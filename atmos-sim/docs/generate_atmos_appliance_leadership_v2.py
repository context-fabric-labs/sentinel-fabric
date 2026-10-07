from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from math import isclose
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import generate_atmos_capacity_portfolio_v2 as visual


DOCS = Path(__file__).resolve().parent
SOURCE = Path.home() / "Downloads" / "Sandisk_ATMOS_AI_Appliance_Portfolio_PRD_v2.docx"
DOC_OUTPUT = DOCS / "Sandisk_ATMOS_AI_Appliance_Portfolio_PRD_v2_Linked.docx"
PPT_OUTPUT = DOCS / "Sandisk_ATMOS_AI_Appliance_Leadership_Proposal_v2.pptx"
STAMP = "6 October 2026"
SOURCE_DIGEST = sha256(SOURCE.read_bytes()).hexdigest()
WHITE, INK, LIGHT = visual.WHITE, visual.INK, visual.LIGHT
RED, TEAL, BLUE, AMBER, MUTED = visual.RED, visual.TEAL, visual.BLUE, visual.AMBER, visual.MUTED

SLIDE_SECTIONS = [
    (1, 3, 20), (2, 15), (2, 12), (3, 4, 7), (4, 16), (5, 6, 8),
    (6, 7, 8, 12), (9, 10), (3, 10, 11), (13, 14, 15), (17, 18), (19, 20), (1, 19, 20),
]


def source_references(document):
    references = {}
    for block in document.tables:
        for row in block.rows:
            values = [cell.text for cell in row.cells]
            if len(values) == 3 and re.fullmatch(r"R\d+", values[0]) and values[2].startswith("http"):
                references[values[0]] = (values[1], values[2])
    return references


def content_signature(document):
    return [
        ("table", [[cell.text for cell in row.cells] for row in block.rows])
        if hasattr(block, "rows") else ("paragraph", block.style.name, block.text)
        for block in document.iter_inner_content()
    ]


def section_map(document):
    sections = {}
    for paragraph in document.paragraphs:
        match = re.match(r"^(\d+)\.\s", paragraph.text)
        if paragraph.style.name == "Title" and match:
            number = int(match.group(1))
            sections[number] = (paragraph, f"PRD_{number:02d}")
        elif paragraph.style.name == "Title" and paragraph.text.startswith("Appendix A"):
            sections["A"] = (paragraph, "PRD_Appendix_A")
    if set(sections) != {*range(1, 21), "A"}:
        raise AssertionError(f"Unexpected PRD section taxonomy: {list(sections)}")
    return sections


def add_navigation_bookmarks(document):
    sections = section_map(document)
    existing = [int(value) for value in document._element.xpath("//w:bookmarkStart/@w:id")]
    identifier = max(existing, default=0) + 1
    for paragraph, name in sections.values():
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(identifier))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(identifier))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        identifier += 1
    return sections


def configure_visuals():
    visual.DOC_OUTPUT = DOC_OUTPUT
    visual.PPT_OUTPUT = PPT_OUTPUT
    visual.STAMP = STAMP


def new_slide(deck, title, section, subtitle="", dark=False):
    return visual.new_slide(deck, len(deck.slides) + 1, title, f"PRD_{section:02d}" if isinstance(section, int)
                            else "PRD_Appendix_A", str(section), subtitle, dark)


def notes(slide, sections, detail, references=()):
    slide.notes_slide.notes_text_frame.text += (
        "\nPRD v2 is the controlling proposal, not an approved product or measured result. "
        "All ATMOS specifications, prices, node counts, supported models and service levels "
        "are planning requirements pending final silicon, OEM qualification, model certification "
        "and matched customer service measurements. The linked companion's text/tables are unchanged.\n"
        + detail + "\n"
    )
    source = Document(SOURCE)
    mapped = section_map(source)
    for section in sections:
        paragraph, bookmark = mapped[section]
        slide.notes_slide.notes_text_frame.text += f"\nPRD: {paragraph.text}\n{DOC_OUTPUT.name}#{bookmark}"
    registered = source_references(source)
    for reference in references:
        title, url = registered[reference]
        slide.notes_slide.notes_text_frame.text += (
            f"\n[{reference}] PRD reference: {title}\n{url}\n"
            "Source supplied with PRD; no ATMOS interoperability, engagement or certification implied."
        )


def label(slide, value, x, y, width=6, color=TEAL):
    visual.section_label(slide, value, x, y, width, color)


def body(slide, value, x, y, width, height, size=16, color=INK, bold=False, link=None):
    return visual.text(slide, value, x, y, width, height, size, color, bold, link)


def footnote(slide, value):
    body(slide, value, 0.55, 6.69, 12.16, 0.25, 10.8, MUTED)


def flow(slide, labels, x, y, widths, height=0.86, colors=None, gap=0.32, font=15):
    cursor = x
    for index, (value, width) in enumerate(zip(labels, widths)):
        visual.flow_node(slide, value, cursor, y, width, height,
                         colors[index] if colors else TEAL, font=font)
        if index + 1 < len(labels):
            visual.arrow(slide, cursor + width + 0.045, y + height / 2 - 0.1, gap - 0.09)
        cursor += width + gap


def slide_charter(deck):
    slide = new_slide(deck, "Sandisk ATMOS AI Appliance Portfolio", 1,
                      "Sell a complete private-AI service, not an E3.S accelerator kit. Leadership product and proof proposal.", True)
    pillars = [
        ("CERTIFIED APPLIANCE", "4 / 8 / 16 / 24-device family\nOEM, topology, thermals and serviceability"),
        ("TURNKEY SOFTWARE", "AI OS + Model Services + Operator\nServing Manager, SDK and Control Center"),
        ("KNOWLEDGE AGENT", "Ingest -> retrieve -> rerank -> cite\nGoverned tools, evaluation and audit"),
        ("ONE SUPPORT ROUTE", "Signed packs, upgrades and rollback\nSandisk triage + coordinated OEM support"),
    ]
    for index, (title, detail) in enumerate(pillars):
        position = 2.05 + index * 0.88
        visual.rect(slide, 0.55, position, 0.055, 0.65, RED if index == 0 else TEAL)
        body(slide, title, 0.79, position + 0.02, 4.35, 0.39, 18, WHITE, True)
        body(slide, detail, 5.36, position, 7.11, 0.7, 16, "D3E1E8")
    visual.rect(slide, 0.55, 5.85, 12.23, 0.79, RED)
    body(slide, "REQUEST: approve 8/16-device POCs, flagship Knowledge Agent, software ownership and parallel OEM feasibility.",
         0.74, 6.0, 11.83, 0.49, 16, WHITE, True)
    notes(slide, (1, 3, 20),
          "Management approves a bounded product/proof program, not a launch date, final bill of "
          "materials, vendor price, supported model or achieved SLA. The PRD's initial wedge is "
          "data-heavy inference and enterprise knowledge; target buyers include CIO/infrastructure, "
          "AI platform, knowledge-management, security, developers and SRE, including constrained "
          "hospital/factory sites. Do not claim training, graphics, coherent terabyte memory, "
          "functional-safety control or universal dense-model TPOT parity. The 24-device SKU "
          "is stretch. Design-to-cost and support lifecycle targets remain proposed.")


def slide_skus(deck):
    slide = new_slide(deck, "Named SKUs. Explicit system and service profiles.", 2,
                      "Five commercial profiles share four device densities; Enterprise 8 adds production integration and serviceability.")
    visual.slide_table(slide, ["Configuration", "Devices / groups", "HBF nominal / study", "DRAM / NVMe", "Service network"], [
        ["Edge Workstation 4", "4 / 1", "512 / 400 GB", "128-256 GB / 2-4 TB", "25/100 GbE"],
        ["Departmental 8", "8 / 2", "1.024 / 0.800 TB", "256-512 GB / 4-8 TB", "100 GbE"],
        ["Enterprise 8", "8 / 2", "Same device envelope", "Qualified 8-device BOM", "Enterprise 8 BOM"],
        ["Enterprise 16", "16 / 4", "2.048 / 1.600 TB", "512 GB-1 TB / 8-16 TB", "2 x 100/200 GbE"],
        ["Enterprise 24 / stretch", "24 / 6", "3.072 / 2.400 TB", "1 TB+ / 16 TB+", "200/400 GbE"],
    ], 0.55, 2.03, [2.4, 1.56, 2.7, 3.29, 2.28], 0.58, 14)
    label(slide, "PROFILE DIFFERENCES / CUSTOMER RECEIVES", 0.55, 5.74, 12)
    body(slide, "Workstation: local catalog + APIs. Departmental: projects, quotas and separate queues. Enterprise: identity, audit, backup and support.",
         0.55, 6.1, 12.12, 0.46, 14.4)
    footnote(slide, "CPU, form factor, redundant power and acoustic/air/DLC qualification are OEM-specific. Aggregate HBF is distributed, not coherent. PRD 2 / 15.")
    notes(slide, (2, 3, 15),
          "The source defines Edge AI Workstation 4, Departmental AI Appliance 8, and Enterprise "
          "AI Appliance 8/16/24. Enterprise 8 has the same device resource envelope as Departmental "
          "8 but a production OEM/service profile; do not invent a separate host/NIC specification. "
          "Study budgets derive from 100 GB/device, not nominal 128 GB. For Workstation 4 the "
          "PRD proposes a pedestal/short-depth 2U, one high-lane socket and acoustic qualification. "
          "Departmental 8 proposes 2U/3U and one high-lane or two balanced sockets. Enterprise "
          "16/24 uses two sockets and qualified 2U/4U rack options with N+1 and high-performance "
          "air or DLC study. Conventional NVMe retains boot, corpus, metadata and rebuild inputs; "
          "ATMOS is not an SSD. No single-node SKU alone promises service HA.")


def slide_device(deck):
    slide = new_slide(deck, "Four-device domains: local work before peer exchange", 2,
                      "One CPU/NUMA root -> x16 uplink -> qualified switch -> four x4 ATMOS modules. Group placement is a software contract.")
    label(slide, "STANDARD SWITCH-LOCAL ALLOCATION DOMAIN", 0.55, 1.98, 6.38)
    visual.flow_node(slide, "CPU / NUMA root\nx16 uplink", 0.55, 2.61, 2.12, 0.88, BLUE, font=15)
    visual.flow_node(slide, "PCIe P2P switch\n4 x x4 ports", 0.55, 3.77, 2.12, 0.9, TEAL, font=15)
    visual.line(slide, 1.61, 3.49, 1.61, 3.77, TEAL)
    visual.line(slide, 2.67, 4.22, 3.01, 4.22, TEAL)
    visual.line(slide, 3.01, 2.51, 3.01, 4.75, TEAL)
    for index in range(4):
        position = 2.35 + index * 0.66
        visual.flow_node(slide, f"ATMOS {index}: HBF + NPU + LPDDR", 3.43, position, 3.36, 0.58, TEAL, font=13.4)
        visual.line(slide, 3.01, position + 0.29, 3.43, position + 0.29, TEAL)
    label(slide, "PER-DEVICE PLANNING ENVELOPE", 7.33, 1.98, 5.44)
    visual.slide_table(slide, ["Resource", "Target / constraint"], [
        ["HBF", "128 GB nominal / 100 GB study"],
        ["Local read / NPU", "~200 GB/s / ~64 TOPS-class"],
        ["Dynamic state", "8-16 GB LPDDR + SRAM"],
        ["Module / link", "E3.S 2T / Gen5 or Gen6 x4"],
        ["Sustained power", "30-40 W; accelerator duty"],
    ], 7.33, 2.36, [1.52, 3.92], 0.48, 14)
    envelopes = [(4, "0.512 TB", "120-160 W"), (8, "1.024 TB", "240-320 W"),
                 (16, "2.048 TB", "480-640 W"), (24, "3.072 TB", "720-960 W")]
    for index, (devices, capacity, power) in enumerate(envelopes):
        position = 0.55 + index * 3.1
        visual.rect(slide, position, 5.62, 2.92, 0.95, LIGHT)
        body(slide, f"{devices} devices / {capacity}", position + 0.12, 5.73, 2.69, 0.33, 17, TEAL, True)
        body(slide, f"Device-only power: {power}", position + 0.12, 6.16, 2.69, 0.26, 11.4, MUTED)
    footnote(slide, "P2P is not coherent memory or local HBF bandwidth. Qualify Gen5/Gen6 separately, per-device fit, IOMMU, safe reset and cross-domain traffic. PRD 2 / 12.")
    notes(slide, (2, 12, 15),
          "All targets are unverified planning assumptions. The source's 64 TOPS-class target "
          "does not define precision/sparsity and is not comparable to other vendors' peak "
          "compute. Arithmetic local-read sums are 0.8/1.6/3.2/4.8 TB/s for 4/8/16/24 "
          "devices, but not coherent or per-request throughput. Each group has four x4 "
          "endpoints and one x16 uplink; 16 and 24 use four and six domains. Large model/vector "
          "objects stay local; compact reductions, replicas/recovery can use qualified peers. "
          "Cross-group/host-root paths require measured service and cooperating-driver support. "
          "Keep IOMMU mappings bounded and invalidate on reset; never trade isolation for "
          "a P2P benchmark. Whole-group allocation, device health, DMA and pack-aware placement "
          "are ATMOS deliverables, not automatically supplied by DRA or a PCIe switch.", ("R15",))


def slide_software(deck):
    slide = new_slide(deck, "The software is the product: deliver an integrated stack", 4,
                      "Packaged components, stable contracts and operations—not a requirement for the customer to assemble an inference platform.")
    label(slide, "EXPERIENCE AND KNOWLEDGE", 0.55, 2.03, 3.75)
    label(slide, "SERVING AND DEVICE EXECUTION", 4.7, 2.03, 4.09)
    label(slide, "TRUST AND FLEET OPERATIONS", 9.2, 2.03, 3.58, RED)
    left = [
        "Knowledge Agent / UI\nSearch, RAG, approved workflows",
        "SDK / API portal / notebooks\nPython, REST, examples and profiling",
        "Knowledge plane\nConnectors, parsers, ACLs and indexes",
        "Certified ModelPacks\nWeights, tokenizer, kernels and tests",
    ]
    middle = [
        "Gateway\nAuth, quota, API policy and accounting",
        "Serving Manager\nKServe/llm-d candidate; replicas, queues",
        "Model Services / ATMOS runtime\nCertified operators and memory plans",
        "AI OS + Device Service\nDRA/CDI/operator, DMA, P2P, firmware",
    ]
    right = [
        "Control Center\nFleet, SLO, updates and cases",
        "Identity / secrets / network\nOIDC/SAML, RBAC, vault, TLS/mTLS",
        "Supply chain / policy / evidence\nSigning, SBOM, policy and audit",
        "Enterprise Support\nProd/LTS, CVEs, OEM escalation",
    ]
    for index, (left_text, middle_text, right_text) in enumerate(zip(left, middle, right)):
        position = 2.48 + index * 0.91
        visual.flow_node(slide, left_text, 0.55, position, 3.75, 0.75, BLUE, font=14)
        visual.flow_node(slide, middle_text, 4.7, position, 4.09, 0.75, TEAL, font=14)
        visual.flow_node(slide, right_text, 9.2, position, 3.58, 0.75, RED, font=13.8)
        if index < 3:
            visual.line(slide, 6.745, position + 0.75, 6.745, position + 0.91, TEAL)
    visual.arrow(slide, 4.35, 2.76, 0.28)
    visual.line(slide, 8.79, 3.77, 9.2, 3.77, RED)
    footnote(slide, "ATMOS runtime/driver/operator, HBF placement and workload adapters must be delivered/qualified; upstream projects do not establish ready ATMOS support. PRD 3-4 / 12.")
    notes(slide, (3, 4, 7, 12),
          "The PRD names nine software products: AI OS, Device Service, Model Services, Certified "
          "ModelPacks, Serving Manager, Knowledge Agent, Control Center, SDK and Enterprise "
          "Support. Candidate gateway/serving components are Envoy AI Gateway or equivalent, "
          "KServe LLMInferenceService, llm-d or Sandisk-packaged alternatives. Those projects "
          "supply patterns/components, not an ATMOS-certified backend. Sandisk must own the "
          "device/runtime and signed distribution and select/name the NPU compiler/runtime "
          "partner. ModelPack binds executable preparation separately from weights-as-data; "
          "no custom compiler CLI is claimed available. Readiness depends on signed pack, "
          "ABI, per-device memory fit, golden correctness and warmup. Customer-safe configuration "
          "and supported SDK extensions are different from firmware/kernel or unchecked binaries.", ("R16", "R17"))


def slide_apis(deck):
    slide = new_slide(deck, "Seven service APIs. A bounded certified workload catalog.", 4,
                      "Support follows model, input, candidate, policy and SLO certification—not compatibility by endpoint name alone.")
    visual.slide_table(slide, ["API family", "Proposed surface / service contract"], [
        ["OpenAI-compatible", "/v1/chat/completions, /responses, /models; certified params"],
        ["Embeddings", "/v1/embeddings; batch/indexing; model and input limits"],
        ["Retrieval", "/v1/retrieve; hybrid search, collections, filters and metadata"],
        ["Reranking", "/v1/rerank; batch; certified candidate limits"],
        ["Documents", "Ingest, status, delete, reindex, version and extraction results"],
        ["Agents", "Sessions/runs, approved tools, citations and policy output"],
        ["Administration", "Projects/users/keys/quotas, models, updates and support"],
    ], 0.55, 2.08, [2.06, 5.76], 0.52, 13.5)
    label(slide, "CERTIFY FIRST / LAUNCH CANDIDATES", 8.81, 2.1, 3.97)
    body(slide, "Embeddings, hybrid retrieval, reranking, cited Knowledge Agent and qualified document services.",
         8.81, 2.56, 3.97, 1.03, 16)
    label(slide, "MODEL-SPECIFIC PILOTS / RESEARCH", 8.81, 3.91, 3.97, AMBER)
    body(slide, "7B-32B private LLM: quality + TTFT/TPOT/context proof. Recommendation features: pilot. MoE prefill/cold experts: research/pilot.",
         8.81, 4.36, 3.97, 1.44, 15.5)
    body(slide, "Not launch targets: long-context dense decode, training, graphics or safety-critical control.",
         0.55, 6.34, 12.09, 0.33, 14, RED, True)
    footnote(slide, "GA candidate is a target, not a released certification. Healthcare/industrial packs provide knowledge assistance, not diagnostic or control claims. PRD 4 / 16.")
    notes(slide, (4, 7, 8, 16),
          "Preserve the actual seven API families. /v1/responses is a proposed supported API "
          "only for certified models/documented parameters, not a claim of full external-provider "
          "equivalence. Retrieval and rerank have separate metadata/candidate contracts. "
          "Multitenancy scopes models, data/indexes, API keys and audit by tenant/project, with "
          "separate bounded interactive/batch queues and resource claims. Hard isolation can use "
          "dedicated nodes/replicas. The catalog treats embeddings/hybrid retrieval/rerank/Knowledge "
          "Agent as GA candidates, documents as GA/vertical candidates, recommendation features "
          "as pilot and each private LLM as model-specific pilot/GA. No aggregate HBF/TOPS "
          "proves full attention/KV/operator support. Vertical templates cover healthcare, "
          "manufacturing, regulated enterprise, research library and support with explicit claims limits.")


def slide_knowledge(deck):
    slide = new_slide(deck, "Out-of-box Knowledge Agent: governed evidence to answers", 5,
                      "The flagship solution ships ingestion, hybrid retrieval, reranking, citations, evaluation and bounded agent workflows.")
    label(slide, "INGESTION / SOURCE, VERSION, TENANT AND ACL LINEAGE", 0.55, 2.0, 12)
    flow(slide, ["Approved sources\nfiles / S3 / enterprise", "Extract + chunk\nparser / OCR profile",
                 "Embed + sparse terms\ncertified models", "Versioned collections\nindexes + lineage"],
         0.55, 2.39, [2.81, 2.81, 2.81, 2.82], 0.82, font=14.5)
    label(slide, "ONLINE QUERY / ACL FILTERING BEFORE EVIDENCE REACHES GENERATION", 0.55, 3.48, 12, RED)
    flow(slide, ["Query / understand", "ACL + hybrid\nretrieve", "Rerank / select", "Generate + cite", "Verify + audit"],
         0.55, 3.87, [2.19, 2.19, 2.19, 2.2, 2.2], 0.83, gap=0.32, font=14)
    label(slide, "SIX INCLUDED MODES / SELECT WITHIN THE CERTIFIED ENVELOPE", 0.55, 5.02, 12)
    modes = [("Search", "Ranked passages only"), ("Standard RAG", "One retrieval / answer pass"),
             ("Deep RAG", "Decompose, broaden, synthesize"), ("Agentic RAG", "Plan, tools, retry, verify"),
             ("Retrieval-only", "Evidence to an existing LLM"), ("Batch knowledge", "Summaries, extract, tag")]
    for index, (title, description) in enumerate(modes):
        column, row = index % 3, index // 3
        left, top = 0.55 + column * 4.16, 5.41 + row * 0.54
        body(slide, title, left, top, 1.64, 0.37, 13.2, TEAL, True)
        body(slide, description, left + 1.68, top, 2.3, 0.42, 12.4)
    footnote(slide, "Agentic mode adds model calls, latency and budget risk. Search/retrieval-only can return evidence without ATMOS generation. Full generation requires model certification. PRD 5-8.")
    notes(slide, (5, 6, 8),
          "The source's six modes are search, standard RAG, deep RAG, agentic RAG, retrieval-only "
          "and batch knowledge jobs. The flow summarizes the standard generated-answer mode; "
          "search/retrieval-only return evidence and do not require the ATMOS generator. "
          "Agentic mode can loop through planning, sub-tasks, tool calls, retries, synthesis "
          "and optional verification, but only under explicit step/tool/time/token/cost limits. "
          "Launch connectors are NFS/SMB read-only, S3-compatible, controlled uploads, supported "
          "SharePoint/OneDrive equivalents, approved database exports and allow-listed web/intranet. "
          "Parsers support certified PDF/text/HTML/Markdown/Office; OCR, images and audio/video "
          "need explicit profiles/licenses. Every chunk retains source page/section, document "
          "version, tenant and ACL. Failed ingestion is quarantined and must not expose partial data.",
          ("R9", "R10", "R12", "R13", "R14"))


def slide_governance(deck):
    slide = new_slide(deck, "Governed by default. Configurable within certified limits.", 7,
                      "Tenant/source ACLs must filter retrieval before generation; post-generation masking is not the security boundary.")
    flow(slide, ["Identity + tenant\nOIDC/SAML, RBAC", "ACL + freshness\napproved evidence",
                 "Tools + budget policy\nscoped credentials", "Output + immutable audit\ncitations, model and decisions"],
         0.55, 2.07, [2.81, 2.81, 2.81, 2.82], 0.9, [BLUE, RED, TEAL, BLUE], font=14.5)
    columns = [
        ("CUSTOMER CONFIGURATION", "Collections, source schedules and retention. Prompts/citations and certified retrieval settings. Approved tools, models, modes, quotas and golden tests.", TEAL),
        ("SUPPORTED EXTENSIONS", "REST/OpenAI + Python SDK, webhooks and approved MCP. Signed restricted connectors and supported workflows. BYO models through Sandisk certification.", BLUE),
        ("LOCKED PRODUCTION", "Firmware, driver, kernels and switch policy. Unsigned binaries, hidden tokenizer/precision changes, unbounded tools and cross-tenant P2P are not self-service.", RED),
    ]
    for index, (title, detail, color) in enumerate(columns):
        position = 0.55 + index * 4.16
        label(slide, title, position, 3.39, 3.96, color)
        body(slide, detail, position, 3.86, 3.96, 1.67, 16)
    visual.rect(slide, 0.55, 5.84, 12.23, 0.72, LIGHT)
    body(slide, "BASE TOOLS: read-only search/fetch/lookup; sandboxed calculator; draft-only messaging. Actions need approval; code execution is not in base.",
         0.7, 5.98, 11.91, 0.47, 14.5)
    footnote(slide, "Trust: signed firmware/images/ModelPacks, SBOM/CVE, TPM/secure boot, vault, TLS/mTLS, retention/deletion and audit. No unearned compliance certification. PRD 6-8 / 12.")
    notes(slide, (6, 7, 8, 12),
          "The retrieval boundary is tenant and source ACL filtering before evidence reaches the "
          "generator. Source/path/schema allow-lists, scoped connector credentials, retention and "
          "deletion apply throughout ingestion and retrieval. The tool catalog defaults search, "
          "document fetch and structured lookup to read-only; calculator is sandboxed with no "
          "external network. Ticket creation needs human confirmation and a scoped credential; "
          "email/message is draft-only; workflow actions are disabled until signed/policy-approved. "
          "Code execution is a separate sandboxed premium profile, not base. Agent limits cover "
          "steps, calls, retrieved bytes, tokens, elapsed time and cost. Evaluate recall/NDCG, "
          "groundedness, citations, task/tool completion, leakage/unauthorized actions and tails. "
          "Developer mode is explicitly nonproduction and isolated from production data. "
          "Secure/measured boot, signed supply chain, encryption, vault, TLS/mTLS and audit are "
          "requirements that support assessments, not an automatic regulated certification.", ("R15",))


def slide_lifecycle(deck):
    slide = new_slide(deck, "Model and platform upgrades are controlled product releases", 9,
                      "ModelPack binds weights, tokenizer, preprocessing, precision, compiler output, runtime, golden tests, APIs and SLO.")
    label(slide, "MODELPACK STATE MACHINE", 0.55, 2.0, 12)
    flow(slide, ["Available\ncatalog", "Staged\nsign + reserve", "Validated\nquality + hardware",
                 "Canary\nshadow / small share", "Production\ndefault version", "Draining\nno new sessions"],
         0.55, 2.4, [1.7, 1.7, 1.7, 1.7, 1.7, 1.7], 0.88, gap=0.4, font=13)
    body(slide, "Failed canary -> Rolled back. Draining -> Retired. Active sessions remain generation-pinned; preserve a rollback generation.",
         0.55, 3.52, 12.13, 0.38, 14.1, RED, True)
    label(slide, "GENERATOR / RERANKER UPDATE", 0.55, 4.05, 5.75)
    body(slide, "Stage -> compatibility/golden tests -> shadow/canary -> compare quality/p99/errors -> promote gradually; retain rollback.",
         0.55, 4.49, 5.78, 0.79, 15.6)
    label(slide, "EMBEDDING OR INDEX-COMPATIBILITY CHANGE", 6.83, 4.05, 5.95, AMBER)
    body(slide, "Build/re-embed a shadow collection -> retrieval evaluation / optional dual-read -> atomic generation switch -> hold old index.",
         6.83, 4.49, 5.95, 0.79, 15.6)
    label(slide, "PLATFORM UPDATE / ADVANCE ONE FAILURE DOMAIN AT A TIME", 0.55, 5.51, 12)
    flow(slide, ["Preflight + backup\nhealth, space, recovery", "Drain + signed bundle\nboot candidate slot",
                 "Smoke + canary\nhardware/P2P/model/SLO", "Promote or roll back\nthen next domain"],
         0.55, 5.89, [2.81, 2.81, 2.81, 2.82], 0.62, font=13)
    footnote(slide, "Offline: signed bundles, customer windows, no cloud required. Breaking parser/index/model/policy changes need migration + rollback. PRD 9-10.")
    notes(slide, (9, 10),
          "All eight PRD ModelPack states appear: Available, Staged, Validated, Canary, Production, "
          "Draining, Rolled back and Retired. A same-embedding generator/reranker update does not "
          "silently reindex the corpus. An embedding-model change creates a shadow collection, "
          "re-embeds incrementally without modifying the active index, evaluates retrieval/end-to-end "
          "quality, optionally dual-reads representative traffic, then atomically switches. "
          "Plan two generations and rollback space before approval. Customer fine-tuned weights, "
          "new quantization and BYO embedding models go through certification; community models "
          "stay developer-only until certified. Platform bundles cover OEM BIOS/BMC, switch/retimer "
          "and ATMOS firmware, Linux/kernel/driver, NPU compiler/runtime/operators, Kubernetes/DRA, "
          "gateway/identity/observability, model compatibility, SBOM/CVEs and rollback evidence. "
          "Offline mode uses signed repository bundles, customer-controlled windows and exportable "
          "evidence without mandatory cloud connectivity. No support duration is an approved contract.")


def slide_operations(deck):
    slide = new_slide(deck, "Day-1 setup. Day-2 operation. One support route.", 11,
                      "First-service setup within one business day is a target; corpus ingestion time depends on data volume, connectors and policy.")
    flow(slide, ["Site readiness", "Secure enrollment", "Identity + sources", "Certified profile", "Ingest + evaluate", "Users + operate"],
         0.55, 2.05, [1.77] * 6, 0.52, gap=0.32, font=12.5)
    label(slide, "CONTROL CENTER / CORRELATED FLEET AND SERVICE VIEWS", 0.55, 2.93, 6.08)
    views = [
        ("Fleet", "BOM, serial, warranty, entitlement"), ("Devices", "HBF/NPU, health, resets, thermal"),
        ("Topology", "NUMA, PCIe width, switch/P2P"), ("Services", "Models, sessions, queues, SLO"),
        ("Knowledge", "Ingestion, collections, generations"), ("Tenants", "Usage, quota, keys, chargeback"),
        ("Updates", "Canary, drift, compatibility, rollback"), ("Support", "Alerts, evidence, case, OEM dispatch"),
    ]
    for index, (title, detail) in enumerate(views):
        column, row = index % 2, index // 2
        left, top = 0.55 + column * 3.13, 3.34 + row * 0.65
        visual.rect(slide, left, top, 2.96, 0.59, LIGHT)
        body(slide, title, left + 0.09, top + 0.055, 2.77, 0.24, 12.3, TEAL, True)
        body(slide, detail, left + 0.09, top + 0.33, 2.77, 0.22, 10.9)
    label(slide, "HA REQUIRES COMPLETE REPLICAS AND PEAK RESERVE", 7.14, 2.93, 5.64, RED)
    flow(slide, ["Active node 1", "Active node 2", "N+1 reserve"], 7.14, 3.4,
         [1.72, 1.72, 1.72], 0.62, [BLUE, BLUE, TEAL], gap=0.24, font=13)
    body(slide, "One node: no service HA. Two nodes: replicas, not guaranteed full peak after a loss. Three-node N+1: two-node active capacity plus reserve.",
         7.14, 4.28, 5.63, 0.93, 15)
    label(slide, "SUPPORT OWNERSHIP", 7.14, 5.41, 5.64)
    body(slide, "Sandisk: first call + OS/device/runtime/packs. OEM: chassis/platform, coordinated entitlement. NPU vendor and certified ISV: named escalation.",
            7.14, 5.81, 5.63, 0.65, 13.8)
    body(slide, "Release targets: Production 3y | LTS 5y | optional +2y.\nVertical 7-10y only with a funded business case; not contracted.",
            0.55, 6.1, 6.07, 0.54, 10.9, BLUE, True)
    footnote(slide, "Support/lifecycle durations are proposed targets, not contracted entitlements. Evidence bundles redact inventory, versions, SLO/error, thermal and audit data. PRD 3 / 10-11.")
    notes(slide, (3, 10, 11),
          "Day 0 supplies rack/office power, network, identity/DNS/NTP and sources. Day 1 initializes "
          "hardware burn-in, enrollment, admin, topology/thermal preflight, identity/TLS, projects "
          "and certified signed packs. Day 2 runs ingestion and users/apps with quota, metering "
          "and audit. Do not promise completion of a million-document corpus in one day. "
          "Single Workstation 4 has no service HA; a single appliance can degrade on a device/group "
          "fault but has reduced maintenance capacity. Two nodes maintain service only with "
          "appropriate replicas and sufficient remaining capacity; N+1 must survive declared "
          "peak demand with complete model/index state. Multi-site data replication is solution-specific. "
          "Sandisk is prime for device/firmware/driver/runtime/operator/packs and appliance distribution; "
          "OEM covers CPU/DRAM/PSU/fans/NIC/BIOS/BMC under coordinated entitlement; NPU partner "
          "covers compiler/runtime internals; connector/application ownership is named in the "
          "solution manifest. Developer/Feature/Production/LTS/Extended/Vertical channels differ "
          "in cadence and support. Final SLAs, geographic coverage and durations need funding/contracts.")


def slide_oem(deck):
    slide = new_slide(deck, "Adapt an established OEM platform—qualify the whole BOM", 13,
                      "Parallel feasibility, not a partner commitment: early engineering flexibility, enterprise service and regulated lifecycle.")
    visual.slide_table(slide, ["Candidate", "PRD rationale / public platform", "Product role", "ATMOS qualification risk"], [
        ["Supermicro", "X14 Petascale / E3.S-rich 24-32 class", "Early 8/16 engineering; flexible configuration", "2T/x4/bay power, P2P backplane and productized support"],
        ["Dell", "PowerEdge R7725 / R7725xd; dense E3.S and root maps", "Enterprise channel, OpenManage and field service", "Direct CPU attachment may need a new switch/P2P backplane"],
        ["HPE", "DL380/DL345 Gen12; iLO, cooling and x4 cabling options", "Regulated edge, lifecycle and management", "Dedicated thermal BOM; GPUs/NICs/DIMMs constrain ambient"],
        ], 0.55, 2.04, [1.5, 3.56, 3.0, 4.17], 0.85, 14.5)
    label(slide, "QUALIFICATION LADDER / DO NOT INFER SUPPORT FROM STORAGE BAY COUNT", 0.55, 5.58, 12)
    flow(slide, ["1 module\npower + throttle", "4-device group\nheat + P2P + removal", "8-device POC\nCPU/NIC + 72h soak",
                 "16 standard\nmax inlet + N+1", "24 stretch\n6 groups + rack proof"],
            0.55, 5.98, [2.19, 2.19, 2.19, 2.2, 2.2], 0.62, gap=0.32, font=11.8)
    footnote(slide, "Sign exact 2T/x4 topology, sustained bay/switch cooling, BIOS/BMC, DMA isolation, safe replacement, spares, NRE and prime support. PRD 13-15.")
    notes(slide, (13, 14, 15),
          "The source proposes Supermicro for early engineering, parallel Dell/HPE feasibility "
          "for enterprise pilot, one primary OEM and second-source roadmap for launch. No vendor "
          "agreement, ATMOS compatibility or launch date is established by public E3.S/storage "
          "specifications. Verify thickness, actual lane width/gen, root/switch layout, power, "
          "airflow and reset/error recovery. Direct CPU attachment is not the proposed switch-local "
          "peer domain. Storage 1T/x2 populations cannot be advertised as 2T/x4 accelerator bays. "
          "HPE ambient restrictions involving DIMMs, NIC/GPU cards and E3.S demonstrate need for "
          "a dedicated thermal configuration. Test switch/retimer and connector temperatures, "
          "one-fan degradation, full CPU/NIC/service load, 72h soak, N+1 power/cooling, inlet "
          "and service stability. The 24-device system has six groups and is a gated stretch SKU. "
          "Evaluate a tower/pedestal OEM for Workstation 4 acoustics; do not force Departmental "
          "8 into a quiet desktop. All quoted chassis candidates remain subject to exact OEM selection.",
          ("R1", "R3", "R4", "R5", "R6", "R7", "R8"))


def counterpart_tco(atmos_systems=1.0):
    baseline = 1 + 0.05 + 3 * (0.12 + 0.07)
    acquisition = 0.70
    atmos = atmos_systems * acquisition * (1 + 0.08 + 3 * (0.14 + 0.08))
    return baseline, atmos


def slide_comparison(deck):
    baseline, atmos = counterpart_tco()
    saving = 1 - atmos / baseline
    break_even = baseline / atmos
    slide = new_slide(deck, "Counterpart positioning and the conditional cost target", 17,
                      "The only NVIDIA comparison page: workload-specific positioning, not a universal replacement or a measured price/performance claim.")
    visual.slide_table(slide, ["PRD reference", "Published reference memory", "Closest SKU / limit"], [
        ["DGX Spark", "128 GB unified; 273 GB/s", "W4 / A8; capacity is not compute parity"],
        ["RTX PRO 6000", "96 GB GDDR7", "A8; no graphics/training parity"],
        ["DGX Station", "748 GB coherent; HBM up to 7.1 TB/s", "A8; much lower local compute/bandwidth"],
        ["RTX PRO Server", "2/4/8 x 96 GB; ~1.6 TB/s per GPU", "E8/16/24; certify the exact service"],
        ["IGX Thor", "128 GB + optional dGPU", "E8/16 backend, not a safety/controller replacement"],
    ], 0.55, 2.09, [2.04, 2.61, 2.73], 0.64, 13)
    label(slide, "70% ACQUISITION: DESIGN-TO-COST OBJECTIVE", 8.38, 2.09, 4.4, AMBER)
    body(slide, f"{saving:.1%}", 8.38, 2.61, 4.18, 0.86, 47, TEAL, True)
    body(slide, f"Illustrative 3-year TCO saving\n{atmos:.3f}N versus {baseline:.2f}N", 8.38, 3.65, 4.36, 0.76, 17)
    counterpart_price = 180_000
    body(slide, "Synthetic 4-GPU-server-class N=$180k:\n$126k ATMOS acquisition; 3-year TCO\n$219,240 versus $291,600.",
         8.38, 4.71, 4.35, 0.97, 15.3)
    body(slide, f"Break-even: <{break_even:.2f} ATMOS systems per counterpart. At 1.25x: ~6% saving. At 2x: ~50% higher TCO.",
         0.55, 6.12, 12.12, 0.44, 15, RED, True)
    footnote(slide, "PRD scenario only: same quality, SLO, peak demand and availability; 3y undiscounted. All prices synthetic; use quotes, measured node counts and removable spend. PRD 17-18.")
    notes(slide, (17, 18),
          "Keep the PRD's consolidated reference section in one deck page. Published reference "
          "memory types, compute precision/sparsity and platform boundaries are not interchangeable "
          "with distributed HBF or the ATMOS 64-TOPS target. Reference configurations and exact OEM "
          "quotes must be frozen before external comparative claims. N is synthetic counterpart "
          "acquisition cost, not a negotiated quote. Baseline TCO=1.00N + 0.05N integration + "
          "3*(0.12N software/support + 0.07N facility/energy)=1.62N. ATMOS acquisition=0.70N; "
          "integration=8%, annual support=14%, facility/energy=8% of ATMOS acquisition; total "
          "0.70*(1.08+3*0.22)N=1.218N. Saving=(1.62-1.218)/1.62=24.8148%. Break-even "
          "system ratio=1.62/1.218=1.33005, under proportional cost and equal accepted service. "
          "1.25 systems: 6.0185% saving; 2 systems: 50.3704% higher TCO. Discrete actual "
          "node counts, N+1, staging/index/model generations, migration, subscription/support, "
          "OEM NRE and measured power can change the result. Freeze demand, quality, p99, "
          "availability and best CPU/managed/service alternative. Customer spend must be avoidable. "
          "Required finance evidence includes ATMOS BOM/yield/warranty/margin, OEM quote/NRE, "
          "software/model certification staffing, accepted-service rates and finance-approved DCF.",
          ("R20", "R21", "R22", "R23", "R24"))


def slide_gates(deck):
    slide = new_slide(deck, "Release gates convert a PRD into a supportable product", 19,
                      "Evidence sequence, not a calendar promise. Approve claims only after certified quality, service, availability and economics.")
    gates = [
        ("G0 / CHARTER", "PRD, personas, SKUs, catalog\nManagement approval"),
        ("G1 / DEVICE", "Service curves, compiler, thermals\nThree launch workloads feasible"),
        ("G2 / GROUP", "4-device P2P, isolation, resets\nQualified hardware building block"),
        ("G3 / 8-DEVICE OOB", "Knowledge Agent pilot\nFirst-service setup within one day"),
        ("G4 / SOFTWARE BASELINE", "Packs, catalog, canary and rollback\nProduction update/support workflow"),
        ("G5 / 16-DEVICE OEM", "HA, security and 72-hour soak\nOrderable qualified pilot SKU"),
        ("G6 / VERTICAL PACKS", "Healthcare/manufacturing + ISVs\nReference-customer evidence"),
        ("G7 / 24-DEVICE STRETCH", "Six groups, thermal, N+1 and cost\nAdvance only after standard proof"),
        ("G8 / COMPARATIVE CLAIMS", "Matched service + TCO evidence\nLegal, product and finance approval"),
    ]
    for index, (title, description) in enumerate(gates):
        column, row = index % 3, index // 3
        left, top = 0.55 + column * 4.16, 2.05 + row * 1.36
        visual.rect(slide, left, top, 3.9, 1.08, LIGHT)
        body(slide, title, left + 0.12, top + 0.12, 3.66, 0.35, 14, TEAL if row < 2 else AMBER, True)
        body(slide, description, left + 0.12, top + 0.6, 3.66, 0.43, 12.3)
        if column < 2:
            visual.arrow(slide, left + 3.94, top + 0.48, 0.15)
    body(slide, "First POCs: 8-device Knowledge Agent / 1-5m pages or chunks; 16-device retrieval/rerank; 8-device private AI; thermal and lifecycle tests.",
         0.55, 6.28, 12.12, 0.36, 13.4)
    footnote(slide, "Fail on quality/ACL/safety, unsafe DMA, missed p99/freshness or peak HA, thermal derating, incomplete software/support, or adverse accepted-service TCO. PRD 19-20.")
    notes(slide, (19, 20),
          "The nine PRD gates are management charter, single device/three launch workloads, "
          "four-device domain, departmental 8 with OOB setup, production software lifecycle, "
          "qualified 16 with HA/security/72-hour soak, vertical/ISV/reference customers, "
          "qualified 24 and comparative claim approval. The 1-5m pages/chunks Knowledge Agent "
          "scope is a POC dataset target, not an achieved index or same-day ingestion promise. "
          "Retrieval/rerank needs candidate-distribution recall/NDCG, p99 and cost/query. "
          "Private AI needs multi-user quotas, batch and one certified full model. Lifecycle "
          "POC exercises two model generations and a platform update through canary/rollback. "
          "Simulation/emulation can reject or narrow assumptions but cannot establish physical "
          "OEM thermals, accepted service or support readiness. Owners must be named across "
          "product, NPU/runtime, platform/OEM, security, applications, SRE/support, sourcing "
          "and finance. 24-device and broad replacement claims remain held until their gates.")


def slide_decisions(deck):
    slide = new_slide(deck, "Approve bounded product proof—not a launch claim", 20,
                      "Leadership needs scope, accountable software/service owners, an OEM qualification path and a finance gate.")
    label(slide, "APPROVE / ASSIGN / HOLD", 0.55, 2.01, 7.04)
    decisions = [
        ("PRODUCT", "Approve 8/16-device POCs and the OOB Knowledge Agent; retain strict workload and claims boundaries."),
        ("SOFTWARE", "Name owners for AI OS, driver/runtime, Model Services/Operator, Control Center, SDK and release/support."),
        ("PLATFORM", "Start Supermicro/Dell/HPE feasibility; freeze 2T/x4/P2P, thermal, HA, firmware and service contracts."),
        ("COMMERCIAL", "Keep 70% price as a target; require quotes, certification staffing, measured unit costs and cashable displacement."),
        ("HOLD", "24-device launch, unsupported models, compliance/diagnostic/safety claims and universal replacement promises."),
    ]
    for index, (title, description) in enumerate(decisions):
        position = 2.45 + index * 0.79
        body(slide, title, 0.55, position, 1.54, 0.4, 12.2, RED if title == "HOLD" else TEAL, True)
        body(slide, description, 2.22, position, 5.5, 0.68, 15)
    visual.rect(slide, 8.15, 2.0, 4.63, 4.61, LIGHT)
    label(slide, "CLICK TO OPEN PRD DETAILS", 8.32, 2.18, 4.29, BLUE)
    navigation = [
        ("2 / 15  Hardware and SKU specs", 2),
        ("3-4  Turnkey stack and APIs", 3),
        ("5-8  Knowledge, data, tools, safety", 5),
        ("9-10  Models and platform upgrades", 9),
        ("11-12  Operations, support, security", 11),
        ("13-15  OEM design and qualification", 13),
        ("16  Workload and vertical catalog", 16),
        ("17-18  Counterparts and TCO", 17),
        ("19-20  Gates, risks and decisions", 19),
        ("A  PRD public source register", "A"),
    ]
    for index, (title, section) in enumerate(navigation):
        bookmark = f"PRD_{section:02d}" if isinstance(section, int) else "PRD_Appendix_A"
        body(slide, title, 8.32, 2.73 + index * 0.34, 4.28, 0.27, 12.6, BLUE,
             link=f"{DOC_OUTPUT.name}#{bookmark}")
    footnote(slide, "Keep PPTX + linked PRD together. Each slide offers a full-document link if the viewer ignores section bookmarks. Original PRD content is preserved.")
    notes(slide, (1, 19, 20),
          "Leadership decisions are those requested in the PRD executive summary: approve "
          "8-device departmental and 16-device enterprise proof; flagship Enterprise Knowledge "
          "Agent; AI OS/Model Services/Operator/Control Center/ModelPack lifecycle investment; "
          "parallel Supermicro/Dell/HPE feasibility; and a conditional design-to-cost target. "
          "No staffing count, R&D budget, partner agreement, silicon availability or launch "
          "date is invented. Required financial evidence precedes funding/product/price approval. "
          "Risk register covers compute/operator deficit, distributed-memory fit, OEM bay mismatch, "
          "software ecosystem gap, agent safety, model/index churn, node-count economics, "
          "fragmented support and regulated claims. Full-document and section/bookmark links "
          "are relative to the linked PRD copy. Keep files in one directory; Office security "
          "policies may prompt and some viewers ignore fragments. The source PRD's 281 "
          "paragraphs and 55 tables are not rewritten; only navigation bookmarks are added.")


def build_deck(document):
    configure_visuals()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for builder in (slide_charter, slide_skus, slide_device, slide_software, slide_apis, slide_knowledge,
                    slide_governance, slide_lifecycle, slide_operations, slide_oem, slide_comparison,
                    slide_gates, slide_decisions):
        builder(deck)
    core = deck.core_properties
    core.title = "Sandisk ATMOS AI Appliance Portfolio: Leadership Proposal"
    core.subject = "Source-controlled PRD requirements, product scope and investment gates"
    core.author = core.last_modified_by = "Sandisk / ATMOS Product and Engineering"
    core.revision = 2
    core.created = core.modified = datetime.now(timezone.utc)
    return deck


def validate(document, deck):
    original = Document(SOURCE)
    if content_signature(original) != content_signature(document):
        raise AssertionError("Linked PRD's substantive content changed")
    if sha256(SOURCE.read_bytes()).hexdigest() != SOURCE_DIGEST:
        raise AssertionError("Original source PRD changed")
    bookmarks = set(document._element.xpath("//w:bookmarkStart/@w:name"))
    if len(deck.slides) != 13:
        raise AssertionError("Expected the thirteen-slide leadership deck")
    baseline, atmos = counterpart_tco()
    if not isclose(baseline, 1.62) or not isclose(atmos, 1.218):
        raise AssertionError("TCO scenario no longer reconciles to the PRD")
    if not isclose(counterpart_tco(baseline / atmos)[1], baseline):
        raise AssertionError("Node-count break-even does not reconcile")
    body_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    for required in ("AI OS", "Model Services", "Control Center", "ModelPacks", "DRA/CDI/operator",
                     "Workstation 4", "Departmental 8", "Enterprise 8", "Enterprise 16", "Enterprise 24",
                     "Knowledge Agent", "OpenAI-compatible", "Embeddings", "Retrieval", "Reranking",
                     "Documents", "Agents", "Administration", "Available", "Staged", "Validated",
                     "Canary", "Production", "Draining", "Rolled back", "Retired", "ACL", "N+1",
                     "3y", "5y", "Supermicro", "Dell", "HPE", "24.8%", "G0", "G8"):
        if required.lower() not in body_text.lower():
            raise AssertionError(f"Visible requirement missing: {required}")
    note_sections = set()
    for slide in deck.slides:
        note_sections.update(int(value) for value in re.findall(r"PRD: (\d+)\.", slide.notes_slide.notes_text_frame.text))
        destinations = [shape.click_action.hyperlink.address for shape in slide.shapes
                        if shape.has_text_frame and shape.click_action.hyperlink.address]
        if DOC_OUTPUT.name not in destinations:
            raise AssertionError("Slide missing portable full-document link")
        for destination in destinations:
            name, separator, bookmark = destination.partition("#")
            if name != DOC_OUTPUT.name or (separator and bookmark not in bookmarks):
                raise AssertionError(f"Unresolved detail link: {destination}")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + Inches(0.01) or shape.top + shape.height > deck.slide_height + Inches(0.01):
                raise AssertionError(f"Shape outside slide: {shape.name}")
    if note_sections != set(range(1, 21)):
        raise AssertionError(f"PRD section coverage incomplete: {set(range(1, 21)) - note_sections}")


def main():
    source_document = Document(SOURCE)
    add_navigation_bookmarks(source_document)
    deck = build_deck(source_document)
    validate(source_document, deck)
    source_document.save(DOC_OUTPUT)
    deck.save(PPT_OUTPUT)
    validate(Document(DOC_OUTPUT), Presentation(PPT_OUTPUT))
    print(f"Generated {PPT_OUTPUT.name}: {len(deck.slides)} slides")
    print(f"Linked companion: {DOC_OUTPUT.name}; 21 section bookmarks; PRD content/source unchanged")


if __name__ == "__main__":
    main()