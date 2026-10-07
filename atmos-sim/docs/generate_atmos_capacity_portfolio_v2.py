from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from math import isclose
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement as WordElement
from docx.oxml.ns import qn as word_qn
from docx.opc.constants import RELATIONSHIP_TYPE as WordRelationship
from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordColor
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement as SlideElement
from pptx.util import Inches, Pt


DOCS = Path(__file__).resolve().parent
DOC_OUTPUT = DOCS / "ATMOS_E3S_Product_Portfolio_and_Rackscale_Integration_Proposal_v2.docx"
PPT_OUTPUT = DOCS / "ATMOS_E3S_AI_Capacity_Portfolio_Executive_v2.pptx"
STAMP = "5 October 2026"

INK = "183246"
MUTED = "536879"
WHITE = "FFFFFF"
LIGHT = "F3F7F9"
RED = "D51F3C"
TEAL = "007B78"
BLUE = "2B6791"
AMBER = "AC6528"
ASSETS = DOCS / "assets" / "atmos_capacity_portfolio"

SLIDE_GUIDE = [
    (1, "Portfolio and market priorities", "1", "Portfolio"),
    (2, "Common hardware and software foundation", "2", "Foundation"),
    (3, "E1: AI Appliance", "3", "Appliance"),
    (4, "E2: Companion Service", "4", "Companion"),
    (5, "E3: Integrated Capacity Pool", "5", "CapacityPool"),
    (6, "H: Hyperscale Neural-Data Sidecar", "6", "HyperscaleSidecar"),
    (7, "Matched-service cost and ROI", "8", "Economics"),
    (8, "Sensitivity and break-even", "8", "Economics"),
    (9, "OEM and evidence gates", "9 / 11", "OEM"),
    (10, "Investment decision and detail navigator", "12", "Decisions"),
]

OFFERINGS = [
    {
        "id": "E1", "name": "ATMOS AI Appliance", "audience": "Enterprise / private neural services",
        "boundary": "A new, independently operated inference endpoint",
        "bookmark": "Appliance", "section": "3", "color": BLUE,
    },
    {
        "id": "E2", "name": "ATMOS Companion Service", "audience": "Existing enterprise GPU estates",
        "boundary": "Choose GPU OR ATMOS before the complete inference request",
        "bookmark": "Companion", "section": "4", "color": TEAL,
    },
    {
        "id": "E3", "name": "ATMOS Integrated Capacity Pool", "audience": "Enterprises with platform operations teams",
        "boundary": "Whole-request hardware selection behind one qualified gateway",
        "bookmark": "CapacityPool", "section": "5", "color": RED,
    },
    {
        "id": "H", "name": "Hyperscale Neural-Data Sidecar", "audience": "Hyperscalers and very large search / recommendation operators",
        "boundary": "Shard-local scoring and reduction; compact results to GPU services",
        "bookmark": "HyperscaleSidecar", "section": "6", "color": AMBER,
    },
]


@dataclass(frozen=True)
class Scenario:
    track: str
    baseline_avoidable_year: float
    retained_year: float
    capex: float
    deployment: float
    support_year: float
    software_year: float
    operations_year: float
    average_it_kw: float
    annual_units: int
    unit: str
    configuration: str
    baseline: str

    @property
    def initial(self):
        return self.capex + self.deployment

    @property
    def energy_year(self):
        return self.average_it_kw * 8760 * 1.4 * 0.12

    @property
    def recurring(self):
        return self.support_year + self.software_year + self.operations_year + self.energy_year

    @property
    def baseline_tco(self):
        return 3 * (self.retained_year + self.baseline_avoidable_year)

    @property
    def project_tco(self):
        return self.initial + 3 * (self.retained_year + self.recurring)

    def net_benefit(self, monetized_fraction=1.0):
        return 3 * (self.baseline_avoidable_year * monetized_fraction - self.recurring) - self.initial

    @property
    def tco_saving(self):
        return self.net_benefit() / self.baseline_tco

    @property
    def roi(self):
        return self.net_benefit() / self.initial

    @property
    def payback_months(self):
        annual_net = self.baseline_avoidable_year - self.recurring
        return self.initial / annual_net * 12 if annual_net > 0 else float("inf")

    @property
    def break_even_fraction(self):
        return (self.initial / 3 + self.recurring) / self.baseline_avoidable_year


SCENARIOS = {
    "E1": Scenario("E1", 180_000, 0, 90_000, 30_000, 12_000, 8_000, 25_000,
                   1.5, 100_000_000, "embedding requests",
                   "Three assumed $30k 8-device nodes; two active plus one N+1 node",
                   "A cancellable cloud/leased embedding-service block costing $180k/year"),
    "E2": Scenario("E2", 120_000, 0, 75_000, 25_000, 10_000, 8_000, 24_000,
                   1.2, 20_000_000, "qualified batch jobs",
                   "Three assumed $25k 8-device nodes; two active plus one N+1 node",
                   "An otherwise required, separately billed $120k/year GPU overflow block"),
    "E3": Scenario("E3", 300_000, 700_000, 120_000, 60_000, 18_000, 12_000, 50_000,
                   2.0, 500_000_000, "requests in the unchanged estate workload mix",
                   "Three assumed $40k 16-device nodes; two active plus one N+1 node",
                   "A $1m/year estate with a $300k/year eligible, genuinely removable capacity block"),
    "H": Scenario("H", 600_000, 900_000, 240_000, 120_000, 36_000, 24_000, 120_000,
                  6.0, 10_000_000_000, "search/recommendation queries",
                  "Four assumed $60k 16-device nodes: two index generations times two service replicas",
                  "A $600k/year replaceable optimized retrieval tier plus $900k/year retained GPU services"),
}

SOURCES = [
    ("S1", "Linux kernel: PCI Peer-to-Peer DMA Support",
     "https://docs.kernel.org/driver-api/pci/p2pdma.html",
     "Topology limits, cooperating provider/client/orchestrator drivers, safe DMA invalidation. Not ATMOS qualification."),
    ("S2", "Kubernetes: Dynamic Resource Allocation",
     "https://kubernetes.io/docs/concepts/scheduling-eviction/dynamic-resource-allocation/",
     "ResourceSlices, DeviceClasses, ResourceClaims and CDI. Version and feature status must be pinned; DRA is not model readiness or failover."),
    ("S3", "llm-d: Architecture",
     "https://llm-d.ai/docs/architecture",
     "Router, endpoint picker, InferencePool, model servers and workload variants. ATMOS still requires its own backend, metrics and qualification."),
    ("S4", "Meta: Scaling the Instagram Explore recommendations system (9 Aug 2023)",
     "https://engineering.fb.com/2023/08/09/ml-applications/scaling-instagram-explore-recommendations-system/",
     "Billions of items; cached two-tower retrieval; multi-stage ranking that narrows thousands of candidates. Market-scale precedent, not ATMOS adoption."),
    ("S5", "Meta authors: Deep Learning Recommendation Model (31 May 2019)",
     "https://arxiv.org/abs/1906.00091",
     "Embedding-table memory constraints and separate dense compute. Older architecture evidence, not a current measured ATMOS business case."),
    ("S6", "Supermicro: ASG-2115S-NE332R product specifications",
     "https://www.supermicro.com/en/products/system/storage/2u/asg-2115s-ne332r",
     "32 front E3.S 1T Gen5 x2 storage bays. Demonstrates why bay count is not proof of x4, 2T, accelerator-class support."),
    ("S7", "NVIDIA: H200 product specifications",
     "https://www.nvidia.com/en-us/data-center/h200/",
     "141 GB HBM3e and 4.8 TB/s per GPU; distinct SXM/NVL configurations. A strong installed-base control, not a price quote."),
    ("S8", "NVIDIA: GB300 NVL72 product specifications",
     "https://www.nvidia.com/en-us/data-center/gb300-nvl72/",
     "72 Blackwell Ultra GPUs; rounded 20 TB GPU memory; 130 TB/s rack NVLink. No ATMOS access to that fabric is implied."),
    ("S9", "Tenstorrent: vLLM TT Plugin",
     "https://github.com/tenstorrent/vllm-tt-plugin",
     "A public Tenstorrent/vLLM integration path with model/device/version limits. Does not establish an ATMOS HBF backend or ATMOS performance."),
    ("S10", "LinkedIn authors: LiRank (2024; revised 7 Aug 2024)",
     "https://arxiv.org/abs/2402.06859",
     "Production-scale feed, jobs and ads ranking; quantization/vocabulary compression. Target-market precedent and incumbent optimization, not ATMOS validation."),
    ("S11", "Meta: Multi-Stage Architecture for Ads Ranking (5 Aug 2026)",
     "https://engineering.fb.com/2026/08/05/ml-applications/from-user-sequences-to-scaling-laws-a-multi-stage-architecture-for-metas-ads-ranking/",
     "Heavy offline user modeling and cached embeddings with latency-constrained online ranking. Co-design and caching must be included in the comparator."),
    ("S12", "Agent Router capabilities (formerly Envoy AI Gateway)",
     "https://theagentrouter.ai/docs/capabilities/",
     "Provider/pool routing, quota, rate limits and metrics. A candidate gateway foundation, not an ATMOS-ready distribution."),
    ("S13", "OpenAI: API pricing, specialized embedding models",
     "https://developers.openai.com/api/docs/pricing/",
     "Published standard input-token prices: text-embedding-3-small $0.02 and text-embedding-3-large $0.13 per million tokens. Policy/quality/SLO and non-token costs differ from private infrastructure."),
]


def money(value: float):
    return f"-${abs(value):,.0f}" if value < 0 else f"${value:,.0f}"


def compact_money(value: float):
    prefix = "-$" if value < 0 else "$"
    return f"{prefix}{abs(value) / 1_000_000:.2f}m" if abs(value) >= 1_000_000 else f"{prefix}{abs(value) / 1000:,.1f}k"


def add_bookmarked_heading(document, text: str, bookmark: str, identifier: int, level: int = 1):
    paragraph = document.add_heading(text, level)
    if level == 1:
        paragraph.paragraph_format.page_break_before = bookmark in {"Portfolio", "Economics", "Sources"}
    start = WordElement("w:bookmarkStart")
    start.set(word_qn("w:id"), str(identifier))
    start.set(word_qn("w:name"), bookmark)
    end = WordElement("w:bookmarkEnd")
    end.set(word_qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    return paragraph


def configure_word(document):
    section = document.sections[0]
    section.page_width, section.page_height = WordInches(8.5), WordInches(11)
    section.top_margin = section.bottom_margin = WordInches(0.65)
    section.left_margin = section.right_margin = WordInches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = WordPt(10)
    normal.font.color.rgb = WordColor.from_string(INK)
    normal.paragraph_format.space_after = WordPt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = WordPt(size)
        style.font.color.rgb = WordColor.from_string(INK)
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.page_break_before = False
    document.styles["Heading 3"].font.size = WordPt(11)
    document.styles["Heading 3"].font.color.rgb = WordColor.from_string(TEAL)
    header = section.header.paragraphs[0]
    header.text = "SANDISK / ATMOS     AI CAPACITY PORTFOLIO"
    for run in header.runs:
        run.font.size = WordPt(8)
        run.font.color.rgb = WordColor.from_string(RED)
    footer = section.footer.paragraphs[0]
    footer.text = f"INTERNAL | v2 | {STAMP}     PAGE "
    field = WordElement("w:fldSimple")
    field.set(word_qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = WordPt(8)
        run.font.color.rgb = WordColor.from_string(MUTED)
    core = document.core_properties
    core.title = "ATMOS E3.S AI Capacity Portfolio"
    core.author = core.last_modified_by = "SanDisk / ATMOS Architecture"
    core.revision = 2
    core.created = core.modified = datetime.now(timezone.utc)


def table(document, headers: list[str], rows: list[list[str]]):
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    repeat = WordElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cell, text in zip(result.rows[0].cells, headers):
        cell.text = text
        properties = cell._tc.get_or_add_tcPr()
        shading = WordElement("w:shd")
        shading.set(word_qn("w:fill"), INK)
        properties.append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = WordColor.from_string(WHITE)
            run.bold = True
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            cell.text = value
    for row_index, row in enumerate(result.rows):
        keep = WordElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for cell in row.cells:
            if row_index and row_index % 2:
                shading = WordElement("w:shd")
                shading.set(word_qn("w:fill"), LIGHT)
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = WordPt(3)
                paragraph.paragraph_format.space_before = WordPt(2)
                for run in paragraph.runs:
                    run.font.size = WordPt(9)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = spacer.paragraph_format.space_after = WordPt(0)
    spacer.paragraph_format.line_spacing = WordPt(3)
    spacer.add_run(" ").font.size = WordPt(3)
    return result


def text(slide, value: str, x: float, y: float, width: float, height: float,
         size: float = 18, color: str = INK, bold: bool = False, link: str | None = None,
         align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.paragraphs[0].alignment = align
    frame.paragraphs[0].space_before = frame.paragraphs[0].space_after = Pt(0)
    frame.paragraphs[0].line_spacing = 1.04
    run = frame.paragraphs[0].add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if link:
        shape.click_action.hyperlink.address = link
    return shape


def new_slide(deck, number: int, title: str, bookmark: str, section: str,
            subtitle: str = "", dark: bool = False):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(INK if dark else WHITE)
    text(slide, "SANDISK / ATMOS", 0.53, 0.21, 3.2, 0.28, 11, RED, True)
    text(slide, f"AI CAPACITY PORTFOLIO / {number:02d}", 9.02, 0.21, 3.78, 0.28,
        10, WHITE if dark else MUTED, False, align=PP_ALIGN.RIGHT)
    rect(slide, 0.55, 0.62, 12.23, 0.027, RED)
    text(slide, title, 0.53, 0.83, 12.23, 0.58, 28, WHITE if dark else INK, True)
    if subtitle:
       text(slide, subtitle, 0.55, 1.44, 12.15, 0.45, 15.5, "C9D7DF" if dark else MUTED)
    rect(slide, 0.55, 6.97, 12.23, 0.012, "41596A" if dark else "D9E2E8")
    text(slide, f"INTERNAL / {STAMP} / PLANNING, NOT QUALIFIED PERFORMANCE", 0.53, 7.12,
        7.0, 0.23, 8.5, "C9D7DF" if dark else MUTED)
    text(slide, "Open Word proposal", 7.85, 7.1, 2.35, 0.27, 10,
        "B8DFE6" if dark else BLUE, False, DOC_OUTPUT.name)
    text(slide, f"Section {section}", 10.45, 7.1, 2.34, 0.27, 10,
        "B8DFE6" if dark else BLUE, False, f"{DOC_OUTPUT.name}#{bookmark}", PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = (
        f"Companion document: {DOC_OUTPUT.name}. Section {section}, bookmark {bookmark}.\n"
        "Keep both files together. Full-document link is the portable fallback if a viewer ignores "
        "bookmark fragments. Office security policies can require confirmation for local links.\n"
        f"Slide {number}: {title}. All ATMOS specifications and economic scenarios require qualification."
    )
    return slide


def word_link(paragraph, label: str, destination: str | None = None, bookmark: str | None = None):
    hyperlink = WordElement("w:hyperlink")
    if destination:
        relationship = paragraph.part.relate_to(destination, WordRelationship.HYPERLINK, is_external=True)
        hyperlink.set(word_qn("r:id"), relationship)
    if bookmark:
        hyperlink.set(word_qn("w:anchor"), bookmark)
    run = WordElement("w:r")
    properties = WordElement("w:rPr")
    style = WordElement("w:rStyle")
    style.set(word_qn("w:val"), "Hyperlink")
    properties.append(style)
    color = WordElement("w:color")
    color.set(word_qn("w:val"), BLUE)
    properties.append(color)
    run.append(properties)
    content = WordElement("w:t")
    content.text = label
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def paragraphs(document, values):
    for value in values:
        document.add_paragraph(value)


def bullets(document, values):
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def figure(document, filename: str, caption: str):
    document.add_picture(str(ASSETS / filename), width=WordInches(7.05))
    paragraph = document.add_paragraph(caption)
    paragraph.paragraph_format.space_after = WordPt(10)
    for run in paragraph.runs:
        run.font.size = WordPt(8.5)
        run.italic = True
        run.font.color.rgb = WordColor.from_string(MUTED)


def build_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    regular = ImageFont.truetype(font_path, 23)
    small = ImageFont.truetype(font_path, 19)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 26)

    def canvas(title):
        image = Image.new("RGB", (1500, 660), "#" + WHITE)
        drawing = ImageDraw.Draw(image)
        drawing.rectangle((0, 0, 1500, 62), fill="#" + INK)
        drawing.text((25, 16), title, font=bold, fill="#" + WHITE)
        return image, drawing

    def node(drawing, bounds, label, color=BLUE):
        drawing.rounded_rectangle(bounds, radius=7, fill="#" + LIGHT, outline="#" + color, width=3)
        drawing.multiline_text((bounds[0] + 16, bounds[1] + 17), label,
                               font=regular, spacing=9, fill="#" + INK)

    def arrow(drawing, start, end, color=MUTED):
        drawing.line((start, end), fill="#" + color, width=4)
        if end[0] == start[0]:
            drawing.polygon([(end[0], end[1]), (end[0] - 9, end[1] - 14),
                             (end[0] + 9, end[1] - 14)], fill="#" + color)
        else:
            drawing.polygon([(end[0], end[1]), (end[0] - 14, end[1] - 9),
                             (end[0] - 14, end[1] + 9)], fill="#" + color)

    image, drawing = canvas("COMMON FOUNDATION  /  topology to be qualified")
    node(drawing, (70, 120, 520, 225), "CPU / NUMA root\nPCIe Gen5 x16 uplink")
    node(drawing, (70, 310, 520, 415), "P2P-capable PCIe switch\nFour Gen5 x4 endpoints", TEAL)
    arrow(drawing, (295, 225), (295, 310))
    for index in range(4):
        y = 103 + index * 125
        node(drawing, (735, y, 1405, y + 101),
             f"ATMOS {index}: local HBF + NPU + LPDDR\n128 GB nominal / 100 GB study budget", TEAL)
        drawing.line((600, 360, 600, y + 50, 735, y + 50), fill="#" + TEAL, width=3)
    drawing.line((520, 360, 600, 360), fill="#" + TEAL, width=4)
    drawing.text((70, 535), "Read locally; exchange compact results. P2P is not shared coherent memory.",
                 font=regular, fill="#" + INK)
    drawing.text((70, 583), "Host staging remains the functional baseline. No unsafe ACS/IOMMU bypass. [S1]",
                 font=small, fill="#" + MUTED)
    image.save(ASSETS / "p2p_foundation.png")

    image, drawing = canvas("ENTERPRISE BOUNDARY  /  choose one complete-request backend")
    node(drawing, (45, 248, 385, 373), "Client / application\nModel + declared SLO")
    node(drawing, (485, 248, 850, 373), "Gateway policy\nIdentity, quota, routing", RED)
    node(drawing, (965, 125, 1455, 250), "GPU endpoint / pool\nExisting runtime and KV", BLUE)
    node(drawing, (965, 391, 1455, 516), "ATMOS endpoint / pool\nQualified operators and SLO", TEAL)
    arrow(drawing, (385, 310), (485, 310))
    drawing.line((850, 310, 907, 310, 907, 187), fill="#" + BLUE, width=4)
    arrow(drawing, (907, 187), (965, 187), BLUE)
    drawing.line((907, 310, 907, 453), fill="#" + TEAL, width=4)
    arrow(drawing, (907, 453), (965, 453), TEAL)
    drawing.text((45, 566), "E2: independent endpoints. E3: unified identity, policy and platform operations.",
                 font=regular, fill="#" + INK)
    drawing.text((45, 611), "No in-layer GPU/ATMOS crossing. Shared gateway, CPU and NIC interference still require proof.",
                 font=small, fill="#" + MUTED)
    image.save(ASSETS / "enterprise_boundary.png")

    image, drawing = canvas("HYPERSCALE SIDECAR  /  neural-data stage, not an LLM execution split")
    node(drawing, (35, 147, 347, 293), "Existing ANN /\ncandidate generator\nVersioned IDs + query")
    node(drawing, (447, 98, 1030, 358), "ATMOS shard-local stage\nPersistent vectors / features\nGather, score, pool, local top-k\nMerge compact group results", TEAL)
    node(drawing, (1130, 147, 1465, 293), "GPU ranker / LLM\nSelected features\nAttention + KV + output", RED)
    arrow(drawing, (347, 220), (447, 220))
    arrow(drawing, (1030, 220), (1130, 220))
    node(drawing, (447, 418, 1030, 525), "Host control: refresh, generations,\naccess policy, replicas and replay", AMBER)
    arrow(drawing, (738, 358), (738, 418), AMBER)
    drawing.text((35, 573), "Qualification: recall/NDCG, freshness, end-to-end p99, failover and cashable tier removal.",
                 font=regular, fill="#" + INK)
    drawing.text((35, 616), "Meta and LinkedIn are target-market archetypes, not customers, partners or validated ATMOS deployments.",
                 font=small, fill="#" + MUTED)
    image.save(ASSETS / "hyperscale_sidecar.png")


def add_document_cover(document):
    document.add_paragraph("SANDISK / ATMOS", style="Subtitle")
    document.add_heading("ATMOS E3.S\nAI Capacity Portfolio", 0)
    document.add_paragraph("Three enterprise deployment modes. One hyperscale neural-data track.", style="Subtitle")
    document.add_paragraph(f"Technical product proposal and conditional business case | v2 | {STAMP}")
    table(document, ["Decision lens", "Position"], [
        ["Enterprise market entry", "Independent appliance, companion service and whole-request capacity pool"],
        ["Hyperscale opportunity", "Neural-Data Sidecar for very large search / recommendation infrastructure"],
        ["Initial product density", "8- and 16-device qualified OEM nodes; 20-22 only after customer and OEM proof"],
        ["Evidence status", "Proposed architecture and planning envelope. No measured ATMOS silicon results or prices supplied."],
        ["Economics status", "Transparent, independent 36-month scenarios; not quoted prices, forecasts or guaranteed savings"],
    ])
    paragraph = document.add_paragraph()
    word_link(paragraph, "Open the companion executive presentation", destination=PPT_OUTPUT.name)
    document.add_paragraph(
        "Keep the presentation and this document in the same folder. Each slide has a full-document "
        "link and a section-specific bookmark link. Local Office security policy may ask for "
        "confirmation; section numbers remain a fallback if an application ignores bookmark fragments."
    )
    document.add_heading("Reader guide", 2)
    guide = table(document, ["Slide", "Executive view", "Document section"], [
        [str(number), title, section] for number, title, section, bookmark in SLIDE_GUIDE
    ])
    for row, entry in zip(guide.rows[1:], SLIDE_GUIDE):
        paragraph = row.cells[2].paragraphs[0]
        paragraph.text = ""
        word_link(paragraph, entry[2], bookmark=entry[3])


def add_portfolio_document(document):
    add_bookmarked_heading(document, "1. Portfolio and decision", "Portfolio", 1)
    paragraphs(document, [
        "Lead with SLO-isolated capacity rather than a monolithic heterogeneous inference "
        "system. E1, E2 and E3 share one proposed device/runtime/operations foundation, "
        "but solve different enterprise adoption and ownership problems. E2 is the first "
        "integration priority because a separate service can be piloted without changing "
        "the existing GPU runtime. E1 requires a turnkey operations/support distribution; "
        "E3 requires the customer's platform and routing integration.",
        "H is a separate design-partner track for hyperscalers and large internet search, "
        "social, recommendation and advertising operators. LinkedIn and Meta are market "
        "archetypes, not claimed customers or endorsers. A general enterprise search or "
        "RAG installation does not establish enough economic or latency pain to justify "
        "another accelerator and data lifecycle. Admit an exceptional enterprise only "
        "after its scale, memory/network profile and removable cost are demonstrated.",
    ])
    table(document, ["ID / product", "Buyer and adoption", "Execution boundary", "First value proof"], [
        ["E1 / " + OFFERINGS[0]["name"], "Private AI / neural services; new supported endpoint",
         "Complete, qualified workload on ATMOS", "Replace a paid capacity block at identical quality/SLO"],
        ["E2 / " + OFFERINGS[1]["name"], "Existing GPU estate; independently operated parallel endpoint",
         "Choose GPU OR ATMOS before execution", "Remove a forecast GPU overflow purchase or bill"],
        ["E3 / " + OFFERINGS[2]["name"], "Mature platform team; one gateway and service catalog",
         "Whole-request hardware selection; no token-level crossing", "Realize demand placement and measurable bill/lease reduction"],
        ["H / " + OFFERINGS[3]["name"], "Hyperscale search / recommendation; sponsored custom integration",
         "Local neural-data operations; compact results to GPU/application", "Retire an optimized retrieval tier or avoid its expansion"],
    ])
    bullets(document, [
        "Measure accepted work at the same quality, SLO, workload mix, precision policy and availability; capacity alone is not service.",
        "Keep latency-critical GPU workloads, attention and KV ownership on their existing qualified runtime unless an entire ATMOS workload is independently proven.",
        "Do not promise native full-LLM execution from HBF capacity, a TT-NN kernel or an OpenAI-compatible HTTP interface.",
        "Use OEM-qualified four-device domains as a proposed schedulable unit; preserve isolation and include service redundancy.",
        "Bind releases to a manifest and a named support prime. None of the tracks is a launch commitment before the evidence gates.",
    ])
    document.add_heading("1.1 Evidence classes and commercial language", 2)
    table(document, ["Class", "Permitted use", "Not permitted"], [
        ["Public component facts [S1-S13]", "Choose candidate stacks, controls and qualification questions", "Claim ATMOS interoperability, certification or a partner commitment"],
        ["Private planning inputs", "Device/power/topology/model sizing studies", "Sell aggregate HBF or TOPS as accepted service"],
        ["Simulation / emulation", "Find bottlenecks, reject infeasible plans, prepare physical tests", "Pass a silicon, OEM thermal or customer SLO gate"],
        ["Matched physical evidence and quotes", "Approve the bounded workload, SKU, support and conditional economics", "Generalize to different models, platforms or tenants without retesting"],
    ])


def add_foundation_document(document):
    add_bookmarked_heading(document, "2. Common hardware and software foundation", "Foundation", 2)
    document.add_paragraph(
        "ATMOS is a proposed E3.S accelerator, not a conventional SSD. Every value below "
        "comes from the supplied planning envelope; physical silicon, final packaging "
        "and an orderable OEM configuration must confirm it. Use decimal GB/TB for HBF "
        "rather than silently changing GB to GiB."
    )
    table(document, ["Resource", "Planning value", "Release condition"], [
        ["Persistent HBF", "128 GB nominal/device; 100 GB application study budget", "Per-device packing, metadata, replication, scratch and skew must fit"],
        ["Local HBF read", "200 GB/s/device target", "Useful sustained rate for actual access pattern; not host/P2P bandwidth"],
        ["Working memory", "8-16 GB LPDDR plus SRAM", "Workspace, activations and any whole-model state must fit with reserve"],
        ["NPU", "Approximately 64 TOPS-class source planning target", "Define precision, sparsity, clock, kernel coverage and measured service first"],
        ["PCIe", "Gen5 x4 baseline; Gen6 a separate qualification option", "Endpoint/backplane/switch/retimer/BIOS lane match and payload proof"],
        ["Module", "E3.S 2T direction; 30-40 W planning envelope", "Sustained accelerator-class thermal, power and mechanical certification"],
    ])
    table(document, ["Population", "Nominal HBF", "Application study budget", "Device-only power"], [
        ["8 / two 4-device domains", "1.024 TB", "0.800 TB", "240-320 W"],
        ["16 / four 4-device domains", "2.048 TB", "1.600 TB", "480-640 W"],
        ["20-22 / stretch only", "2.560-2.816 TB", "2.000-2.200 TB", "600-880 W"],
    ])
    document.add_paragraph(
        "All totals are distributed arithmetic sums, not one coherent memory or an "
        "application-bandwidth promise. At 8 devices, nominal LPDDR is 64-128 GB distributed; "
        "state cannot be pooled without an explicit runtime and communication design. "
        "Availability requires complete eligible-service replicas, not just spare memory."
    )
    document.add_heading("2.1 Four-device PCIe/P2P group", 2)
    figure(document, "p2p_foundation.png", "Proposed switch-local domain. P2P capability, safe teardown and payload rate require physical proof. [S1]")
    paragraphs(document, [
        "Gen5 x4 carries approximately 15.75 GB/s per direction before transaction overhead; "
        "Gen5 x16 is approximately 63 GB/s. Four x4 endpoints roughly balance an x16 "
        "uplink. Eight x4 endpoints behind one x16 uplink introduce about 2:1 aggregate "
        "host oversubscription. Two four-device domains use two x16 uplinks, or 32 host "
        "lanes: this is a locality/fault-containment choice, not automatic lane savings.",
        "Local HBF processing should dominate transferred bytes. A downstream peer path "
        "does not become shared coherent memory, a collective implementation, NVLink, "
        "NVSwitch or an AMD scale-up fabric. Cross-root P2P and GPU-to-ATMOS transport "
        "are separate platform and cooperating-driver qualifications [S1].",
        "Use host staging as the functional baseline. Enable only documented, supported "
        "peer paths with safe buffer ownership, timeout, cancellation and DMA invalidation "
        "on reset/removal. Do not bypass IOMMU, ACS, DPC or tenant isolation for a benchmark. "
        "P2P is a desired subsystem capability; service economics must still be evaluated "
        "with the qualified fallback if P2P is unavailable.",
    ])
    document.add_heading("2.2 Software architecture and delivery boundary", 2)
    table(document, ["Layer", "Candidate building block", "ATMOS-specific obligation"], [
        ["Northbound service", "OpenAI-compatible chat/completions/embeddings where supported", "Prove semantics, model catalog and workload envelope; rerank/search are separate API contracts"],
        ["Gateway", "Customer gateway; Agent Router (formerly Envoy AI Gateway) candidate [S12]", "Identity, quota, tenant policy, audit, cost metrics and safe endpoint selection"],
        ["Serving control", "KServe or llm-d / Gateway API Inference Extension [S3]", "Backend metrics, compatible pool/model contracts, readiness and overload handling"],
        ["Kubernetes devices", "DRA, ResourceSlices, DeviceClasses, ResourceClaims and CDI [S2]", "Build/qualify the ATMOS DRA driver; whole-group allocation and tested version matrix"],
        ["Model runtime", "Selected NPU SDK; TT-NN / TT-Metal / vLLM TT plugin candidate [S9]", "HBF address/layout, operators, loader, memory planner, batching and complete workload integration"],
        ["Device plane", "Linux PCIe/DMA interfaces [S1]", "Driver/firmware/queues/health, reset, safe mapping and peer lifecycle"],
        ["Release and operations", "Container/Kubernetes/OpenTelemetry/Prometheus patterns", "Operator, signed ModelPacks, ABI checks, canary/drain/rollback, diagnostics and one support route"],
    ])
    document.add_paragraph(
        "The public TT plugin is evidence that vendor integration can use vLLM extension "
        "points, not evidence that ATMOS already works [S9]. Pin selected versions and "
        "per-model capabilities. ModelPack/ExpertPack is proposed packaging of weights, "
        "layout, graph/kernels, metadata and compatibility; a custom 'expertpack' compiler "
        "is not an available off-the-shelf ATMOS executable. Weights are data; runtime "
        "graph/kernel preparation is a distinct compiler concern."
    )


def product_heading(document, track):
    item = next(value for value in OFFERINGS if value["id"] == track)
    add_bookmarked_heading(document, f"{item['section']}. {item['name']}", item["bookmark"], int(item["section"]))
    document.add_paragraph(f"Target buyer: {item['audience']}. Execution boundary: {item['boundary']}.")


def add_appliance_document(document):
    product_heading(document, "E1")
    paragraphs(document, [
        "A new private endpoint packaged as an OEM-supported node or small cluster. ATMOS "
        "is the primary accelerator only for qualified workloads. Start with embedding, "
        "reranking, retrieval or a bounded document-batch operator set; general interactive "
        "LLM support is conditional, not implied by device capacity.",
        "Request path: application -> identity/quota gateway -> workload queue and model "
        "catalog -> ATMOS model server -> allocated device/group -> result and metering. "
        "Streaming or complete responses follow the declared API; batch jobs use explicit "
        "job state, retries and cancellation. A GPU fallback, if sold, remains a separate "
        "qualified endpoint and is included in both support and cost.",
    ])
    table(document, ["Package", "Scope", "Qualification focus"], [
        ["Entry / 8 devices", "Two four-device domains; 1.024 TB nominal HBF", "One frozen model/workload, two active plus reserve service configuration"],
        ["Standard / 16 devices", "Four domains; 2.048 TB nominal HBF", "Multi-tenant isolation, per-domain fit, sustained power and complete-service HA"],
        ["Stretch / 20-22", "Capacity-driven custom configuration only", "Customer need, bay mechanics, x4 lanes, cooling, CPU/NIC/storage and support quote"],
    ])
    document.add_heading("3.1 Turnkey service contract", 2)
    table(document, ["Function", "Minimum release evidence"], [
        ["Identity / tenant", "OIDC or approved API identity, tenant-scoped credentials, audit and DMA isolation"],
        ["Admission / queues", "Context/batch/concurrency caps, bounded queues, overload response and noisy-neighbor tests"],
        ["Model / API", "Signed model/pack digest, tokenizer/precision/operator matrix and golden correctness"],
        ["Lifecycle", "Canary, warmup before readiness, drain, cancel, reset, rollback and controlled replacement"],
        ["Metering", "Accepted/rejected work, end-to-end latency, device time, usage and chargeback; not raw submissions"],
        ["Availability", "Complete replicas and measured failover/restart behavior at required peak demand"],
        ["Support", "One prime, signed BOM and runbook for OEM/device/runtime/control-plane escalation"],
    ])
    document.add_heading("3.2 Efficiency mechanism and cost scenario", 2)
    paragraphs(document, [
        "The hypothesis is lower fully loaded cost for a capacity-dominant eligible service, "
        "not GPU-equivalent single-stream token latency. Validate against the cheapest "
        "policy-compliant CPU, GPU or managed API alternative; a small enterprise embedding "
        "service may already be cheaper on CPU or pay-per-token APIs.",
        "Section 8 models a replaceable $180k/year service block, three assumed 8-device "
        "nodes including N+1, and $120k total upfront spend. The hypothetical demand is "
        "100m accepted embedding requests/year with an unchanged input distribution. "
        "This is a financial scenario, not a measured rate or a price quote. Exact model, "
        "tokens/request, arrival trace, p99, privacy requirement and availability must "
        "justify both the comparator and the selected node count.",
    ])


def add_companion_document(document):
    product_heading(document, "E2")
    figure(document, "enterprise_boundary.png", "E2 and E3 share a complete-request boundary, but differ in gateway and operational ownership.")
    paragraphs(document, [
        "Deploy an independently operated ATMOS endpoint beside the customer's existing "
        "NVIDIA or AMD service. The customer, portal or policy router selects one backend "
        "before the request starts. GPU firmware, CUDA/ROCm, model runtime and native "
        "collectives stay unchanged. GPU-to-ATMOS peer DMA is not a prerequisite.",
        "Separate endpoint choice avoids a token-level dependency, but does not guarantee "
        "unchanged GPU tail latency. Shared CPU, NIC, gateway, storage, power and monitoring "
        "can still interfere. Reserve and test these resources under the mixed arrival "
        "trace. Reject or queue excess eligible work within the contracted policy rather "
        "than silently moving an unqualified workload.",
    ])
    table(document, ["Route to existing GPU service", "Route to ATMOS only when qualified"], [
        ["Strict streaming TTFT/TPOT and long-context interactive generation", "Embedding, rerank and retrieval service contracts"],
        ["Dense/VLM workloads or unsupported operators", "Document/batch jobs with the declared deadline and model quality"],
        ["Training/fine-tuning and CUDA/ROCm-dependent work", "A complete small/moderate model that passed the operator/state/SLO gate"],
        ["ATMOS overload/degraded state if compatible fallback is contracted", "Healthy ready model/pack and available qualified service capacity"],
    ])
    document.add_heading("4.1 Deployment and failure contract", 2)
    bullets(document, [
        "Start with separate endpoints or a parallel cluster; reuse corporate identity and observability without coupling device lifecycles.",
        "Freeze request semantics, sampling/precision, retry and idempotency rules. A partial stream cannot silently resume on a different model/backend.",
        "Use health/readiness and admission signals to route new complete requests. An in-flight failure follows the declared restart/fail policy.",
        "Keep tenant quotas and shared-resource reserves visible. Validate the existing GPU SLO under simultaneous ATMOS saturation and fault injection.",
    ])
    document.add_heading("4.2 Efficiency mechanism and cost scenario", 2)
    document.add_paragraph(
        "Use the companion to avoid an identifiable future GPU overflow block, not to "
        "claim a refund for installed GPUs. Section 8 assumes a separately billed "
        "$120k/year capacity block, $100k upfront ATMOS deployment, and 20m qualified "
        "batch jobs/year. Routeable request share is not equal to reclaimable GPU cost: "
        "residency, peak demand, batch efficiency and existing lease terms determine whether "
        "the expansion is actually avoided. If only 50% of that budget is removable, this "
        "scenario has a negative three-year return."
    )


def add_capacity_document(document):
    product_heading(document, "E3")
    paragraphs(document, [
        "Expose ATMOS service replicas through the customer's established gateway, identity, "
        "quotas, audit and chargeback. The platform chooses a ready, compatible inference "
        "pool by model/operator support, workload tier, live queue and measured SLO/unit cost. "
        "This is a complete-request resource-pool integration, not an internal Transformer split.",
        "Use separate pools for incompatible model or API contracts. An InferencePool/variant "
        "must retain a consistent base model and qualified response semantics [S3]. Do not "
        "route merely because two servers accept the same JSON or advertise a similar model "
        "name. Precision changes and model versions need explicit quality qualification.",
    ])
    document.add_heading("5.1 Kubernetes ownership and constraints", 2)
    table(document, ["Artifact", "Required behavior"], [
        ["ATMOS DRA driver", "Publish capacity/topology via ResourceSlices, allocate/configure claims and inject devices with CDI [S2]"],
        ["DeviceClass / ResourceClaim", "Select qualified generation, NUMA/switch domain and workload capability; prevent partial-group conflicts"],
        ["Atomic group representation", "Vendor driver represents a whole qualified group or enforces equivalent atomic placement; DRA does not invent P2P connectivity"],
        ["ATMOS Operator", "Firmware/pack compatibility, warmup readiness, drain, replacement, rollback and fault quarantine"],
        ["Gateway / endpoint picker", "Model/SLO/tier routing, admission, safe fallback and queue/cost/quality/thermal signals"],
        ["Model readiness", "Only register an endpoint after ABI, pack hash, memory fit, correctness and warmup pass"],
        ["Version / scheduling", "Pin kernel/Kubernetes/DRA/API versions and supported features; do not assume DRA preemption or automatic application failover"],
    ])
    document.add_paragraph(
        "Example metadata concepts, not a deployable API schema: device-class=atmos-e3s; "
        "hbf-capacity=128 GB nominal; application-budget=100 GB study; switch-domain=A; "
        "group-size=4; numa=0; qualified-p2p-profile=<manifest-id>; pack-abi=<pinned-version>; "
        "workload-class=<tested-embedding/rerank/batch-contract>. The ATMOS driver, naming "
        "namespace and claim semantics must be implemented and tested."
    )
    document.add_heading("5.2 Difference from E2 and proof", 2)
    table(document, ["Dimension", "E2 companion", "E3 integrated pool"], [
        ["Front door", "Separate service / explicit selection", "Unified gateway and catalog"],
        ["Device/cluster operations", "Independent appliance/cluster acceptable", "Customer platform, DRA/CDI/operator and release matrix"],
        ["Routing responsibility", "Customer or bounded policy router", "Platform endpoint picker and qualified workload policies"],
        ["Integration burden", "Lowest first-pilot scope", "Higher multi-tenant, telemetry, admission and lifecycle scope"],
        ["Execution invariant", "One whole-request backend", "Same; no in-layer crossing"],
    ])
    document.add_heading("5.3 Efficiency mechanism and cost scenario", 2)
    document.add_paragraph(
        "Pooling may improve placement and avoid overbuying, but it does not manufacture "
        "GPU savings. Section 8 assumes a $1m/year estate with $300k/year of genuinely "
        "removable eligible service, $700k/year retained GPU service, and $180k upfront "
        "ATMOS/platform deployment. The 500m/year request denominator is an unchanged "
        "workload mix, not an equivalence between embedding, batch and chat requests. "
        "Report full-estate savings separately from the eligible-slice benefit and verify "
        "that request routing, cache/residency and peak reserve allow the specific lease "
        "or expansion block to be removed."
    )


def add_hyperscale_document(document):
    product_heading(document, "H")
    paragraphs(document, [
        "Primary audience: hyperscalers and internet-scale search, social/recommendation, "
        "ads, commerce and content platforms with very large persistent vector/feature "
        "sets, sustained serving demand, a measured memory/network/latency bottleneck and "
        "a custom infrastructure team. Meta's Explore engineering account documents a "
        "billions-of-items, multi-stage retrieval/ranking funnel with cacheable embeddings "
        "[S4]. The DLRM paper separates embedding-table memory constraints from dense "
        "compute [S5]. LinkedIn's LiRank describes production feed, jobs and ads ranking "
        "with compression/quantization [S10]. Meta's 2026 sequence-model architecture "
        "decouples heavy offline modeling from cached, latency-constrained online ranking "
        "[S11]. These are architectural precedents, not evidence of ATMOS adoption "
        "or a transferable performance gain. LinkedIn and Meta are target-market examples.",
        "Ordinary enterprise RAG/search is not the default buyer. A small index, low QPS, "
        "an optimized CPU/DRAM ANN service, effective caching or a cheap managed endpoint "
        "can make the added accelerator and synchronization work uneconomic. An exceptional "
        "enterprise must show the same bottleneck and financial evidence as a hyperscaler; "
        "a large corpus alone is insufficient.",
    ])
    figure(document, "hyperscale_sidecar.png", "H operates at the neural-data stage. The GPU retains its model runtime, attention, KV and generation path.")
    document.add_heading("6.1 Launch scope and ownership", 2)
    table(document, ["Stage", "Owner / contract", "Required proof"], [
        ["Candidate generation", "Existing ANN/search stack; query plus versioned candidate IDs", "Candidate distribution, filters, ACLs and freshness trace frozen"],
        ["ATMOS local stage", "Persistent vector/feature shards; gather, dot/MaxSim, pooling, lightweight scoring and local top-k", "Kernel coverage, useful random-read rate, skew and LPDDR workspace"],
        ["Group / service merge", "Compact IDs/scores/features; deterministic merge and deadline", "Fan-out/fan-in, queueing, complete result semantics and failure behavior"],
        ["GPU/application stage", "Existing heavy ranker/LLM/VLM receives qualified compact inputs", "No quality, relevance, personalization, safety or end-to-end latency regression"],
        ["Data lifecycle", "Host control plane plus qualified import/refresh and complete service replicas", "Atomic index/model generations, replication, deletion/ACL compliance and stale-data rejection"],
    ])
    document.add_paragraph(
        "This scope is neural-data processing, not cold-expert/per-layer GPU offload. "
        "Keeping compact results at an application boundary preserves a testable product "
        "contract. Semantic-cache reuse requires version/ACL/expiry correctness; a cache "
        "hit is not assumed safe. Token-level late interaction is a candidate workload "
        "only after actual token-vector expansion, MaxSim operators and relevance "
        "quality pass. A simple vector dot product is not automatically equivalent to "
        "a cross-encoder reranker or a full recommendation model."
    )
    document.add_heading("6.2 Transparent sizing and traffic example", 2)
    table(document, ["Hypothetical input / derivation", "Result / interpretation"], [
        ["500m vectors x 768 dimensions x 2 bytes", "768 GB raw vectors; decimal units"],
        ["25% hypothetical packing/IDs/metadata allowance", "960 GB per generation; actual index/layout metadata must be measured"],
        ["One 16-device node x 100 GB application study budget", "1,600 GB distributed budget; one generation can fit only if every shard fits"],
        ["Two generations x two complete replicas", "3,840 GB data budget; four 16-device nodes / 64 devices assumed for live plus staging replicas"],
        ["20,000 candidate vectors x 768 x 2 bytes/query", "30.72 MB candidate-vector payload if the baseline actually transfers these vectors"],
        ["200 selected items x (8-byte ID + 4-byte score)", "2.4 kB result payload; excludes query/candidate IDs, features, headers, refresh and replicas"],
        ["20,000 candidates -> 200 heavy-ranker candidates", "99% fewer downstream candidates only if the same quality/relevance is retained; not 99% total cost saving"],
        ["10,000 queries/s x 30.72 MB local read/query", "307.2 GB/s useful local vector demand before metadata/cache effects; not an achieved rate"],
    ])
    paragraphs(document, [
        "The sizing illustration uses two complete live replicas plus two full staging "
        "replicas for a controlled generation switch. A surviving live replica must "
        "still meet the peak SLA, or more capacity is required. Four independent PCIe "
        "domains in a node do not each contain a complete index. HBF write/import "
        "bandwidth, endurance, delete handling and update cost are unproven and may "
        "change this node count. Host ANN metadata, journals, CPU DRAM, SSDs and NICs "
        "are separately budgeted where they are not included in the HBF layout.",
        "The apparent traffic reduction is valid only for the described vector-moving "
        "baseline. An ANN service already returning compact IDs may see little or no "
        "inter-stage byte reduction. Measure bytes on the actual bottleneck, including "
        "candidate input IDs, scatter/gather, selected features, group merge and refresh. "
        "Do not full-scan 768 GB for every query merely because HBF can hold the corpus.",
        "A critical-path illustration, not a p99 prediction: if the removable stage "
        "occupies 35 ms of a 100 ms serial budget and is twice as fast, the budget "
        "becomes 82.5 ms, or 17.5% lower. End-to-end p99, overlapping work, queueing "
        "and fan-out tails must be measured; percentile latencies cannot be added as "
        "though they were deterministic components.",
    ])
    document.add_heading("6.3 Hyperscale admission and rejection charter", 2)
    bullets(document, [
        "Sponsor provides index/model size, precision, candidate histogram, QPS/arrival trace, p95/p99 budget, relevance/freshness/ACL requirements and removable spend.",
        "Benchmark the best optimized CPU/DRAM ANN + cache/compression/SSD tier, GPU-resident variant and incumbent stack; do not choose a weak GPU-only strawman.",
        "Freeze recall@k/NDCG or recommendation quality, filters, freshness, update rate and complete-service availability before measuring throughput.",
        "Replay updates, hot/skewed shards, tenant contention, a lost node/group, stale generations and restart while measuring the complete pipeline.",
        "Proceed only if measured local useful work outweighs additional DMA/RPC/merge, the full result meets p99, and support/replica/refresh costs leave positive unit economics.",
        "Reject if the data fits cheaply in the incumbent tier, the baseline already sends compact results, relevance drops, updates cannot meet freshness, or no capacity block can be retired/avoided.",
    ])
    document.add_paragraph(
        "Section 8's H scenario is an assumed $600k/year replaceable retrieval tier, "
        "$900k/year retained GPU service and $360k upfront deployment of four assumed "
        "16-device nodes. No GPU generation savings or monetized relevance/revenue uplift "
        "is included. At only 50% realized tier reduction, the three-year return is negative."
    )


def add_slo_document(document):
    add_bookmarked_heading(document, "7. Workload fit and matched-service measurement", "SloFit", 7)
    table(document, ["Workload", "Track / control", "Required acceptance metric", "Position"], [
        ["Embeddings / rerank", "E1/E2/E3; CPU/managed/GPU control", "Model quality, input length, batch, accepted rate and p99", "First bounded service candidates"],
        ["Document / asynchronous batch", "E1/E2/E3; incumbent runtime", "Accuracy, completion deadline, cancellation and energy/job", "Conditional operator coverage"],
        ["Complete small/moderate LLM", "E1/E2/E3 only after full model proof", "TTFT, TPOT/ITL, context, concurrency, sampling and state fit", "Not implied by HBF/NPU capacity"],
        ["Vector/feature refinement / late interaction", "H; best CPU/DRAM ANN and GPU variants", "Recall/NDCG, freshness, p99, update load and cost/query", "Hyperscale design-partner track"],
        ["Long-context interactive / strict token SLO / VLM", "Existing qualified NVIDIA/AMD service", "Matched end-to-end and per-token tails", "GPU preference unless whole ATMOS service independently passes"],
        ["Training, fine-tuning and token-synchronous heterogeneous execution", "Existing GPU platform / outside scope", "No claimed ATMOS product result", "Not portfolio launch workloads"],
    ])
    document.add_heading("7.1 Freeze a workload card before the POC", 2)
    table(document, ["Field", "Required value"], [
        ["Workload identity", "Immutable model/dataset/index/tokenizer/pack hashes; licenses and tenant/ACL policy"],
        ["Arrival and input", "Length/candidate/update distributions, batching, concurrency, burst/skew and peak demand"],
        ["Quality", "Golden output / accepted domain score / recall@k or NDCG threshold against the same reference"],
        ["Service", "Sponsor-owned p50/p95/p99, TTFT/TPOT where relevant, deadline, timeout and overload budget"],
        ["Availability", "Complete replica policy, N+1 peak service, restart/failover semantics and maintenance"],
        ["Comparator", "Cheapest qualified policy-compliant incumbent/CPU/managed/GPU configuration and approved contract price"],
        ["Cash proof", "Named removable bill/lease/expansion block, realized fraction, approved quotes and 36-month horizon"],
    ])
    document.add_heading("7.2 Measurement procedure", 2)
    for value in [
        "Freeze manifests and the above workload card; warm up both paths and record actual software/hardware versions.",
        "Run identical steady, burst, hot-shard and mixed-tenant traces. Vary offered load until each qualified service reaches its SLO-limited capacity.",
        "Count only accepted, quality-correct work that meets the declared SLO. Retain all rejects, errors, timeouts and stale results in the denominator/report.",
        "Correlate queue, host, device, DMA, merge, runtime and GPU traces. Record useful HBF bytes, peer/host/NIC bytes, CPU/LPDDR, thermal and facility power.",
        "Inject cancellation, module/group/node faults, reboot, pack/version mismatch and data refresh; demonstrate isolation and the contracted failure policy.",
        "Apply measured rates and reserve to actual node counts. Obtain OEM/support quotes and verify which baseline spend is truly avoided before recomputing ROI.",
    ]:
        document.add_paragraph(value, style="List Number")
    document.add_paragraph(
        "Use current qualified baselines, not an old GPU memory shortage alone. H200's "
        "published 141 GB HBM3e / 4.8 TB/s per GPU [S7] and GB300 NVL72's 72 GPUs, "
        "rounded 20 TB GPU memory and 130 TB/s rack NVLink [S8] show how high-end "
        "controls can remove the simple capacity-fit thesis. A rack and an 8/16-device "
        "node are not equal service units. Match the workload, service envelope, facility "
        "and total contracted cost; GPU memory/BW totals do not imply ATMOS-accessible memory."
    )


def add_oem_document(document):
    add_bookmarked_heading(document, "9. OEM, thermal, availability and support", "OEM", 9)
    document.add_paragraph(
        "Bay count is not accelerator support. The verified Supermicro ASG-2115S-NE332R "
        "example has 32 front E3.S 1T Gen5 x2 storage bays [S6]; that is not a drop-in "
        "match for the proposed x4, 2T, 30-40 W ATMOS module. Seek named, orderable "
        "OEM configurations and a signed accelerator-class qualification, not a percentage "
        "of a storage bay population. Dell/HPE and other OEM options require the same "
        "BOM/topology/thermal proof; no unsupported compatibility assertion is made here."
    )
    table(document, ["Stage", "Test scope", "Advance evidence"], [
        ["1 module", "Power, useful HBF/NPU service, inlet/throttle, voltage and reset", "Sustained qualified service and mechanical/electrical envelope"],
        ["4-module group", "Adjacent heat, switch/retimer load, P2P, security and removal", "Safe lifecycle, lane map and complete host-staged/peer-path comparison"],
        ["8-module node", "NUMA, CPU/NIC load, multi-tenant saturation, full soak and one-domain failure", "Accepted service, fault containment and complete replica policy"],
        ["16-module standard", "Maximum inlet, N+1 PSU/fan, maintenance, spares and support", "Orderable BOM, full service under failure/maintenance and cost quote"],
        ["20-22 stretch", "Customer-specific fit, all bays, NIC/storage reserve and cooling", "Only after standard SKU proof; no initial launch promise"],
    ])
    table(document, ["16-device power item", "Budget treatment"], [
        ["Devices", "480-640 W planning arithmetic; measure sustained workload/idle/peak"],
        ["CPU / DRAM", "OEM-specific; source planning 300-700+ W is not a certified BOM"],
        ["Switch / retimer / backplane / NIC / SSD", "Measure rather than omit; maintain lane and thermal headroom"],
        ["Fans / conversion / facility", "Wall power, N+1/inlet condition and customer PUE; module wattage is not system energy"],
    ])
    document.add_heading("9.1 Required configuration contract", 2)
    bullets(document, [
        "Mechanics: exact E3.S thickness/carrier, bay spacing, retention, connector rating, service access and hot-replacement rules.",
        "Electrical: root/switch/NUMA map, actual x4/Gen5 or separately qualified Gen6, BAR/IOMMU, AER/DPC/FLR and safe peer-path enablement.",
        "Thermal: continuous qualified workload at maximum inlet, adjacent modules, one-fan failure, connector temperatures and throttling behavior.",
        "Firmware/security: signed BIOS/BMC/device updates, rollback compatibility, health telemetry, DMA isolation and failure containment.",
        "Availability: complete model/index service replicas with enough surviving peak capacity; disjoint unique shards are not HA.",
        "Commercial: all-in BOM, NRE, device/switch/runtime support, lead times, spares, geography, escalation and liability boundary.",
    ])
    document.add_heading("9.2 One accountable support route", 2)
    table(document, ["Layer", "Accountable owner / required interface"], [
        ["Chassis, power/cooling, BIOS/BMC", "OEM/integrator; named SKU and field-service SLA"],
        ["Switch/backplane and P2P", "OEM + switch vendor + SanDisk; one root-cause/firmware route"],
        ["ATMOS module, firmware, SDK/driver", "SanDisk; diagnostics, reset/isolation and replacement contract"],
        ["NPU compiler/runtime/operators", "SanDisk + selected NPU partner; signed scope and version/test matrix"],
        ["DRA/operator/serving adapters", "SanDisk distribution owner; readiness, drain/rollback and issue response"],
        ["Customer gateway/data pipeline", "Named customer/software partner; routing, refresh, ACL and support boundary"],
        ["First call / economic case", "One prime per released package; support included in approved ROI ledger"],
    ])


def add_delivery_document(document):
    add_bookmarked_heading(document, "10. Software deliverability and release package", "Delivery", 10)
    table(document, ["Workstream", "Existing foundation", "Build or qualify before release"], [
        ["Device/SDK", "Linux DMA/PCIe and selected NPU APIs", "ATMOS queue ABI, safe mappings, HBF/LPDDR planner, health and recovery"],
        ["Workload runtime", "Candidate TT-NN/TT-Metal/vLLM plugin [S9]", "Actual ATMOS operators, graphs/kernels, model loader, batching and memory/service curve"],
        ["Packs", "Container/signing/model artifact patterns", "ModelPack/feature-pack manifest, precision/layout, immutable hashes, verification and rollback"],
        ["Kubernetes", "DRA/CDI and standard controllers [S2]", "ATMOS DRA driver, atomic group policy, operator, readiness/drain and taint/quarantine"],
        ["Serving/gateway", "KServe/llm-d/customer gateway patterns [S3]", "Qualified API/model semantics, endpoint metrics, route/admission and tenant contracts"],
        ["Hyperscale data", "Incumbent ANN/feature pipeline", "Local kernels, shard/merge, generation protocol, update/delete/ACL and complete replicas"],
        ["Evidence/support", "OpenTelemetry/Prometheus/CI conventions", "Correlated traces, golden suites, fault harness, signed BOM, support/runbooks and cost worksheet"],
    ])
    document.add_heading("10.1 Minimum compatibility manifest", 2)
    bullets(document, [
        "Hardware/device firmware/NPU generation; switch, retimer, backplane, carrier and per-root lane topology.",
        "OEM, CPU/NUMA, BIOS/BMC, inlet/power policy, NIC/SSD/host-memory budget and qualified security controls.",
        "OS/kernel, driver/SDK/runtime, containers, Kubernetes/DRA/CDI/operator and gateway/serving versions.",
        "Model/tokenizer/pack/index hashes, operator/precision/layout, per-device memory fit, supported API and workload envelope.",
        "Peak arrival/load, quality and freshness thresholds, SLO, timeout/cancel/retry policy, replica/reserve and failure response.",
        "Quotes and cost inputs, known limits, first-call support, replacement/rollback procedure and evidence bundle IDs.",
    ])
    document.add_paragraph(
        "Pin actual upstream releases; neither a DRA object nor a model family in a "
        "vendor README proves the ATMOS device path. Do not expose a pod as ready until "
        "device checks, ABI/pack match, memory plan, golden correctness and warmup pass. "
        "Treat security and update/delete semantics as release gates, not post-pilot cleanup."
    )


def add_evidence_document(document):
    add_bookmarked_heading(document, "11. Gated proof and adoption program", "Evidence", 11)
    document.add_paragraph(
        "The gates are evidence order, not a promised silicon or launch schedule. "
        "Simulation/emulation can refine or reject the plan while hardware is unavailable, "
        "but cannot replace the physical service, thermal, support or customer-economic gates."
    )
    table(document, ["Gate", "Scope / owner", "Required exit package"], [
        ["G0 / workload and economics", "Product + customer sponsor + finance", "One immutable workload card, best comparator, bill/expansion block, SLO/quality, baseline traces and quote plan"],
        ["G1 / one device", "SanDisk + NPU partner", "Operators/correctness, HBF useful read, LPDDR fit, power/thermal, queue/DMA/reset and service curve"],
        ["G2 / four-device domain", "SanDisk + OEM/switch vendor", "Host/peer comparison, fan-out/merge, isolation, safe teardown, skew and domain-fault evidence"],
        ["G3 / 8 then 16-device appliance", "SanDisk distribution + OEM/support prime", "API/admission, tenant interference, full soak, complete replicas/peak failover, signed BOM/support"],
        ["G4 / E2 or E3 customer integration", "Customer platform + SanDisk", "Whole-request routing, GPU tails protected, ready/drain/rollback, measured costs and approved removable capacity"],
        ["G5 / H design partner", "Hyperscale search/data sponsor + SanDisk", "Optimized incumbent comparison, recall/NDCG/freshness, refresh/delete/ACL, full-pipeline p99 and actual tier avoidance"],
    ])
    document.add_heading("11.1 Smallest useful first pilots", 2)
    bullets(document, [
        "Enterprise: one existing paid workload and one ATMOS endpoint using E2. Do not expand to every model/operator or a unified gateway before matched service is demonstrated.",
        "Appliance: turn that proven workload into an OEM/runbook/support package, then qualify the 8/16-device product envelope.",
        "Integrated pool: add E3 only for a customer with DRA/operator/gateway ownership and a measurable placement/removable-cost problem.",
        "Hyperscale: one neural-data operator stage, one trace and one index-generation protocol. No assumption that vector scoring, late interaction and recommendations all share one kernel contract.",
    ])
    document.add_heading("11.2 Stop, narrow or return to baseline", 2)
    bullets(document, [
        "Stop on incorrect outputs, relevance/freshness/ACL regression, unsafe DMA/reset or a missing required operator.",
        "Narrow when lower-level evidence passes but complete service fails its p99, queue, context, concurrency or availability envelope.",
        "Use the qualified staged path or redesign the OEM fixture if peer access is unsupported; do not relax security to keep a claim.",
        "Reject density or revise node counts if thermal derating, shared CPU/NIC load, replicas, staging or repair reserve changes accepted capacity.",
        "Fail the business gate if the realized fraction falls below Section 8's break-even floor, quotes/NRE rise, or incumbent optimization wins.",
    ])


def add_decisions_document(document):
    add_bookmarked_heading(document, "12. Investment decision and customer discovery", "Decisions", 12)
    table(document, ["Decision", "Recommendation / required next evidence"], [
        ["Enterprise product charter", "Approve E1/E2/E3 definitions; lead one E2 matched-service pilot and a supported 8-device path"],
        ["OEM scope", "Approve a four-device fixture and 8->16 qualification; hold 20-22 until standard-node proof"],
        ["Software ownership", "Name driver/runtime/operator/gateway/support owners and signed NPU/OEM deliverables"],
        ["Hyperscale charter", "Sponsor H only with large-search bottleneck, update/quality/SLO trace and an actually removable retrieval tier"],
        ["Commercial claims", "Hold all device specifications, throughput, efficiency, savings and launch dates pending matched physical evidence and quotes"],
        ["Finance gate", "Publish complete cash and noncash benefits separately; approve only the recomputed, bounded business case"],
    ])
    document.add_heading("12.1 Discovery questions that decide fit", 2)
    for value in [
        "Which exact service, model/index and bill or future capacity purchase could ATMOS replace? Is that spend contractually avoidable?",
        "What are current accepted throughput, peak arrival, p99/TTFT/TPOT, quality and freshness? Which stage actually limits the service?",
        "How large is the persistent data/model and working state per shard? What is the measured byte movement, cache behavior and update/delete rate?",
        "What is the cheapest qualified CPU/managed/GPU alternative after normal caching, batching, ANN/compression and routing improvements?",
        "Who owns API/model semantics, Kubernetes/DRA/gateway, data refresh, security, OEM lifecycle and first-call support?",
        "What complete replica and peak reserve survives a module/domain/node fault, and what change would invalidate the cost model?",
    ]:
        document.add_paragraph(value, style="List Number")
    document.add_paragraph(
        "Recommended proof sequence: matched baseline -> one operator/device -> qualified "
        "group -> supported whole service -> customer routing and actual avoided spend. "
        "Make H a separate hyperscale workload and data-lifecycle charter rather than "
        "a generic enterprise add-on."
    )
    paragraph = document.add_paragraph()
    word_link(paragraph, "Return to the executive deck", destination=PPT_OUTPUT.name)


def add_sources_document(document):
    add_bookmarked_heading(document, "Appendix A. Public research and claim boundaries", "Sources", 30)
    document.add_paragraph(
        "Primary public references below were reviewed for architectural constraints and "
        "market context. Dates/specifications are source-specific. They do not validate "
        "ATMOS hardware, customers, partnership, prices or measured performance. Private "
        "device values originate in the supplied planning proposal; economic inputs in "
        "Section 8 are explicitly hypothetical. No unresolved search-result identifiers "
        "are used as citations."
    )
    for reference, title, url, relevance in SOURCES:
        document.add_heading(f"{reference}. {title}", 2)
        paragraph = document.add_paragraph()
        word_link(paragraph, url, destination=url)
        document.add_paragraph(f"Basis: {relevance} Reviewed {STAMP}.")


def build_document():
    document = Document()
    configure_word(document)
    add_document_cover(document)
    add_portfolio_document(document)
    add_foundation_document(document)
    add_appliance_document(document)
    add_companion_document(document)
    add_capacity_document(document)
    add_hyperscale_document(document)
    add_slo_document(document)
    add_economics_document(document)
    add_oem_document(document)
    add_delivery_document(document)
    add_evidence_document(document)
    add_decisions_document(document)
    add_sources_document(document)
    return document


def add_economics_document(document):
    add_bookmarked_heading(document, "8. Cost, efficiency and conditional ROI", "Economics", 20)
    document.add_paragraph(
        "ILLUSTRATIVE USD CASH MODEL, NOT A QUOTE OR BENCHMARK. Every device price, "
        "service budget, accepted throughput, node count and wattage below is an explicit "
        "hypothetical input. Replace them with customer demand, measured whole-service "
        "performance and binding OEM/support quotes before investment approval. Scenarios "
        "are independent and must not be added together. No ATMOS speedup is assumed."
    )
    document.add_heading("8.1 Method and financial boundary", 2)
    document.add_paragraph(
        "Use an undiscounted 36-month cash horizon in constant USD. Baseline avoidable "
        "budgets are recurring cloud/lease/service commitments that can actually be "
        "cancelled or not purchased; they already include their provider's power. "
        "Retained GPU costs are identical on both sides. Existing sunk GPU capital is "
        "not reclaimed, and freed capacity is only a noncash benefit unless it avoids "
        "a specific future purchase or bill. Taxes, financing, residual value and vendor "
        "silicon R&D are not separately modeled; customer procurement quotes must include "
        "supplier margin and any allocated program NRE."
    )
    table(document, ["Metric", "Calculation"], [
        ["Initial investment I", "Quoted hardware/system/spares + one-time integration and validation"],
        ["Annual ATMOS cost A", "Support + software + operations + measured facility-adjusted electricity"],
        ["Annual electricity", "Average IT kW x 8,760 hours x PUE 1.4 x $0.12/kWh; PUE and tariff are assumptions"],
        ["Three-year net benefit", "3 x (avoidable annual baseline x realized fraction - A) - I"],
        ["Three-year ROI", "Three-year net benefit / I; not a vendor margin or measured product return"],
        ["Simple payback", "12 x I / (avoidable annual baseline - A), if annual net saving is positive"],
        ["Break-even monetization", "(I / 3 + A) / avoidable annual baseline for a 36-month investment"],
        ["Unit economics", "Three-year cash cost / three-year accepted service units at unchanged quality, demand, SLO and availability"],
    ])
    document.add_heading("8.2 Worked scenario results", 2)
    table(document, ["Track", "Initial", "ATMOS/year", "3y baseline", "3y project", "Net benefit", "TCO saving"], [
        [track, money(model.initial), money(model.recurring), money(model.baseline_tco),
         money(model.project_tco), money(model.net_benefit()), f"{model.tco_saving:.1%}"]
        for track, model in SCENARIOS.items()
    ])
    table(document, ["Track", "3y ROI", "Payback", "Baseline / new per 1,000 units", "Annual accepted demand"], [
        [track, f"{model.roi:.1%}", f"{model.payback_months:.1f} months",
         f"${model.baseline_tco / (3 * model.annual_units) * 1000:.4f} / "
         f"${model.project_tco / (3 * model.annual_units) * 1000:.4f}",
         f"{model.annual_units:,} {model.unit}"]
        for track, model in SCENARIOS.items()
    ])
    document.add_paragraph(
        "E3's headline percentage applies to the full $1m/year estate, including $700k/year "
        "retained GPU services. H includes $900k/year retained GPU services. The E3 denominator "
        "uses the identical, frozen workload mix; it is not a claim that different request "
        "types have equal cost. Service units differ by track and are not cross-comparable."
    )
    document.add_heading("8.3 Inputs, configuration and realization risk", 2)
    for track, model in SCENARIOS.items():
        document.add_heading(f"{track}: scenario ledger", 3)
        document.add_paragraph(f"Baseline: {model.baseline}.")
        document.add_paragraph(f"Configuration: {model.configuration}.")
        table(document, ["Input", "Hypothetical budget"], [
            ["Hardware and included reserve", money(model.capex)],
            ["One-time rollout, migration, qualification and integration", money(model.deployment)],
            ["Annual support / software / operations", f"{money(model.support_year)} / {money(model.software_year)} / {money(model.operations_year)}"],
            ["Average complete-block IT power / annual facility electricity", f"{model.average_it_kw:g} kW / {money(model.energy_year)}"],
            ["Annual avoidable baseline / retained service", f"{money(model.baseline_avoidable_year)} / {money(model.retained_year)}"],
        ])
        document.add_paragraph(
            f"36-month break-even requires more than {model.break_even_fraction:.1%} of the "
            f"specified baseline spend to be realized as avoided bills. If only half is "
            f"monetized, net three-year benefit becomes {money(model.net_benefit(0.5))}. "
            "Revenue growth, latency improvement and capacity credit are excluded from this cash return."
        )
    document.add_heading("8.4 Sensitivity and stop rule", 2)
    table(document, ["Track", "0% realized", "50% realized", "75% realized", "100% realized", "Break-even floor"], [
        [track, *[money(model.net_benefit(fraction)) for fraction in (0, 0.5, 0.75, 1)],
         f"{model.break_even_fraction:.1%}"] for track, model in SCENARIOS.items()
    ])
    document.add_paragraph(
        "Fail the economic gate if the matched-SLO measured node count increases, qualified "
        "capacity or fleet demand is insufficient, baseline leases cannot be reduced, or "
        "index refresh, replicas, CPU/NIC load, OEM NRE and support erase the margin. "
        "Recalculate for a 25% capex increase, 25% lower goodput, higher support cost and "
        "the customer's tariff/PUE; do not substitute nominal HBF bandwidth or TOPS for "
        "accepted service. Directly compare H with optimized CPU/DRAM plus ANN, compression, "
        "caching and SSD-tiering alternatives as well as GPU-resident implementations."
    )
    document.add_heading("8.5 Published-price counterexample: a managed API can win", 2)
    table(document, ["Public reference / hypothetical demand", "Input-token charge"], [
        ["text-embedding-3-small: $0.02 per million input tokens [S13]", "$1,024/year at 100m requests x 512 tokens"],
        ["text-embedding-3-large: $0.13 per million input tokens [S13]", "$6,656/year at the same token count"],
    ])
    document.add_paragraph(
        "These published standard prices were reviewed on 5 October 2026 and can change. "
        "They are input-token charges only, not full retrieval/storage/network/support "
        "costs or a private deployment quote. The model/embedding dimension and quality "
        "may differ. If a permitted managed API meets the same privacy, quality, data "
        "residency, availability and latency contract, E1's assumed $180k/year private "
        "baseline is not the right alternative and its positive example does not apply. "
        f"The small-model token charge is $3,072 over three years versus E1's {money(SCENARIOS['E1'].project_tco)} "
        "hypothetical deployment cost before any API-side additions. Do not justify an "
        "appliance for ordinary embedding calls using an artificially expensive baseline."
    )


def rect(slide, x, y, width, height, fill=LIGHT, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        shape.line.color.rgb = RGBColor.from_string(border)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(SlideElement("a:effectLst"))
    for effect in shape._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    return shape


def line(slide, start_x, start_y, end_x, end_y, color=MUTED, width=1.5):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(start_x), Inches(start_y),
                                       Inches(end_x), Inches(end_y))
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    return shape


def flow_node(slide, label, x, y, width, height=0.92, color=BLUE, fill=LIGHT, font=16):
    rect(slide, x, y, width, height, fill, color)
    text(slide, label, x + 0.1, y + 0.055, width - 0.2, height - 0.1, font,
         WHITE if fill == INK else INK, True, align=PP_ALIGN.CENTER)


def arrow(slide, x, y, width=0.34, color=MUTED):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(width), Inches(0.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    shape._element.spPr.append(SlideElement("a:effectLst"))
    for effect in shape._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")


def section_label(slide, label, x, y, width=6.5, color=TEAL):
    text(slide, label, x, y, width, 0.31, 11.5, color, True)


def bullet_block(slide, values, x, y, width, size=17, row_height=0.78):
    for index, value in enumerate(values):
        position = y + index * row_height
        rect(slide, x, position + 0.12, 0.045, 0.17, TEAL)
        text(slide, value, x + 0.16, position, width - 0.16, row_height - 0.07, size)


def slide_table(slide, headers, rows, x, y, widths, row_height=0.65, size=15, row_colors=None):
    for row_index, values in enumerate([headers, *rows]):
        cursor = x
        for column_index, (value, width) in enumerate(zip(values, widths)):
            fill = INK if row_index == 0 else (LIGHT if row_index % 2 else WHITE)
            rect(slide, cursor, y + row_index * row_height, width, row_height, fill)
            color = WHITE if row_index == 0 else INK
            if row_index and row_colors and column_index:
                color = row_colors[row_index - 1]
            text(slide, str(value), cursor + 0.08, y + row_index * row_height + 0.1,
                 width - 0.16, row_height - 0.15, size if row_index else 11.5,
                 color, row_index == 0)
            cursor += width


def add_research_notes(slide, references=(), details=""):
    values = {reference: (title, url, relevance) for reference, title, url, relevance in SOURCES}
    extra = "\n" + details
    for reference in references:
        title, url, relevance = values[reference]
        extra += f"\n[{reference}] {title}\n{url}\nUse: {relevance} Reviewed {STAMP}."
    slide.notes_slide.notes_text_frame.text += extra


def add_finance_notes(slide, tracks):
    for track in tracks:
        model = SCENARIOS[track]
        slide.notes_slide.notes_text_frame.text += (
            f"\nILLUSTRATIVE {track}, constant USD, 36 months, no discounting or residual value. "
            f"Initial hardware {money(model.capex)} + deployment {money(model.deployment)} = {money(model.initial)}. "
            f"Annual support {money(model.support_year)}, software {money(model.software_year)}, "
            f"operations {money(model.operations_year)}, average IT {model.average_it_kw} kW. "
            f"Power cost = kW x 8760 x assumed PUE 1.4 x $0.12/kWh = {money(model.energy_year)}/year. "
            f"Annual recurring {money(model.recurring)}; avoidable baseline {money(model.baseline_avoidable_year)}; "
            f"retained {money(model.retained_year)}. Baseline 3y {money(model.baseline_tco)}; "
            f"project 3y {money(model.project_tco)}; net {money(model.net_benefit())}. "
            f"TCO saving {model.tco_saving:.1%}, upfront-investment ROI {model.roi:.1%}, "
            f"simple payback {model.payback_months:.1f} months. Break-even monetization "
            f"{model.break_even_fraction:.1%}. At 50% realization net benefit {money(model.net_benefit(0.5))}. "
            f"Annual demand {model.annual_units:,} {model.unit}; {model.configuration}. "
            "Node count and accepted rates must be measured under identical quality, SLO and HA. "
            "Freed installed-GPU capacity is noncash unless a named bill or future purchase is avoided. "
            f"Scenarios are independent/non-additive. Details: {DOC_OUTPUT.name}#Economics."
        )


def roi_panel(slide, track, x=8.02, y=3.02, width=4.72):
    item = next(value for value in OFFERINGS if value["id"] == track)
    model = SCENARIOS[track]
    rect(slide, x, y, width, 2.72, LIGHT)
    text(slide, "ILLUSTRATIVE 3-YEAR TCO SAVING", x + 0.19, y + 0.16, width - 0.38, 0.31,
         11.5, item["color"], True)
    text(slide, f"{model.tco_saving:.1%}", x + 0.17, y + 0.59, 2.33, 0.81, 47, item["color"], True)
    text(slide, f"{model.payback_months:.1f} months\nsimple payback", x + 2.6, y + 0.72,
         width - 2.79, 0.73, 15, INK, True)
    text(slide, f"{money(model.initial)} upfront | {money(model.recurring)}/year ATMOS",
         x + 0.18, y + 1.59, width - 0.36, 0.6, 15)
    text(slide, f"Cash benefit needs >{model.break_even_fraction:.1%} baseline-spend realization.",
         x + 0.18, y + 2.2, width - 0.36, 0.4, 12.5, MUTED)


def slide_portfolio(deck):
    slide = new_slide(deck, 1, "ATMOS E3.S AI Capacity Portfolio", "Portfolio", "1",
                      "Three enterprise deployment modes. One separate hyperscale neural-data track.", dark=True)
    for index, item in enumerate(OFFERINGS):
        position = 2.05 + index * 0.92
        rect(slide, 0.55, position, 0.05, 0.69, item["color"])
        text(slide, item["id"], 0.76, position + 0.08, 0.66, 0.42, 25, "F4B9C2", True)
        text(slide, item["name"], 1.65, position + 0.04, 5.04, 0.72, 19.5, WHITE, True,
             f"{DOC_OUTPUT.name}#{item['bookmark']}")
        text(slide, item["audience"], 7.1, position + 0.04, 5.25, 0.65, 16, "C9D7DF")
        if index < 3:
            line(slide, 1.65, position + 0.8, 12.62, position + 0.8, "41596A", 0.6)
    rect(slide, 0.55, 6.03, 12.23, 0.68, RED)
    text(slide, "LEAD E2: prove one complete service. Productize E1; integrate E3. Sponsor H only at scale.",
         0.73, 6.17, 11.9, 0.37, 16.5, WHITE, True)
    add_research_notes(slide, ["S4", "S10", "S11"],
                       "E1/E2/E3 are enterprise deployment modes, not additive pools of savings. "
                       "H is hyperscale-first. The supplied portfolio's removed research product "
                       "is absent from product scope, roadmap, pricing and claims. No launch is approved here.")


def slide_foundation(deck):
    slide = new_slide(deck, 2, "Localize bytes. Exchange compact results.", "Foundation", "2",
                      "One qualified four-device group is the atomic hardware and scheduling unit, not shared memory.")
    section_label(slide, "FOUR-DEVICE SWITCH / NUMA DOMAIN", 0.55, 1.97)
    flow_node(slide, "CPU / NUMA root\nGen5 x16", 0.76, 2.56, 2.1, 0.91, BLUE)
    flow_node(slide, "P2P-capable\nPCIe switch", 0.76, 3.79, 2.1, 0.91, TEAL)
    line(slide, 1.81, 3.47, 1.81, 3.79, TEAL)
    line(slide, 2.86, 4.23, 3.25, 4.23, TEAL)
    line(slide, 3.25, 2.56, 3.25, 4.64, TEAL)
    for index in range(4):
        position = 2.25 + index * 0.68
        flow_node(slide, f"ATMOS {index}  /  local HBF + NPU", 3.68, position, 3.58, 0.57, TEAL, font=14.5)
        line(slide, 3.25, position + 0.285, 3.68, position + 0.285, TEAL)
    section_label(slide, "STACK: REUSE UPSTREAM; QUALIFY THE ATMOS PATH", 7.78, 1.97, 5)
    stack = [
        "Gateway / identity / quota / complete-request routing",
        "KServe or llm-d / model API and live endpoint metrics",
        "DRA / CDI + ATMOS Operator / ready-drain-rollback",
        "ATMOS workload runtime / selected NPU SDK / packs",
        "ATMOS driver, firmware, DMA, P2P and telemetry",
    ]
    for index, label in enumerate(stack):
        position = 2.38 + index * 0.49
        rect(slide, 7.78, position, 5, 0.44, LIGHT if index % 2 == 0 else WHITE)
        text(slide, label, 7.88, position + 0.05, 4.8, 0.34, 13.2, INK, index == 3)
    metrics = [
        ("128 / 100 GB", "Nominal HBF / study budget per device"),
        ("200 GB/s", "Local HBF read target; not peer bandwidth"),
        ("~15.75 GB/s", "Gen5 x4 / direction before overhead"),
        ("8 -> 16", "Qualified node densities; 20-22 stretch"),
    ]
    for index, (value, label) in enumerate(metrics):
        position = 0.55 + index * 3.09
        rect(slide, position, 5.32, 2.94, 1.0, LIGHT)
        text(slide, value, position + 0.12, 5.43, 2.7, 0.37, 21, TEAL, True)
        text(slide, label, position + 0.12, 5.89, 2.7, 0.35, 11.3, MUTED)
    text(slide, "Host staging first. No unsafe isolation bypass. P2P/TT plugin/DRA is not ATMOS-ready model support. [S1-S3, S9]",
         0.55, 6.54, 12.15, 0.33, 12.2, MUTED)
    add_research_notes(slide, ["S1", "S2", "S3", "S9"],
                       "All ATMOS values are source planning assumptions. Per-device LPDDR 8-16 GB, "
                       "30-40 W device power, 64 TOPS-class source target with undefined precision/sparsity. "
                       "Two four-device x16 domains consume 32 host lanes. Peer DMA is a transport "
                       "capability, not a collective, coherent pool or connection to GPU native fabric. "
                       "ModelPack/ExpertPack and ATMOS-specific driver/runtime/operator require delivery.")


def slide_appliance(deck):
    slide = new_slide(deck, 3, "E1 / A turnkey private neural-service endpoint", "Appliance", "3",
                      "Sell a complete, supported workload first; full interactive LLM support is a separate qualification.")
    labels = ["Client / job\nmodel + tier", "Identity / quota\nAPI + queues", "ATMOS runtime\nqualified operators", "8 / 16 devices\nOEM + support"]
    for index, label in enumerate(labels):
        position = 0.55 + index * 3.09
        flow_node(slide, label, position, 2.02, 2.79, 0.87, BLUE, font=16)
        if index < 3:
            arrow(slide, position + 2.83, 2.36, 0.22)
    section_label(slide, "PRODUCT AND RELEASE ENVELOPE", 0.55, 3.18)
    bullet_block(slide, [
        "Lead with embeddings, rerank, retrieval or bounded document/batch operators.",
        "8 devices: 1.024 TB nominal HBF. 16: 2.048 TB, distributed across qualified domains.",
        "Turnkey means identity, tenant isolation, packs, canary/drain, metering, HA and one support prime.",
    ], 0.55, 3.61, 7.02, 16.8, 0.72)
    roi_panel(slide, "E1", y=3.18)
    rect(slide, 0.55, 6.08, 12.23, 0.58, LIGHT)
    text(slide, "SCENARIO: replace $180k/year private service; $120k upfront; 3 x 8-device nodes including N+1.",
         0.7, 6.2, 11.92, 0.32, 14, INK, True)
    text(slide, "Compare with CPU / managed API too. Positive ROI is conditional on the best permitted baseline. [S13]",
         0.55, 6.68, 12.14, 0.23, 10.8, MUTED)
    add_finance_notes(slide, ["E1"])
    add_research_notes(slide, ["S13"],
                       "Published embedding API prices can defeat the private-appliance business case. "
                       "100m requests x hypothetical 512 tokens = 51.2 billion input tokens/year. "
                       "At $0.02/million, small-model token charges are $1,024/year, excluding other "
                       "costs and requiring permitted privacy/quality/SLO. No direct API equivalence is assumed.")


def slide_companion(deck):
    slide = new_slide(deck, 4, "E2 / Add capacity without splitting an inference request", "Companion", "4",
                      "First enterprise integration priority: separate endpoint, failure domain and device lifecycle.")
    flow_node(slide, "Application / portal\nchoose before execution", 0.55, 2.06, 3.5, 1.0, TEAL)
    flow_node(slide, "Existing GPU endpoint\nstrict token SLO / VLM / long context", 5.03, 2.02, 7.74, 0.65, BLUE, font=16)
    flow_node(slide, "Independent ATMOS endpoint\nqualified neural services / batch", 5.03, 2.88, 7.74, 0.65, TEAL, font=16)
    line(slide, 4.05, 2.56, 4.55, 2.56)
    line(slide, 4.55, 2.345, 4.55, 3.205)
    arrow(slide, 4.59, 2.245, 0.32, BLUE)
    arrow(slide, 4.59, 3.105, 0.32, TEAL)
    section_label(slide, "LOW-INTEGRATION ADOPTION / NON-NEGOTIABLE TESTS", 0.55, 3.87)
    bullet_block(slide, [
        "No CUDA/ROCm, GPU firmware or native-collective modification. No peer-DMA dependency.",
        "Shared gateway, CPU, NIC, storage and power can still hurt GPU tails: test mixed saturation.",
        "Freed installed GPU capacity is not cash. Avoid a specific future purchase or bill.",
    ], 0.55, 4.24, 7.02, 16.2, 0.66)
    roi_panel(slide, "E2", y=3.87)
    text(slide, "SCENARIO: avoid $120k/year GPU overflow; $100k upfront; 3 x 8-device nodes, including N+1.",
         0.55, 6.72, 12.15, 0.23, 11.5, MUTED)
    add_finance_notes(slide, ["E2"])
    add_research_notes(slide, ["S3"],
                       "Request count share is not equal to GPU cost share. Peak, model residency, batch "
                       "efficiency and lease terms control removability. At only half the avoidable "
                       "baseline budget monetized, this scenario loses money over 36 months.")


def slide_capacity(deck):
    slide = new_slide(deck, 5, "E3 / One gateway. Two qualified hardware pools.", "CapacityPool", "5",
                      "Reuse the enterprise catalog and policy; select one compatible, ready backend for the whole request.")
    flow_node(slide, "Enterprise gateway\nidentity / quota / SLO / cost", 0.55, 2.04, 3.22, 0.91, RED)
    flow_node(slide, "GPU InferencePool\nexisting runtime", 4.35, 2.02, 3.1, 0.62, BLUE, font=15)
    flow_node(slide, "ATMOS InferencePool\nqualified model / pack", 4.35, 2.92, 3.1, 0.62, TEAL, font=15)
    line(slide, 3.77, 2.51, 4.03, 2.51)
    line(slide, 4.03, 2.33, 4.03, 3.23)
    arrow(slide, 4.05, 2.23, 0.23, BLUE)
    arrow(slide, 4.05, 3.13, 0.23, TEAL)
    section_label(slide, "DRA ALLOCATES DEVICES; THE OPERATOR MAKES A SERVICE READY", 0.55, 3.79, 7.05)
    slide_table(slide, ["Contract", "ATMOS responsibility"], [
        ["DRA / CDI", "Group, NUMA, topology, allocation and isolation"],
        ["Operator", "ABI / pack / warmup -> ready -> drain / rollback"],
        ["Gateway", "Same model/API quality; live queue and safe overload"],
    ], 0.55, 4.19, [1.58, 5.6], 0.55, 14.5)
    roi_panel(slide, "E3", y=2.03)
    text(slide, "vs E2: unified platform, identity and chargeback; more lifecycle/routing integration.",
         8.21, 5.04, 4.38, 0.93, 16)
    text(slide, "SCENARIO: $1m/year estate; $300k removable, $700k retained; $180k upfront. Full-estate denominator.",
         0.55, 6.68, 12.15, 0.25, 11.5, MUTED)
    add_finance_notes(slide, ["E3"])
    add_research_notes(slide, ["S2", "S3", "S12"],
                       "DRA does not imply atomic P2P groups unless represented/enforced by the vendor "
                       "driver, model readiness, application failover or device preemption. Pin features "
                       "to selected Kubernetes/API versions. E3 assumes three 16-device nodes, two active "
                       "and one N+1; measure surviving peak capacity. 500m requests/year use a frozen estate mix.")


def slide_hyperscale(deck):
    slide = new_slide(deck, 6, "H / Hyperscale Neural-Data Sidecar", "HyperscaleSidecar", "6",
                      "Hyperscalers and Meta / LinkedIn-like operators; enterprise only with exceptional scale and proven pain.")
    labels = ["Existing ANN / candidates\nquery + versioned IDs", "ATMOS local vector / feature shards\ngather, score, pool, top-k", "Compact results -> GPU / app\nranker, attention, KV, generation"]
    widths = [3.34, 4.21, 3.95]
    positions = [0.55, 4.32, 8.83]
    for index, (label, position, width) in enumerate(zip(labels, positions, widths)):
        flow_node(slide, label, position, 2.01, width, 0.99, AMBER if index == 0 else TEAL, font=15.5)
        if index < 2:
            arrow(slide, position + width + 0.06, 2.38, 0.24)
    line(slide, 0.55, 3.29, 12.78, 3.29, "D9E2E8", 1)
    section_label(slide, "MEMORY AND LIFECYCLE / HYPOTHETICAL", 0.55, 3.54, 4.06, AMBER)
    text(slide, "500m x 768 x FP16\n= 768 GB raw vectors", 0.55, 3.99, 4.02, 0.89, 20, INK, True)
    text(slide, "960 GB with a 25% layout allowance.\nFour 16-device nodes: 2 generations x 2 complete replicas; each shard must fit.",
         0.55, 5.07, 4.03, 1.14, 15.5)
    section_label(slide, "REDUCE BEFORE EXPENSIVE RANKING", 4.85, 3.54, 3.79)
    text(slide, "20,000 -> 200", 4.85, 4.02, 3.62, 0.54, 26, TEAL, True)
    text(slide, "99% fewer candidates, only if relevance is retained; not 99% total savings.",
         4.85, 4.73, 3.65, 0.83, 16)
    text(slide, "30.72 MB vectors -> 2.4 kB IDs/scores ONLY for a vector-moving baseline. Input IDs / refresh / features excluded.",
         4.85, 5.64, 3.65, 0.95, 13.8, MUTED)
    model = SCENARIOS["H"]
    rect(slide, 8.83, 3.54, 3.95, 2.88, LIGHT)
    text(slide, "ILLUSTRATIVE WHOLE-PIPELINE TCO", 9.01, 3.7, 3.59, 0.35, 11.2, AMBER, True)
    text(slide, f"{model.tco_saving:.1%}", 8.99, 4.2, 3.35, 0.84, 46, AMBER, True)
    text(slide, "$600k/year tier removable; $900k/year GPU retained. $360k upfront, including replicas / staging nodes.",
         9.01, 5.17, 3.54, 1.02, 15.5)
    text(slide, "Gate: best CPU/DRAM ANN control + recall/NDCG + freshness + end-to-end p99 + deletions/ACL + actual tier removal. [S4, S10-S11]",
         0.55, 6.66, 12.15, 0.27, 11.3, MUTED)
    add_finance_notes(slide, ["H"])
    add_research_notes(slide, ["S4", "S5", "S10", "S11"],
                       "Companies named are market archetypes, not partners/customers/endorsers. "
                       "No cold-expert or Transformer-layer execution split is in the sidecar product. "
                       "500m*768*2=768e9 bytes; 25% allowance gives 960e9; two generations x two "
                       "replicas=3.84e12 bytes, four separately qualified 16-device nodes at 100GB "
                       "study budget/device. A surviving complete live node must meet the peak SLO. "
                       "20k*768*2=30.72e6 bytes; 200*(8+4)=2400 bytes. Actual input IDs, query, "
                       "headers, selected features, merge, update/replica traffic are excluded. "
                       "If incumbent ANN already emits compact IDs, this vector-movement benefit "
                       "may be zero. No full scan of the corpus per query. HBF ingest/update is unproven.")


def slide_economics(deck):
    slide = new_slide(deck, 7, "Cost advantage needs accepted service AND avoided spend", "Economics", "8",
                      "Illustrative constant-USD 36-month cash models. Same quality, SLO, demand and availability; not ATMOS quotes or benchmarks.")
    section_label(slide, "3-YEAR TCO / EACH BASELINE NORMALIZED TO 100", 0.55, 2.02, 6.0)
    for index, item in enumerate(OFFERINGS):
        model = SCENARIOS[item["id"]]
        position = 2.55 + index * 0.91
        text(slide, item["id"], 0.55, position + 0.04, 0.48, 0.34, 19, item["color"], True)
        rect(slide, 1.28, position + 0.01, 4.48, 0.16, "D9E2E8")
        rect(slide, 1.28, position + 0.26, 4.48 * (1 - model.tco_saving), 0.23, item["color"])
        text(slide, f"{model.tco_saving:.1%} lower", 1.28, position + 0.56, 2.53, 0.28, 13.5, item["color"], True)
        text(slide, f"{compact_money(model.baseline_tco)} -> {compact_money(model.project_tco)}",
             3.93, position + 0.53, 2.4, 0.34, 12.8, MUTED, align=PP_ALIGN.RIGHT)
    section_label(slide, "100% OF SPECIFIED AVOIDABLE BUDGET REALIZED", 6.82, 2.02, 5.95)
    slide_table(slide, ["Track", "Upfront", "3y net benefit", "Payback", "3y ROI"], [
        [track, compact_money(model.initial), money(model.net_benefit()),
         f"{model.payback_months:.1f} mo", f"{model.roi:.0%}"]
        for track, model in SCENARIOS.items()
    ], 6.82, 2.48, [0.72, 1.24, 1.50, 1.24, 1.26], 0.61, 14)
    text(slide, "E3 retains $700k/year GPU services. H retains $900k/year. Denominators cover the whole unchanged service/estate.",
         6.9, 5.8, 5.75, 0.76, 15)
    text(slide, "Includes hardware, rollout, support, software, operations, electricity and reserve. No double-counting / no additive portfolio savings.",
         0.55, 6.63, 12.15, 0.29, 11.7, MUTED)
    add_finance_notes(slide, SCENARIOS)
    add_research_notes(slide, [],
                       "Simple ROI is cumulative 36-month net cash benefit divided by upfront investment, "
                       "not vendor gross margin. Baseline provider power is already in cloud/lease/service "
                       "budgets. No sunk-capital refund, tax, financing, residual value or revenue uplift. "
                       "Customer quotes must include allocated supplier program NRE. All assumed service "
                       "rate and node count inputs remain to be proven; see per-track ledgers and unit rates.")


def slide_sensitivity(deck):
    slide = new_slide(deck, 8, "Where the business case breaks", "Economics", "8",
                      "Realized fraction means baseline bills or future purchases actually avoided, not just requests routed elsewhere.")
    section_label(slide, "NET 3-YEAR BENEFIT AT DIFFERENT REALIZATION LEVELS", 0.55, 2.04, 7.95)
    slide_table(slide, ["Track", "50% realized", "75% realized", "100% realized", "Break-even"], [
        [track, money(model.net_benefit(0.5)), money(model.net_benefit(0.75)),
         money(model.net_benefit()), f">{model.break_even_fraction:.1%}"]
        for track, model in SCENARIOS.items()
    ], 0.55, 2.53, [0.68, 1.76, 1.78, 1.78, 1.65], 0.68, 15,
                row_colors=[TEAL, RED, TEAL, RED])
    section_label(slide, "THE CHEAPER ALTERNATIVE CAN WIN", 8.78, 2.04, 3.97, AMBER)
    rect(slide, 8.78, 2.53, 4.0, 2.62, LIGHT)
    text(slide, "$1,024/year", 8.97, 2.79, 3.59, 0.64, 28, AMBER, True)
    text(slide, "Published small-embedding API token charges at 100m x 512 tokens. $0.02 per million tokens. [S13]",
         8.97, 3.63, 3.58, 1.18, 16)
    bullet_block(slide, [
        "Freed GPU capacity is noncash until a bill or future purchase is avoided.",
        "Stress-test goodput, node count, refresh/HA, support/NRE and energy.",
    ], 0.55, 6.08, 7.62, 14.5, 0.37)
    text(slide, "API is not an equal private-service quote. If privacy, quality and SLO permit it, the assumed appliance baseline must change.",
         8.82, 5.49, 3.9, 1.03, 15.5)
    add_finance_notes(slide, SCENARIOS)
    add_research_notes(slide, ["S13"],
                       "51.2 billion tokens / one million * $0.02 = $1,024 yearly token charge; "
                       "at $0.13 text-embedding-3-large charge=$6,656. These published standard "
                       "prices reviewed 5 Oct 2026 exclude retrieval, network, storage, support and "
                       "private-policy requirements. No quality/model identity assumption is made. "
                       "E2 and H net return are negative at 50% realized avoided cost. E1 and E3 "
                       "positive cases are small at that level and remain subject to measured rates/quotes.")


def slide_qualification(deck):
    slide = new_slide(deck, 9, "OEM density and service are evidence gates, not bay counts", "OEM", "9 / 11",
                      "A four-device fixture -> supported 8-device pilot -> 16-device standard. Hold 20-22 until customer/OEM proof.")
    stages = [
        ("1 DEVICE", "Correctness / useful rate\npower / thermal / reset"),
        ("4 DEVICES", "P2P + staging / skew\nisolation / safe removal"),
        ("8 DEVICES", "Whole service / tenants\nNUMA / complete replicas"),
        ("16 DEVICES", "OEM BOM / full soak\nN+1 peak service / support"),
        ("CUSTOMER", "E2 / E3 routing or H\nquality + cost gate"),
    ]
    for index, (title, label) in enumerate(stages):
        position = 0.55 + index * 2.5
        rect(slide, position, 2.08, 2.23, 1.14, LIGHT)
        text(slide, title, position + 0.12, 2.22, 2.0, 0.32, 15, TEAL, True)
        text(slide, label, position + 0.12, 2.68, 2.0, 0.45, 11.7)
        if index < 4:
            arrow(slide, position + 2.27, 2.58, 0.19)
    section_label(slide, "WHAT THE OEM / PRIME MUST SIGN", 0.55, 3.64, 5.84)
    bullet_block(slide, [
        "E3.S 2T / x4 lanes / BAR and security / BIOS-BMC / sustained bay power and cooling.",
        "Wall power at maximum inlet; CPU/NIC/storage load; safe replacement and full replica reserve.",
        "Named BOM, NRE, firmware/driver/runtime matrix, spares and one support escalation route.",
    ], 0.55, 4.1, 5.85, 16.1, 0.73)
    section_label(slide, "WHAT THE SERVICE / FINANCE GATE MUST PROVE", 6.85, 3.64, 5.93, RED)
    bullet_block(slide, [
        "Same quality, p99 / TTFT / TPOT where relevant, freshness and accepted service under failures.",
        "Protect GPU tails and tenant isolation; no forced unqualified routing or unsafe DMA bypass.",
        "Best incumbent/CPU/GPU comparison + measured node count + actual quotes + removable spend.",
    ], 6.85, 4.1, 5.93, 16.1, 0.73)
    text(slide, "Public 32-bay E3.S 1T x2 storage density does not certify ATMOS x4/2T accelerators. Simulation cannot pass physical gates. [S1, S6]",
         0.55, 6.65, 12.15, 0.29, 11.1, MUTED)
    add_research_notes(slide, ["S1", "S6", "S7", "S8"],
                       "Eight-module power is 240-320 W device-only; sixteen is 480-640 W. "
                       "CPU/DRAM, switch/retimer/backplane, NIC/SSD, fans and conversion loss are additional. "
                       "Disjoint shards in separate groups do not provide full-model/index HA. "
                       "High-end GPU controls have substantially larger/faster HBM than older capacity "
                       "arguments assume; compare exact service units, not a rack versus a node.")


def slide_decisions(deck):
    slide = new_slide(deck, 10, "Fund bounded proof. Hold launch and savings claims.", "Decisions", "12",
                      "One shared foundation, clear ownership, four retained tracks and a linked technical/economic decision package.")
    section_label(slide, "NEXT DECISIONS / ACCOUNTABLE OWNERS", 0.55, 2.01, 7.25)
    decisions = [
        ("E2 FIRST", "Product + customer: freeze one paid workload, best control and removable overflow block."),
        ("E1 PRODUCT", "OEM + SanDisk: qualify 4 -> 8 -> 16 devices, full-service reserve and one support prime."),
        ("E3 INTEGRATE", "Platform owner: DRA/operator/gateway contract, GPU tail protection and cost-aware policy."),
        ("H AT SCALE", "Hyperscale sponsor: index/update trace, relevance/freshness, p99 and tier-removal charter."),
        ("FINANCE GATE", "Finance + engineering: substitute matched physical rates and binding quotes; reject below break-even."),
    ]
    for index, (label, body) in enumerate(decisions):
        position = 2.51 + index * 0.78
        text(slide, label, 0.55, position, 1.53, 0.44, 12.3, RED if index == 4 else TEAL, True)
        text(slide, body, 2.29, position, 5.43, 0.64, 16)
    rect(slide, 8.22, 2.03, 4.56, 4.65, LIGHT)
    section_label(slide, "LINKED WORD DETAIL NAVIGATOR", 8.41, 2.23, 4.16, BLUE)
    destinations = [
        ("Foundation / software boundary", "2", "Foundation"),
        ("Enterprise E1 / E2 / E3", "3-5", "Appliance"),
        ("Hyperscale H / sizing / refresh", "6", "HyperscaleSidecar"),
        ("Cost / ROI / downside / API control", "8", "Economics"),
        ("OEM / support / gates", "9-11", "OEM"),
        ("Research / source and claim register", "A", "Sources"),
    ]
    for index, (label, section, bookmark) in enumerate(destinations):
        position = 2.83 + index * 0.53
        text(slide, f"{section}   {label}", 8.4, position, 4.15, 0.43, 13.6, BLUE,
             False, f"{DOC_OUTPUT.name}#{bookmark}")
    text(slide, "Keep PPT + DOCX together.\nUse 'Open Word proposal' if a viewer ignores section bookmarks.",
         8.4, 6.16, 4.12, 0.48, 10.4, MUTED)
    add_research_notes(slide, [],
                       "Each slide provides a relative full-document link and an explicit bookmark link. "
                       "All destinations refer to the paired v2 DOCX in the same directory. Viewer/Office "
                       "security policies can require confirmation; section numbers are a human-readable "
                       "fallback. Word contains reciprocal deck links and an internal slide-to-section guide. "
                       "No silicon launch date, benchmark, quote or named-company engagement is invented.")


def build_presentation():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for builder in (slide_portfolio, slide_foundation, slide_appliance, slide_companion,
                    slide_capacity, slide_hyperscale, slide_economics, slide_sensitivity,
                    slide_qualification, slide_decisions):
        builder(deck)
    deck.core_properties.title = "ATMOS E3.S AI Capacity Portfolio"
    deck.core_properties.author = deck.core_properties.last_modified_by = "SanDisk / ATMOS Architecture"
    deck.core_properties.revision = 2
    deck.core_properties.created = deck.core_properties.modified = datetime.now(timezone.utc)
    return deck


def validate_pair(document, deck):
    blocks = [paragraph.text for paragraph in document.paragraphs]
    blocks.extend(cell.text for item in document.tables for row in item.rows for cell in row.cells)
    doc_text = "\n".join(blocks)
    slide_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    for item in OFFERINGS:
        if item["name"] not in doc_text or item["name"] not in slide_text:
            raise AssertionError(f"Offering missing from companion outputs: {item['id']}")
    for model in SCENARIOS.values():
        if not isclose(model.baseline_tco - model.project_tco, model.net_benefit(), abs_tol=0.0001):
            raise AssertionError(f"Economic reconciliation failed: {model.track}")
        if not isclose(model.net_benefit(model.break_even_fraction), 0, abs_tol=0.0001):
            raise AssertionError(f"Break-even arithmetic failed: {model.track}")
        if money(model.net_benefit()) not in doc_text or money(model.net_benefit()) not in slide_text:
            raise AssertionError(f"Economic values are not synchronized: {model.track}")
    for forbidden in ("ATMOS Decode Tier", "Decode Research Kit", "Offering 4", "Offering 5"):
        if forbidden.lower() in (doc_text + slide_text).lower():
            raise AssertionError(f"Retired product reference remains: {forbidden}")
    bookmarks = set(document._element.xpath("//w:bookmarkStart/@w:name"))
    for slide in deck.slides:
        links = [run.hyperlink.address for shape in slide.shapes if shape.has_text_frame
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                 if run.hyperlink.address]
        links.extend(shape.click_action.hyperlink.address for shape in slide.shapes
                     if shape.has_text_frame and shape.click_action.hyperlink.address)
        if not links:
            raise AssertionError("Slide has no companion-document detail link")
        for destination in links:
            filename, separator, bookmark = destination.partition("#")
            if filename != DOC_OUTPUT.name or (separator and bookmark not in bookmarks):
                raise AssertionError(f"Unresolved companion link: {destination}")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + Inches(0.01) or shape.top + shape.height > deck.slide_height + Inches(0.01):
                raise AssertionError(f"Shape outside slide: {shape.name}")
    if len(deck.slides) != len(SLIDE_GUIDE):
        raise AssertionError("Deck and document reader guide have different slide counts")
    for bookmark in document._element.xpath("//w:hyperlink/@w:anchor"):
        if bookmark not in bookmarks:
            raise AssertionError(f"Unresolved Word navigation bookmark: {bookmark}")


def main():
    build_diagrams()
    document = build_document()
    deck = build_presentation()
    validate_pair(document, deck)
    document.save(DOC_OUTPUT)
    deck.save(PPT_OUTPUT)
    validate_pair(Document(DOC_OUTPUT), Presentation(PPT_OUTPUT))
    print(f"Generated and reopened {DOC_OUTPUT.name} and {PPT_OUTPUT.name}")
    print(f"Retained tracks={len(OFFERINGS)}; slides={len(deck.slides)}; companion links validated")


if __name__ == "__main__":
    main()