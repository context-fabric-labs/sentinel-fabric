from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


DOCS_DIR = Path(__file__).resolve().parent
OUTPUT = DOCS_DIR / "SanDisk_ATMOS_Three_Offerings_Technical_Value_Proposal_v4.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = "F6F8FA"
DARK = "111827"
NAVY = "102A43"
INK = "24384A"
MUTED = "64748B"
LINE = "D7E0E7"
WHITE = "FFFFFF"
RED = "E31837"
BLUE = "2E6F9E"
TEAL = "008C82"
GREEN = "4D8B61"
ORANGE = "D9823B"
PURPLE = "6B5AA6"
PALE_BLUE = "E7F0F7"
PALE_TEAL = "E4F3F0"
PALE_ORANGE = "FBEDE1"
PALE_RED = "F7E7E7"

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"
FONT_MONO = "Aptos Mono"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_background(slide, color: str = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def set_shape_fill(shape, color: str, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def set_shape_line(shape, color: str, width: float = 1.0, transparency: int = 0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_text(slide, text: str, x: float, y: float, w: float, h: float, *,
             size: float = 16, color: str = INK, bold: bool = False,
             font: str = FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin: float = 0.03, name: str | None = None) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs: list[tuple[str, bool, str]], x: float, y: float,
                  w: float, h: float, *, size: float = 16,
                  align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
                  name: str | None = None) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    for text, bold, color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x: float, y: float, w: float, h: float, *, fill: str = WHITE,
             line: str = LINE, radius: bool = True, line_width: float = 1.0,
             name: str | None = None) -> object:
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, line_width)
    return shape


def add_badge(slide, text: str, x: float, y: float, w: float, *, fill: str,
              color: str = WHITE, size: float = 10.5) -> None:
    add_rect(slide, x, y, w, 0.30, fill=fill, line=fill, radius=True)
    add_text(slide, text, x + 0.04, y + 0.02, w - 0.08, 0.25, size=size,
             color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *,
             color: str = LINE, width: float = 1.2, arrow: bool = False) -> object:
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        connector.line._get_or_add_ln().append(end)
    return connector


def add_header(slide, number: int, section: str, title: str, subtitle: str = "") -> None:
    add_text(slide, "SANDISK / ATMOS", 0.55, 0.28, 1.62, 0.28, size=10, color=RED, bold=True)
    add_text(slide, section.upper(), 2.35, 0.28, 4.2, 0.28, size=10, color=MUTED, bold=True)
    add_text(slide, title, 0.55, 0.72, 12.1, 0.52, size=27, color=NAVY,
             bold=True, font=FONT_DISPLAY)
    if subtitle:
        add_text(slide, subtitle, 0.57, 1.27, 12.0, 0.42, size=13.5, color=MUTED)
    add_line(slide, 0.55, 1.75, 12.78, 1.75, color=LINE, width=1)
    add_text(slide, f"{number:02d}", 12.25, 0.28, 0.50, 0.28, size=10,
             color=MUTED, bold=True, align=PP_ALIGN.RIGHT)


def add_footer(slide, number: int, note: str = "") -> None:
    add_line(slide, 0.55, 7.12, 12.78, 7.12, color=LINE, width=0.8)
    add_text(slide, "INTERNAL TECHNICAL PROPOSAL • 02 OCT 2026 • NO MEASURED ATMOS PERFORMANCE CLAIM",
             0.55, 7.18, 8.90, 0.18, size=7.5, color=MUTED, bold=True)
    if note:
        add_text(slide, note, 9.72, 7.18, 2.43, 0.18, size=7.3, color=MUTED,
                 align=PP_ALIGN.RIGHT)
    add_text(slide, str(number), 12.43, 7.18, 0.35, 0.18, size=7.5, color=MUTED,
             bold=True, align=PP_ALIGN.RIGHT)


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float,
                *, size: float = 14, color: str = INK, bullet_color: str = RED,
                gap: float = 0.45) -> None:
    for index, item in enumerate(items):
        yy = y + index * gap
        add_text(slide, "•", x, yy, 0.20, gap, size=size + 1, color=bullet_color,
                 bold=True, valign=MSO_ANCHOR.TOP)
        add_text(slide, item, x + 0.25, yy, w - 0.25, gap, size=size, color=color)


def add_metric(slide, value: str, label: str, x: float, y: float, w: float,
               *, accent: str = BLUE, value_size: float = 25) -> None:
    add_rect(slide, x, y, w, 0.95, fill=WHITE, line=LINE, radius=True)
    add_text(slide, value, x + 0.14, y + 0.10, w - 0.28, 0.38, size=value_size,
             color=accent, bold=True, font=FONT_DISPLAY)
    add_text(slide, label, x + 0.14, y + 0.52, w - 0.28, 0.28, size=10.5,
             color=MUTED, bold=True)


def add_cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DARK)
    add_rect(slide, 0, 0, 0.18, 7.5, fill=RED, line=RED, radius=False)
    add_text(slide, "SANDISK / ATMOS", 0.65, 0.45, 2.2, 0.32, size=11,
             color=RED, bold=True)
    add_text(slide, "THREE TECHNICAL OFFERINGS", 0.65, 0.88, 5.4, 0.30,
             size=11, color="9FB3C8", bold=True)
    add_text(slide, "ATMOS in a Blackwell / Hopper\ninference architecture", 0.65,
             1.35, 11.5, 1.35, size=35, color=WHITE, bold=True, font=FONT_DISPLAY)
    add_text(slide, "Execution boundaries, technical value, and the evidence required before productization",
             0.68, 2.82, 10.8, 0.38, size=16, color="C7D3DF")

    cards = [
        ("A", "NATIVE SERVICE", "Request boundary", "Complete prefill + decode on ATMOS", BLUE),
        ("B", "DISAGGREGATED", "KV boundary", "GPU prefill → ATMOS decode", TEAL),
        ("C", "SPECULATIVE", "Token-block boundary", "GPU drafts → ATMOS verifies", ORANGE),
    ]
    for index, (letter, name, boundary, desc, accent) in enumerate(cards):
        x = 0.68 + index * 4.05
        add_rect(slide, x, 3.72, 3.72, 1.62, fill="172235", line="334155", radius=True)
        add_text(slide, letter, x + 0.18, 3.88, 0.46, 0.48, size=27, color=accent,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, name, x + 0.72, 3.88, 2.72, 0.30, size=12, color=WHITE, bold=True)
        add_text(slide, boundary, x + 0.72, 4.22, 2.72, 0.24, size=10.5,
                 color="9FB3C8", bold=True)
        add_text(slide, desc, x + 0.18, 4.70, 3.30, 0.38, size=13, color="DDE5ED")

    add_rect(slide, 0.68, 5.78, 11.82, 0.72, fill="2A1D25", line="5D2734", radius=True)
    add_rich_text(slide, [
        ("RECOMMENDATION  ", True, RED),
        ("Lead with A; qualify B only after a measured KV handoff; retain C as a research option with an automatic normal-decode fallback.", False, WHITE),
    ], 0.90, 5.94, 11.35, 0.37, size=14)
    add_text(slide, "INTERNAL • 02 OCT 2026 • PRODUCT AND PERFORMANCE SUBJECT TO QUALIFICATION",
             0.68, 7.07, 9.8, 0.18, size=7.5, color="8294A8", bold=True)
    add_text(slide, "01", 12.08, 7.07, 0.40, 0.18, size=7.5, color="8294A8",
             bold=True, align=PP_ALIGN.RIGHT)
    add_notes(slide, """
Open with execution boundaries, not customer names or cost forecasts. Offering A is an independent full-model service. Offering B performs complete target-model prefill on NVIDIA and transfers target KV plus continuation state once to an ATMOS-native decoder. Offering C uses a separate small draft model on GPU; ATMOS owns the full target model and verifies candidate token blocks with target-correct acceptance/correction. C is not voting, a quality classifier, or KV handoff from B.

Recommendation: A is the least intrusive integration because requests and responses cross the boundary. B is conditional on a compatible KV layout/connector and measured TTFT/TPOT benefit; public vLLM documentation explicitly states disaggregated prefilling does not inherently improve throughput. C is conditional on acceptance rate and verification economics and must fall back to normal target decode when it loses.

No measured ATMOS performance is asserted. All ATMOS device capacities, local bandwidth, runtime, power and support remain internal planning targets until qualified.
""")


def add_nvidia_baseline(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 2, "NVIDIA baseline", "ATMOS competes with a moving GPU memory + fabric target",
               "Hopper is the installed-base control; Blackwell Ultra is the strategic rack-scale control.")
    columns = [
        (0.58, "RTX PRO 6000", "BLACKWELL / PCIe", RED, [
            ("96 GB", "GDDR7 / GPU"), ("1.597 TB/s", "memory BW / GPU"), ("PCIe Gen5", "host interface")
        ], "Best development and OEM integration control", "8-GPU MGX examples exist; 768 GB aggregate GDDR is calculated."),
        (4.48, "H200", "HOPPER / SXM or NVL", BLUE, [
            ("141 GB", "HBM3e / GPU"), ("4.8 TB/s", "memory BW / GPU"), ("900 GB/s", "NVLink / GPU")
        ], "Installed-base high-bandwidth control", "4/8-GPU HGX and enterprise H200 NVL systems are available."),
        (8.38, "GB300 NVL72", "BLACKWELL ULTRA / RACK", TEAL, [
            ("72", "Blackwell Ultra GPUs"), ("20 TB", "aggregate GPU memory"), ("130 TB/s", "rack NVLink")
        ], "Strategic rack product control", "≈288 GB HBM3e/GPU; 576 TB/s aggregate memory bandwidth; liquid cooled."),
    ]
    for x, name, arch, accent, metrics, role, implication in columns:
        add_rect(slide, x, 2.02, 3.55, 4.48, fill=WHITE, line=LINE, radius=True)
        add_badge(slide, arch, x + 0.18, 2.18, 1.75, fill=accent, size=9.5)
        add_text(slide, name, x + 0.18, 2.62, 3.15, 0.36, size=21, color=NAVY,
                 bold=True, font=FONT_DISPLAY)
        for idx, (value, label) in enumerate(metrics):
            yy = 3.17 + idx * 0.69
            add_text(slide, value, x + 0.20, yy, 1.30, 0.34, size=20, color=accent,
                     bold=True, font=FONT_DISPLAY)
            add_text(slide, label, x + 1.48, yy + 0.05, 1.75, 0.26, size=10.5,
                     color=MUTED, bold=True)
        add_line(slide, x + 0.18, 5.28, x + 3.37, 5.28, color=LINE, width=0.8)
        add_text(slide, role, x + 0.20, 5.43, 3.12, 0.36, size=12.5, color=INK, bold=True)
        add_text(slide, implication, x + 0.20, 5.88, 3.12, 0.46, size=10.5, color=MUTED)
    add_rect(slide, 0.58, 6.66, 11.35, 0.34, fill=PALE_RED, line=PALE_RED, radius=True)
    add_rich_text(slide, [
        ("ATMOS implication  ", True, RED),
        ("As native GPU memory and fabric improve, ATMOS cannot rely on 'the model does not fit.' It must prove workload separation, phase disaggregation, or more accepted work per rack/power/support envelope.", False, INK),
    ], 0.72, 6.69, 11.05, 0.26, size=11.2)
    add_footer(slide, 2, "NVIDIA public specs • calculations labeled")
    add_notes(slide, """
Sources checked 2 Oct 2026:
- NVIDIA RTX PRO 6000 Blackwell Server Edition: 96 GB GDDR7, 1,597 GB/s, PCIe Gen5; available now. https://www.nvidia.com/en-us/data-center/rtx-pro-6000-blackwell-server-edition/
- NVIDIA RTX PRO Server: MGX 4U/6U examples support 8 RTX PRO 6000 GPUs. https://www.nvidia.com/en-us/data-center/products/rtx-pro-server/
- NVIDIA H200: 141 GB HBM3e, 4.8 TB/s, up to 700 W SXM / 600 W NVL, NVLink 900 GB/s per GPU, HGX 4/8 GPU or H200 NVL options. https://www.nvidia.com/en-us/data-center/h200/
- NVIDIA GB300 NVL72: 72 Blackwell Ultra GPUs, 36 Grace CPUs, 20 TB rounded aggregate GPU memory, up to 576 TB/s aggregate GPU-memory bandwidth and 130 TB/s rack NVLink; available now and liquid cooled. https://www.nvidia.com/en-us/data-center/gb300-nvl72/
- NVIDIA Blackwell Ultra datasheet: 288 GB HBM3e per B300 GPU. https://resources.nvidia.com/en-us-blackwell-architecture/blackwell-ultra-datasheet

Calculated arithmetic: 8 × 96 GB RTX PRO = 768 GB aggregate GDDR. 8 × 141 GB H200 = 1,128 GB aggregate HBM and 8 × 4.8 TB/s = 38.4 TB/s arithmetic local-memory bandwidth. These sums are not uniform shared memory and do not represent application throughput.

Technical interpretation: RTX PRO is the most accessible Blackwell development/OEM control. H200 represents a mature Hopper installed base with very high HBM bandwidth. GB300 NVL72 is the strategic rack control whose native memory and NVLink reduce any pure capacity argument for ATMOS. Qualification must hold model, quality, latency and failure reserve constant.
""")


def add_hypothetical_workload(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 3, "Hypothetical study", "A bounded MoE service with numbers we can defend",
               "No customer reference: this is a technical sizing profile, not a deployment claim.")

    add_rect(slide, 0.58, 2.02, 3.25, 4.74, fill=NAVY, line=NAVY, radius=True)
    add_text(slide, "REGULATED DOCUMENT\nSYNTHESIS", 0.83, 2.25, 2.75, 0.72,
             size=19, color=WHITE, bold=True, font=FONT_DISPLAY)
    add_text(slide, "Qwen3-235B-A22B class", 0.83, 3.17, 2.70, 0.30,
             size=13.5, color="BFD7EA", bold=True)
    workload = [
        ("4,096", "input tokens"), ("≤512", "output tokens"),
        ("8K", "total-context cap"), ("2 RPS", "peak study load"),
        ("≤120 s", "proposed p95 completion")
    ]
    for idx, (value, label) in enumerate(workload):
        yy = 3.72 + idx * 0.53
        add_text(slide, value, 0.86, yy, 0.85, 0.28, size=16.5, color=WHITE,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, label, 1.78, yy + 0.03, 1.55, 0.24, size=10.5,
                 color="C7D3DF", bold=True)

    add_rect(slide, 4.12, 2.02, 4.18, 4.74, fill=WHITE, line=LINE, radius=True)
    add_text(slide, "CAPACITY SCREEN", 4.37, 2.24, 2.2, 0.28, size=11,
             color=RED, bold=True)
    add_metric(slide, "≈235 GB", "raw FP8 model-weight floor", 4.37, 2.68, 1.72,
               accent=RED, value_size=21)
    add_metric(slide, "0.734 GiB", "BF16 logical KV at 4K", 6.28, 2.68, 1.72,
               accent=BLUE, value_size=20)
    add_metric(slide, "1.469 GiB", "BF16 logical KV at 8K", 4.37, 3.82, 1.72,
               accent=BLUE, value_size=20)
    add_metric(slide, "47 GiB", "32 sequences × 8K KV", 6.28, 3.82, 1.72,
               accent=TEAL, value_size=21)
    add_text(slide, "ATMOS study envelope", 4.37, 5.15, 2.5, 0.28, size=12,
             color=NAVY, bold=True)
    add_bullets(slide, [
        "8 × 128 GB HBF = 1,024 GB nominal",
        "800 GB application weight budget",
        "128 GB LPDDR nominal; 96 GB after provisional reserve",
    ], 4.37, 5.50, 3.62, 1.05, size=10.8, gap=0.34)

    add_rect(slide, 8.58, 2.02, 4.17, 4.74, fill=WHITE, line=LINE, radius=True)
    add_text(slide, "WHAT THE STUDY MUST PROVE", 8.83, 2.24, 3.50, 0.28,
             size=11, color=TEAL, bold=True)
    add_bullets(slide, [
        "Full-model compiler + attention/KV + MoE runtime",
        "Per-device working-memory fit—not aggregate arithmetic",
        "Native device fabric, queueing, failure and recovery",
        "Quality and p95 deadline at the declared load",
        "Accepted work versus optimized GPU controls",
    ], 8.83, 2.73, 3.55, 2.35, size=12.1, gap=0.47)
    add_rect(slide, 8.83, 5.48, 3.48, 0.92, fill=PALE_ORANGE, line="E9C8A8", radius=True)
    add_text(slide, "Capacity fit is only the admission ticket.\nCompute, active KV and fabric decide the service.",
             9.02, 5.66, 3.12, 0.52, size=11.2, color=INK, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_footer(slide, 3, "hypothetical inputs • arithmetic disclosed")
    add_notes(slide, """
Hypothetical use case, not tied to a named customer: regulated document synthesis with 4,096 input tokens, up to 512 generated tokens, 8K total-context cap, 2 requests/s peak and proposed p95 completion <=120 s. These are study inputs, not a measured SLA.

Model: Qwen3-235B-A22B-class. A raw FP8 weight floor is approximately 235 GB by one byte per parameter; real packages include scales, mixed-precision tensors, alignment, metadata, runtime and replicas. Qwen model configuration used for KV arithmetic: 94 layers, 4 KV heads, head dimension 128. BF16 logical KV bytes = 2 (K,V) × 94 × 4 × 128 × 2 bytes × cached tokens. At 4,096 tokens: 788,529,152 bytes = 0.734375 GiB/request. At 8,192 tokens: 1.46875 GiB/request. Thirty-two 8K sequences: 47 GiB logical KV. Excludes block rounding, replication, staging and workspace. Source: https://huggingface.co/Qwen/Qwen3-235B-A22B/blob/main/config.json

Internal ATMOS planning envelope retained from supplied materials: up to 128 GB HBF and 8–16 GB LPDDR per E3.S, 100 GB application weight budget per device, 200 GB/s local-read target. For eight devices this is 1,024 GB nominal HBF, 800 GB weight study budget and 128 GB LPDDR at 16 GB/device; a provisional 4 GB/device reserve leaves 96 GB. These are unqualified planning targets. Per-device sharding and fabric matter; aggregate memory is not unified.

GPU controls: 4 RTX PRO 6000 GPUs provide 384 GB aggregate GDDR7 and an accessible Blackwell development control. An H200 or Blackwell Ultra control must be sized by the actual runtime; raw memory fit does not establish accepted service. No cost savings are assumed.
""")


def add_portfolio_map(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 4, "Portfolio", "Three offerings, three execution boundaries",
               "Option C is now explicit: speculative decoding, not another prefill/decode split.")
    cards = [
        (0.58, "A", "NATIVE ATMOS SERVICE", BLUE, "REQUEST / RESPONSE", "ATMOS", "Full target model: prefill + decode", "Prompt in; response out", "LEAD"),
        (4.48, "B", "GPU PREFILL → ATMOS DECODE", TEAL, "TARGET KV + STATE", "GPU + ATMOS", "GPU runs full prefill; ATMOS runs full target decode", "One handoff per request", "CONDITIONAL"),
        (8.38, "C", "GPU DRAFT → ATMOS VERIFY", ORANGE, "CANDIDATE TOKEN BLOCK", "SMALL GPU + ATMOS", "GPU draft model proposes; ATMOS target accepts/corrects", "Repeated verification cycles", "RESEARCH"),
    ]
    for x, letter, title, accent, boundary, engines, execution, crossing, maturity in cards:
        add_rect(slide, x, 2.03, 3.55, 4.67, fill=WHITE, line=LINE, radius=True)
        add_text(slide, letter, x + 0.18, 2.18, 0.42, 0.46, size=27, color=accent,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, title, x + 0.70, 2.20, 2.60, 0.58, size=13.3, color=NAVY,
                 bold=True)
        add_badge(slide, maturity, x + 2.25, 2.86, 1.07, fill=accent, size=8.8)
        add_text(slide, "BOUNDARY", x + 0.20, 3.36, 1.10, 0.22, size=9.5,
                 color=MUTED, bold=True)
        add_text(slide, boundary, x + 0.20, 3.62, 3.05, 0.42, size=15,
                 color=accent, bold=True, font=FONT_DISPLAY)
        add_text(slide, "EXECUTION OWNERS", x + 0.20, 4.24, 1.55, 0.22, size=9.5,
                 color=MUTED, bold=True)
        add_text(slide, engines, x + 0.20, 4.50, 3.05, 0.32, size=13, color=INK, bold=True)
        add_text(slide, execution, x + 0.20, 5.02, 3.10, 0.69, size=11.8, color=INK)
        add_line(slide, x + 0.20, 5.88, x + 3.35, 5.88, color=LINE, width=0.8)
        add_text(slide, "WHAT CROSSES", x + 0.20, 6.01, 1.18, 0.20, size=9.3,
                 color=MUTED, bold=True)
        add_text(slide, crossing, x + 1.35, 5.99, 1.95, 0.33, size=10.6,
                 color=INK, bold=True)
    add_footer(slide, 4, "same target quality • different integration work")
    add_notes(slide, """
Offering A: one request is routed either to the existing GPU service or to a complete ATMOS-native target-model service. Only request/response crosses. ATMOS must implement full model loading, attention, KV, router, experts, combine and sampling. This is not GPU expert offload.

Offering B: the GPU executes the complete target model during prefill and creates target-model KV. The service transfers exact KV blocks plus model identity, positions, dtype/scales, head/shard layout, fences and continuation state to an ATMOS-native decoder holding the same target weights. ATMOS then executes the complete decode loop without returning per layer. Both pools need target weights. vLLM describes disaggregated prefill as useful for separately tuning TTFT and inter-token latency and controlling tail ITL, but explicitly states it does not inherently improve throughput. Public connectors are precedents, not an ATMOS KV-compatibility guarantee.

Offering C: a separate small draft model on GPU proposes a token block using its own KV. ATMOS owns the large target model and target KV and scores the candidates using a distribution-preserving speculative decoding algorithm. Candidate token IDs plus required probability/protocol metadata cross every verification cycle; committed prefix/correction returns to the drafter. Draft KV is never imported as target KV. C must automatically use normal target decoding when expected accepted output does not repay draft + verification overhead.

Speculative-decoding algorithm precedent: Leviathan, Kalman & Matias, ICML 2023, https://proceedings.mlr.press/v202/leviathan23a.html
Disaggregated prefill reference: https://docs.vllm.ai/en/latest/features/disagg_prefill/
NVIDIA Dynamo reference: https://docs.nvidia.com/dynamo/components/router/disaggregated-serving
""")


def add_flow_box(slide, title: str, detail: str, x: float, y: float, w: float,
                 h: float, *, accent: str, fill: str = WHITE) -> None:
    add_rect(slide, x, y, w, h, fill=fill, line=accent, radius=True, line_width=1.5)
    add_rect(slide, x, y, 0.10, h, fill=accent, line=accent, radius=False)
    add_text(slide, title, x + 0.24, y + 0.16, w - 0.38, 0.30, size=14,
             color=NAVY, bold=True)
    add_text(slide, detail, x + 0.24, y + 0.51, w - 0.38, h - 0.63, size=11.3,
             color=INK)


def add_offering_a(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 5, "Offering A", "Native ATMOS service: isolate a complete request",
               "The lowest-coupling offer—ATMOS is a full model endpoint, not a GPU weight cache.")

    add_flow_box(slide, "Admission / policy", "Choose one complete service by model, context, deadline, tenant and current queue.",
                 0.60, 2.10, 2.35, 1.25, accent=RED, fill=PALE_RED)
    add_line(slide, 2.95, 2.72, 3.53, 2.72, color=RED, width=2.2, arrow=True)
    add_flow_box(slide, "GPU service", "Existing CUDA model + GPU KV\nInteractive or ineligible work",
                 3.57, 2.10, 2.55, 1.25, accent=BLUE, fill=PALE_BLUE)
    add_text(slide, "OR", 6.30, 2.50, 0.42, 0.30, size=12, color=MUTED,
             bold=True, align=PP_ALIGN.CENTER)
    add_flow_box(slide, "ATMOS-native service group", "Full target model: attention → router → experts → combine → sampling",
                 6.90, 2.10, 3.55, 1.25, accent=TEAL, fill=PALE_TEAL)
    add_line(slide, 10.45, 2.72, 11.04, 2.72, color=TEAL, width=2.2, arrow=True)
    add_flow_box(slide, "Response", "One target-approved output\nNo GPU state handoff",
                 11.08, 2.10, 1.62, 1.25, accent=GREEN, fill="EAF4EC")

    add_rect(slide, 0.60, 3.72, 7.42, 2.67, fill=WHITE, line=LINE, radius=True)
    add_text(slide, "ATMOS CLUSTER RESPONSIBILITIES", 0.84, 3.92, 3.20, 0.27,
             size=11, color=TEAL, bold=True)
    responsibilities = [
        ("HBF", "Full target weights; versioned native pack"),
        ("LPDDR / SRAM", "Active target KV, scratch and runtime state"),
        ("NPU + fabric", "Attention, MoE, collectives, sampling and inter-device execution"),
        ("Service plane", "Admission, readiness, isolation, update, rollback and recovery"),
    ]
    for idx, (label, detail) in enumerate(responsibilities):
        yy = 4.36 + idx * 0.47
        add_text(slide, label, 0.86, yy, 1.15, 0.25, size=11.2, color=TEAL, bold=True)
        add_text(slide, detail, 2.07, yy, 5.50, 0.29, size=11.2, color=INK)

    add_rect(slide, 8.30, 3.72, 4.40, 2.67, fill=NAVY, line=NAVY, radius=True)
    add_text(slide, "TECHNICAL VALUE", 8.57, 3.94, 2.10, 0.28, size=11,
             color="BFD7EA", bold=True)
    add_bullets(slide, [
        "No KV, hidden-state or per-layer return to NVIDIA",
        "Existing Hopper / Blackwell service can remain unchanged",
        "Independent queue, failure and scaling boundary",
    ], 8.55, 4.38, 3.80, 1.36, size=11.4, bullet_color=TEAL, gap=0.45)
    add_text(slide, "GATE", 8.57, 5.93, 0.55, 0.22, size=9.5, color="BFD7EA", bold=True)
    add_text(slide, "Full-model quality + jobs/hour + p95 completion + recovery",
             9.18, 5.88, 3.18, 0.38, size=10.8, color=WHITE, bold=True)
    add_footer(slide, 5, "A is the recommended lead offer")
    add_notes(slide, """
Offering A routes each eligible request to exactly one complete service. The NVIDIA service and its KV cache remain independent. ATMOS must load and execute the complete target model, including prompt prefill, attention, active KV, router, routed/shared experts, combination and output sampling. If a model spans several E3.S modules, the ATMOS-native fabric and failure policy are part of the product.

What crosses the integration boundary: prompt/token input, model/version, generation limits, tenant identity, deadline/cancellation and the response. No target KV or per-layer activation crosses from NVIDIA.

Why it can help in a Hopper/Blackwell environment: it creates an independently managed service pool for bounded, lower-urgency or private jobs without modifying the protected NVIDIA inference engine. The value is not that ATMOS is faster than H200/B300 memory. It is that eligible work can be isolated and scaled under a different capacity/power/support envelope if the complete service meets quality and deadline.

Primary technical risks: full-model compiler/runtime breadth, attention and active-KV fit, native inter-device fabric, small-batch performance, queue stability, complete-service availability, platform operations and support. Capacity fit alone does not pass.
""")


def add_offering_b(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 6, "Offering B", "GPU prefill → ATMOS decode: one target-KV handoff",
               "Use NVIDIA for prompt parallelism; move the serial token-generation phase only after exact cache compatibility is proven.")

    add_flow_box(slide, "NVIDIA prefill pool", "Complete target model weights\nAttention + experts on GPU\nCreates target-model KV",
                 0.60, 2.18, 3.20, 1.65, accent=BLUE, fill=PALE_BLUE)
    add_line(slide, 3.80, 3.00, 5.13, 3.00, color=TEAL, width=2.5, arrow=True)
    add_rect(slide, 4.08, 2.22, 0.78, 0.30, fill=TEAL, line=TEAL, radius=True)
    add_text(slide, "ONCE", 4.11, 2.25, 0.72, 0.23, size=9.2, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Target KV blocks + model/hash + positions +\nhead/shard layout + dtype/scales + fences + continuation",
             3.92, 3.24, 1.18, 1.02, size=9.3, color=MUTED,
             align=PP_ALIGN.CENTER)
    add_flow_box(slide, "ATMOS decode group", "Same target weights in HBF\nImported target KV\nAttention + experts + next-token loop",
                 5.18, 2.18, 3.35, 1.65, accent=TEAL, fill=PALE_TEAL)
    add_line(slide, 8.53, 3.00, 9.20, 3.00, color=GREEN, width=2.5, arrow=True)
    add_flow_box(slide, "Streaming response", "ATMOS owns decode until completion or declared fallback",
                 9.25, 2.18, 3.44, 1.65, accent=GREEN, fill="EAF4EC")

    add_metric(slide, "0.734 GiB", "logical BF16 KV / 4K prompt", 0.60, 4.23, 2.30,
               accent=BLUE, value_size=22)
    add_metric(slide, "1.47 GiB/s", "logical handoff at 2 RPS", 3.12, 4.23, 2.30,
               accent=TEAL, value_size=22)
    add_metric(slide, "2× weights", "target weights live in both pools", 5.64, 4.23, 2.30,
               accent=ORANGE, value_size=22)

    add_rect(slide, 8.35, 4.23, 4.34, 2.05, fill=NAVY, line=NAVY, radius=True)
    add_text(slide, "PASS ONLY IF", 8.61, 4.45, 1.45, 0.25, size=10.5,
             color="BFD7EA", bold=True)
    add_bullets(slide, [
        "KV conversion/import is numerically exact",
        "Combined TTFT + TPOT improves the service",
        "Paired admission prevents decode backlog",
        "Reset/stale ownership is safe",
    ], 8.58, 4.82, 3.72, 1.28, size=10.8, bullet_color=TEAL, gap=0.32)
    add_rect(slide, 0.60, 5.47, 7.34, 0.81, fill=PALE_ORANGE, line="E7C7A6", radius=True)
    add_rich_text(slide, [
        ("Important  ", True, ORANGE),
        ("Disaggregated prefill separates TTFT and ITL tuning; it does not inherently improve throughput. B earns value only from the measured combined service.", False, INK),
    ], 0.82, 5.66, 6.90, 0.41, size=11.2)
    add_footer(slide, 6, "B duplicates target weights and exposes a real state-transfer contract")
    add_notes(slide, """
Offering B uses NVIDIA where its parallel prompt processing is strongest and ATMOS for the complete target decode loop. Both phases execute the same target model and therefore both hold target weights. The primary benefit is separate scaling and queue isolation of prefill versus decode—not GPU weight reclamation.

The handoff is a model-state transfer, not a pointer. It includes exact model/checkpoint identity, positions, token IDs, layer/head/shard ownership, KV dtype/scales/layout, transfer fences, completion/ownership, first-token handling, cancellation and reset semantics. A connector may transfer blocks layer-by-layer; 'one handoff' means one service-phase boundary per request, not one DMA operation.

Qwen3-235B example: BF16 logical KV at a 4,096-token prompt is 0.734375 GiB. At 2 requests/s the logical handoff is 1.46875 GiB/s before block rounding, duplication, staging, protocol and burst reserve. These are arithmetic inputs, not measured ATMOS traffic.

vLLM documentation states disaggregated prefilling is useful for independently tuning TTFT/ITL and controlling tail ITL, and explicitly says it does not improve throughput. Public NIXL/LMCache/Dynamo connectors do not prove GPU-to-ATMOS cache compatibility. Sources:
- https://docs.vllm.ai/en/latest/features/disagg_prefill/
- https://docs.nvidia.com/dynamo/components/router/disaggregated-serving

Adoption gate: same model/quality, combined TTFT, handoff, TPOT, accepted rate, duplicated model capacity, failure reserve and operational cost must beat the optimized control. If ATMOS decode backs up, GPU prefill cannot create value by itself.
""")


def add_offering_c(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 7, "Offering C", "Speculative node: GPU drafts; ATMOS target verifies",
               "A distinct algorithmic offer with separate model state and a measurable acceptance break-even.")

    add_flow_box(slide, "Small draft model / GPU", "Fast candidate generation\nDraft weights + draft KV\nProposes k tokens",
                 0.60, 2.13, 3.15, 1.52, accent=BLUE, fill=PALE_BLUE)
    add_line(slide, 3.75, 2.89, 5.03, 2.89, color=ORANGE, width=2.4, arrow=True)
    add_text(slide, "candidate IDs + required\nprobability/protocol metadata",
             3.78, 3.11, 1.20, 0.62, size=9.4, color=MUTED, align=PP_ALIGN.CENTER)
    add_flow_box(slide, "Large target MoE / ATMOS", "Target weights + target KV\nScores block in parallel\nAccepts prefix or corrects",
                 5.08, 2.13, 3.40, 1.52, accent=ORANGE, fill=PALE_ORANGE)
    add_line(slide, 8.48, 2.89, 9.62, 2.89, color=GREEN, width=2.4, arrow=True)
    add_flow_box(slide, "Commit + resynchronize", "Publish target-approved tokens\nRollback rejected target state\nDrafter follows committed prefix",
                 9.67, 2.13, 3.02, 1.52, accent=GREEN, fill="EAF4EC")

    add_rect(slide, 0.60, 4.03, 7.32, 2.32, fill=NAVY, line=NAVY, radius=True)
    add_text(slide, "BREAK-EVEN", 0.86, 4.25, 1.25, 0.27, size=11, color="BFD7EA", bold=True)
    add_text(slide, "effective ms / committed token  =  cycle time ÷ E[committed tokens]",
             0.86, 4.69, 6.65, 0.43, size=18, color=WHITE, bold=True, font=FONT_MONO)
    add_text(slide, "Win when  E[committed]  >  speculative cycle time ÷ native target time/token",
             0.86, 5.24, 6.62, 0.42, size=13.5, color="C7D3DF", bold=True)
    add_rect(slide, 0.86, 5.79, 6.58, 0.35, fill="2A3A50", line="2A3A50", radius=True)
    add_rich_text(slide, [
        ("Illustration: ", True, ORANGE),
        ("native = 80 ms/token; speculative cycle = 120 ms → threshold >1.5 commits/cycle. 3 commits = 40 ms/token; 1 commit = 120 ms/token.", False, WHITE),
    ], 1.02, 5.82, 6.24, 0.28, size=10.2)

    add_rect(slide, 8.22, 4.03, 4.47, 2.32, fill=WHITE, line=LINE, radius=True)
    add_text(slide, "STATE + CORRECTNESS", 8.48, 4.25, 2.25, 0.27, size=11,
             color=ORANGE, bold=True)
    add_bullets(slide, [
        "Draft KV and target KV remain separate",
        "Tokenizer and sampling protocol must align",
        "Target acceptance/correction preserves distribution",
        "Rejected state is reclaimed safely",
        "Auto-disable below rolling break-even",
    ], 8.44, 4.66, 3.78, 1.46, size=10.8, bullet_color=ORANGE, gap=0.29)
    add_footer(slide, 7, "C is a research option—not a default product promise")
    add_notes(slide, """
This slide clarifies Option C. It is speculative decoding: a small, fast GPU draft model proposes a block; ATMOS executes the large target model to score the candidates and applies a distribution-preserving acceptance/correction algorithm. It is not GPU prefill of the target, not majority voting and not a quality classifier.

State: the draft model and target model have separate weights and separate KV caches. Both initialize prompt state unless a separately qualified optimization exists. The draft KV is never imported into the target. The wire contract may require draft probabilities/logits or algorithm-specific metadata in addition to token IDs. The target commits accepted output and corrects according to the chosen exact algorithm. Both sides roll forward only from the committed prefix.

Break-even derivation: let t_native be ordinary target time per committed token. Let t_cycle(k) include drafting k candidates, target verification and protocol/communication. Let C be committed target-approved tokens per cycle. Effective time/token = t_cycle / E[C]. Speculation wins only if E[C] > t_cycle / t_native. In the illustration, 120/80 = 1.5, so the rolling expected committed count must exceed 1.5. Three commits gives 40 ms/token (2x); one gives 120 ms/token (slower). These numbers are illustrative, not an ATMOS benchmark.

Measure acceptance by workload and context, accepted tokens/s, target cycles, draft/verification time, communication, rejected work, target KV rollback, quality/distribution and tail latency. Automatically fall back to normal ATMOS target decoding when the rolling confidence-adjusted benefit is negative.

Algorithm reference: Leviathan, Kalman & Matias, Fast Inference from Transformers via Speculative Decoding, ICML 2023. https://proceedings.mlr.press/v202/leviathan23a.html
""")


def add_common_platform(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 8, "Common platform", "One ATMOS cluster; three traffic and state contracts",
               "Do not assume E3.S capacity, a PCIe switch, or NVIDIA fabric membership automatically creates a serving system.")

    add_rect(slide, 0.60, 2.03, 5.10, 4.68, fill=NAVY, line=NAVY, radius=True)
    add_text(slide, "ATMOS EXECUTION CLUSTER — 8-DEVICE STUDY", 0.88, 2.26,
             4.55, 0.30, size=12, color="BFD7EA", bold=True)
    for row in range(2):
        for col in range(4):
            idx = row * 4 + col
            x = 0.88 + col * 1.08
            y = 2.88 + row * 0.90
            add_rect(slide, x, y, 0.88, 0.67, fill="183B50", line=TEAL, radius=True)
            add_text(slide, f"A{idx}", x, y + 0.10, 0.88, 0.24, size=13,
                     color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            add_text(slide, "HBF + NPU", x, y + 0.38, 0.88, 0.17, size=7.8,
                     color="BFD7EA", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HBF", 0.90, 4.81, 0.62, 0.24, size=10.5, color=TEAL, bold=True)
    add_text(slide, "1,024 GB nominal / 800 GB application study budget", 1.52, 4.81,
             3.77, 0.25, size=10.4, color=WHITE)
    add_text(slide, "WORK", 0.90, 5.21, 0.62, 0.24, size=10.5, color=TEAL, bold=True)
    add_text(slide, "128 GB LPDDR nominal / 96 GB after provisional reserve", 1.52, 5.21,
             3.77, 0.25, size=10.4, color=WHITE)
    add_text(slide, "FABRIC", 0.90, 5.61, 0.62, 0.24, size=10.5, color=TEAL, bold=True)
    add_text(slide, "Native NPU collectives + failure domains remain qualification gates", 1.52, 5.61,
             3.77, 0.38, size=10.4, color=WHITE)
    add_text(slide, "These are internal targets—not measured product specifications.",
             0.90, 6.27, 4.20, 0.23, size=9.2, color="9FB3C8", bold=True)

    add_rect(slide, 6.02, 2.03, 6.68, 4.68, fill=WHITE, line=LINE, radius=True)
    add_text(slide, "NVIDIA INTEGRATION BY OFFERING", 6.28, 2.26, 3.60, 0.28,
             size=12, color=RED, bold=True)
    rows = [
        ("A", BLUE, "Network / service API", "Lowest coupling", "Request + response", "RTX PRO dev → H200 / GB300 control"),
        ("B", TEAL, "KV connector / transport", "High state fidelity", "Target KV + continuation", "TTFT + handoff + TPOT must win"),
        ("C", ORANGE, "Verification protocol", "Low bytes, frequent sync", "Candidate block + metadata", "Acceptance and verification latency decide"),
    ]
    for idx, (letter, accent, interface, property_, payload, target) in enumerate(rows):
        yy = 2.82 + idx * 1.10
        add_text(slide, letter, 6.30, yy, 0.40, 0.42, size=24, color=accent,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, interface, 6.83, yy + 0.02, 2.25, 0.26, size=12.2,
                 color=NAVY, bold=True)
        add_text(slide, property_, 9.20, yy + 0.02, 1.32, 0.24, size=10.0,
                 color=MUTED, bold=True)
        add_text(slide, payload, 6.83, yy + 0.39, 2.60, 0.25, size=10.7, color=INK)
        add_text(slide, target, 9.45, yy + 0.39, 2.85, 0.32, size=10.3,
                 color=INK, bold=True)
        if idx < 2:
            add_line(slide, 6.28, yy + 0.90, 12.35, yy + 0.90, color=LINE, width=0.7)
    add_rect(slide, 6.28, 6.15, 6.05, 0.36, fill=PALE_RED, line=PALE_RED, radius=True)
    add_rich_text(slide, [
        ("Invariant  ", True, RED),
        ("ATMOS does not become an NVLink / NVSwitch endpoint. Preserve stock CUDA, NCCL and NVIDIA firmware/driver boundaries.", False, INK),
    ], 6.44, 6.19, 5.72, 0.26, size=10.2)
    add_footer(slide, 8, "physical topology and support are configuration-specific")
    add_notes(slide, """
The ATMOS cluster is a common study platform, not a validated BOM. Internal planning inputs: eight E3.S devices, up to 128 GB HBF/device and 8–16 GB LPDDR/device, 100 GB application HBF budget/device and 200 GB/s local-read target/device. The displayed 1,024/800/128/96 GB arithmetic uses 16 GB LPDDR/device and a provisional 4 GB/device working-memory reserve. Per-device placement, scratch, updates, replicas and failure reserve reduce usable capacity.

A: integrate through a normal service API over the qualified network. It is the least invasive path for RTX PRO, H200/HGX or GB300 environments because the existing NVIDIA service can remain unchanged.

B: requires a target-KV connector and transport. Depending on physical productization this may use a supported network/RDMA/NIXL-like path or an OEM-approved local path. Exact layout conversion and ownership are more important than headline bandwidth. The target weights exist in both pools.

C: transfers less state per cycle than B but synchronizes repeatedly, making latency and protocol efficiency critical. The small draft GPU can be an RTX PRO-class or another supported GPU selected by fit; it does not run target layers.

NVIDIA target sequence: RTX PRO 6000 Blackwell Server Edition for development/integration; H200/HGX as an installed-base high-bandwidth control; GB300 NVL72 for strategic rack qualification. Every platform is remeasured. ATMOS remains outside NVLink/NVSwitch and stock NVIDIA firmware/driver/NCCL are preserved.
""")


def add_selection_matrix(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 9, "Technical selection", "Choose the offering by the bottleneck—not novelty",
               "Cost is one output metric; the primary decision is whether the execution boundary improves the service.")

    columns = [0.58, 2.80, 5.25, 7.70, 10.05]
    widths = [2.08, 2.31, 2.31, 2.21, 2.66]
    headers = ["OFFER", "WHEN IT FITS", "STATE CROSSING", "ATMOS MUST PROVE", "GO / NO-GO KPI"]
    for x, w, head in zip(columns, widths, headers):
        add_rect(slide, x, 2.04, w, 0.52, fill=NAVY, line=NAVY, radius=False)
        add_text(slide, head, x + 0.10, 2.18, w - 0.20, 0.22, size=9.5,
                 color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    data = [
        ("A\nNATIVE", BLUE, "Bounded jobs can use a separate complete service", "Request / response", "Full model, KV, fabric, availability", "Accepted jobs/hour at quality + deadline"),
        ("B\nPREFILL / DECODE", TEAL, "GPU prefill is valuable; decode is the remaining bottleneck", "Target KV + continuation once/request", "KV compatibility + full decoder + paired queues", "Combined TTFT + TPOT + accepted rate"),
        ("C\nSPECULATIVE", ORANGE, "Target decode is serial and draft acceptance stays high", "Candidates + protocol metadata every cycle", "Exact target verify + rollback + adaptive fallback", "Accepted tokens/s after all overhead"),
    ]
    for row_idx, (offer, accent, fit, state, prove, kpi) in enumerate(data):
        yy = 2.56 + row_idx * 1.20
        values = [offer, fit, state, prove, kpi]
        for col_idx, (x, w, value) in enumerate(zip(columns, widths, values)):
            fill = "F8FAFC" if row_idx % 2 == 0 else "EDF2F6"
            add_rect(slide, x, yy, w, 1.20, fill=fill, line=LINE, radius=False)
            if col_idx == 0:
                add_text(slide, value, x + 0.08, yy + 0.22, w - 0.16, 0.62,
                         size=13, color=accent, bold=True, font=FONT_DISPLAY,
                         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
            else:
                add_text(slide, value, x + 0.12, yy + 0.18, w - 0.24, 0.80,
                         size=10.7, color=INK, align=PP_ALIGN.CENTER,
                         valign=MSO_ANCHOR.MIDDLE)

    add_rect(slide, 0.58, 6.38, 12.13, 0.45, fill=PALE_ORANGE, line="E7C7A6", radius=True)
    add_rich_text(slide, [
        ("Bounded commercial test  ", True, ORANGE),
        ("Compare total delivered cost ÷ accepted work at matched quality, SLO and reserve. B includes duplicated target weights + KV transfer; C includes draft compute + rejected work. No fleet-savings forecast before measurement.", False, INK),
    ], 0.78, 6.46, 11.72, 0.29, size=10.5)
    add_footer(slide, 9, "A first • B/C only after their specific technical gate")
    add_notes(slide, """
The table is the selection rule. A is preferred when the request can be assigned entirely to an independent ATMOS service. B is appropriate only when ATMOS decode is already competitive and prompt processing or prefill/decode interference is the measured bottleneck. C is appropriate only when target decoding is the bottleneck and the selected draft model has a high, stable, workload-specific acceptance rate.

Control requirements: each offering must be compared with an optimized NVIDIA control at the same checkpoint/target quality, prompt/output distribution, TTFT/TPOT or completion deadline, availability and reserve. H200 or GB300 can make memory capacity less scarce; ATMOS still needs to show useful accepted work or a deployment boundary worth operating.

Cost attribute: use total delivered cost per accepted request/output token. Include hardware, network, rack, power/cooling, licenses/support, spares/failure reserve, engineering/qualification and operations. Offering B includes full target weights in both pools, GPU prefill and KV transfer. Offering C includes the draft GPU, rejected draft work, target verification and protocol/rollback. Do not equate low GPU utilization or nominal ATMOS HBF with savings.
""")


def add_validation_plan(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, 10, "Qualification", "Four gates turn concepts into a supported offering",
               "Simulator and emulator can overlap; only physical and named-platform evidence can earn product claims.")
    stages = [
        (0.62, "1", "SIMULATE", BLUE, "Full native model", "Weights + active KV\nNPU/fabric queues\nHopper/Blackwell controls", "Architecture fit"),
        (3.75, "2", "EMULATE", TEAL, "Real interfaces", "Pack/loader + API\nB: KV connector\nC: verify/rollback", "Correct contracts"),
        (6.88, "3", "PHYSICAL NODE", ORANGE, "Prove A first", "Quality + latency\nPower/thermal/faults\nThen B and C", "Silicon gate"),
        (10.01, "4", "NAMED PLATFORM", PURPLE, "NVIDIA control", "RTX/H200/GB300\nSupport + lifecycle\nAccepted-work result", "Product gate"),
    ]
    for idx, (x, number, name, accent, title, detail, gate) in enumerate(stages):
        add_rect(slide, x, 2.20, 2.55, 3.42, fill=WHITE, line=accent, radius=True, line_width=1.5)
        add_text(slide, number, x + 0.18, 2.34, 0.46, 0.44, size=25, color=accent,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, name, x + 0.70, 2.39, 1.60, 0.26, size=11, color=NAVY, bold=True)
        add_text(slide, title, x + 0.20, 3.06, 2.15, 0.34, size=15, color=accent,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.25, 3.60, 2.05, 0.88, size=11.5,
                 color=INK, align=PP_ALIGN.CENTER)
        add_rect(slide, x + 0.25, 4.79, 2.05, 0.48, fill=accent, line=accent, radius=True)
        add_text(slide, gate, x + 0.31, 4.91, 1.93, 0.22, size=10, color=WHITE,
                 bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_line(slide, x + 2.55, 3.90, stages[idx + 1][0] - 0.06, 3.90,
                     color=MUTED, width=1.8, arrow=True)
    add_rect(slide, 0.62, 5.97, 11.94, 0.76, fill=NAVY, line=NAVY, radius=True)
    add_rich_text(slide, [
        ("DECISION SEQUENCE  ", True, RED),
        ("A native service must pass before B or C. B then proves target-KV ownership and combined TTFT/TPOT. C separately proves target-correct speculative acceptance, rollback and break-even fallback.", False, WHITE),
    ], 0.88, 6.15, 11.42, 0.39, size=12.2)
    add_footer(slide, 10, "configuration-specific evidence • no transfer across GPU generations")
    add_notes(slide, """
Gate 1 — simulation: represent the complete native target model, not only expert capacity. Model full weights, attention, active KV, NPU service, native ATMOS fabric, queues, device failure and control GPU service. Use RTX PRO/H200/GB300 specs only as control inputs, not ATMOS predictions.

Gate 2 — emulator: operate the real software contracts and produce correct target results on surrogate/partner compute. A needs full model pack/loader/runtime and request API. B adds exact KV export/import, ownership and continuation. C adds draft/target scheduler, exact acceptance/correction and KV commit/rollback. Emulator timing remains provisional.

Gate 3 — physical node: prove offering A first on ATMOS silicon/prototype with quality, completion/TTFT/TPOT, NPU/HBF/LPDDR/fabric, all-device load, power/thermal, faults, updates and recovery. Only then execute B handoff and C verification experiments. A successful demo without failure/lifecycle evidence is not a gate pass.

Gate 4 — named NVIDIA platform: compare each selected offering against its own optimized control. Development target: RTX PRO 6000 Blackwell Server Edition. Installed-base production control: H200/HGX or H200 NVL. Strategic rack control: one exact GB300 NVL72 OEM configuration. Include support, lifecycle, supply, security, availability reserve and accepted-work result. Requalify every OEM and generation.
""")


def add_recommendation(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, DARK)
    add_text(slide, "SANDISK / ATMOS", 0.62, 0.36, 1.55, 0.28, size=10,
             color=RED, bold=True)
    add_text(slide, "RECOMMENDATION", 2.35, 0.36, 2.5, 0.28, size=10,
             color="9FB3C8", bold=True)
    add_text(slide, "Lead with a complete native service.\nEarn the other two boundaries.",
             0.62, 0.92, 11.7, 1.12, size=32, color=WHITE, bold=True, font=FONT_DISPLAY)

    decisions = [
        ("APPROVE", GREEN, "Offering A native backend", "One model, full attention/KV/MoE runtime, RTX PRO development control and OEM cluster study."),
        ("FEASIBILITY", TEAL, "Offering B KV bridge", "Freeze exact Qwen target cache schema; prototype export/import; measure combined TTFT + TPOT."),
        ("RESEARCH", ORANGE, "Offering C speculative loop", "Select draft pair; simulate acceptance; implement target-correct verify/rollback; retain auto fallback."),
    ]
    for idx, (status, accent, title, detail) in enumerate(decisions):
        x = 0.65 + idx * 4.13
        add_rect(slide, x, 2.52, 3.78, 1.58, fill="172235", line="334155", radius=True)
        add_badge(slide, status, x + 0.18, 2.70, 1.10, fill=accent, size=8.8)
        add_text(slide, title, x + 0.18, 3.14, 3.30, 0.30, size=13.8,
                 color=WHITE, bold=True)
        add_text(slide, detail, x + 0.18, 3.51, 3.34, 0.45, size=10.2, color="C7D3DF")

    add_text(slide, "FIRST 90 DAYS", 0.65, 4.62, 1.80, 0.27, size=11,
             color=RED, bold=True)
    timeline = [
        ("0–30", "Freeze Qwen profile, native model contract and NVIDIA controls"),
        ("31–60", "Compile/load complete native pack; emulate A and KV/verify interfaces"),
        ("61–90", "Run physical A primitive; return B/C go/no-go and OEM path"),
    ]
    for idx, (window, action) in enumerate(timeline):
        x = 0.65 + idx * 4.13
        add_text(slide, window, x, 5.07, 0.72, 0.32, size=18, color=RED,
                 bold=True, font=FONT_DISPLAY)
        add_text(slide, action, x + 0.82, 5.05, 2.96, 0.58, size=11.2,
                 color=WHITE, bold=True)
        if idx < 2:
            add_line(slide, x + 3.84, 5.30, x + 4.05, 5.30, color="64748B", width=1.5, arrow=True)

    add_rect(slide, 0.65, 6.12, 12.00, 0.55, fill="2A1D25", line="5D2734", radius=True)
    add_rich_text(slide, [
        ("HOLD  ", True, RED),
        ("Performance, product date, GPU reduction and TCO claims until a named configuration passes quality, service, failure, lifecycle and support gates.", False, WHITE),
    ], 0.88, 6.26, 11.50, 0.28, size=11.5)
    add_text(slide, "INTERNAL • 02 OCT 2026 • SOURCES AND CALCULATIONS IN SPEAKER NOTES",
             0.65, 7.08, 10.5, 0.18, size=7.5, color="8294A8", bold=True)
    add_text(slide, "11", 12.10, 7.08, 0.42, 0.18, size=7.5, color="8294A8",
             bold=True, align=PP_ALIGN.RIGHT)
    add_notes(slide, """
Decision request:
1. Approve a bounded native ATMOS full-model backend and Offering A study for one Qwen3-235B-A22B-class profile.
2. Use RTX PRO 6000 Blackwell Server Edition as the most accessible development control; retain H200 as installed-base production control and GB300 NVL72 as strategic rack control.
3. Fund B only to close target-KV schema/connector/ownership and combined-service feasibility. Public disaggregated serving does not deliver an ATMOS decoder or throughput guarantee.
4. Keep C as research until a selected small draft/large target pair shows a confidence-adjusted committed-token rate above the measured break-even, with exact distribution-preserving acceptance/correction and safe target/draft KV state.
5. Request one OEM physical/fabric study and explicit SanDisk/Tenstorrent/runtime ownership. Do not imply partner commitment from public software.

First 90 days:
- Days 0–30: freeze checkpoint, precision, workload, quality/deadline, GPU controls, full-model memory/active-state/fabric ledger and interface ownership.
- Days 31–60: build/verify the complete native pack/runtime on emulator; implement A API; define and unit-test B KV schema and C verification protocol; calibrate simulator.
- Days 61–90: run physical primitive/native service tests as evidence permits; return A go/pivot/stop, B KV feasibility and C acceptance simulation, plus one named OEM path and risk/cost boundary.

Cost is subordinate to technical proof. Later compare total delivered cost per accepted work at equal quality/SLO/availability. No measured ATMOS rates, power, price, OEM support or launch date are asserted.

Principal public references:
- NVIDIA RTX PRO 6000: https://www.nvidia.com/en-us/data-center/rtx-pro-6000-blackwell-server-edition/
- NVIDIA RTX PRO Server: https://www.nvidia.com/en-us/data-center/products/rtx-pro-server/
- NVIDIA H200: https://www.nvidia.com/en-us/data-center/h200/
- NVIDIA GB300 NVL72: https://www.nvidia.com/en-us/data-center/gb300-nvl72/
- NVIDIA Dynamo disaggregated serving: https://docs.nvidia.com/dynamo/components/router/disaggregated-serving
- vLLM disaggregated prefill: https://docs.vllm.ai/en/latest/features/disagg_prefill/
- Speculative decoding: https://proceedings.mlr.press/v202/leviathan23a.html
- Qwen3-235B configuration: https://huggingface.co/Qwen/Qwen3-235B-A22B/blob/main/config.json
""")


def validate(prs: Presentation, expected_slides: int) -> None:
    if len(prs.slides) != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} slides, found {len(prs.slides)}")
    for index, slide in enumerate(prs.slides, 1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if not texts:
            raise RuntimeError(f"Slide {index} has no visible text")
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            raise RuntimeError(f"Slide {index} has no speaker notes")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size and run.font.size.pt < 7.0:
                        raise RuntimeError(f"Slide {index} contains text below 7 pt in {shape.name}")
    all_text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text.strip()
    )
    required = [
        "NATIVE SERVICE", "GPU prefill → ATMOS decode", "GPU DRAFT → ATMOS VERIFY",
        "RTX PRO 6000", "H200", "GB300 NVL72", "BREAK-EVEN",
        "Disaggregated prefill", "Four gates", "Lead with a complete native service",
    ]
    for item in required:
        if item not in all_text:
            raise RuntimeError(f"Missing required deck content: {item}")


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    add_cover(prs)
    add_nvidia_baseline(prs)
    add_hypothetical_workload(prs)
    add_portfolio_map(prs)
    add_offering_a(prs)
    add_offering_b(prs)
    add_offering_c(prs)
    add_common_platform(prs)
    add_selection_matrix(prs)
    add_validation_plan(prs)
    add_recommendation(prs)
    validate(prs, 11)
    prs.core_properties.title = "SanDisk ATMOS — Three Technical Offerings"
    prs.core_properties.subject = "ATMOS technical value against NVIDIA Hopper and Blackwell architectures"
    prs.core_properties.author = "SanDisk ATMOS Strategy"
    prs.core_properties.last_modified_by = "SanDisk ATMOS Strategy"
    prs.core_properties.comments = "Internal technical proposal; no measured ATMOS performance claims."
    prs.core_properties.revision = 4
    prs.core_properties.created = datetime.now(timezone.utc)
    prs.core_properties.modified = datetime.now(timezone.utc)
    prs.save(OUTPUT)
    reopened = Presentation(OUTPUT)
    validate(reopened, 11)
    print(f"Wrote {OUTPUT} | slides={len(reopened.slides)} size={OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()