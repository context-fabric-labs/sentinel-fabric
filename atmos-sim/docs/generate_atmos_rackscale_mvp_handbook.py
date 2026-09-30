from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCS_DIR = Path(__file__).resolve().parent
ASSET_DIR = DOCS_DIR / "assets" / "atmos_rackscale_mvp_handbook"
OUTPUT = DOCS_DIR / "ATMOS_RackScale_MVP_Engineering_and_POC_Execution_Handbook.docx"

NAVY = "102A43"
INK = "24384A"
BLUE = "2E6F9E"
TEAL = "008C82"
GREEN = "4D8B61"
ORANGE = "D9823B"
RED = "B64B4B"
PURPLE = "6A5A98"
GRAY = "647382"
WHITE = "FFFFFF"
MIST = "E9F1F5"
LIGHT_GRAY = "F4F6F7"
MID_GRAY = "D8E0E5"
PALE_BLUE = "E7F0F7"
PALE_TEAL = "E4F3F0"
PALE_ORANGE = "FBEDE1"
PALE_RED = "F7E7E7"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def pil_rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def font_path() -> str:
    for candidate in (
        "/System/Library/Fonts/Avenir Next.ttc",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ):
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No supported font found")


FONT_PATH = font_path()


def image_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(FONT_PATH, size=size, index=1 if bold else 0)
    except OSError:
        return ImageFont.truetype(FONT_PATH, size=size)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 85, start: int = 95, bottom: int = 85,
                     end: int = 95) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = MID_GRAY, size: int = 6) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.05

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(12.5)
    subtitle.font.color.rgb = rgb(BLUE)

    for name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 13, 7),
        ("Heading 2", 13.5, BLUE, 10, 5),
        ("Heading 3", 10.5, TEAL, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.1)
    styles["List Bullet"].paragraph_format.space_after = Pt(2.5)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9.1)
    styles["List Number"].paragraph_format.space_after = Pt(2.5)

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(7.8)
    caption.font.italic = True
    caption.font.color.rgb = rgb(GRAY)

    for name, size, color, bold in (
        ("Kicker", 8.2, TEAL, True),
        ("Source Note", 7.3, GRAY, False),
        ("Small", 7.8, GRAY, False),
        ("Command", 8.1, INK, False),
    ):
        if name not in styles:
            styles.add_style(name, 1)
        style = styles[name]
        style.font.name = "Aptos" if name != "Command" else "Menlo"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_after = Pt(3)
        if name == "Command":
            style.paragraph_format.left_indent = Inches(0.2)
            style.paragraph_format.right_indent = Inches(0.2)
            style.paragraph_format.keep_together = True


def configure_document(document: Document) -> None:
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.24)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ATMOS RACK-SCALE MVP | ENGINEERING & POC HANDBOOK")
    run.font.name = "Aptos"
    run.font.size = Pt(7.3)
    run.font.bold = True
    run.font.color.rgb = rgb(GRAY)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(7.1))
    left = table.cell(0, 0).paragraphs[0]
    left.add_run("INTERNAL EXECUTION GUIDE").font.size = Pt(7.2)
    left.runs[0].font.color.rgb = rgb(GRAY)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right.add_run("PAGE ")
    run.font.size = Pt(7.2)
    run.font.color.rgb = rgb(GRAY)
    add_field(right, "PAGE")

    props = document.core_properties
    props.title = "ATMOS Rack-Scale MVP Engineering and POC Execution Handbook"
    props.subject = "Step-by-step software, hardware, productization, and four-gate validation guide"
    props.author = "Sandisk / ATMOS Architecture"
    props.last_modified_by = "Sandisk / ATMOS Architecture"
    props.comments = "Targets and procedures require configuration-specific approval and measured evidence."
    props.created = datetime.now(timezone.utc)
    props.modified = datetime.now(timezone.utc)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_paragraph(text, style=f"Heading {level}")


def add_body(document: Document, text: str) -> None:
    document.add_paragraph(text)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: Iterable[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK,
                  size: float = 8.0, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]], *,
              font_size: float = 7.8, header_fill: str = NAVY,
              banded: bool = True) -> object:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE,
                      size=font_size)
        set_cell_shading(table.rows[0].cells[index], header_fill)
        set_cell_border(table.rows[0].cells[index], header_fill)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value, size=font_size)
            set_cell_border(row.cells[index])
            if banded and row_index % 2:
                set_cell_shading(row.cells[index], LIGHT_GRAY)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(document: Document, label: str, text: str, *, fill: str = PALE_TEAL,
                accent: str = TEAL) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.75)
    set_cell_text(table.cell(0, 0), label, bold=True, color=WHITE, size=8.3,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_border(table.cell(0, 0), accent)
    set_cell_text(table.cell(0, 1), text, size=8.6)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_border(table.cell(0, 1), accent)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_command(document: Document, command: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, MID_GRAY)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = "Command"
    paragraph.add_run(command)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_part_opener(document: Document, number: str, title: str, promise: str) -> None:
    add_page_break(document)
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(6.0)
    set_cell_text(table.cell(0, 0), number, bold=True, color=WHITE, size=20,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.cell(0, 0), TEAL)
    set_cell_border(table.cell(0, 0), TEAL)
    cell = table.cell(0, 1)
    cell.text = ""
    set_cell_shading(cell, NAVY)
    set_cell_border(cell, NAVY)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = rgb(WHITE)
    paragraph = cell.add_paragraph()
    run = paragraph.add_run(promise)
    run.font.name = "Aptos"
    run.font.size = Pt(8.8)
    run.font.color.rgb = rgb(MIST)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def draw_wrapped(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
                 font: ImageFont.FreeTypeFont, fill: tuple[int, int, int],
                 align: str = "center") -> None:
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 24
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + max(0, len(lines) - 1) * 5
    y = y1 + max(0, (y2 - y1 - total) // 2)
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        x = x1 + 12 if align == "left" else x1 + (x2 - x1 - width) // 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + 5


def draw_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
             fill: str, size: int = 24) -> None:
    draw.rounded_rectangle(box, radius=18, fill=pil_rgb(fill), outline=pil_rgb(WHITE), width=3)
    draw_wrapped(draw, box, text, image_font(size, True), pil_rgb(WHITE))


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
               color: str) -> None:
    draw.line([start, end], fill=pil_rgb(color), width=6)
    x2, y2 = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x2, y2), (x2 - direction * 18, y2 - 10), (x2 - direction * 18, y2 + 10)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x2, y2), (x2 - 10, y2 - direction * 18), (x2 + 10, y2 - direction * 18)]
    draw.polygon(points, fill=pil_rgb(color))


def make_delivery_map(path: Path) -> None:
    image = Image.new("RGB", (1800, 950), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "RACK-SCALE MVP DELIVERY MAP", font=image_font(36, True), fill=pil_rgb(WHITE))
    layers = [
        ("PRODUCT", "Compatibility manifest | economic/support contract", PURPLE),
        ("OPERATIONS", "Kubernetes allocation | operator/lifecycle | telemetry/evidence", GRAY),
        ("SERVING", "Model contract | vLLM | placement | framework seam | collectives", BLUE),
        ("TRANSPORT", "ATMOS SDK/driver | registered buffers | DMA | completion | reset", ORANGE),
        ("HARDWARE", "NPU/HBF/LPDDR | PCIe endpoint | switch/backplane | server/rack", GREEN),
    ]
    y = 135
    for title, description, color in layers:
        draw_box(draw, (70, y, 380, y + 120), title, color, 27)
        draw.rounded_rectangle((410, y, 1730, y + 120), radius=18, fill=pil_rgb(LIGHT_GRAY),
                               outline=pil_rgb(color), width=4)
        draw_wrapped(draw, (445, y + 8, 1700, y + 112), description, image_font(25),
                     pil_rgb(INK), "left")
        y += 145
    draw.rounded_rectangle((70, 875, 1730, 930), radius=14, fill=pil_rgb(PALE_RED),
                           outline=pil_rgb(RED), width=3)
    draw_wrapped(draw, (85, 878, 1715, 927),
                 "Every layer must pass function, service, failure, operations, and support evidence before the rack claim is earned.",
                 image_font(21, True), pil_rgb(RED))
    image.save(path, quality=95)


def make_poc_ladder(path: Path) -> None:
    image = Image.new("RGB", (1800, 850), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "FOUR GROUPED POCs", font=image_font(36, True), fill=pil_rgb(WHITE))
    phases = [
        ("POC 1", "Vendor-neutral expert primitive", BLUE),
        ("POC 2", "Switched E3.S subsystem", TEAL),
        ("POC 3", "8 GPU + 8 ATMOS service", ORANGE),
        ("POC 4", "Named rack-scale product", PURPLE),
    ]
    x_positions = [80, 505, 930, 1355]
    for index, ((code, label, color), x) in enumerate(zip(phases, x_positions)):
        draw_box(draw, (x, 180, x + 360, 390), f"{code}\n{label}", color, 26)
        if index < 3:
            draw_arrow(draw, (x + 360, 285), (x_positions[index + 1] - 15, 285), color)
    gates = [
        "Numerical + ABI + lifecycle",
        "BOM + route + DMA + faults",
        "Quality + SLO + break-even",
        "Platform support + rack ROI",
    ]
    for text, (_, _, color), x in zip(gates, phases, x_positions):
        draw.rounded_rectangle((x, 470, x + 360, 590), radius=16, fill=pil_rgb(LIGHT_GRAY),
                               outline=pil_rgb(color), width=3)
        draw_wrapped(draw, (x + 10, 480, x + 350, 580), text, image_font(22, True), pil_rgb(INK))
    draw.rounded_rectangle((80, 680, 1715, 795), radius=18, fill=pil_rgb(PALE_ORANGE),
                           outline=pil_rgb(ORANGE), width=3)
    draw_wrapped(draw, (100, 690, 1695, 785),
                 "A gate consumes signed artifacts from the work packages. A demo, meeting, capacity fit, or successful enumeration is not a gate pass.",
                 image_font(24, True), pil_rgb(INK))
    image.save(path, quality=95)


def make_dependency_map(path: Path) -> None:
    image = Image.new("RGB", (1800, 1060), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "WORK-PACKAGE DEPENDENCY AND PROMOTION FLOW",
              font=image_font(34, True), fill=pil_rgb(WHITE))
    rows = [
        ("CONTROL", "SW-01 model contract | SW-10 evidence", BLUE),
        ("PRIMITIVE", "SW-02 packs | SW-07 SDK/driver | HW-01 NPU/memory | HW-02 endpoint", GREEN),
        ("HYBRID PATH", "SW-03 vLLM | SW-04 placement | SW-05 GPU seam | HW-03 transport", TEAL),
        ("PLATFORM", "SW-06 collectives | HW-04 native fabric | HW-05 switch | HW-06 server/rack", ORANGE),
        ("OPERATIONS", "SW-08 Kubernetes | SW-09 operator/lifecycle", GRAY),
        ("PRODUCT", "PR-01 manifest | PR-02 economics/support", PURPLE),
    ]
    y = 135
    centers: list[tuple[int, int]] = []
    for title, detail, color in rows:
        draw_box(draw, (75, y, 390, y + 115), title, color, 25)
        draw.rounded_rectangle((425, y, 1725, y + 115), radius=17, fill=pil_rgb(LIGHT_GRAY),
                               outline=pil_rgb(color), width=4)
        draw_wrapped(draw, (450, y + 5, 1700, y + 110), detail, image_font(23), pil_rgb(INK), "left")
        centers.append((900, y + 115))
        y += 145
    for index in range(len(centers) - 1):
        draw_arrow(draw, centers[index], (900, centers[index + 1][1] - 115), rows[index][2])
    draw.rounded_rectangle((75, 1000, 1725, 1045), radius=12, fill=pil_rgb(PALE_RED),
                           outline=pil_rgb(RED), width=2)
    draw_wrapped(draw, (90, 1002, 1710, 1043),
                 "Promotion moves downward only after upstream identities, negative tests, evidence and owners are stable.",
                 image_font(20, True), pil_rgb(RED))
    image.save(path, quality=95)


def build_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "delivery": ASSET_DIR / "delivery_map.png",
        "poc": ASSET_DIR / "poc_ladder.png",
        "dependencies": ASSET_DIR / "work_package_dependencies.png",
    }
    make_delivery_map(assets["delivery"])
    make_poc_ladder(assets["poc"])
    make_dependency_map(assets["dependencies"])
    return assets


def add_figure(document: Document, path: Path, caption: str, width: float = 7.0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


@dataclass(frozen=True)
class WorkPackageSummary:
    identifier: str
    domain: str
    artifact: str
    action: str
    owner: str
    purpose: str
    consumed_by: str


WORK_PACKAGES = [
    WorkPackageSummary("SW-01", "Software", "Model contract / learned router", "V", "Sandisk model", "Freeze unchanged model semantics and export reproducible routing traces.", "POC 1, 3, 4"),
    WorkPackageSummary("SW-02", "Software", "ExpertPack compiler", "N", "Sandisk + Tenstorrent", "Build versioned NPU packs for every ATMOS-owned expert.", "POC 1, 3, 4"),
    WorkPackageSummary("SW-03", "Software", "vLLM MoE integration", "M/N", "Sandisk + Tenstorrent", "Add the ATMOS prepare/execute/finalize seam without changing customer APIs.", "POC 1, 3, 4"),
    WorkPackageSummary("SW-04", "Software", "Heterogeneous adapter / placement", "N", "Sandisk + Tenstorrent", "Map experts to healthy ATMOS owners and preserve identity, gates, epochs, and bounds.", "POC 1-4"),
    WorkPackageSummary("SW-05", "Software", "GPU framework seam", "M", "Sandisk + Tenstorrent", "Pack, synchronize, restore, combine, and resume through stock GPU APIs.", "POC 1, 3, 4"),
    WorkPackageSummary("SW-06", "Software", "GPU collectives", "V", "Vendor runtime", "Keep ATMOS outside NCCL/RCCL and revalidate GPU rank/group behavior.", "POC 3, 4"),
    WorkPackageSummary("SW-07", "Software", "ATMOS SDK + Linux driver", "N", "Sandisk silicon/runtime", "Expose discovery, buffers, DMA queues, completion, errors, reset, and telemetry.", "POC 1-4"),
    WorkPackageSummary("SW-08", "Software", "Kubernetes allocation", "M/N", "Sandisk platform", "Co-allocate qualified GPU+ATMOS resources while preserving the GPU operator.", "POC 3, 4"),
    WorkPackageSummary("SW-09", "Software", "Operator / lifecycle", "N", "SRE / release", "Stage packs, publish complete readiness, drain, update, rollback, and recover.", "POC 3, 4"),
    WorkPackageSummary("SW-10", "Software", "Telemetry / evidence", "N", "Performance / validation", "Correlate request-to-NPU traces and publish stable service metrics.", "POC 1-4"),
    WorkPackageSummary("HW-01", "Hardware", "ATMOS NPU + HBF/LPDDR path", "V/M", "NPU + memory", "Qualify real FFN service curves, useful HBF, staging, thermals, and small batch.", "POC 1, 3, 4"),
    WorkPackageSummary("HW-02", "Hardware", "PCIe endpoint / DMA", "V/M", "Silicon / firmware", "Freeze endpoint/DMA capabilities, ordering, completion, errors, and reset.", "POC 1-4"),
    WorkPackageSummary("HW-03", "Hardware", "GPU-to-ATMOS transport", "V/M", "Sandisk + OEM", "Prove host-staged baseline and qualify any documented peer route.", "POC 2-4"),
    WorkPackageSummary("HW-04", "Hardware", "Native GPU scale-up fabric", "V", "OEM / platform", "Preserve NVLink/NVSwitch or UALink and verify no ATMOS rank/fabric intrusion.", "POC 3, 4"),
    WorkPackageSummary("HW-05", "Hardware", "Switch / backplane", "V/M", "OEM + switch FAE", "Qualify map, cables, security, diagnostics, containment, and firmware lifecycle.", "POC 2-4"),
    WorkPackageSummary("HW-06", "Hardware", "Server / rack envelope", "V/M", "OEM engineering", "Close lanes, storage displacement, power, cooling, BMC, service, and failure coupling.", "POC 2-4"),
    WorkPackageSummary("PR-01", "Product", "Compatibility manifest", "N", "Sandisk release", "Bind every released hardware, software, model, topology, and workload limit.", "POC 3, 4"),
    WorkPackageSummary("PR-02", "Product", "Economic / support contract", "N", "Product + solution architecture", "Report equal-SLO rate, reserve, avoidable cost, supply, and one support route.", "POC 3, 4"),
]


@dataclass(frozen=True)
class WorkPackageDetail:
    identifier: str
    title: str
    action: str
    owner: str
    purpose: str
    ready: tuple[str, ...]
    steps: tuple[str, ...]
    commands: tuple[str, ...]
    verification: tuple[tuple[str, str, str], ...]
    deliverables: tuple[str, ...]
    done: tuple[str, ...]
    failure: tuple[str, ...]
    platform_seam: str
    dependencies: str


SOFTWARE_PACKAGES = [
    WorkPackageDetail(
        "SW-01", "Model contract and learned-router baseline", "V", "Sandisk model/runtime",
        "Freeze the exact model semantics that ATMOS must preserve, create the numerical oracle, and "
        "export router traces that drive expert placement, service curves, and POC load generation.",
        (
            "Approved checkpoint, tokenizer, model code, license/custody route, and GPU-only reference environment are accessible.",
            "Product owner has named the first workload, quality measures, TTFT/TPOT objectives, context/output distribution, and data classification.",
            "A reproducible GPU-only run completes without ATMOS and produces stable model/task outputs.",
        ),
        (
            "Create a model configuration ID. Record checkpoint/tokenizer hashes, repository commits, model configuration, precision recipe, framework/container, and license/custody restrictions.",
            "Inventory every MoE layer: hidden size, routed/shared expert count, expert intermediate size, activation, top-k, gate normalization, accumulation, tensor-parallel partition, and any expert-specific exceptions.",
            "Define the unchanged semantic boundary: the learned router chooses the same expert IDs and weights; ATMOS placement does not predict or replace neural routing; the learned gate is applied exactly once.",
            "Instrument the reference path to export request, sequence, token, layer, selected expert IDs, gate weights, tensor shape/dtype, timing, and rank identity. Scrub or tokenize sensitive payloads while retaining routing shape.",
            "Capture representative traces at low, target, burst, and saturation load; include prefill/decode, context/output buckets, skew, hot/cold experts, and at least one failure/retry-free control run.",
            "Generate golden expert vectors for representative first/middle/last MoE layers, rare and hot experts, tokens-per-expert 1/2/4/8/16/32/64, boundary values, and the selected precision/accumulation rules.",
            "Freeze operator-level and end-task quality tolerances before ATMOS output is inspected. Include deterministic seeds, judge versions, and a rule for nondeterministic sampling comparisons.",
            "Run the reference three times from a clean process. Compare hashes/configuration, router identities, goldens, quality, memory residency, and latency distributions; explain every unstable field.",
            "Publish model-contract.yaml, router-trace.schema.json, golden vector bundle, workload-profile.yaml, quality-contract.yaml, and the signed baseline report.",
        ),
        (
            "shasum -a 256 <checkpoint-files> <tokenizer-files> > evidence/.../model/checksums.sha256",
            "python -m vllm.entrypoints.openai_api_server --model <checkpoint> --port 8001 <pinned-options>",
            "# TEMPLATE: python tools/export_moe_trace.py --endpoint http://localhost:8001 --workload <profile> --out evidence/.../raw/router.jsonl",
        ),
        (
            ("Configuration identity", "Recreate from hashes/commits/container on a clean host", "Manifest and checksums match; deviations fail closed"),
            ("Router semantics", "Compare selected IDs, gate weights, normalization and shared-expert behavior", "Exact IDs and weights within declared numeric tolerance"),
            ("Expert oracle", "Replay golden tensors on the GPU reference", "Every vector passes shape/dtype/status and numerical tolerance"),
            ("Workload representativeness", "Compare trace buckets to the approved service profile", "Coverage table includes target and tail/skew cases"),
            ("Quality baseline", "Run frozen task/operator evaluation three times", "Variance and acceptance threshold are signed before hybrid tests"),
        ),
        (
            "Model contract and compatibility fields", "Router trace schema and versioned trace corpus",
            "Golden expert-vector bundle", "Quality/SLO/workload contract", "GPU-only baseline evidence package",
        ),
        (
            "Another engineer reproduces the reference from the manifest.",
            "All ATMOS-owned expert shapes required by the first checkpoint are enumerated.",
            "Router traces and goldens pass schema, privacy, checksum, and replay tests.",
            "Quality and SLO acceptance are approved before any ATMOS comparison.",
        ),
        (
            "If the GPU reference is unstable, stop and fix the reference; ATMOS cannot be judged against a moving oracle.",
            "If a model operator, license, or precision cannot be supported, narrow the first checkpoint/operator scope and issue a new configuration ID.",
            "If production traces cannot be shared, build an approved synthetic trace with matching bucket/skew statistics and label the evidence limitation.",
        ),
        "Common semantics on NVIDIA and AMD. Re-run the reference quality, routing, memory, and latency baseline on each pinned CUDA or ROCm stack; never assume equal framework kernels imply equal service behavior.",
        "Starts the program. Feeds SW-02/03/04/05/10, HW-01, POC 1, POC 3, and every platform control.",
    ),
    WorkPackageDetail(
        "SW-02", "ExpertPack compiler and artifact pipeline", "N", "Sandisk + Tenstorrent",
        "Compile every ATMOS-owned (model, layer, expert) FFN into immutable, signed, versioned NPU "
        "artifacts with complete memory, precision, compatibility, and numerical metadata.",
        (
            "SW-01 model contract, tensor oracle, precision recipe, and complete expert inventory are approved.",
            "NPU compiler/kernel capability matrix and target firmware/runtime interface are versioned.",
            "Application-usable HBF, scratch, LPDDR/SRAM, alignment, and dual-generation update budgets are defined.",
        ),
        (
            "Define the ExpertPack schema: model/checkpoint hash, layer/expert identity, tensor shapes, weight layout, scale/zero-point data, activation, accumulation, kernel ID, scratch/tile budget, firmware/runtime minimums, checksum, signature, and golden-output references.",
            "Build a converter from the pinned checkpoint representation to a canonical intermediate form. Reject unknown operators, missing scales, unsupported shapes, duplicate identities, and ambiguous gate ownership.",
            "Implement the NPU compiler adapter using a pinned Tenstorrent/partner compiler and kernel set. Treat public TT-Metal/TT-NN support as reference only; ATMOS HBF/controller support requires the signed partner scope.",
            "Compile a smoke set covering one expert per shape family; load in the functional emulator/NPU; compare metadata, memory allocation, and golden outputs before compiling the full bank.",
            "Compile every ATMOS-owned expert. Produce a coverage report mapping each model layer/expert to a pack object, layout, precision, bytes, scratch, and compile result. A missing expert keeps the generation unready.",
            "Run operator goldens across token counts and edge values. Compare error distribution by layer/expert, not only an aggregate model score; quarantine any outlier pack.",
            "Calculate nominal, usable, and survivable placement capacity including scales, metadata, alignment, firmware reservation, scratch, fragmentation, replicas/spares, and two generations during rolling update.",
            "Make the build reproducible: rebuild from a clean environment, compare content hashes, capture compiler logs, lock tool/container versions, and document any nondeterministic field excluded from signing.",
            "Sign and publish the complete generation to an artifact repository with immutable retention, access control, SBOM/provenance, promotion state, and rollback references.",
            "Execute negative lifecycle tests: corrupt object, wrong checkpoint, wrong firmware, partial generation, revoked signature, insufficient capacity, and rollback to the prior complete generation.",
        ),
        (
            "# TEMPLATE: expertpack compile --model-contract model-contract.yaml --target atmos-<rev> --out packs/<generation>",
            "# TEMPLATE: expertpack verify packs/<generation> --goldens goldens/ --require-complete",
            "shasum -a 256 packs/<generation>/**/* > packs/<generation>/checksums.sha256",
        ),
        (
            ("Coverage", "Join model expert inventory with pack index", "100% of required ATMOS-owned experts map to exactly one valid object"),
            ("Numerics", "Run golden tensors for every shape family and sampled/all experts", "All outputs satisfy signed per-operator and task tolerance"),
            ("Memory", "Load generation under exact firmware reservation and scratch", "Usable and update-time fit pass; no nominal-capacity shortcut"),
            ("Reproducibility", "Clean rebuild and hash comparison", "Identical signed content or reviewed deterministic exclusions"),
            ("Lifecycle", "Stage/activate/rollback/corrupt/partial tests", "Only complete, authentic, compatible generations report ready"),
        ),
        (
            "ExpertPack schema/specification", "Compiler/converter and pinned toolchain", "Complete signed generation",
            "Coverage, numerical, capacity and reproducibility reports", "Artifact-repository and rollback runbook",
        ),
        (
            "Every first-checkpoint expert is compiled, verified, signed, and addressable by stable identity.",
            "Pack memory/scratch figures feed placement and hardware service-curve tests.",
            "Two generations can coexist or the operator has an approved stop-the-world alternative.",
            "The lifecycle refuses incomplete, corrupt, wrong-model, or wrong-firmware packs.",
        ),
        (
            "Unsupported shape/activation/precision: stop that operator or add a funded compiler/kernel work item; do not substitute a silent approximation.",
            "Pack bank does not fit usable HBF: change precision/placement/device count/failure policy and re-run economics; raw capacity is not a waiver.",
            "Numerical failure concentrated in selected layers: retain those experts on GPU or change the accumulation recipe under a new contract.",
        ),
        "Pack identity and Expert ABI are common. Compiler output targets ATMOS, not CUDA or ROCm; revalidate only framework-side precision/combination and end-task quality per GPU platform.",
        "Requires SW-01 and HW-01 capability inputs. Feeds SW-03/04/09, HW-01, PR-01, and POCs 1/3/4.",
    ),
    WorkPackageDetail(
        "SW-03", "vLLM MoE integration", "M/N", "Sandisk + Tenstorrent",
        "Integrate ATMOS expert ownership into vLLM/PyTorch through a maintained prepare/execute/finalize "
        "seam while preserving the public serving API, GPU attention/KV path, batching, and model semantics.",
        (
            "SW-01 model contract and SW-02 ExpertPack interface are stable.",
            "A pinned upstream vLLM/PyTorch build works on the GPU-only reference.",
            "Supported feature envelope is explicit: continuous batching required; speculative decoding, LoRA, graph capture, prefix caching, and distributed layouts each have a declared initial state.",
        ),
        (
            "Pin vLLM, PyTorch, CUDA/ROCm, Python, compiler/container, and model commits. Create a compatibility branch and a minimal patch inventory; do not begin from a deprecated private fork.",
            "Locate the MoE dispatch/execute/combine and weight-allocation extension points for the pinned release. Write a design note showing which code remains upstream/vendor-owned and which ATMOS plugin/backend code Sandisk owns.",
            "Add model-load ownership: build the placement map before weight allocation; omit only ATMOS-owned experts from GPU residency; keep shared/dense/GPU-local experts and an explicit fallback reserve where required.",
            "Implement prepare(): capture router outputs, validate model/placement generation, group compatible tokens, reserve output slots, enforce queue/deadline limits, and emit trace context.",
            "Implement execute(): submit groups through the heterogeneous adapter/transport without blocking the scheduler thread; retain cancellation and dependency handles.",
            "Implement finalize(): wait asynchronously for all required contributions, reject stale/partial results, restore token/rank order, apply the learned gates exactly once, combine, and resume the GPU continuation stream.",
            "Preserve continuous batching and admission behavior. Test mixed prefill/decode, request churn, cancellation, timeout, backpressure, small/large groups, and a slow/missing destination.",
            "Feature-gate unsupported paths at startup/request validation. Never accept a feature and silently run incorrect semantics; publish the supported matrix in the compatibility manifest.",
            "Add GPU-only fallback and ATMOS-disabled modes using the same service API. Ensure a failed ATMOS initialization cannot leave partially freed GPU expert weights in a ready service.",
            "Build unit, golden operator, scheduler, API, distributed, upgrade, and end-to-end tests. Re-run against every supported vLLM release; classify required upstream hooks and maintain bounded backports.",
        ),
        (
            "python -m vllm.entrypoints.openai_api_server --model <checkpoint> --port 8001 <gpu-control-options>",
            "# TEMPLATE: VLLM_PLUGINS=atmos python -m vllm.entrypoints.openai_api_server --model <checkpoint> --port 8002 --additional-config '{\"atmos\":{\"manifest\":\"<path>\"}}'",
            "# TEMPLATE: pytest tests/vllm_atmos -q --manifest <compatibility-manifest>",
        ),
        (
            ("API compatibility", "Run frozen OpenAI API contract and error tests", "Customer request/response semantics remain unchanged"),
            ("Model residency", "Compare GPU allocations to placement manifest", "Only declared ATMOS-owned tensors are absent; fallback reserve is visible"),
            ("Numerical path", "GPU-only versus hybrid operator/end-task comparison", "SW-01 tolerances pass with gates applied once"),
            ("Scheduler behavior", "Mixed batching, cancel, deadline, overload, slow device", "No deadlock, unbounded queue, cross-request result, or hidden partial output"),
            ("Release compatibility", "Test pinned and next supported vLLM version", "Patch set remains bounded and owned; unsupported combinations fail early"),
        ),
        (
            "ATMOS vLLM plugin/backend and bounded patch set", "Feature/support matrix", "Model-load ownership implementation",
            "Prepare/execute/finalize conformance tests", "Upgrade/backport/upstream plan", "GPU-only and ATMOS-disabled recovery mode",
        ),
        (
            "Frozen API, quality, continuous batching, cancellation, and overload tests pass.",
            "No proprietary NVIDIA/AMD driver or firmware changes are required.",
            "Unsupported features fail at configuration/request validation, not during execution.",
            "A clean install from the manifest starts GPU control and hybrid services reproducibly.",
        ),
        (
            "Required upstream hook absent: propose a generic upstream interface or maintain a reviewed, budgeted patch; otherwise narrow the feature/release envelope.",
            "Dynamic scheduling prevents bounded latency: cap grouping wait/queue, change admission, or keep latency-critical experts GPU-local.",
            "Feature interaction changes semantics: disable it in the first release and create a separately gated expansion package.",
        ),
        "Common plugin semantics; separate pinned CUDA and ROCm builds. GPU stream/event, allocator, graph, peer-memory, and distributed behavior are tested independently; no separate model fork.",
        "Requires SW-01/02/04/05/07. Feeds POC 1 integration, POC 3 full service, POC 4 platform releases and PR-01.",
    ),
    WorkPackageDetail(
        "SW-04", "Heterogeneous expert adapter and placement", "N", "Sandisk + Tenstorrent",
        "Own stable expert identity, physical placement, grouping, queue bounds, replicas/fallback, and "
        "immutable placement epochs without changing the learned router.",
        (
            "SW-01 router trace and expert inventory, SW-02 pack index, and usable/survivable device capacity are available.",
            "Topology/failure domains and first availability policy are defined, even if provisional.",
            "SW-07 transport exposes device health, queue credits, completion, cancellation, and generation identity.",
        ),
        (
            "Define stable logical ExpertID = model generation + layer + expert and physical OwnerID = platform/configuration + device + pack generation. Keep logical identity independent of GPU rank and PCIe address.",
            "Define placement-manifest fields: owners/replicas, GPU-local flag, affinity, capacity/scratch, switch/NUMA/failure domain, readiness, fallback, activation epoch, and drain state.",
            "Implement a static placement loader first. Validate complete coverage, no duplicate primary without policy, memory/scratch fit, complete surviving service, and pack readiness before publishing an epoch.",
            "Implement token grouping by model generation, placement epoch, layer, expert, precision, physical destination, and deadline class. Preserve original token/rank/output slots and learned gate values.",
            "Implement per-device and global queue bounds, credits, oldest-work age, grouping-wait cap, deadline-aware admission, fairness, and backpressure. Define whether rejected ATMOS work falls back or fails the request.",
            "Replay SW-01 traces through baseline placement strategies: hash/random control, co-activation-aware, GPU/ATMOS locality, hot-expert replica, failure-domain-aware, and selected GPU-local experts.",
            "Measure mean/p95 destinations per token/layer, co-activation edge cut, bytes per switch/root, queue imbalance, slowest contribution, head-of-line blocking, HBF capacity, replica cost, and surviving complete-service rate.",
            "Implement epoch lifecycle: stage owners, verify complete readiness, atomically publish a new epoch for new requests, drain old requests, and retire only after no live request/fallback references it.",
            "Inject owner loss, stale completion, pack mismatch, capacity loss, switch-domain loss, skew burst, and partial rollout. Prove no in-flight request mixes epochs or continues with missing contributions.",
            "Publish the first placement policy and a trace-to-placement report; keep optimization policy replaceable behind the stable manifest/epoch contract.",
        ),
        (
            "# TEMPLATE: atmos-placement validate --model model-contract.yaml --packs pack-index.json --topology topology.yaml --placement placement.yaml",
            "# TEMPLATE: atmos-sim/build/test_model_graph --gtest_filter='*Placement*:*Routing*'",
            "# TEMPLATE: atmos-placement replay --trace router.jsonl --policies hash,coactivation,replicated --out evidence/.../analysis/placement",
        ),
        (
            ("Coverage/capacity", "Validate every expert, owner, replica and memory/scratch sum", "Complete service fits nominally and after declared failure"),
            ("Identity/order", "Golden token routing and shuffled completion tests", "Expert/gate/output identities are exact across reorder/fan-in"),
            ("Queue stability", "Low/target/burst/saturation trace replay", "Bounded depth/age; target plus reserve has no unbounded growth"),
            ("Epoch safety", "Concurrent rollout/drain/reset/cancellation", "No request mixes placement or pack generations"),
            ("Placement value", "Compare candidate policies to hash/random control", "Chosen policy improves signed metric without violating failure/capacity/SLO"),
        ),
        (
            "Placement manifest/schema and validator", "Grouping/queue/epoch library", "Trace replay and policy comparator",
            "First placement report", "Failure and rollout tests", "GPU-local/fallback policy",
        ),
        (
            "Static placement and epoch lifecycle are correct before adaptive optimization is enabled.",
            "All target traces have complete coverage, bounded queues, and signed placement metrics.",
            "Declared endpoint/switch failure preserves correctness and either complete service or explicit fail-closed behavior.",
            "Placement output is portable across NVIDIA/AMD because ATMOS is not a GPU collective rank.",
        ),
        (
            "Fan-out/tail too high: co-locate co-activated experts, add replicas, retain hot experts on GPU, reduce offload scope, or change execution boundary.",
            "Capacity fails after reserve/update: increase devices, reduce model/precision/replicas, or reject the topology; do not consume safety reserve silently.",
            "Queue instability: reduce admission, cap grouping delay, add service capacity, or change workload class before full-service POC.",
        ),
        "Vendor-neutral logical placement. Physical topology fields differ: NVIDIA/AMD root/NUMA, host architecture and peer capability are manifest data, not branches in the routing algorithm.",
        "Requires SW-01/02/07/10 and HW-01/05/06 inputs. Feeds SW-03/05/09, POC 1 trace tests, POC 2 load/fault tests, POC 3/4.",
    ),
    WorkPackageDetail(
        "SW-05", "GPU framework seam", "M", "Sandisk + Tenstorrent",
        "Implement activation packing, asynchronous synchronization, result restoration, learned-gate "
        "combination, and continuation using stock PyTorch/CUDA or PyTorch/ROCm interfaces.",
        (
            "SW-01 tensor/gate contract and goldens, SW-03 integration points, SW-04 grouping identity, and SW-07 transport API are stable.",
            "Host-staged buffer ownership and lifetime are defined; peer DMA is a separately qualified adapter.",
            "Pinned GPU framework, driver, compiler, and profiling tools are available.",
        ),
        (
            "Define canonical activation/result wire layouts, alignment, dtype/scales, token/expert metadata, maximum group sizes, and integrity/status fields. Keep the wire contract GPU-vendor neutral.",
            "Allocate reusable pinned host buffers for the baseline. Implement size classes/pools, ownership generations, registration cache, bounded memory, NUMA affinity, and cleanup on process/device reset.",
            "Implement GPU packing/gather: group hidden states using SW-04 mappings, preserve original token/rank/output slots and gates, avoid scheduler-thread blocking, and record GPU events/dependencies.",
            "Connect dependency events to transport submission. Define when the host may read, when ATMOS may DMA, when the output slot is writable, and what reset/cancel does to every outstanding buffer.",
            "Implement return scatter/reorder and top-k fan-in. Reject duplicate, missing, corrupt, stale-generation, wrong-shape, and wrong-expert contributions before combination.",
            "Implement learned-gate application and accumulation in the specified precision exactly once, then record an event the GPU continuation stream waits upon.",
            "Build a transport adapter boundary. Host staging is mandatory; optional NVIDIA and AMD peer adapters implement the same buffer/submit/completion semantics and fall back only through an explicit supported policy.",
            "Measure pack, submit, DMA, wait, scatter, combine, and continuation separately with synchronized clocks/events. Include message sizes and tokens/expert from real traces, not only large bandwidth transfers.",
            "Test cancellation, timeout, late completion, reset during DMA, process exit, backpressure, pool exhaustion, mixed precision, multi-GPU rank ordering, and all top-k contributions arriving out of order.",
            "Optimize only after the correctness trace passes: fused gather/scatter, persistent buffers, vectorized metadata, overlap, and graph-compatible portions. Re-run all error and quality tests after each optimization.",
        ),
        (
            "# NVIDIA TEMPLATE: CUDA_LAUNCH_BLOCKING=0 NSYS_NVTX_PROFILER_REGISTER_ONLY=0 nsys profile -o evidence/.../traces/gpu_seam <hybrid-command>",
            "# AMD TEMPLATE: rocprofv3 --output-format json --output-file evidence/.../traces/gpu_seam -- <hybrid-command>",
            "data-plane/scripts/run_benchmark.sh --sizes 1024,4096,16384,65536,262144 --warmup 20 --iterations 200",
        ),
        (
            ("Wire correctness", "Round-trip patterned tensors/metadata across size/precision/top-k", "Exact identity/order/status and signed numerical tolerance"),
            ("Synchronization", "Out-of-order streams/completions and dependency tracing", "No early read/write, global device sync, deadlock or hidden serialization"),
            ("Buffer safety", "Cancel/reset/exit/pool exhaustion/stale handle", "No use-after-free, cross-request corruption or stale DMA"),
            ("Latency", "Real trace message/group sweep with synchronized decomposition", "Complete contribution-to-continuation fits the POC budget"),
            ("Portability", "Same vectors/API on CUDA and ROCm builds", "Common semantics pass; platform deltas stay in adapters/manifest"),
        ),
        (
            "Wire-format and buffer-lifecycle specification", "CUDA and ROCm framework adapters", "Host-staged implementation",
            "Optional peer-adapter interface", "Pack/scatter/combine kernels", "Correctness, fault and profiling reports",
        ),
        (
            "No global synchronization is required in the steady-state critical path unless explicitly budgeted.",
            "Host-staged path is correct and bounded before any peer route is promoted.",
            "Every error/reset/cancel path retires or invalidates buffers safely.",
            "Complete GPU continuation latency is measured and consumed by POC gate calculations.",
        ),
        (
            "Host staging misses budget: optimize grouping/pools/overlap, then qualify peer DMA; if both fail, pivot the execution boundary.",
            "Graph capture conflicts: exclude the ATMOS segment initially and quantify the regression; do not claim graph support without tests.",
            "Cross-rank reorder/collective conflict: simplify the first distributed topology and revalidate SW-06 before expanding.",
        ),
        "One wire/transport contract. CUDA uses streams/events and supported pinned/peer APIs; ROCm uses HIP streams/events and supported amdgpu/peer APIs. Build and test independently.",
        "Requires SW-01/03/04/07 and HW-03. Feeds POC 1 end-to-end primitive, POC 3/4 and SW-10 latency evidence.",
    ),
    WorkPackageDetail(
        "SW-06", "GPU collectives invariant", "V", "NVIDIA/AMD vendor runtime + Sandisk validation",
        "Prove that moving selected experts to ATMOS does not make ATMOS a GPU rank, alter native "
        "collective libraries, or break rank/group configuration and ordering.",
        (
            "GPU-only tensor/pipeline/expert-parallel topology and collective traces are reproducible.",
            "SW-03 placement determines which GPU expert weights/ranks remain and which move to ATMOS.",
            "Pinned NCCL or RCCL versions, topology files, environment, and multi-node network are recorded.",
        ),
        (
            "Capture the GPU-only process/rank/group topology, collective sequence, message sizes, algorithms/protocols, topology discovery, and latency/bandwidth baseline.",
            "Document the invariant: ATMOS has no CUDA/HIP device identity, NCCL/RCCL rank, NVLink/NVSwitch/UALink membership, or collective-library patch.",
            "Reconfigure GPU expert/tensor/pipeline groups after offloaded weights are omitted. Ensure each remaining GPU rank performs the same ordered collectives expected by the model graph.",
            "Run single-node and, where in scope, multi-node correctness with ATMOS disabled, host-staged hybrid, and each supported failure/fallback mode. Compare collective order and outputs.",
            "Collect NCCL/RCCL debug/topology logs and GPU traces. Attribute any changed collective size/count to the model placement contract, not to ATMOS joining the fabric.",
            "Inject slow ATMOS completion, request cancellation, endpoint failure, GPU rank failure, and timeout. Prove GPU ranks do not enter mismatched collectives or deadlock while waiting for heterogeneous fan-in.",
            "Test rolling pack/placement epochs while GPU process groups remain stable, or document a controlled restart if group topology must change.",
            "Publish the platform-specific collective manifest and invariant report; repeat for every named NVIDIA and AMD platform/configuration.",
        ),
        (
            "# NVIDIA TEMPLATE: NCCL_DEBUG=INFO NCCL_DEBUG_SUBSYS=INIT,COLL,GRAPH <gpu-control-or-hybrid-command> 2> evidence/.../raw/nccl.log",
            "# AMD TEMPLATE: NCCL_DEBUG=INFO RCCL_ENABLE_TUNING_TRACE=1 <gpu-control-or-hybrid-command> 2> evidence/.../raw/rccl.log",
        ),
        (
            ("No ATMOS rank", "Inspect process groups, device discovery and collective logs", "Only native GPUs appear in NCCL/RCCL groups"),
            ("Collective order", "Compare rank traces under control/hybrid/cancel/fault", "No mismatch, deadlock, timeout cascade or hidden patch"),
            ("Numerical equivalence", "Distributed GPU-only versus hybrid output", "SW-01 quality contract passes"),
            ("Performance regression", "Collective latency/bytes/count at equal workload", "Changes are explained and within signed service budget"),
            ("Failure containment", "Slow/missing ATMOS and GPU-rank fault", "Declared failover/fail-closed path avoids collective hang"),
        ),
        (
            "NVIDIA NCCL and AMD RCCL invariant reports", "Pinned rank/group/topology manifests",
            "Distributed correctness/performance/fault evidence", "Supported restart/reconfiguration rule",
        ),
        (
            "Stock supported NCCL/RCCL pass without modification.",
            "No collective mismatch/deadlock across declared heterogeneous failures.",
            "Platform-specific group topology is captured in PR-01.",
        ),
        (
            "Hybrid path changes collective ordering: correct the model graph/integration; do not patch vendor collectives as a shortcut.",
            "Tail wait causes rank stalls: improve placement/admission/fan-in or move the critical expert path; record the service impact.",
            "A platform requires proprietary collective changes: reject that path for the first product.",
        ),
        "Separate verification on stock NCCL and RCCL. NVLink/NVSwitch and UALink/Infinity Fabric remain platform-owned boundaries.",
        "Requires SW-01/03/04/05 and HW-04. Feeds POC 3 and each POC 4 platform release.",
    ),
    WorkPackageDetail(
        "SW-07", "ATMOS SDK and Linux driver", "N", "Sandisk silicon/firmware/runtime",
        "Provide the supported host contract for discovery, isolation, registered buffers, DMA queues, "
        "completion, cancellation, reset, telemetry, firmware, and error recovery.",
        (
            "HW-02 endpoint/DMA register and capability contract exists; security/IOMMU assumptions are reviewed.",
            "SW-02 pack and HW-01 execution requirements define queue, memory, completion, and telemetry needs.",
            "Supported kernels/distros/host architectures and device/firmware update model are selected for the first fixture.",
        ),
        (
            "Define device discovery and capability ABI: hardware/FW revision, endpoint/link, queue count/depth, address width, buffer rules, pack capacity, telemetry, reset, error, and optional peer capabilities.",
            "Define the stable userspace API: open/close, query, register/unregister, create/destroy queue, submit dependencies/work, poll/wait completion, cancel, error detail, reset/recover, telemetry, and firmware/pack operations.",
            "Implement least-privilege BAR/MMIO mapping and DMA memory registration with bounded addresses, pin/accounting limits, ownership/process/tenant generation, IOMMU mappings, and cleanup on close/crash/reset.",
            "Implement queue/ring allocation, descriptor validation, doorbells, MSI-X or polling, completion status, sequence numbers, dependency fences, queue credits, timeout/cancel, and backpressure.",
            "Keep DMA and NPU completion distinct. Define operation states and ambiguity: submitted, DMA-in complete, admitted, executing, DMA-out complete, failed/cancelled/indeterminate; make buffer reuse safe in every state.",
            "Implement reset invalidation. FLR/device/switch/root reset revokes registrations, queues, pack readiness, and placement ownership; rejoin requires rediscovery, self-test, pack verify, and a new generation.",
            "Implement telemetry and diagnostics: link/errors/AER, queue age/depth, bytes, operations, timeouts, resets, temperature/power/throttle, firmware, pack and fatal snapshots with stable IDs.",
            "Implement secure firmware update/rollback, signature/version policy, audit, SBOM, failure recovery, and field diagnostics bundle. Integrate with OEM BMC only through a documented boundary.",
            "Build a functional emulator/stub backend for CI that fails loudly on unimplemented hardware behavior. Use real endpoint/prototype tests for DMA/reset/timing claims.",
            "Execute unit, ABI, stress, fuzz, race, leak, suspend/reboot, reset-during-DMA, malformed descriptor, out-of-range DMA, stale handle, cross-process/tenant, queue exhaustion, and kernel upgrade tests.",
            "Package version-matched SDK/driver/firmware with DKMS/kmod or approved distro method, install/uninstall/upgrade/rollback docs, diagnostics, support matrix, and security response owner.",
        ),
        (
            "cmake -S data-plane -B data-plane/build -DCMAKE_BUILD_TYPE=Release -DBUILD_TESTING=ON -DBUILD_BENCHMARKS=ON && cmake --build data-plane/build --parallel",
            "ctest --test-dir data-plane/build -N  # If Total Tests is nonzero, run: ctest --test-dir data-plane/build --output-on-failure; otherwise provide GoogleTest and reconfigure",
            "data-plane/scripts/run_benchmark.sh --sizes 1024,4096,16384,65536,262144 --warmup 20 --iterations 200",
            "# HARDWARE TEMPLATE: sudo -n ./tools/atmos-diag --manifest <id> --self-test --json evidence/.../raw/diag.json",
        ),
        (
            ("ABI/API", "Version and negative compatibility tests", "Supported versions interoperate; unsupported versions fail early"),
            ("Memory isolation", "Bounds, IOMMU, stale/cross-process handles, crash/reset", "No arbitrary/cross-owner DMA; all resources reclaimed"),
            ("Queue/completion", "Concurrency, reorder, wrap, backpressure, timeout/cancel", "No lost/duplicate/ambiguous completion without explicit status"),
            ("Reset/rejoin", "Reset at every operation state and repeated recovery", "No stale access; readiness only after full requalification"),
            ("Lifecycle", "Fresh install, upgrade, rollback, uninstall, kernel update", "Supported manifest reproduces and diagnostics identify drift"),
            ("Security", "Fuzz/threat-model negative suite and signed update", "No critical open finding; audit/SBOM/response owner ready"),
        ),
        (
            "Public SDK/UAPI and internal driver/FW specifications", "Userspace library and language bindings as required",
            "Linux driver, firmware package and emulator/stub", "Conformance/stress/security/reset suites",
            "Diagnostics, telemetry and support bundle", "Install/upgrade/rollback and compatibility documentation",
        ),
        (
            "All required API and negative tests pass on the exact fixture manifest.",
            "No reset/error can permit stale DMA or falsely ready service.",
            "Driver/FW package has an owner, support matrix, security process, and rollback path.",
            "Performance counters needed by SW-10 and POC decomposition are available and calibrated.",
        ),
        (
            "Endpoint behavior is ambiguous: block higher POCs until firmware/driver state and ownership are explicit.",
            "Peer aperture cannot meet isolation: retain host staging; do not weaken security controls.",
            "Kernel/distro maintenance cost is unowned: narrow supported matrix or secure an enterprise distro/OEM support route.",
        ),
        "Same userspace contract; separate x86-64/Arm64 and kernel/OEM builds. NVIDIA/AMD peer registration is an optional transport backend, not the core driver API.",
        "Requires HW-01/02 and security/platform inputs. Feeds every POC, SW-03/04/05/08/09/10 and PR-01.",
    ),
    WorkPackageDetail(
        "SW-08", "Kubernetes allocation and topology-aware scheduling", "M/N", "Sandisk platform",
        "Expose ATMOS as allocatable, health-aware, topology-described resources and co-schedule only "
        "qualified GPU+ATMOS groups while preserving NVIDIA/AMD GPU operators.",
        (
            "SW-07 discovery/health API and HW-06 topology/failure-domain data are stable.",
            "SW-09 defines complete-service readiness and PR-01 defines compatible combinations.",
            "One supported Kubernetes distribution/version and existing NVIDIA or AMD GPU operator are selected.",
        ),
        (
            "Choose the initial allocation model: whole ATMOS devices or an explicitly qualified complete service group. Defer fractional sharing until isolation, queue QoS, accounting, and noisy-neighbor tests pass.",
            "Define resource attributes: device/service-group ID, hardware/FW, capacity, NUMA/root, switch/tray/rack failure domain, health, thermals, pack generation, placement epoch, and compatibility-manifest ID.",
            "Implement the Kubernetes device plugin or DRA driver using supported upstream APIs. Discovery comes from SW-07; never infer readiness solely from a present PCIe function.",
            "Publish topology and health without replacing the NVIDIA GPU Operator or AMD GPU Operator. Define labels/claims/affinity for qualified GPU+ATMOS combinations and prevent unsupported cross-root/rack scheduling.",
            "Implement allocation preparation: verify manifest, health, complete pack/service-group readiness, capacity and failure policy before mounting device nodes/config/credentials into the workload.",
            "Implement health transitions and quarantine. A device becoming unhealthy stops new allocation; workload-aware recovery comes from SW-09 and the application, not automatic device-plugin migration.",
            "Implement deallocation/cleanup: drain or cancel work, revoke buffers/queues, clear tenant state, release claims, and prevent immediate reuse until the driver/operator confirms safety.",
            "Test scheduler pressure, insufficient groups, NUMA/root affinity, mixed GPU models, pack generation drift, rolling node reboot, plugin/operator restart, endpoint/switch failure, and network partition.",
            "Add policy and admission checks so a pod cannot request an unsupported GPU/ATMOS/model/topology tuple or bypass the compatibility manifest.",
            "Package manifests/Helm/operator dependencies with install/upgrade/rollback, RBAC, network/security policy, observability, and support documentation.",
        ),
        (
            "kubectl get nodes -L topology.kubernetes.io/zone,kubernetes.io/arch",
            "kubectl get pods,nodes -o wide > evidence/.../environment/k8s-inventory.txt",
            "# TEMPLATE: kubectl get resourceclaims,resourceslices -A -o yaml > evidence/.../raw/atmos-resources.yaml",
            "# TEMPLATE: kubectl apply -f deploy/atmos-dra/ && kubectl wait --for=condition=Ready pod -l app=atmos-dra -n atmos-system --timeout=180s",
        ),
        (
            ("Resource discovery", "Compare Kubernetes objects to driver/OEM inventory", "Exact devices/groups, topology and manifest IDs match"),
            ("Co-allocation", "Schedule qualified/unqualified GPU+ATMOS requests", "Qualified requests land correctly; invalid topology fails with clear reason"),
            ("Health", "Endpoint/switch/plugin/node failures", "No new allocation to failed domain; running workload follows declared recovery"),
            ("Isolation/cleanup", "Multi-tenant allocate/deallocate/crash/reset", "No cross-tenant state or stale resource access"),
            ("Upgrade", "Plugin/DRA/operator/Kubernetes upgrade/rollback", "Claims and workloads preserve or follow documented drain/restart behavior"),
        ),
        (
            "Device plugin/DRA driver", "Resource schema/topology labels and policies", "GPU+ATMOS co-allocation examples",
            "Kubernetes conformance/fault/upgrade tests", "RBAC/security/install/support documentation",
        ),
        (
            "Only compatible, healthy, complete groups can be allocated.",
            "GPU operators remain stock and separately owned.",
            "Health, topology, cleanup, and upgrade tests pass on the supported distro/version.",
            "Fractional tenancy is absent or separately qualified and documented.",
        ),
        (
            "DRA unavailable in selected distro: use whole-device plugin plus affinity/admission policy and document migration path.",
            "Topology cannot express the qualified group: package it as a node/service-level resource or narrow platform scope.",
            "Automatic health does not restore complete model: block readiness and invoke SW-09 recovery; never rely on pod restart alone.",
        ),
        "Common ATMOS resource semantics. Preserve NVIDIA and AMD GPU Operators; platform topology and CPU architecture are manifest attributes and independently tested.",
        "Requires SW-07/09/10, HW-06 and PR-01. Feeds POC 3 operational tests and POC 4 product qualification.",
    ),
    WorkPackageDetail(
        "SW-09", "ATMOS operator and lifecycle", "N", "SRE / release",
        "Manage complete ExpertPack generations, readiness, placement activation, drain, update, rollback, "
        "repair, failover, and recovery at the model-service boundary.",
        (
            "SW-02 pack repository, SW-04 placement epochs, SW-07 health/reset, SW-08 allocation, and PR-01 manifest schema are available.",
            "Failure policy, RTO/RPO, surviving capacity, spares, and support escalation are declared for the first service.",
            "Supported Kubernetes/host deployment and secret/signing infrastructure are ready.",
        ),
        (
            "Define custom resources or service state for ATMOS hardware groups, pack generations, placement epochs, compatibility manifests, rollout policy, replica/failure domain, degraded envelope, and support condition.",
            "Implement reconciliation from observed hardware/driver/Kubernetes state to desired service state. Use idempotent operations and durable status; retries must not duplicate activation or destroy the rollback generation.",
            "Implement pack staging: reserve capacity, fetch immutable objects, verify signature/checksum/compatibility, load every required owner/replica, run device goldens, and keep service ownership unchanged.",
            "Implement readiness: device healthy is insufficient. Publish service-group readiness only when the complete pack generation, placement, dependencies, surviving capacity, and compatibility manifest pass.",
            "Implement activation: atomically publish a new placement epoch for new requests, observe quality/SLO canary, and retain old-generation rollback. Prevent mixed generation within a request.",
            "Implement drain/retire: stop new requests for the old epoch, wait/cancel according to deadline policy, verify zero references, remove old packs only after rollback policy permits, and record audit evidence.",
            "Implement rollback for quality/SLO/error/health regression. Restore prior pack/placement and verify complete readiness; define stop-the-world behavior if dual generations cannot coexist.",
            "Implement recovery: quarantine failed endpoint/domain, choose complete replica/GPU fallback/fail-closed policy, rebuild replacement, verify self-test and packs, publish a new epoch, and meet RTO/RPO.",
            "Implement firmware/driver/platform maintenance coordination: cordon/drain, support bundle, update, reset/reboot, rediscovery, pack verify, placement rejoin, and rollback.",
            "Add status/conditions/events/metrics, audit, RBAC, alerts, runbooks, support bundle collection, and escalation routing. Test controller restarts and control-plane partitions.",
            "Run lifecycle/fault matrix under live target load: normal update, partial stage, corrupt pack, failed canary, endpoint/switch/node loss, insufficient surviving capacity, operator restart, network partition, and repeated recovery.",
        ),
        (
            "# TEMPLATE: kubectl apply -f deploy/atmos-operator/",
            "# TEMPLATE: kubectl apply -f examples/atmosservice-qwen3.yaml && kubectl wait --for=condition=Ready atmosservice/qwen3 --timeout=1800s",
            "# TEMPLATE: kubectl patch atmosservice qwen3 --type=merge -p '{\"spec\":{\"packGeneration\":\"<new>\"}}'",
            "# TEMPLATE: kubectl get atmosservice qwen3 -o yaml > evidence/.../raw/operator-state.yaml",
        ),
        (
            ("Complete readiness", "Remove one expert/replica/pack/device", "Service refuses ready or enters only the signed degraded mode"),
            ("Generation safety", "Concurrent traffic during stage/activate/drain/rollback", "No request mixes pack/placement generations"),
            ("Recovery", "Endpoint/switch/node loss and replacement", "Correct policy, RTO/RPO and surviving accepted rate pass"),
            ("Controller resilience", "Operator/API restart and partition", "Reconciliation is idempotent; desired/observed state converges safely"),
            ("Maintenance", "Driver/FW/OS/Kubernetes upgrade and rollback", "Manifest and readiness update correctly; diagnostics retained"),
        ),
        (
            "Operator/controller and APIs/CRDs", "Pack/placement rollout implementation", "Readiness/failure/recovery state machine",
            "Lifecycle/fault/upgrade test suite", "SRE runbooks, alerts and diagnostics bundle", "RTO/RPO and degraded-envelope report",
        ),
        (
            "Complete service, not PCIe presence, controls readiness.",
            "All lifecycle and declared failure tests pass under load.",
            "Fresh install, upgrade and rollback are reproducible from PR-01.",
            "Support can diagnose and escalate from the generated bundle without developer-only access.",
        ),
        (
            "Dual-generation capacity unavailable: approve stop-the-world rollout or change capacity; never advertise rolling update implicitly.",
            "Recovery misses RTO or capacity: add complete replicas/spares, narrow SLO, or reject the availability claim.",
            "Ownership across OEM/Sandisk/vendor is ambiguous: block product readiness until one first-response and escalation path is contracted.",
        ),
        "Common lifecycle state machine; package/test per CPU architecture, distro and GPU platform. GPU vendor operators remain separate dependencies.",
        "Requires SW-02/04/07/08/10, HW-06, PR-01/02. Feeds POC 3 lifecycle/fault and POC 4 release/support qualification.",
    ),
    WorkPackageDetail(
        "SW-10", "Telemetry and evidence system", "N", "Performance / validation",
        "Correlate request, layer, token, expert, transport, NPU, GPU, switch, thermal, failure, and cost "
        "signals into reproducible evidence used consistently from RTX development through rack release.",
        (
            "SW-01 identifiers/workload/SLO, SW-04 placement identity, SW-07 counters, and POC evidence repository are defined.",
            "Clock sources and synchronization method for GPU/CPU/ATMOS/BMC/switch are identified.",
            "Data classification, retention, access, and customer-trace scrubbing rules are approved.",
        ),
        (
            "Define one correlation context: program/config/run, request/sequence/token range, model/pack/placement generation, layer/expert, GPU rank, ATMOS device/queue/work, transport buffer, and retry/cancellation generation.",
            "Define metric semantics and units before instrumentation: offered/admitted/accepted/rejected/failed/out-of-SLO work; TTFT/TPOT/output rate; queue wait/age/depth; logical/physical bytes; useful HBF; NPU/kernel; power/thermal; resets/errors/retries; RTO/degraded time.",
            "Instrument vLLM/service, placement, GPU seam, transport/driver, firmware/NPU, switch/OEM, Kubernetes/operator, and power/BMC. Preserve raw events plus derived metrics; never retain only dashboards.",
            "Establish clock correlation with PTP or measured offsets where supported. Include source clock, timestamp quality, calibration event, uncertainty, and drift in the run manifest; avoid fake sub-microsecond decomposition across unsynchronized clocks.",
            "Implement trace propagation and sampling. Capture 100% errors/faults/gate runs; use deterministic sampling for high-volume soak while retaining aggregate counters and exemplars.",
            "Implement the evidence writer: immutable run directory, manifest, environment inventory, configs, logs, traces, metrics, outputs, analysis, checksums, deviations, and tool versions.",
            "Build standard analyses: latency critical path, queue/service curves, placement/fan-out, route matrix, quality, accepted-load curve, fault/RTO, power/thermal, and TCO input extraction.",
            "Build dashboards and alerts for operations, but derive gate decisions from archived evidence and scripted analysis. Version queries and dashboard definitions with the manifest.",
            "Validate telemetry under known synthetic events: one request, known bytes, queue delay, injected error/reset, thermal threshold, cancellation, and accepted/rejected work. Reconcile counters across layers and explain loss/duplication.",
            "Run privacy/security review: redact prompts/activations/weights/secrets, control evidence access, sign/checksum packages, and define support-bundle content safe for each recipient.",
            "Publish the evidence schema, instrumentation matrix, analysis scripts, dashboard/alert pack, and a complete sample gate package another engineer can reproduce.",
        ),
        (
            "ctest --test-dir atmos-sim/build -R 'metrics|golden|e2e' --output-on-failure",
            "cargo test --manifest-path control-plane/Cargo.toml telemetry -- --nocapture",
            "# TEMPLATE: python tools/validate_evidence.py evidence/<program>/<poc>/<run-id> --require-checksums --require-manifest",
        ),
        (
            ("Correlation", "Single known request across every layer", "IDs join without ambiguity from API through continuation"),
            ("Counter reconciliation", "Known messages/bytes/work/errors", "Layer totals reconcile or documented loss/sampling is bounded"),
            ("Clock quality", "Calibration event and drift test", "Uncertainty is smaller than claimed decomposition or explicitly limits it"),
            ("Evidence reproducibility", "Re-run scripted analysis from archived raw data", "Tables/plots/decision inputs reproduce without live systems"),
            ("Privacy/security", "Inspect sample/support packages", "No prohibited data/secret; access, checksum and retention controls pass"),
        ),
        (
            "Stable telemetry/evidence schema", "Instrumentation and correlation libraries", "Standard analysis scripts",
            "Dashboards/alerts with versioned queries", "Evidence validator and sample package", "Privacy/retention/support-bundle policy",
        ),
        (
            "Every POC metric has one definition, source, unit and owner.",
            "A slow token can be decomposed from admission through GPU continuation within known clock uncertainty.",
            "Raw evidence and scripted analysis reproduce each gate result.",
            "RTX, NVIDIA rack and AMD rack results use the same metric semantics without reinterpreting definitions.",
        ),
        (
            "Clocks cannot correlate finely: report component distributions and end-to-end latency without unsupported critical-path attribution; improve instrumentation before optimization claims.",
            "Counters disagree: block gate decision until loss, duplication, sampling or retry semantics are resolved.",
            "Sensitive trace cannot leave customer boundary: run analysis in place and export approved aggregates/checksums/decision evidence only.",
        ),
        "Common schema and analyses; instrumentation adapters differ for Nsight/DCGM/NCCL versus ROCm/rocprof/RCCL and for OEM BMC/switch tools.",
        "Requires all identity/metric producers. Starts in wave 0 and is mandatory for every work package, POC and PR-02 economics.",
    ),
]


HARDWARE_PACKAGES = [
    WorkPackageDetail(
        "HW-01", "ATMOS NPU and HBF/LPDDR execution path", "V/M", "NPU + memory engineering",
        "Qualify real routed-FFN correctness and service curves across HBF, LPDDR/SRAM staging, NPU "
        "kernels, small batches, contention, power, and thermal states.",
        (
            "SW-01 expert shapes/goldens and SW-02 pack compiler/schema are available.",
            "Silicon, emulator or prototype revision and evidence class are recorded; planning targets are not labeled as measurements.",
            "Approved board, cooling, power, firmware, telemetry and hardware safety procedure are ready.",
        ),
        (
            "Create the device capability manifest: NPU revision, supported operators/precisions/shapes, HBF capacity/rate target, LPDDR/SRAM, DMA/queue interfaces, firmware, power/thermal limits, sensors, and known errata.",
            "Select representative expert families from the first checkpoint: first/middle/last layer, min/median/max expert size, hot/cold placement, tokens-per-expert 1/2/4/8/16/32/64, and the selected FP8/accumulation recipe.",
            "Load signed smoke ExpertPacks and run deterministic goldens. Verify tensor layout, scales, activation, accumulation, output status and error handling before performance mode.",
            "Instrument physical bytes at each memory boundary: HBF reads/writes, LPDDR/SRAM staging, tile reload, output writes, descriptor/metadata, and NPU idle/stall. Report useful HBF bytes separately from interface traffic.",
            "Sweep tokens/expert, queue depth, concurrent experts, hot/cold pack state, tile size, ring/buffer configuration, precision, and arrival distribution. Include low/target/burst/saturation and recovery after burst.",
            "Build service curves for service time, expert/s, token/s, p50/p95/p99, useful HBF GB/s, read amplification, NPU utilization, queue stability, scratch, power, temperature, frequency and throttle state.",
            "Test LPDDR/SRAM contention and staging alternatives. Detect any pattern where a nominal 200 GB/s HBF target is constrained by narrower read/write staging or repeated small weight reads.",
            "Run sustained thermal/power tests at representative mixed load until steady state. Exercise approved inlet/cooling conditions and throttle/shutdown transitions; do not bypass hardware safety limits.",
            "Inject malformed/unsupported pack, ECC/RAS event where safely supported, queue exhaustion, kernel timeout, firmware error and reset. Verify explicit status, bounded cleanup and no incorrect output.",
            "Calibrate the digital twin using measured service distributions, useful/physical bytes, queues and thermal states. Preserve pre- and post-calibration parameters and uncertainty.",
            "Publish qualified shape/precision/service envelopes and unsupported combinations. Re-run under each host/rack traffic pattern that changes queueing, thermal or power behavior.",
        ),
        (
            "ctest --test-dir atmos-sim/build -R 'virtual_devices|model_graph|golden|e2e' --output-on-failure",
            "# TEMPLATE: atmos-npu-bench --packs <generation> --tokens 1,2,4,8,16,32,64 --queues 1,4,8,16 --duration 600 --json evidence/.../raw/npu.json",
            "# TEMPLATE: atmos-diag --telemetry 1s --duration 1800 --out evidence/.../power-thermal/device.jsonl",
        ),
        (
            ("Numerics", "All representative/full expert goldens and end-task sample", "SW-01 tolerance; no silent unsupported operator/precision"),
            ("Service curve", "Repeated tokens/expert and load sweep", "Stable distributions, confidence intervals and saturation point"),
            ("Useful bandwidth", "Reconcile weight/output work with memory counters", "Useful versus physical bytes and amplification explained"),
            ("Queue stability", "Target/burst/saturation/recovery", "No unbounded target queue; overload rejection/backpressure explicit"),
            ("Thermal/power", "Steady mixed load and approved high-condition test", "No hidden throttle at supported envelope; safety transitions work"),
            ("Fault/RAS", "Supported negative/fault/reset matrix", "No incorrect output; bounded status, cleanup and rejoin"),
        ),
        (
            "Device capability and errata manifest", "Qualified ExpertPack/operator matrix", "Expert service-curve dataset/model",
            "Useful-HBF/staging analysis", "Power/thermal/throttle envelope", "Fault/RAS/reset report", "Calibrated simulator parameters",
        ),
        (
            "First checkpoint's required expert families pass numerics and service demand plus reserve.",
            "Useful local bandwidth and staging bottlenecks are measured, not inferred from interface rate.",
            "Supported sustained envelope is bounded by power/thermal and failure evidence.",
            "SW-02, SW-04 and POC load models consume the measured service curves.",
        ),
        (
            "Small-batch service misses demand: optimize layouts/kernels/tiling, increase grouping, retain hot experts on GPU, or narrow to prefill/offline.",
            "Staging bottleneck dominates: change memory movement architecture before claiming HBF value.",
            "Thermal throttle breaks rate: lower population/power, redesign cooling or reject the server envelope.",
        ),
        "NPU/HBF function is GPU-vendor neutral, but service curves must be replayed under each platform's actual traffic, thermal and power pattern.",
        "Requires SW-01/02 and silicon/partner evidence. Feeds SW-04/07/10, POC 1, POC 3/4 and PR-01/02.",
    ),
    WorkPackageDetail(
        "HW-02", "PCIe endpoint and DMA qualification", "V/M", "Sandisk silicon / firmware",
        "Freeze and prove the ATMOS PCIe endpoint, addressability, DMA descriptors, ordering, completion, "
        "interrupts, errors, isolation, reset, and recovery before adding GPU/platform-specific routes.",
        (
            "Endpoint/FW revision, approved lab host, kernel/driver, IOMMU policy and hardware safety procedure are available.",
            "SW-07 UAPI/queue/buffer state machine and expected address/queue/telemetry requirements are reviewed.",
            "A direct-attached endpoint fixture exists so switch behavior is not in the first debug path.",
        ),
        (
            "Document PCIe generation/width, IDs/class, BARs/apertures, address widths, DMA masks, MSI-X vectors, queue count/depth, ATS/PASID/PRI capability if any, FLR/reset, AER/DPC and firmware-update behavior.",
            "Enumerate on a known root. Capture lspci/config space, kernel/IOMMU logs, NUMA/root path, BAR allocation and negotiated link. Compare every field with the capability manifest.",
            "Validate bounded BAR/MMIO access and register behavior. Test invalid offsets/widths, privilege, concurrent access, device removal/reset and firmware mismatch.",
            "Validate buffer registration for aligned/unaligned sizes, page sizes, scatter-gather limits, address boundaries, IOMMU on, process exit and memory pressure. Reject unsupported addresses before submission.",
            "Validate DMA host-to-device, device-to-host and bidirectional concurrency across representative activation/result sizes. Record logical/physical bytes, latency, throughput, CPU use and link counters.",
            "Validate descriptor ordering, dependencies, memory visibility/fences, queue wrap, MSI-X/polling, completion sequence/status, duplicate/lost completion detection, cancel and timeout.",
            "Run all queues/endpoints active with low/target/burst/saturation load. Prove bounded backpressure and fairness; capture oldest-work age, not only aggregate bandwidth.",
            "Inject malformed descriptors, out-of-range DMA, stale/unregistered buffers, wrong owner, poisoned/error completion, link degradation, AER, queue hang and firmware error using approved mechanisms.",
            "Reset during each operation state: registered idle, queued, DMA-in, executing, DMA-out and completing. Verify all registrations/queues/readiness invalidate and no stale write occurs after memory reuse.",
            "Repeat across supported cold/warm boot, FLR, secondary bus/switch reset where applicable, kernel reload, driver/FW update and host reboot. Publish exact rejoin prerequisites.",
            "Run soak and race/stress tests; audit errors/leaks and compare driver, firmware and hardware counters. Sign the endpoint qualification report for the exact configuration.",
        ),
        (
            "# READ-ONLY TEMPLATE: lspci -vv -s <BDF> > evidence/.../environment/lspci.txt",
            "# READ-ONLY TEMPLATE: dmesg | grep -Ei 'pcie|aer|iommu|atmos' > evidence/.../environment/kernel-pcie.txt",
            "ctest --test-dir data-plane/build -N  # Run --output-on-failure only when tests are registered; current builds may lack GoogleTest",
            "# HARDWARE TEMPLATE: atmos-dma-bench --device <BDF> --sizes 1024,4096,16384,65536,262144 --queues 1,4,8 --json evidence/.../raw/dma.json",
        ),
        (
            ("Enumeration/config", "Compare PCIe/config/BAR/link/IOMMU inventory", "Manifest exact; unsupported root/config fails clearly"),
            ("DMA correctness", "Pattern/checksum both directions and concurrency", "No corruption, stale access, lost/duplicate completion"),
            ("Ordering/visibility", "Dependency/fence and CPU/device visibility tests", "Documented memory model passes every queue mode"),
            ("Isolation", "Out-of-range/stale/cross-owner and IOMMU tests", "No unauthorized DMA/BAR access"),
            ("Reset", "Reset in every operation state and memory reuse", "No late write; all state invalidated; bounded recovery"),
            ("Performance/soak", "Message/load sweep and long run", "Qualified envelope, errors and queue behavior stable"),
        ),
        (
            "Endpoint capability/configuration manifest", "DMA/ordering/completion specification", "Direct-attach qualification dataset",
            "Security/isolation/fault/reset matrix", "Performance/soak report", "Driver/FW rejoin/update requirements",
        ),
        (
            "Direct endpoint passes function, isolation, reset and required load before switch qualification.",
            "All operation states and ambiguous failure behavior are explicit to SW-07/SW-05.",
            "Exact endpoint/FW/driver/host configuration enters PR-01.",
        ),
        (
            "Address/BAR window cannot fit: engage OEM/BIOS before switch work; do not assume another server fixes it.",
            "Ordering/completion ambiguous: block higher POCs until protocol and retry safety are fixed.",
            "IOMMU/isolation conflicts with performance: preserve security, optimize supported route, or reject peer/direct path.",
        ),
        "Same endpoint contract; every root complex, host CPU architecture, IOMMU/BIOS and peer path is requalified per NVIDIA/AMD OEM configuration.",
        "Requires SW-07 and HW-01. Feeds HW-03/05/06, POC 1/2/3/4 and PR-01.",
    ),
    WorkPackageDetail(
        "HW-03", "GPU-to-ATMOS transport qualification", "V/M", "Sandisk runtime/driver + OEM",
        "Prove the mandatory host-staged route and any optional documented GPU peer route end-to-end, "
        "including topology, memory registration, synchronization, faults, security and complete continuation latency.",
        (
            "HW-02 endpoint and SW-05/SW-07 buffer/transport contracts pass direct tests.",
            "Pinned CUDA or ROCm platform, GPU, driver, kernel, IOMMU/ACS policy and exact root/NUMA map are available.",
            "The product does not depend on proprietary GPU driver/firmware patches; peer mode has a documented vendor/OEM route or remains disabled.",
        ),
        (
            "Create the route inventory: GPU BDF/rank, ATMOS BDF, CPU/root/NUMA, switch/retimer/cable, PCIe generation/width, ACS/IOMMU, memory type, API, and every root/switch/rack crossing.",
            "Implement and qualify host staging first: GPU-to-pinned-host, ATMOS DMA-in, NPU work, ATMOS DMA-out, host-to-GPU, dependencies and continuation. Use persistent registered pools and explicit stream/event synchronization.",
            "Measure each stage and complete producer-to-expert-to-continuation for real hidden/result sizes, top-k, tokens/expert, warm/loaded queues, single/concurrent GPU and low/target/burst/saturation.",
            "Validate CPU/NUMA affinity, pinned allocation locality, memory bandwidth/CPU use and cross-root behavior. Record actual payload route; root traversal does not automatically mean an extra application copy.",
            "For NVIDIA peer inquiry, verify documented peer-memory/GPUDirect requirements, topology, memory registration, DMA addressability, callbacks/invalidation, synchronization, security and OEM support on the exact configuration.",
            "For AMD peer inquiry, verify documented ROCm/HIP/amdgpu/DMA-BUF or approved mechanism, IOMMU/topology, registration/invalidation, synchronization, security and OEM support on the exact configuration.",
            "Implement peer as a transport backend with the same operation/status semantics. Never silently select peer; expose route in the manifest and retain host-staged fallback only when its SLO behavior is qualified.",
            "Inject cancellation, process exit, GPU reset, ATMOS reset, switch/root error, memory invalidation, timeout, stale mapping, and concurrent tenant attempts. Prove no stale DMA and safe ownership callbacks.",
            "Compare routes at identical model/quality/SLO/load. Report complete latency, p50/p95/p99, queueing, bytes, CPU/memory use, errors and power; link bandwidth alone cannot select the product route.",
            "Publish a measured route matrix listing supported and unsupported GPU/ATMOS pairs, payload/latency envelope, root/rack crossings, required settings and first-response owner.",
        ),
        (
            "data-plane/scripts/run_benchmark.sh --sizes 1024,4096,16384,65536,262144 --warmup 20 --iterations 200",
            "# NVIDIA TEMPLATE: nvidia-smi topo -m > evidence/.../environment/nvidia-topology.txt",
            "# AMD TEMPLATE: rocm-smi --showtopo > evidence/.../environment/amd-topology.txt",
            "# TEMPLATE: atmos-roundtrip-bench --transport host-staged,peer-if-supported --trace <router-trace> --json evidence/.../raw/routes.json",
        ),
        (
            ("Host-staged correctness", "Full activation/expert/result/continuation vectors", "Mandatory baseline passes all sizes/top-k/faults"),
            ("Route proof", "Topology, API logs and endpoint/GPU counters", "Actual payload path and crossings are evidenced"),
            ("Peer support", "Vendor/OEM-documented registration and invalidation", "Exact configuration supported; no private patch/security bypass"),
            ("Complete latency", "Real trace/load sweep", "Meets POC budget including synchronization/fan-in/continuation"),
            ("Fault/ownership", "Reset/exit/invalidation/cross-tenant", "No stale DMA or memory reuse ambiguity"),
        ),
        (
            "Host-staged implementation and qualification", "Optional NVIDIA and AMD peer backend results",
            "Measured route matrix", "Topology/NUMA maps", "Complete latency/fault/security report", "Fallback/selection policy",
        ),
        (
            "Host staging is always functionally supported for released configurations.",
            "Any promoted peer route is documented, measured, secure and configuration-specific.",
            "Complete round-trip, not copy bandwidth, satisfies the POC latency decision.",
            "Unsupported combinations are explicit in PR-01 and fail early.",
        ),
        (
            "Host staging misses budget and peer unavailable: pivot batching/execution boundary or reject platform; do not modify proprietary stack.",
            "Peer works only with weakened ACS/IOMMU: reject peer and retain secure baseline.",
            "Cross-root tail dominates: change physical placement/OEM route or companion topology.",
        ),
        "Separate CUDA and ROCm transport adapters and route matrices; common wire/status/buffer ownership semantics.",
        "Requires SW-05/07/10, HW-02/05/06. Feeds POC 2/3/4 and PR-01/02.",
    ),
    WorkPackageDetail(
        "HW-04", "Native GPU scale-up fabric invariant", "V", "OEM / GPU platform",
        "Preserve the platform's native GPU-to-GPU fabric and verify ATMOS neither joins nor destabilizes "
        "NVLink/NVSwitch, UALink or the existing GPU collective/network path.",
        (
            "GPU-only platform topology, firmware, NCCL/RCCL and workload baseline are captured.",
            "ATMOS physical route and SW-06 collective invariant are defined.",
            "OEM/platform owner confirms supported instrumentation and fault-test boundaries.",
        ),
        (
            "Inventory native GPU fabric: GPU/rank/tray groups, NVLink/NVSwitch or UALink/Infinity links, bandwidth/error counters, collective topology, host connections, NIC/DPU paths, firmware and management ownership.",
            "State the boundary in architecture and manifest: ATMOS is a PCIe endpoint/service branch, never advertised as a native GPU-fabric endpoint or GPU rank.",
            "Run GPU-only bandwidth/latency/collective/model controls and archive topology/counter/firmware data at idle, target and saturation load.",
            "Install/enable ATMOS without offload and repeat controls to isolate physical/firmware/resource side effects such as lane sharing, address windows, NUMA, power/cooling or NIC displacement.",
            "Run hybrid workload and compare native GPU fabric traffic, collective latency/order, errors, link state, rank groups and network behavior at equal accepted load/SLO.",
            "Exercise declared ATMOS endpoint/switch failure and verify GPU fabric remains healthy; then exercise approved GPU fabric/rank fault and verify ATMOS service follows the declared application policy without collective deadlock.",
            "Check BIOS/BMC/platform updates and maintenance: ATMOS addition must not require private GPU firmware or unowned fabric configuration.",
            "Publish per-platform invariant evidence and any resource displacement/coupling that changes baseline economics or support.",
        ),
        (
            "# NVIDIA TEMPLATE: nvidia-smi topo -m && NCCL_DEBUG=INFO <collective-test>",
            "# AMD TEMPLATE: rocm-smi --showtopo && NCCL_DEBUG=INFO <rccl-test>",
        ),
        (
            ("Fabric identity", "Topology/rank/device inventory", "No ATMOS endpoint/rank in native fabric"),
            ("No-regression", "GPU-only before/after install and hybrid at equal workload", "Fabric/collective errors and performance within signed tolerance"),
            ("Failure isolation", "ATMOS domain and native GPU domain faults", "No unintended cross-domain reset/deadlock"),
            ("Lifecycle", "BIOS/BMC/GPU stack maintenance", "Stock vendor-supported ownership preserved"),
        ),
        (
            "NVIDIA NVLink/NVSwitch invariant report", "AMD UALink/Infinity invariant report",
            "Before/after topology and performance controls", "Failure-coupling and resource-displacement record",
        ),
        (
            "Stock GPU fabric and collective stack pass without modification.",
            "ATMOS failure does not reset/corrupt the native GPU domain outside the approved boundary.",
            "Every platform delta is reflected in PR-01 and PR-02.",
        ),
        (
            "ATMOS consumes/changes a critical GPU/NIC path: redesign lane/platform integration or reject configuration.",
            "Private GPU firmware/collective change required: stop that first-product path.",
            "Power/cooling coupling throttles GPU fabric: reduce population/redesign cooling or use companion tray.",
        ),
        "NVIDIA and AMD are entirely separate qualification records. Preserve stock NCCL/NVLink/NVSwitch versus RCCL/UALink/Infinity ownership.",
        "Requires SW-06, HW-03/06. Feeds POC 3/4 and PR-01/02.",
    ),
    WorkPackageDetail(
        "HW-05", "PCIe switch and backplane subsystem", "V/M", "OEM + switch FAE",
        "Qualify one exact switch/board/backplane/cable/retimer/firmware subsystem for the required port "
        "map, sustained traffic, security, diagnostics, containment, reset and lifecycle.",
        (
            "HW-02 direct endpoint passes and required endpoint count/link mode are fixed.",
            "OEM returns candidate upstream/root connector/lane map and physical E3.S path.",
            "Switch vendor/board FAE confirms generic non-NVMe endpoint inquiry, board access, firmware/debug and support route.",
        ),
        (
            "Freeze requirements before vendor selection: x16 upstream + target downstream x4 ports, generations, endpoint count, oversubscription, board form, cables/retimers, power/thermal, ACS/IOMMU, DPC/AER, hot-plug/reset, diagnostics, firmware, security, lifecycle and support.",
            "Obtain written candidate maps from Microchip PM50052-class and Broadcom PEX89048-class or OEM-selected alternatives. Treat lane count as necessary, not sufficient.",
            "Prefer an orderable reference/AIC/board fixture. Design a custom switch card only after an existing board fails a documented requirement and NRE/lifecycle ownership is approved.",
            "Build the exact subsystem BOM/configuration: switch/board revision, firmware profile, upstream/downstream map, connectors, backplane, MCIO/cables, retimers, E3.S slots, clock/reset/power and debug tools.",
            "Enumerate all endpoints. Capture port/link state, negotiated rate/width, BAR/address windows, ACS/IOMMU, NUMA/root, switch diagnostics and unsupported combinations.",
            "Run one and all endpoints active for host-bound and any approved peer-bound patterns. Sweep message sizes/queues/directions, independent and synchronized traffic, load/update overlap and upstream saturation; measure fairness and head-of-line blocking.",
            "Validate security/routing: ACS settings, peer blocking/allow rules, IOMMU, address isolation, switch management access and firmware authenticity. Never disable containment/isolation merely to obtain peer traffic.",
            "Inject endpoint/link failure, malformed completion, AER/DPC, downstream FLR, hot reset, switch reset/firmware failure, upstream loss and re-enumeration. Measure blast radius, unaffected traffic, lost/ambiguous work and recovery.",
            "Run cable/retimer signal-integrity and error-counter study across supported environmental/thermal conditions. Qualify only the exact cable/channel lengths and revisions tested.",
            "Validate switch FW update/rollback, diagnostics bundle, secure management, BMC integration, field replacement, spares, FAE escalation, lead time and EOL notice.",
            "Compare one eight-device domain with two balanced four-device domains on lanes, throughput, tail, NUMA, power/cost and availability. Record the production decision separately from the lab fixture.",
            "Publish qualification package and configuration ID. Requalify affected evidence whenever switch, FW, board, cable, retimer, root or endpoint generation changes.",
        ),
        (
            "ctest --test-dir atmos-sim/build -R topology --output-on-failure",
            "# READ-ONLY TEMPLATE: lspci -tv > evidence/.../environment/pcie-tree.txt",
            "# TEMPLATE: <vendor-switch-tool> inventory --json > evidence/.../raw/switch-inventory.json",
            "# TEMPLATE: atmos-dma-bench --all-devices --queues 1,4,8 --traffic independent,correlated --duration 900 --json evidence/.../raw/switch-load.json",
        ),
        (
            ("Port/BOM", "Physical/config/firmware inventory", "Exact supported map and every component/revision identified"),
            ("All-device load", "Direction/message/queue/correlation sweep", "Target plus reserve passes with bounded fairness/tail"),
            ("Security", "ACS/IOMMU/routing/management negative tests", "Isolation enabled; only approved routes pass"),
            ("Containment/reset", "Endpoint/link/switch/upstream fault matrix", "Blast radius and recovery satisfy policy; no silent corruption"),
            ("SI/thermal", "Error counters/channel/temperature sustained test", "Qualified channel has bounded errors/no hidden downtrain"),
            ("Lifecycle/support", "FW/update/diagnostics/spares/escalation review", "One accountable support route and supply plan"),
        ),
        (
            "Qualified subsystem BOM/configuration ID", "Lane/port/NUMA/route map", "Traffic/fairness/tail report",
            "Security/ACS/IOMMU/DPC matrix", "Fault/reset/recovery report", "SI/thermal qualification",
            "Firmware/diagnostics/update/support/supply contract", "Single versus dual-domain decision",
        ),
        (
            "One exact switch subsystem passes function, load, security, faults, lifecycle and support gates.",
            "All unsupported routes/configurations and shared failure domains are explicit.",
            "Subsystem evidence feeds OEM design record and PR-01/02.",
        ),
        (
            "No candidate board supports generic endpoint/map: document requirement gap and approve custom NRE or redesign topology.",
            "One-switch blast radius/oversubscription fails: use two four-device domains or direct paths.",
            "FW/support/supply owner absent: lab use only; cannot promote to product.",
        ),
        "Common downstream ATMOS subsystem is desirable, but upstream root/NUMA, board/channel, host architecture and OEM support differ per NVIDIA/AMD platform.",
        "Requires HW-02/03/06 and SW-07/10. Core of POC 2; feeds POC 3/4 and PR-01/02.",
    ),
    WorkPackageDetail(
        "HW-06", "Server and rack integration envelope", "V/M", "OEM engineering",
        "Close the exact server/rack electrical, mechanical, lane, storage/NIC, power, cooling, BMC, "
        "serviceability, fault-coupling, support, supply and qualification boundary.",
        (
            "ATMOS device envelope and qualified endpoint/switch options are available.",
            "One exact OEM host/tray/rack candidate and engineering/support owners are named.",
            "The workload/SLO, device population, failure policy and native GPU/network/storage baseline are frozen for comparison.",
        ),
        (
            "Assign a configuration ID and inventory exact server/tray/rack, motherboard/baseboard, CPUs, GPUs, NICs/DPUs, boot/cache/storage, risers, PCIe switches, BMC/BIOS, PSUs/PDBs, fans/CDUs/manifolds and firmware/software revisions.",
            "Return a complete lane/NUMA/root map: native GPU fabric boundary, host CPU/C2C, ATMOS upstream/downstream, endpoints, generation/width, cables/retimers, NIC/storage paths, oversubscription and every root/rack crossing.",
            "Close mechanics: E3.S 2T positions, keep-outs, retention, connector/cable routing, insertion/removal, service access and what storage/NIC/other component is displaced. A printed bay count is not compatibility.",
            "Close electrical/power: per-device/switch budget, connector/PDB current, PSU/rack/PDU reserve, transient behavior, sequencing and simultaneous GPU+ATMOS+network worst case.",
            "Close thermal: airflow/cold plate/heatsink, inlet/coolant, pressure/flow, fan/CDU response, sensors, BMC policy, throttling/shutdown and adjacent GPU/NIC/storage interaction.",
            "Close BIOS/OS/security: enumeration/address windows, above-4G/MMIO, IOMMU/ACS/ATS/PASID, AER/DPC, secure boot, firmware update, device assignment and kernel/distro support.",
            "Define reset/failure domains and recovery: endpoint FLR, switch/downstream/upstream reset, root/host reboot, BMC action, GPU/native fabric effects, lost work, re-enumeration and service restart/rejoin.",
            "Define field service: safe drain/power/removal, FRU, replacement, pack/firmware restore, diagnostics, spares, repair time, escalation and warranty boundaries across Sandisk/OEM/GPU/switch/NPU vendors.",
            "Run platform tests under simultaneous target workload: enumeration, all-device DMA, GPU fabric/network/storage baseline, power/thermal steady/burst, throttle, endpoint/switch faults, service replacement and firmware/BIOS/driver update/rollback.",
            "Compare local integrated and approved companion-tray paths where local fit fails. Include cable latency/SI, independent power/cooling, cross-tray reset and rack service/support.",
            "Obtain prototype/production BOM, quote, NRE, qualification cost, MOQ, lead time, spares, support SKU, geographies, lifecycle/EOL and lab/production access terms.",
            "Publish the OEM design record and unresolved-owner list. No platform advances to product POC with an unowned lane, reset, cooling, security, service or support requirement.",
        ),
        (
            "# READ-ONLY TEMPLATE: lspci -tv > evidence/.../environment/pcie-tree.txt",
            "# NVIDIA TEMPLATE: nvidia-smi topo -m > evidence/.../environment/gpu-topology.txt",
            "# AMD TEMPLATE: rocm-smi --showtopo > evidence/.../environment/gpu-topology.txt",
            "# OEM TEMPLATE: <bmc-tool> sensor export --json > evidence/.../power-thermal/bmc.json",
        ),
        (
            ("Configuration identity", "BOM/inventory/firmware/software reconciliation", "Every component/revision and support owner recorded"),
            ("Lane/route", "Topology tools plus physical schematic/design record", "All paths/crossings/resources and unsupported routes explicit"),
            ("Power/thermal", "Worst representative simultaneous load and transients", "No unsupported throttle/safety/facility violation"),
            ("Platform regression", "GPU/network/storage controls before/after/hybrid", "Native services preserved within signed tolerance"),
            ("Fault/service", "Declared failures and physical replacement drill", "Containment, RTO, complete service and runbook pass"),
            ("Commercial/support", "Quote/NRE/supply/support review", "Accountable lifecycle and feasible economics"),
        ),
        (
            "Qualified BOM/configuration ID", "Lane/NUMA/root and native-fabric boundary map", "Mechanical/service design",
            "Power/cooling/BMC envelope", "BIOS/security/reset/failure contract", "Platform regression/fault/replacement results",
            "Support/spares/escalation and supply plan", "Prototype/production quote and NRE", "Local versus companion decision",
        ),
        (
            "One exact platform passes physical, electrical, thermal, firmware, service, support and workload tests.",
            "Displaced storage/NIC and every incremental cost/resource are included in PR-02.",
            "Configuration-specific manifest and OEM design record are signed.",
        ),
        (
            "No local mechanical/thermal/lane fit: evaluate OEM-approved companion tray; do not insert into storage bays by assumption.",
            "Native GPU/NIC/storage regression: redesign route/population or reject platform.",
            "No field support/warranty owner: engineering fixture only, not product.",
        ),
        "NVIDIA NVL/HGX/MGX and AMD OAM/Helios/PCIe systems require separate design records and qualification. Do not transfer RTX or another OEM result.",
        "Requires all hardware and operational packages. Core of POC 2 physical record and POC 4 product qualification; feeds PR-01/02.",
    ),
]


PRODUCT_PACKAGES = [
    WorkPackageDetail(
        "PR-01", "Compatibility manifest and release package", "N", "Sandisk release",
        "Bind one supported hardware, firmware, software, model, topology, workload and operational "
        "combination into a machine-readable, signed, reproducible release identity.",
        (
            "All included work packages have signed definitions of done for the target configuration.",
            "OEM design record, package repositories, test evidence, known limits and support ownership exist.",
            "Release, security, legal/open-source, product and support approvers are named.",
        ),
        (
            "Define the manifest schema and compatibility rules: release/config ID, hardware/FW, endpoint/switch/board/cables, OEM/BIOS/BMC, host/GPU/network, OS/kernel/driver, container, vLLM/PyTorch/CUDA/ROCm, Kubernetes/operator, model/checkpoint/tokenizer, ExpertPack/placement, topology/failure policy, workload/SLO envelope and known limits.",
            "Ingest package outputs by immutable identifier and checksum. Reject latest/unversioned dependencies, missing support owner, incomplete evidence or ambiguous hardware family names.",
            "Encode positive and negative compatibility: supported exact tuples, optional capabilities such as peer DMA, excluded features/topologies, maximum populations and required settings.",
            "Attach evidence IDs for function, quality, service, fault, security, power/thermal, lifecycle, support and economics. State evidence class and expiration/review trigger.",
            "Generate deployable release artifacts: driver/FW/SDK, containers, packs, operator/Kubernetes manifests, configuration, diagnostics, dashboards, SBOMs, signatures, docs, runbooks and support contacts.",
            "Implement preflight validation on host/cluster: inventory actual state, compare to manifest, block unsafe/unsupported drift, and produce a clear remediation report.",
            "Run clean-room installation on a fresh supported platform by an engineer/operator who did not build it. Execute smoke, quality, load, fault, update, rollback and support-bundle tests.",
            "Run upgrade from the prior supported manifest and rollback after a failed canary. Verify pack/placement generations, driver/FW compatibility and workload continuity/declared downtime.",
            "Sign, publish and retain the release; issue release notes, known limitations, security/support duration and EOL policy. Archive superseded manifests and approved claims.",
            "Define change-impact rules: model, precision, device/FW, switch, OEM/BIOS, driver/kernel, serving stack, topology, failure policy or SLO changes identify exactly which evidence must be repeated.",
        ),
        (
            "# TEMPLATE: atmos-manifest validate manifest.yaml --require-evidence --require-support-owner",
            "# TEMPLATE: atmos-preflight --manifest manifest.yaml --json evidence/.../environment/preflight.json",
            "shasum -a 256 <release-artifacts> > checksums.sha256",
        ),
        (
            ("Schema/completeness", "Machine validation and owner/evidence checks", "No floating/missing dependency or unsupported implicit value"),
            ("Clean install", "Independent fresh platform install and smoke/full checks", "Another operator reproduces release from package"),
            ("Drift rejection", "Change supported/unsupported fields", "Preflight blocks and identifies remediation"),
            ("Upgrade/rollback", "Prior-to-new and failed-canary rollback", "Declared continuity/downtime and generations remain safe"),
            ("Claims", "Map every approved statement to evidence/config ID", "No platform-family or transferred economic overclaim"),
        ),
        (
            "Machine-readable compatibility manifest/schema/validator", "Signed release bundle and checksums/SBOM",
            "Preflight/drift tool", "Clean-install and upgrade/rollback evidence", "Release notes/limits/support/EOL",
            "Evidence-to-claim and change-impact matrix",
        ),
        (
            "A new operator installs and operates the exact release from the package.",
            "Unsupported drift fails before workload readiness.",
            "All artifacts, evidence, support and claims resolve from one configuration ID.",
        ),
        (
            "Manifest cannot bind an unversioned dependency: pin it or remove it from supported release.",
            "Clean install depends on developer knowledge: improve packaging/runbook until independent reproduction passes.",
            "Evidence invalidated by change: mark release unsupported until affected qualification reruns.",
        ),
        "Separate manifests for RTX development, each NVIDIA rack and each AMD rack. Common schemas/API do not merge qualification evidence.",
        "Consumes all work packages. Mandatory output of POC 3 and final release artifact of POC 4.",
    ),
    WorkPackageDetail(
        "PR-02", "Economic and support contract", "N", "Product + solution architecture",
        "Prove platform-specific value at equal quality, SLO and availability, and define one accountable "
        "customer support, supply, escalation and lifecycle path.",
        (
            "SW-01 workload/quality/SLO, measured POC service rates, HW-06 quotes/BOM, failure reserve and PR-01 configuration ID are available.",
            "Finance/product has approved the cost horizon, discount/energy/facility assumptions and hurdle before hybrid results are reviewed.",
            "OEM/Sandisk/GPU/switch/NPU/software support and supply owners have returned terms or explicit gaps.",
        ),
        (
            "Freeze equal-service comparison: exact model/precision/quality, prompt/output mix, arrivals, TTFT/TPOT, accepted-work definition, availability/failure reserve, tuning process, measurement window and demand scenarios.",
            "Measure GPU-only and hybrid offered/admitted/accepted/rejected/failed/out-of-SLO work, TTFT/TPOT, output tokens, utilization, GPU/ATMOS/network/rack occupancy, power/cooling and degraded/failure rate with confidence intervals.",
            "Build the fully burdened three-year control and hybrid cost model: GPUs/hosts/ATMOS/switch/cables/backplane, displaced storage/NIC, racks/network/optics, power/cooling/facility, licenses/support, spares/reserve/repair/downtime, operations, NRE/qualification, migration and dual-running.",
            "Calculate TCO per accepted request and output token, required rate uplift, complete group count for demand, ROI/NPV/payback as approved, and sensitivity. Keep quotes, measurements, contracts and assumptions visibly distinct.",
            "Sweep accepted rate, configured ATMOS cost, GPU/rack baseline, demand/utilization, failure policy, energy/PUE, support/operations, NRE allocation and next GPU generation. Show the feasible region, not one favorable cell.",
            "Credit only removable cost or useful capacity: fewer complete GPU groups/racks, more accepted work within an existing constraint, avoidable network ports/racks, measured energy/facility reduction, or measured operational benefit.",
            "Define the support RACI: customer intake, first response, triage/diagnostics, Sandisk endpoint/runtime, OEM platform, GPU stack, switch, NPU/compiler, Kubernetes/distro, spares/RMA, security incidents and escalation/SLA.",
            "Define supply/lifecycle: production BOM/quote, MOQ, lead time, forecast, spares, geographic availability, firmware/software support duration, EOL notice, qualification of substitutions and warranty.",
            "Run a support drill from customer symptom through bundle collection, owner routing, reproduction, workaround/fix, communication and closure. Measure handoff time and unresolved ownership.",
            "Publish only configuration-specific claims tied to PR-01 and POC evidence. Separate feasibility, funded POC, design win, supported release, purchase commitment and recognized revenue.",
            "Return a product decision: release, conditional release with dated closures, narrower workload/platform, or no-fit. Preserve negative economics/support outcomes.",
        ),
        (
            "# TEMPLATE: python tools/tco_model.py --control control.yaml --hybrid hybrid.yaml --demand demand.yaml --out evidence/.../economics/",
            "# TEMPLATE: python tools/check_claims.py --manifest manifest.yaml --claims approved-claims.yaml --evidence evidence-index.yaml",
        ),
        (
            ("Equal service", "Reconcile model/quality/SLO/HA/workload/tuning", "No comparison mismatch or transferred rate"),
            ("Rate/capacity", "Repeated accepted-load curves incl. failure", "Qualified rate and reserve are measured with confidence"),
            ("Cost completeness", "BOM/quote/operations/NRE/support review", "All incremental/displaced/availability costs included"),
            ("Sensitivity", "Approved low/base/high and future-platform sweep", "Positive region satisfies sponsor hurdle, or no-fit recorded"),
            ("Support", "RACI and incident/escalation drill", "One first-response route; no unowned critical handoff"),
            ("Claims", "Evidence-to-claim audit", "Every statement is configuration-specific and supportable"),
        ),
        (
            "Equal-service measurement contract", "Quote-backed TCO/ROI workbook and sensitivity", "Demand/capacity plan",
            "Support RACI/SLA/escalation and drill result", "Supply/lifecycle/spares plan", "Approved claims and product decision",
        ),
        (
            "Measured named-platform result exceeds the predeclared hurdle at equal quality/SLO/HA.",
            "Configured quote, reserve, NRE, support and supply are included.",
            "One accountable support route and lifecycle exists.",
            "Claims are approved only for the exact PR-01 configuration.",
        ),
        (
            "Economics fail: optimize rate/cost/reserve, choose another OEM/workload, narrow to strategic capacity benefit, or stop.",
            "Support ownership fails: engineering fixture only; no customer product claim.",
            "Future GPU control erases value: rebaseline before launch; do not preserve obsolete savings.",
        ),
        "RTX economics never transfer. Build separate quote/control/hybrid/reserve/support models for each NVIDIA and AMD platform/OEM.",
        "Requires PR-01 and measured POC 3/4 outputs. Final business and support gate for POC 4.",
    ),
]


@dataclass(frozen=True)
class PocDetail:
    identifier: str
    title: str
    decision: str
    purpose: str
    packages: str
    entry: tuple[str, ...]
    setup: tuple[tuple[str, str], ...]
    steps: tuple[str, ...]
    matrix: tuple[tuple[str, str, str], ...]
    commands: tuple[str, ...]
    evidence: tuple[tuple[str, str], ...]
    pass_criteria: tuple[str, ...]
    failure: tuple[str, ...]
    gate_package: tuple[str, ...]


POCS = [
    PocDetail(
        "POC 1", "Vendor-neutral expert primitive", "Can ATMOS execute real routed experts correctly and fast enough behind a stable ABI?",
        "Prove model semantics, ExpertPack compilation, one-device execution, small-batch service curves, "
        "endpoint/runtime behavior, and an early complete GPU-to-ATMOS-to-GPU continuation before switch investment.",
        "SW-01, SW-02, SW-04 (identity/static placement), SW-05 (single path), SW-07, SW-10, HW-01, HW-02",
        (
            "POC charter names sponsor, decision owner, first checkpoint, precision, operator scope, evidence class, schedule and gate reviewers.",
            "SW-01 GPU oracle, model contract, router trace, quality tolerance and provisional TTFT/TPOT workload contract are signed.",
            "SW-02 smoke compiler and SW-07 SDK/driver API are versioned; one direct endpoint or explicit partner emulator is available.",
            "Approved lab safety, firmware/driver, cooling/power, data/IP and evidence-retention procedures are in place.",
            "Analytical per-layer expert budget is calculated from the GPU control and labeled provisional until end-to-end replay.",
        ),
        (
            ("Model", "One Qwen3-235B-A22B-class checkpoint first; record exact variant/hash/license. Select representative first/middle/last MoE layers and hot/cold experts."),
            ("Precision", "One approved FP8 weight/activation/accumulation recipe. BF16 or other formats are expansion cells, not first-gate blockers unless required by quality."),
            ("Execution target", "Production-intent ATMOS device preferred; otherwise versioned NPU functional/performance emulator with timing uncertainty. Direct PCIe endpoint for physical DMA evidence."),
            ("GPU control", "One pinned NVIDIA/CUDA reference first; AMD/ROCm repeats later. Host-staged registered buffers are the mandatory round-trip baseline."),
            ("Instrumentation", "SW-10 run IDs and synchronized/offset-calibrated GPU/CPU/driver/NPU events; power/thermal sensors; raw outputs and counters."),
            ("Stop conditions", "Thermal/power safety threshold, repeated data corruption, unbounded queue, stale DMA, firmware fatal, or unsupported operator/precision."),
        ),
        (
            "Create POC1 configuration/run manifests and evidence directories. Record every hardware/software/tool revision and unresolved assumption before execution.",
            "Reproduce the GPU-only expert oracle and selected end-task sample three times. Archive golden hidden states, routed IDs/weights, expert outputs, quality and timing.",
            "Compile one smoke ExpertPack per shape family. Verify schema, signature, memory/scratch budget and clean-build reproducibility; reject incomplete or wrong-model generations.",
            "Bring up the ATMOS target in direct mode. Run discovery, self-test, capability comparison, buffer register/unregister, queue create/destroy, known-pattern DMA and reset/rejoin smoke tests.",
            "Execute one expert with one token and one queue. Compare every output element/status/identity to the GPU oracle before enabling asynchronous or performance paths.",
            "Expand correctness across first/middle/last layers, hot/cold experts, tokens/expert 1-64, boundary values, selected precision, queue depths and cold/warm pack state.",
            "Compile/load the complete first-checkpoint expert generation or an approved representative subset if emulator capacity is limited. Publish exact coverage and what remains provisional.",
            "Run one-device service curves: low/target/burst/saturation arrivals, queue depths, concurrent experts and thermal steady state. Capture useful HBF, physical staging traffic, NPU stalls/utilization and p50/p95/p99.",
            "Run negative/lifecycle tests: unsupported/corrupt/wrong pack, malformed descriptor, buffer bounds, queue exhaustion, timeout/cancel, firmware mismatch, reset during each operation state and process exit.",
            "Calibrate atmos-sim with measured/emulated service distributions, bytes, queues and uncertainty; replay the approved router trace to determine target demand plus burst reserve.",
            "Add the early complete continuation test: GPU router output -> pack/group -> host-staged transfer -> ATMOS expert -> return -> reorder/gate/combine -> measured GPU continuation. Use real hidden/result sizes, top-1/top-2/top-8 and warm/loaded queues.",
            "Compare complete p99 continuation with the signed expert-layer budget and replay the result through the model trace. Do not sum component p99 values; measure end-to-end directly.",
            "Run a short repeated soak, inspect errors/leaks/thermal drift and reproduce the analysis from archived raw evidence.",
            "Hold the POC1 gate review. Classify results as production-silicon or provisional emulator/prototype evidence and decide pass, conditional, pivot or stop.",
        ),
        (
            ("Expert shape", "First/middle/last MoE layer; min/median/max FFN; hot/cold expert", "Numerics, pack bytes/scratch and service curve per family"),
            ("Tokens/expert", "1, 2, 4, 8, 16, 32, 64", "p50/p95/p99, expert/s, token/s, useful HBF and NPU utilization"),
            ("Queue/load", "Depth 1/4/8/16; low/target/1.2x burst/saturation/recovery", "Queue age/depth stability, rejects, errors and saturation knee"),
            ("Pack state", "Cold load, warm resident, second generation staged", "Load/activation latency, capacity and update behavior"),
            ("Top-k/destinations", "Top-1/top-2/top-8; 1/2/4/8 logical destinations", "Complete continuation, fan-in and slowest contribution"),
            ("Fault", "Cancel, timeout, malformed pack/descriptor, reset each state, process exit", "No incorrect output/stale DMA; explicit status and bounded recovery"),
            ("Thermal", "Cold start, steady target, qualified high condition, throttle", "Supported rate/power/temperature and stop response"),
        ),
        (
            "cmake -S atmos-sim -B atmos-sim/build -DCMAKE_BUILD_TYPE=Release && cmake --build atmos-sim/build --parallel",
            "ctest --test-dir atmos-sim/build -R 'virtual_devices|model_graph|metrics|golden|e2e' --output-on-failure",
            "# TEMPLATE: expertpack verify packs/<generation> --goldens <bundle> --require-complete",
            "# TEMPLATE: atmos-npu-bench --tokens 1,2,4,8,16,32,64 --queues 1,4,8,16 --json evidence/.../raw/poc1-npu.json",
            "# TEMPLATE: atmos-roundtrip-bench --transport host-staged --top-k 1,2,8 --trace <router-trace> --json evidence/.../raw/poc1-roundtrip.json",
        ),
        (
            ("Configuration", "Manifest, checksums, environment, device/compiler/FW/driver and evidence class"),
            ("Function/quality", "Golden outputs, coverage, error distributions, end-task sample and failed vectors"),
            ("Performance", "Raw service/round-trip events, p50/p95/p99, accepted rate, queue stability and uncertainty"),
            ("Memory/compute", "Pack/capacity/scratch, useful and physical memory bytes, NPU utilization/stalls"),
            ("Reliability/security", "Negative/reset/cancel/process-exit logs, stale-access checks and recovery"),
            ("Power/thermal", "Power, temperature, frequency/throttle and sustained envelope"),
            ("Decision", "Signed budget replay, calibrated parameters, deviations, unsupported scope and gate decision"),
        ),
        (
            "All required first-checkpoint expert shape families compile and pass numerical tolerance; coverage gaps are explicitly excluded or closed.",
            "One-device service curve sustains trace-derived target demand plus approved burst reserve without unbounded queue growth.",
            "Host-staged complete continuation is correct, bounded and plausibly fits the model's signed per-layer/end-to-end TPOT budget.",
            "No negative/reset/cancel case produces silent incorrect output, stale DMA or false readiness.",
            "Evidence is production-intent or the gate is explicitly conditional pending named silicon evidence; emulator timing alone cannot unlock product claims.",
        ),
        (
            "Numerics fail: fix compiler/layout/scales/accumulation or retain unsupported experts on GPU; create a new pack/model contract before rerun.",
            "Small-batch rate/HBF utilization fails: optimize kernels/tiling/grouping, use prefill/offline/cold-expert scope, or stop the model/operator path.",
            "Complete round trip misses budget: evaluate grouped-layer/router-combine relocation or stop fine-grained FFN offload before POC2/custom switch NRE.",
            "Endpoint/reset safety fails: block all later hardware POCs until HW-02/SW-07 behavior is corrected.",
        ),
        (
            "POC1 manifest and signed charter", "Model/ABI/pack/golden and coverage bundle", "NPU service curves and capacity model",
            "Complete continuation latency/budget replay", "Negative/fault/reset report", "Calibrated simulation parameters",
            "Pass/conditional/pivot/stop decision with owners/dates and explicit POC2 spending authorization",
        ),
    ),
    PocDetail(
        "POC 2", "Switched E3.S subsystem", "Can eight E3.S devices operate as one manageable, bounded, secure and recoverable switched subsystem?",
        "Qualify one exact switch/BOM/route/security/failure/lifecycle configuration, prove all-device DMA "
        "and queue behavior, and select the production topology before full-model service work.",
        "SW-07, SW-10, HW-02, HW-03, HW-05, HW-06 (fixture subset); SW-04 trace traffic for correlated load",
        (
            "POC1 function and endpoint/reset gates pass on qualifying evidence; exact open conditions are approved.",
            "One switch/board/backplane/cable/retimer/firmware candidate and eight ATMOS endpoint revisions are inventoried.",
            "OEM/host returns upstream/root/NUMA/address-window/security/reset path for the fixture.",
            "Switch/OEM/Sandisk owners approve fault-injection, firmware, thermal and hardware safety plan.",
            "Host-staged transport is implemented; peer path is optional and separately documented.",
        ),
        (
            ("Primary fixture", "Exact x16 upstream + 8 x x4 downstream target, or the vendor-confirmed supported equivalent. Capture oversubscription and endpoint generations."),
            ("Control", "Same eight endpoints direct-attached where possible, or one direct endpoint plus calibrated analytical/simulation control; limitations explicit."),
            ("Alternative", "Two x16-to-4x4 domains for availability/NUMA comparison when host lanes and boards permit."),
            ("Security", "IOMMU/ACS/DPC/AER enabled; approved management access; no isolation bypass for peer tests."),
            ("Instrumentation", "Host/endpoint/switch/retimer/BMC counters and synchronized run IDs; cable/channel and thermal inventory."),
            ("Stop conditions", "Repeated corruption, uncontained reset/error, unsafe thermal/power/SI, unsupported FW, or stale DMA."),
        ),
        (
            "Freeze POC2 BOM/configuration ID, physical topology, firmware/settings, cable/channel, endpoint population, driver/kernel/IOMMU and evidence plan.",
            "Run preflight and direct endpoint smoke for every ATMOS device before connecting the switch; quarantine any device that cannot reproduce POC1 basics.",
            "Assemble the switched fixture under approved procedure. Capture physical cabling/retimers/power/cooling and vendor/OEM configuration backups.",
            "Enumerate all endpoints. Reconcile BDF, port, link rate/width, BAR/address windows, MSI-X, ACS/IOMMU, NUMA/root and switch diagnostics against the map.",
            "Run one endpoint at a time through every downstream port. Compare DMA correctness/latency/counters to direct control and detect bad port/channel/device combinations.",
            "Run all endpoints independently in both directions across the message/queue sweep. Measure aggregate/uplink utilization, fairness, queue age, CPU/memory, errors and per-device tail.",
            "Replay correlated MoE traffic from SW-01/SW-04: top-k fan-out, synchronized bursts, skew/hot expert, fan-in and pack-load overlap. Record head-of-line blocking and slowest destination.",
            "Test upstream saturation and recovery. Verify bounded backpressure/admission and that congestion does not corrupt or indefinitely starve one endpoint.",
            "Qualify host-staged routes for every endpoint/GPU/root combination available. If peer DMA is investigated, prove documented supported routes and leave unsupported combinations disabled.",
            "Run endpoint fault matrix: link degradation/downtrain, endpoint fatal/nonfatal error, malformed completion, queue hang, downstream FLR, device reset and removal/rejoin under traffic.",
            "Run switch fault matrix: DPC containment, AER, downstream port reset, hot/switch reset, switch firmware failure/update/rollback, upstream loss and host re-enumeration. Capture lost/ambiguous work and blast radius.",
            "Run power/thermal/SI sustained test with all devices, switch/retimers and representative GPU/host load. Capture errors, downtrain, temperature, power, throttle and fan/BMC behavior.",
            "Run a sustained soak with periodic burst and approved endpoint reset. Reconcile host, switch and endpoint errors/counters; repeat any nondeterministic fault.",
            "Where possible, repeat critical load/fault tests on two four-device domains. Compare lanes, rate, p99, failure containment, NUMA, power, cost and support.",
            "Validate switch diagnostics, support bundle, firmware lifecycle, spares, FAE escalation and replacement/recovery procedure.",
            "Publish route matrix, qualified BOM/design record and topology decision. Hold POC2 gate review before POC3 fixture lock.",
        ),
        (
            ("Message size", "1K, 4K, 16K, 64K, 256K plus real activation/result sizes", "Per-device and aggregate latency/throughput/errors"),
            ("Active devices", "1, 2, 4, 8; independent and synchronized", "Fairness, uplink utilization, queue age and p99"),
            ("Queues", "1, 4, 8 per device; low/target/burst/saturation", "Backpressure, starvation and saturation knee"),
            ("Traffic", "H2D, D2H, bidirectional, correlated top-k, pack-load + inference", "Route/fan-out/head-of-line behavior"),
            ("Route", "Direct control, switched host-staged, peer only if supported, root/NUMA local/crossing", "Measured route matrix"),
            ("Fault", "Endpoint/link/AER/DPC/FLR/switch reset/FW/upstream/re-enumeration", "Containment, lost work, stale access and RTO"),
            ("Topology", "One 8-device domain versus two 4-device domains", "Rate/tail/lanes/power/cost/failure comparison"),
        ),
        (
            "ctest --test-dir atmos-sim/build -R topology --output-on-failure",
            "# READ-ONLY TEMPLATE: lspci -tv > evidence/.../environment/poc2-tree.txt",
            "# TEMPLATE: <switch-tool> inventory --json > evidence/.../raw/poc2-switch.json",
            "# TEMPLATE: atmos-dma-bench --all-devices --sizes 1024,4096,16384,65536,262144 --queues 1,4,8 --traffic independent,correlated --json evidence/.../raw/poc2-load.json",
        ),
        (
            ("BOM/topology", "Configuration ID, photos/schematics, lane/port/NUMA map, FW/settings and checksums"),
            ("Function/routes", "Enumeration, port/device tests, host-staged/peer matrix and unsupported pairs"),
            ("Load", "Raw per-device/switch/host events, fairness, queueing, p50/p95/p99 and saturation"),
            ("Fault/reset", "Injection times, AER/DPC/status/counters, lost work, blast radius, recovery/RTO"),
            ("SI/power/thermal", "Channel revisions/errors/downtrain, temperatures, power, fans/throttle under sustained load"),
            ("Lifecycle/support", "FW update/rollback, diagnostics bundle, spares, replacement, FAE/OEM escalation"),
            ("Decision", "Single/dual domain comparison and signed production/fixture topology decision"),
        ),
        (
            "All eight endpoints enumerate and pass correct bounded DMA under simultaneous target load plus reserve.",
            "Security settings remain enabled; no stale/cross-owner access or silent corruption occurs.",
            "Endpoint failures are contained as declared; switch reset/blast radius and recovery are acceptable for the chosen availability policy.",
            "Power, thermal and SI remain within the exact qualified channel/envelope without hidden downtrain/throttle.",
            "One supported BOM, firmware/diagnostics lifecycle and accountable support owner are returned.",
        ),
        (
            "Lane map/board unsupported: use another orderable board/vendor or approve documented custom NRE; do not infer from lane count.",
            "One-switch contention/blast radius fails: split into two balanced domains or retain direct/reference lab topology.",
            "Peer route fails: keep host staging; POC2 can pass without peer if service budget remains credible.",
            "No firmware/support owner: classify subsystem as lab-only and block product promotion.",
        ),
        (
            "Qualified subsystem BOM/configuration and lane/NUMA/route maps", "All-device traffic and route matrix",
            "Security/fault/reset/recovery and soak package", "Power/thermal/SI envelope", "Firmware/diagnostics/support/supply record",
            "Single versus dual-domain decision", "Pass/conditional/pivot/stop decision and POC3 fixture authorization",
        ),
    ),
    PocDetail(
        "POC 3", "Eight-GPU plus eight-ATMOS development service", "Does the complete hybrid service create useful headroom and beat the cost-adjusted threshold at equal quality, SLO and availability?",
        "Integrate every software/hardware package on the development fixture, run the full model and workload, "
        "prove lifecycle/failure behavior, and populate the first complete-service economics without transferring it to rack scale.",
        "All SW packages, HW-01 through HW-05, HW-06 development fixture, PR-01 skeleton, PR-02 method",
        (
            "POC1 and POC2 gates pass or have dated conditions that cannot affect POC3 conclusions.",
            "Exact 8-GPU/8-ATMOS fixture, preferably two balanced four-device domains if POC2 selects it, is configured and inventoried.",
            "First checkpoint/precision/workload/quality/SLO/availability/economic hurdle are frozen before hybrid result review.",
            "Complete signed ExpertPack, placement/failure policy, vLLM integration, driver/runtime, telemetry and operational lifecycle are installable.",
            "GPU-only control has equivalent model/quality/SLO/tuning and reproducible accepted-load curves.",
        ),
        (
            ("GPU", "8 x RTX PRO 6000 Blackwell Server Edition-class development server as current reference; exact OEM, CPU, memory, GPU, NIC and driver required."),
            ("ATMOS", "8 identical qualified E3.S endpoints with complete pack bank, selected switch topology and declared reserve/fallback."),
            ("First service", "Qwen3-235B-A22B-class FP8 routed FFNs; claims/knowledge-assistant profile. Exact checkpoint and license frozen."),
            ("Starting objectives", "Planning evaluation only: TTFT <=2.5 s, TPOT <=45 ms, target demand profile frozen in charter; replace with sponsor-approved values."),
            ("Control arms", "GPU-only; hybrid host-staged; optional hybrid peer only if POC2/HW-03 qualified. Same API, model, quality, workload and HA."),
            ("Operations", "Pinned Kubernetes or host deployment, complete PR-01 draft, pack/operator rollout, monitoring, support and hardware safety procedures."),
        ),
        (
            "Freeze POC3 charter, manifest, demand/load matrix, quality/SLO, availability policy, accepted-work definition, cost boundary and predeclared economic hurdle. Do not reveal hybrid rates before signatures.",
            "Build or refresh the GPU-only control. Run quality, memory inventory, GPU/topology/collectives, low/target/burst/saturation, contexts/outputs/sequences and fault/availability baseline three times.",
            "Install ATMOS driver/FW/SDK, switch/endpoint configuration, packs, placement, vLLM integration, telemetry and operator from versioned artifacts. Run PR-01 preflight and archive environment inventory.",
            "Run hardware smoke: all devices/links/queues, route matrix, self-tests, pack completeness, placement coverage, complete service-group readiness and a single known request.",
            "Run numerical/model smoke at increasing scope: one expert, one MoE layer, several layers, full short decode, full reference prompts. Compare router IDs/gates and task quality at each step.",
            "Run the complete latency decomposition for representative decode tokens. Validate per-layer expert budget, fan-out/fan-in, slowest contribution and GPU continuation under warm and loaded queues.",
            "Run the full offered-load sweep for GPU-only and hybrid: low, target, 1.2x burst, increasing saturation and recovery; repeat three times after stable warmup/temperature.",
            "Sweep contexts 4K/8K/16K/32K, outputs 128/256/1,024 and active sequences 1/4/8/16/32/64. Select decision-critical cells using the approved workload distribution; do not imply every Cartesian cell is supported.",
            "Measure GPU memory relief and how it is used: weight residency removed, KV/workspace/headroom, admitted sequences and OOM/rejection. Do not monetize free bytes without accepted-capacity change.",
            "Run placement experiments: hash/random control, co-activation-aware, hot expert replicas, selected GPU-local experts and failure-domain-aware placement. Compare destinations, skew, queue tail, capacity and surviving service.",
            "Run feature matrix: continuous batching required; prefix cache hit/miss if supported; graph behavior; fixed distributed topology; cancellation/timeout/backpressure; speculative decoding/LoRA disabled unless separately qualified.",
            "Run pack lifecycle: stage second generation, verify complete readiness, canary activate, drain, retire and rollback after an injected quality/SLO failure. Prove no mixed generation.",
            "Run live fault matrix at target load: endpoint/link/switch domain, timeout, thermal throttle, pack corruption, driver/operator restart, node/service restart and chosen failover/fallback. Capture RTO, lost/retried work and surviving accepted rate.",
            "Run 24-hour representative soak with periodic burst, pack health checks and approved fault/reset; track memory/resource leaks, queue drift, thermal behavior and error accumulation.",
            "Populate PR-02 with measured accepted-load curves, failure reserve, configured BOM/quotes, power, support/operations and NRE assumptions. Report TCO per accepted request/token and sensitivity; label all planning values.",
            "Complete PR-01 clean install and rollback by an independent operator; collect support bundle and run one support-triage drill.",
            "Hold POC3 gate review. Approve only the development configuration and POC4 spending; never transfer RTX rate/node-reduction to NVL, HGX, Helios or another OEM.",
        ),
        (
            ("Context/output", "4K/8K/16K/32K; 128/256/1,024", "TTFT/TPOT/output rate/quality and memory by bucket"),
            ("Active sequences", "1/4/8/16/32/64", "Accepted rate, queue stability, KV/headroom and saturation"),
            ("Load", "Low/target/1.2x burst/saturation/recovery", "Offered/admitted/accepted/rejected/error/out-of-SLO distributions"),
            ("Transport", "GPU-only, host-staged hybrid, qualified peer arm", "Complete service and route-specific latency/cost"),
            ("Placement", "Hash, co-activation, replicas, GPU-local hot, failure-aware", "Fan-out, slowest contribution, queues, capacity and survivability"),
            ("Lifecycle", "Fresh install, second pack, canary, rollback, driver/operator restart", "Generation/readiness/RTO/support evidence"),
            ("Fault", "Endpoint/link/switch/thermal/pack/process/node", "Correctness, containment, RTO, surviving accepted rate"),
            ("Duration", "Short benchmark repetitions + 24-hour soak", "Confidence, leaks/drift/errors/thermal stability"),
        ),
        (
            "cargo test --manifest-path control-plane/Cargo.toml",
            "ctest --test-dir atmos-sim/build --output-on-failure && ctest --test-dir data-plane/build -N  # Run data-plane tests only if registered",
            "# CONTROL TEMPLATE: python -m vllm.entrypoints.openai_api_server --model <checkpoint> --port 8001 <control-options>",
            "# HYBRID TEMPLATE: VLLM_PLUGINS=atmos python -m vllm.entrypoints.openai_api_server --model <checkpoint> --port 8002 --additional-config '{\"atmos\":{\"manifest\":\"<id>\"}}'",
            "# TEMPLATE: vllm bench serve --base-url http://localhost:8001 --model <checkpoint> <frozen-workload-options>",
            "# TEMPLATE: vllm bench serve --base-url http://localhost:8002 --model <checkpoint> <same-frozen-workload-options>",
        ),
        (
            ("Manifest/control", "Complete PR-01 drafts, configs/checksums, GPU-only repeated baseline and tuning log"),
            ("Quality", "Expert and end-task results, router/gates, failed cases and feature matrix"),
            ("Service", "Raw requests/traces/metrics, load curves, TTFT/TPOT, output and accepted rate with confidence"),
            ("Resources", "GPU/ATMOS memory, packs, queueing, bytes, NPU/GPU, network, power/thermal and utilization"),
            ("Placement", "Trace/replay/live fan-out/skew/slowest/replica/GPU-local and capacity evidence"),
            ("Lifecycle/fault", "Install/update/rollback/fault/RTO/soak/support drill and all errors/retries"),
            ("Economics", "Configured cost inputs, rate/reserve, TCO/ROI sensitivity and evidence labels"),
            ("Decision", "Supported development envelope, known limits, no-fit cells and signed POC4 authorization"),
        ),
        (
            "Quality passes the frozen operator/end-task contract; no missing/partial expert can silently continue.",
            "Hybrid accepted service rate at the frozen SLO and availability exceeds the predeclared cost-adjusted hurdle or approved strategic-capacity hurdle.",
            "Queues remain bounded at target and approved burst; p99, cancellation, overload and recovery behavior are supportable.",
            "Chosen failure policy preserves correctness and declared complete-service capacity/RTO; lifecycle and soak pass.",
            "Independent clean install, compatibility manifest and support path reproduce the configuration.",
            "All economics are explicitly development-fixture-only and no rack claim is made.",
        ),
        (
            "Quality fails: correct model/pack/gate/combine semantics or retain failing experts/layers on GPU; invalidate affected performance runs.",
            "SLO/rate fails despite capacity fit: do not infer rack value; narrow to prefill/offline/cold experts, change execution boundary or stop.",
            "Availability reserve erases value: redesign replicas/fallback/device group or record no-fit economics.",
            "Operations/support fail: remain engineering POC; do not promote to product POC until lifecycle and ownership close.",
        ),
        (
            "Signed POC3 charter and GPU-only/hybrid manifests", "Reproducible quality and full-service benchmark package",
            "Latency/load/memory/placement/resource reports", "Lifecycle/fault/RTO/soak/support package",
            "PR-01 development release", "PR-02 development economics and sensitivity",
            "Pass/conditional/pivot/stop decision listing what may and may not transfer to POC4",
        ),
    ),
    PocDetail(
        "POC 4", "Named rack-scale product qualification", "Does one exact NVIDIA or AMD rack configuration earn a supported product and economic claim?",
        "Repeat the complete service, lifecycle, fault, security and economics study on a named OEM rack "
        "against its own GPU-only control; qualify NVIDIA and AMD platforms independently.",
        "All 18 work packages, production-intent PR-01, quote/support-ready PR-02, named OEM/switch/GPU/NPU/platform owners",
        (
            "POC3 passes and identifies common contracts versus development-fixture-only evidence.",
            "One exact orderable/production-intent OEM platform per qualification has a returned design record, lab access, configured BOM/quote, engineering and field-support owners.",
            "Production target is selected explicitly: NVIDIA candidates may include an HGX B300 server then GB300 NVL72; AMD candidates may include MI355X server then MI455X Helios. PCIe development systems do not substitute for rack qualification.",
            "GPU-only rack control, hybrid physical route, power/cooling/facilities, network, support and availability policy are approved.",
            "Security, legal, supply, serviceability and fault procedures are authorized for the named facility/configuration.",
        ),
        (
            ("NVIDIA path", "One named OEM and exact HGX/MGX/NVL tray/rack. Grace/Vera Arm64 or x86 host, CUDA/NCCL/NVLink/NVSwitch and network are pinned. ATMOS stays on an OEM-supported PCIe branch/companion path."),
            ("AMD path", "One named OEM/OCP and exact MI355X/MI455X/Helios or approved server. EPYC x86, ROCm/HIP/RCCL and UALink/Infinity/network are pinned. ATMOS remains a separately qualified PCIe endpoint."),
            ("ATMOS geometry", "Start from the OEM-returned balanced 4 GPU + 4 ATMOS local tray where feasible; approved companion PCIe tray is fallback. Exact population follows measured traffic, not a fixed one-to-one rule."),
            ("Control", "Same named rack without ATMOS or with ATMOS disabled, same checkpoint/quality/workload/SLO/HA and equivalent tuning process."),
            ("Facility", "Production-like power, cooling, CDU/fans, network, BMC/management, storage, orchestration, security and support operations."),
            ("Release", "Platform-specific PR-01 and PR-02; no combined NVIDIA/AMD claim and no transfer between generations/OEMs."),
        ),
        (
            "Freeze platform POC charter and configuration ID. Attach OEM design record, schematic/lane/NUMA, mechanics, power/cooling, BIOS/BMC/reset, support, supply, BOM/quote and unresolved-owner list.",
            "Build/reproduce the named GPU-only rack control. Measure quality, accepted-load curves, GPU memory/KV, native fabric/network, power/cooling, faults/availability and operations using the same definitions as POC3.",
            "Install ATMOS physical subsystem under OEM procedure. Inventory every device/board/cable/retimer/firmware and confirm storage/NIC/resource displacement against BOM.",
            "Bring up platform in stages: BIOS/address/IOMMU/ACS; endpoint enumeration; switch ports; driver/FW; one-device DMA; all-device DMA; route matrix; power/thermal sensors; BMC/diagnostics; pack/service readiness.",
            "Port/build the common software stack for host architecture and GPU ecosystem. Re-run SW-01 through SW-10 conformance; close CUDA versus ROCm, Arm64 versus x86, distro/kernel, operator and vendor runtime deltas in PR-01.",
            "Qualify host-staged routes for every intended GPU/tray/ATMOS path. Promote peer DMA only with documented OEM/vendor support, exact topology, enabled security and complete fault/invalidation tests.",
            "Load complete packs and placement/failure policy. Run one request/layer/full short decode smoke, then the complete quality suite before any performance result.",
            "Repeat the POC3 workload/load/context/output/concurrency matrix on GPU-only and hybrid at identical model/quality/SLO/availability. Do not multiply or reuse POC3 rates.",
            "Run rack placement/locality experiments and measure destinations, root/tray/rack crossings, native fabric/network traffic, slowest contribution, queues and accepted capacity.",
            "Run platform power/thermal testing at simultaneous GPU+ATMOS+network/storage target/burst and approved facility conditions. Capture rack/PDU/CDU/BMC and per-component behavior, throttle and safety response.",
            "Run complete fault matrix: endpoint, link/retimer, switch/domain, tray/root/host, native GPU fabric/rank, network, power/cooling sensor, pack, driver/operator and control-plane failures. Validate containment, RTO/RPO, surviving rate and support bundle.",
            "Run firmware/BIOS/driver/Kubernetes/operator/pack update and rollback, physical replacement, spare restore, re-enumeration/rejoin and customer support escalation drill.",
            "Run security/tenancy qualification: IOMMU/ACS/DPC, DMA bounds, stale/cross-tenant handles, secure boot/update, pack/FW authenticity, RBAC, secrets, evidence privacy and vulnerability response.",
            "Run 72-hour representative production-intent soak with scheduled bursts, maintenance/fault event and recovery. Track errors, resets, leaks, thermal/power drift, queue stability, accepted work and operator intervention.",
            "Populate quote-backed PR-02: complete GPU-only/hybrid rack counts, reserve/spares, rack/network/power/cooling/licenses/support/NRE/migration and sensitivity. Calculate cost per accepted request/token and sponsor hurdle.",
            "Perform independent fresh install from PR-01 by operations/support, reproduce smoke and one benchmark cell, and audit every approved product statement against evidence/config ID.",
            "Hold platform-specific release review. Issue release, conditional closures, narrower workload/platform, or no-fit. Repeat the entire relevant qualification for the other GPU vendor/OEM/generation.",
        ),
        (
            ("Platform", "Named NVIDIA and named AMD configuration separately", "No family-level or cross-vendor transfer"),
            ("Workload", "Same POC3 decision-critical matrix plus rack-specific demand/locality", "Independent accepted-load/quality/SLO curves"),
            ("Transport", "Every intended host-staged route; exact peer routes only", "Route matrix by GPU/tray/root/rack crossing"),
            ("Power/thermal", "Idle/low/target/burst/saturation/steady/high approved condition", "Rack/PDU/CDU/BMC/device rates and safety"),
            ("Failure", "Endpoint/switch/tray/root/GPU fabric/network/power/cooling/software", "Containment, RTO/RPO, surviving accepted rate and support"),
            ("Lifecycle", "Fresh install, updates/rollback, physical replacement/spares and EOL substitute rules", "Operational reproducibility and support ownership"),
            ("Duration", "Repeated benchmark cells + 72-hour representative soak", "Confidence, drift, incident and human intervention"),
        ),
        (
            "# NVIDIA TEMPLATE: nvidia-smi topo -m && NCCL_DEBUG=INFO <control-or-hybrid-service>",
            "# AMD TEMPLATE: rocm-smi --showtopo && NCCL_DEBUG=INFO <control-or-hybrid-service>",
            "# TEMPLATE: kubectl get nodes,pods -o wide > evidence/.../environment/cluster.txt",
            "# TEMPLATE: atmos-preflight --manifest <platform-manifest> --json evidence/.../environment/preflight.json",
            "# TEMPLATE: vllm bench serve --base-url <control-or-hybrid> --model <checkpoint> <identical-frozen-workload-options>",
        ),
        (
            ("OEM design/config", "Signed configuration/BOM/lane/mechanics/power/cooling/BIOS/BMC/reset/support/supply package"),
            ("Control/hybrid service", "Independent quality/load/resource/native-fabric/network datasets with confidence"),
            ("Transport/locality", "Route matrix and root/tray/rack crossings with counters and complete latency"),
            ("Facility", "Rack/PDU/CDU/BMC/device power/thermal and safety/throttle evidence"),
            ("Reliability/security", "Fault/RTO/RPO/surviving rate, isolation, updates, replacement, soak and security review"),
            ("Operations/support", "Clean install, runbooks, diagnostics, support drill, spares, supply, SLA/escalation"),
            ("Economics", "Quote-backed equal-service TCO/ROI/sensitivity and demand plan"),
            ("Release/claims", "Signed PR-01/PR-02, known limits, approved claims and platform-specific decision"),
        ),
        (
            "Named hybrid rack preserves frozen model quality, TTFT/TPOT and availability while meeting accepted-load and queue stability criteria.",
            "Declared hardware/software/facility failures are contained and recover within approved RTO/RPO with sufficient surviving capacity.",
            "Security, lifecycle, fresh install, update/rollback, replacement, support, supply and 72-hour soak pass.",
            "Quote-backed three-year TCO per accepted request/token or approved strategic-capacity measure exceeds the predeclared product hurdle over the signed sensitivity region.",
            "Every product claim names the exact platform/configuration and maps to evidence; the other GPU vendor/platform is not implied.",
        ),
        (
            "No supported physical route: use OEM-approved companion tray or next vendor; otherwise no product on that platform.",
            "Rack GPU memory/compute makes ATMOS uneconomic: record no-fit or target a different workload/value mechanism; capacity pressure alone is insufficient.",
            "SLO/availability fails: narrow workload/expert scope, change placement/topology/transport or stop the platform claim.",
            "Support/supply/security owner absent: no release even if performance passes.",
        ),
        (
            "Named OEM design/configuration and full raw qualification package", "Platform-specific PR-01 release bundle",
            "Platform-specific PR-02 economics/support/supply contract", "Independent quality/service/fault/security/facility/soak reports",
            "Approved known limits and evidence-to-claim map", "Release/conditional/narrow/no-fit decision",
        ),
    ),
]


def add_poc_runbook(document: Document, poc: PocDetail) -> None:
    add_page_break(document)
    document.add_paragraph(f"{poc.identifier} | INTEGRATION & DECISION GATE", style="Kicker")
    add_heading(document, poc.title, 1)
    add_table(document, ["Decision question", "Purpose", "Required work packages"], [[
        poc.decision, poc.purpose, poc.packages,
    ]], font_size=7.2, header_fill=PURPLE, banded=False)
    add_heading(document, "Entry checklist", 2)
    add_bullets(document, poc.entry)
    add_heading(document, "Lab / platform setup", 2)
    add_table(document, ["Area", "Required setup"], [[a, b] for a, b in poc.setup], font_size=7.35)
    add_heading(document, "Ordered execution procedure", 2)
    add_numbered(document, poc.steps)
    add_heading(document, "Experiment matrix", 2)
    add_table(document, ["Axis", "Required cells", "Report"], [[a, b, c] for a, b, c in poc.matrix], font_size=7.15)
    add_heading(document, "Commands and automation templates", 2)
    add_body(document, "Replace placeholders with the signed manifest. Physical commands require approved "
             "vendor/OEM procedures and trained operators. Archive stdout, stderr, exit status and tool versions.")
    for command in poc.commands:
        add_command(document, command)
    add_heading(document, "Evidence that must be captured", 2)
    add_table(document, ["Evidence area", "Required content"], [[a, b] for a, b in poc.evidence], font_size=7.35)
    add_heading(document, "Pass criteria", 2)
    add_bullets(document, poc.pass_criteria)
    add_heading(document, "Failure / pivot procedure", 2)
    add_bullets(document, poc.failure)
    add_heading(document, "Gate package", 2)
    add_bullets(document, poc.gate_package)
    add_callout(document, "GATE RULE", "The gate chair records Pass, Conditional Pass with dated mandatory "
                "closures, Pivot with a new hypothesis and invalidated evidence, or Stop. 'Promising' and "
                "'demo complete' are not decision states.", fill=PALE_ORANGE, accent=ORANGE)


def add_poc_section(document: Document) -> None:
    add_part_opener(document, "V", "Four grouped POC execution runbooks",
                    "Consume signed work-package artifacts and finish every gate with a reproducible decision package.")
    add_heading(document, "Common POC control procedure", 1)
    add_numbered(document, (
        "Appoint gate chair, technical lead, evidence owner, safety/reliability owner, product/SLO owner and independent reviewers.",
        "Freeze decision question, entry artifacts, evidence class, manifest, workload/quality/SLO, accepted-work definition, failure policy, cost boundary, pass criteria and stop conditions.",
        "Create immutable run IDs and preflight every tool/configuration. Record deviations before execution; do not repair the manifest after seeing results.",
        "Run control/smoke first, then correctness, performance/load, fault/recovery, lifecycle/operations, soak and economics in that order. A failed earlier layer invalidates later conclusions.",
        "Use stable warmup and at least three independent measured runs for decision-critical cells. Preserve all errors, timeouts, rejects and out-of-SLO work.",
        "Analyze from archived raw evidence with versioned scripts. Review anomalies and repeat only with a documented reason and new run ID.",
        "Hold a pre-read review, then gate meeting. Sign decision, conditions, invalidated evidence, unlocked spending and next owner/date.",
        "Archive the package read-only and update PR-01/PR-02, risk register, compatibility, approved claims and the next POC entry checklist.",
    ))
    add_heading(document, "Gate signers", 2)
    add_table(document, ["Gate dimension", "Required signer", "Cannot be delegated to"], [
        ["Architecture/function", "Chief architect + owning implementation leads", "POC author alone"],
        ["Performance/evidence", "Independent performance/validation lead", "Feature implementation owner alone"],
        ["Reliability/security/safety", "Reliability/security and lab/platform authority", "Benchmark operator alone"],
        ["Product/SLO/economics", "Product owner + finance/solution architecture for POC3/4", "Engineering estimate alone"],
        ["OEM/support/supply", "OEM/platform, procurement and support owners for POC2/4", "Vendor marketing statement"],
    ], font_size=7.4)
    for poc in POCS:
        add_poc_runbook(document, poc)


def add_execution_toolkit(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(document, "VI", "Execution toolkit and reusable templates",
                    "Use these controls to plan work, create reproducible runs, triage failures, conduct reviews, and release one exact configuration.")
    add_figure(document, assets["dependencies"],
               "Figure 3. Dependency and promotion flow. SW-10 evidence spans every layer; PR-01 and PR-02 consume only signed upstream results.")

    add_heading(document, "1. Program RACI", 1)
    add_table(document, ["Function", "Accountable decisions", "Consulted / required evidence"], [
        ["Product / service owner", "Checkpoint/use case, quality, SLO, availability, economic hurdle and permitted claims", "Customer/account, quality, finance, support, legal"],
        ["Chief architect / program technical lead", "Architecture boundary, package dependencies, change impact, gate technical recommendation", "Every work-package owner and independent reviewers"],
        ["Model / compiler / NPU", "Expert ABI, packs, kernels, numerics, useful HBF/NPU service envelope", "SW-01 oracle, HW-01 evidence, partner compiler/runtime"],
        ["Firmware / driver / transport", "Endpoint/DMA/queues, memory ownership, isolation, reset, host/peer routes", "Security, OEM/OS, switch FAE, GPU platform"],
        ["Serving / placement", "vLLM seam, grouping, fan-in, placement epochs, GPU fallback and feature matrix", "Model owner, GPU runtime, performance, SRE"],
        ["Hardware / OEM", "Switch, lane/NUMA, mechanics, power/cooling, BIOS/BMC, serviceability and BOM", "Silicon/FW, switch/board vendors, facility, procurement"],
        ["Platform / SRE", "Kubernetes allocation, lifecycle, readiness, update/rollback, recovery and runbooks", "Driver/runtime, security, support, customer operations"],
        ["Performance / validation", "Run controls, instrumentation, statistics, evidence integrity and independent gate report", "All data producers; implementation owner cannot self-approve"],
        ["Security / reliability / safety", "Threat model, isolation, fault authorization, RTO/RPO, lab/facility safety and release blockers", "OEM, firmware/driver, SRE, support"],
        ["Product finance / procurement / support", "Quote/TCO/hurdle, supply/lifecycle, SLA, escalation and product decision", "OEM, Sandisk, GPU/switch/NPU/software owners"],
    ], font_size=7.0)

    add_heading(document, "2. Work-package status workflow", 1)
    add_table(document, ["Status", "Entry", "Required exit"], [
        ["Not started", "Owner/backlog exists", "Charter, dependencies, target manifest and ready checklist assigned"],
        ["Blocked", "Named dependency/evidence missing", "Blocker owner/date/escalation and workaround decision recorded"],
        ["In design", "Definition of ready passes", "API/design/threat/test/evidence plans reviewed"],
        ["Implementing", "Design review closes mandatory issues", "Code/hardware/config plus unit/negative tests and docs"],
        ["Package verification", "Implementation complete for target manifest", "Verification matrix, artifacts, lifecycle and failure paths pass"],
        ["POC-ready", "Definition of done signed", "Immutable artifacts/evidence index consumed by named POC"],
        ["Platform-qualified", "Package passes on one named production-intent configuration", "PR-01 evidence and limitations updated"],
        ["Released", "POC4 and product reviews pass", "Published signed bundle, support/supply, claims and EOL policy"],
        ["No-fit / retired", "Stop decision or superseded configuration", "Reason, evidence, reopening trigger and migration/EOL recorded"],
    ], font_size=7.45)

    add_heading(document, "3. Run manifest template", 1)
    add_command(document, "run_id: <program>-<poc>-<config>-<UTC timestamp>-<sequence>\nevidence_class: analytic|simulated|emulated|prototype|production-silicon|oem-rack\nconfiguration_id: <PR-01 id or pre-release id>\nmodel:\n  checkpoint: <name>\n  checkpoint_sha256: <hash>\n  tokenizer_sha256: <hash>\n  model_code_commit: <commit>\n  precision_recipe: <id>\n  expertpack_generation: <id>\n  placement_epoch: <id>\nplatform:\n  oem_system: <exact SKU/revision>\n  cpu_arch: x86_64|aarch64\n  gpu: <exact model/count/FW/driver>\n  atmos: <device count/revision/FW>\n  switch_board: <part/revision/FW/port map>\n  bios_bmc_kernel_os: <versions>\nsoftware:\n  container_digest: <digest>\n  vllm_pytorch_cuda_rocm: <versions>\n  atmos_sdk_driver_operator: <versions>\nworkload:\n  profile_id: <id>\n  quality_contract: <id>\n  ttft_tpot_slo: <values/percentiles>\n  load_cell: <low|target|burst|saturation + parameters>\n  failure_policy: <id>\nrun_control:\n  warmup: <rule>\n  duration_or_requests: <value>\n  repetitions: <value>\n  stop_conditions: [<conditions>]\n  deviations: [<ids>]\nowners:\n  operator: <name/function>\n  evidence_owner: <name/function>\n  safety_authority: <name/function>\n  approvers: [<functions>]")

    add_heading(document, "4. Gate decision template", 1)
    add_table(document, ["Section", "Required entry"], [
        ["Decision", "Pass | Conditional Pass | Pivot | Stop"],
        ["Configuration / evidence class", "Exact IDs; no platform-family shorthand"],
        ["Question answered", "One sentence matching the approved charter"],
        ["Mandatory criteria", "Each criterion: pass/fail/not-run with evidence link and reviewer"],
        ["Result summary", "Quality, service, fault, lifecycle, power/thermal, security, economics/support as applicable"],
        ["Deviations / uncertainty", "Missing telemetry, provisional evidence, unsupported cells and effect on conclusion"],
        ["Conditions", "Owner, closure artifact, deadline and what remains locked"],
        ["Invalidated evidence", "Prior results made obsolete by pivot/configuration change"],
        ["Unlocked scope / spend", "Specific next package/POC/NRE; everything else remains held"],
        ["Claims permitted / prohibited", "Exact wording tied to configuration ID"],
        ["Signatures", "Architecture, performance, reliability/security, product/finance, OEM/support where required"],
    ], font_size=7.4)

    add_heading(document, "5. Deviation and defect severity", 1)
    add_table(document, ["Class", "Examples", "Run / gate action"], [
        ["D0 - metadata", "Typo or non-decision label; raw evidence unaffected", "Correct with audit trail; no rerun if independently reviewed"],
        ["D1 - bounded measurement", "One optional counter absent; primary metric intact", "Record uncertainty; reviewer decides whether cell remains usable"],
        ["D2 - configuration drift", "Version/setting/topology differs from manifest", "Invalidate run unless preapproved equivalence is proven"],
        ["D3 - correctness/safety", "Wrong result, stale DMA, isolation failure, uncontained reset, unsafe thermal", "Stop immediately; quarantine; affected performance/economics invalid"],
        ["D4 - evidence integrity", "Missing raw data, changed analysis, clock/counter inconsistency, untracked rerun", "No gate use until reconstructed/repeated with clean run ID"],
        ["D5 - support/supply", "Unowned critical path, unsupported vendor route, no lifecycle/quote", "Engineering result may stand; product gate cannot pass"],
    ], font_size=7.25)

    add_heading(document, "6. Core metric dictionary", 1)
    add_table(document, ["Metric", "Definition", "Required dimensions / caution"], [
        ["Offered work", "Requests/tokens presented to admission during the measurement window", "Workload bucket, tenant/priority, prefill/decode; not throughput"],
        ["Accepted work", "Work completed within frozen quality and latency SLO", "Exclude errors, timeouts, overload cancellations and late completions"],
        ["TTFT", "Request arrival to first accepted output token", "p50/p95/p99 by prompt bucket and load"],
        ["TPOT / inter-token latency", "Time between accepted output tokens after first token", "Distribution and worst sustained periods; define aggregation"],
        ["Expert continuation", "Router result ready to dependent GPU continuation after valid fan-in/combine", "Layer/top-k/destinations/tokens-expert/route/queue state"],
        ["Tokens per expert", "Tokens grouped for one selected expert invocation", "Not API batch or active sequences"],
        ["Destination fan-out", "Distinct ATMOS devices required per token/layer", "Mean/p95 and slowest contribution"],
        ["Queue age/depth", "Waiting work and age of oldest work at each queue", "Depth alone can hide starvation"],
        ["Useful HBF bytes", "Weight/data bytes contributing to executed expert work", "Separate physical/repeated/staging traffic and amplification"],
        ["Physical transport bytes", "Actual bytes crossing GPU/host/PCIe/switch boundaries", "Direction/route/retries/metadata; reconcile counters"],
        ["Surviving accepted rate", "Accepted work after declared failure and recovery/degraded policy", "Include RTO, lost/retried work and complete-service capacity"],
        ["TCO per accepted unit", "Fully burdened cost divided by accepted requests or output tokens", "Same model/quality/SLO/HA/demand; quote/measure/assumption labels"],
    ], font_size=6.95)

    add_heading(document, "7. Troubleshooting and next action", 1)
    add_table(document, ["Symptom", "First checks", "Likely owning packages", "Do next"], [
        ["Wrong expert output", "Router IDs/gates, pack identity/scales/layout, gate applied once, reorder", "SW-01/02/03/04/05, HW-01", "Stop performance; replay one golden; bisect GPU seam versus NPU"],
        ["High p99 with normal mean", "Destinations/token, slowest contribution, queue age, skew, thermal, root crossing", "SW-04/05/10, HW-01/03/05/06", "Trace one tail token; compare placement/route/loaded queue"],
        ["Queue grows at target", "Arrival/service curves, tokens/expert, grouping wait, admission, throttle", "SW-04/03, HW-01", "Reduce admission or add capacity; fix service before soak/economics"],
        ["Low PCIe bandwidth", "Message size, negotiated width/rate, NUMA, pinning, link errors, direction", "HW-02/03/05/06, SW-05/07", "Run direct control; isolate endpoint/port/channel/root"],
        ["Good copies but slow continuation", "GPU stream events, host scheduling, fan-in/combine, global sync, rank stalls", "SW-03/05/06/10", "Profile complete critical path; remove hidden synchronization"],
        ["One endpoint failure loses service", "Unique shards, replicas/fallback, surviving capacity, epoch readiness", "SW-04/09, PR-02", "Fail closed; redesign reserve/placement before HA claim"],
        ["Reset causes stale write", "Buffer generation/invalidation, completion ambiguity, IOMMU mapping", "SW-07, HW-02/03", "Quarantine platform; fix state machine; invalidate related results"],
        ["Switch reset impacts all devices", "DPC/reset domain/FW, one versus dual domain", "HW-05/06, SW-09", "Measure blast radius; split domains or change availability policy"],
        ["Thermal throttling", "BMC/device sensors, inlet/coolant, airflow, fan/CDU, simultaneous GPU load", "HW-01/05/06", "Lower population/power or redesign cooling; rerun sustained envelope"],
        ["Peer works only with IOMMU/ACS off", "Supported vendor/OEM route and security policy", "HW-03/05/06", "Reject peer route; retain secure host staging"],
        ["Hybrid fits but is not cheaper", "Accepted rate, reserve, configured BOM/NRE/support, removable GPU/network cost", "PR-02 + all service packages", "Record no-fit or target another workload/value mechanism"],
        ["Kubernetes pod restarts but service stays incomplete", "Pack/placement/replica readiness and operator recovery", "SW-08/09", "Restore complete service group; pod health alone is insufficient"],
    ], font_size=6.65)

    add_heading(document, "8. OEM and supplier request checklist", 1)
    add_bullets(document, (
        "Exact orderable server/tray/rack configuration and revision; named engineering, support and commercial owners.",
        "CPU/root/NUMA connector/lane map; GPU/native fabric boundary; switch/board/backplane/cable/retimer and all crossings.",
        "E3.S 2T mechanics, retention, cabling, service access and displaced storage/NIC/other resources.",
        "Power/current/transient/PSU/PDB/PDU allocation and simultaneous GPU+ATMOS+network reserve.",
        "Cooling/airflow/cold plate/CDU/fan/sensor/BMC behavior, supported conditions and throttle/shutdown policy.",
        "BIOS address windows, IOMMU/ACS/ATS/PASID, AER/DPC, secure boot/update, FLR/reset/re-enumeration and failure domains.",
        "Supported OS/kernel/driver/container/CUDA/ROCm/Kubernetes, peer-DMA status and unsupported combinations.",
        "Diagnostics, firmware update/rollback, field replacement, spares, RMA, warranty, escalation and SLA.",
        "Prototype and production BOM/quote, NRE/qualification, MOQ, lead time, geographies, lifecycle/EOL and substitution rules.",
        "Lab access, fault/thermal/security test permissions, data handling and returned signed configuration/design record.",
    ))

    add_heading(document, "9. GPU platform target sequence", 1)
    add_table(document, ["Stage", "NVIDIA target", "AMD target", "Why / required conclusion"], [
        ["Development", "RTX PRO 6000 Blackwell Server Edition; 96 GB GDDR7, PCIe Gen5, OEM/MGX servers", "Instinct MI350P; 144 GB HBM3E, PCIe Gen5 add-in card", "Lowest-friction framework/PCIe integration and strong expert-memory pressure; not rack claim"],
        ["High-end server", "8-GPU HGX B300-class OEM server", "8-GPU MI355X OAM OEM server", "Test stronger HBM/control, native scale-up and OEM server lifecycle before full rack"],
        ["Rack product", "One exact GB300 NVL72 OEM configuration", "One exact MI455X Helios/OCP configuration", "Independent rack control/hybrid qualification and quote-backed economics"],
        ["Future rebaseline", "HGX Rubin NVL8 / Vera Rubin NVL72", "Future MI400 variants as orderable", "New memory/compute/fabric generation is a new control and qualification"],
    ], font_size=7.1)
    add_callout(document, "SELECTION RULE", "Target GPUs by integration learning and economic pressure, "
                "not peak specifications alone. Large native HBM can weaken the ATMOS capacity thesis; rack "
                "targets must prove cost per accepted token, concurrency, consolidation or avoidable fabric cost.",
                fill=PALE_ORANGE, accent=ORANGE)

    add_heading(document, "10. Final release checklist", 1)
    add_table(document, ["Dimension", "Required release evidence", "Signer"], [
        ["Function/quality", "Complete checkpoint/operator goldens, end-task quality, fallback/update paths", "Model/quality + architecture"],
        ["Service", "Accepted-load curves, TTFT/TPOT, queues, power/thermal, degraded mode and soak", "Performance/validation"],
        ["Hardware/platform", "OEM design record, endpoint/switch/route, SI, power/cooling, BIOS/BMC and service", "OEM/hardware authority"],
        ["Availability/reliability", "Complete-service reserve, fault containment, RTO/RPO, replacement and recovery", "Reliability/SRE"],
        ["Security", "DMA/isolation, signed lifecycle, tenancy, threat/negative tests, vulnerability response", "Security"],
        ["Operations", "Fresh install, allocation, packs, rollout/rollback, monitoring, incident and support bundle", "Platform/SRE/support"],
        ["Compatibility", "Signed PR-01, preflight/drift, known limits, upgrade/EOL", "Release"],
        ["Economics/supply/support", "Quote-backed PR-02, hurdle/sensitivity, demand, spares, supply, SLA/escalation", "Product/finance/procurement/support"],
        ["Claims", "Evidence-to-claim map naming exact platform; prohibited broad claims listed", "Product/legal/architecture"],
    ], font_size=7.0)
    add_callout(document, "PRODUCT CLAIM", "For the named checkpoint, workload, availability policy and "
                "OEM rack configuration, the qualified ATMOS E3.S expert tier increases accepted service "
                "capacity or reduces fully burdened comparable infrastructure cost while preserving model "
                "quality, TTFT/TPOT, security, support and lifecycle.", fill=PALE_TEAL, accent=TEAL)


def add_command_index(document: Document) -> None:
    add_part_opener(document, "VII", "Command index and environment preflight",
                    "Run repository checks today; run hardware and platform templates only after the required implementation and authorization exist.")
    add_heading(document, "1. Local repository validation", 1)
    add_table(document, ["Path", "Publication-time check (30 September 2026)", "Reader action"], [
        ["atmos-sim", "ctest passed 7/7: core, virtual devices, topology, model graph, metrics, golden traces and end-to-end.", "Reconfigure/build and rerun for every simulator change or calibration update."],
        ["control-plane", "cargo test passed 64/64. Compiler reported existing unused import/dead-code warnings.", "Treat warnings separately; rerun tests for routing/admission/telemetry changes."],
        ["data-plane", "Existing build reported Total Tests: 0 because CMake only registers tests when GoogleTest is found. Benchmark executable was not present in that build.", "Configure with BUILD_TESTING/BUILD_BENCHMARKS; provide GoogleTest; inspect ctest -N; build/run benchmark on an approved CUDA host."],
    ], font_size=7.25)
    add_command(document, "cmake -S atmos-sim -B atmos-sim/build -DCMAKE_BUILD_TYPE=Release\ncmake --build atmos-sim/build --parallel\nctest --test-dir atmos-sim/build --output-on-failure")
    add_command(document, "cmake -S data-plane -B data-plane/build -DCMAKE_BUILD_TYPE=Release -DBUILD_TESTING=ON -DBUILD_BENCHMARKS=ON\ncmake --build data-plane/build --parallel\nctest --test-dir data-plane/build -N\n# If tests are registered: ctest --test-dir data-plane/build --output-on-failure\n# If Total Tests is 0, provide GoogleTest and reconfigure; benchmark remains available:\ndata-plane/scripts/run_benchmark.sh --sizes 1024,4096,16384,65536,262144 --warmup 20 --iterations 200")
    add_command(document, "cargo test --manifest-path control-plane/Cargo.toml\n# Optional existing scenario suite; validate environment and expected-vs-measured labels first:\nexperiments/scripts/run_all_experiments.sh")
    add_heading(document, "2. Preflight inventory templates", 1)
    add_command(document, "uname -a\n<os-release command>\n<kernel package inventory>\nlspci -tv\nlspci -vv -s <ATMOS-BDF>\n<driver/FW/SDK version commands>\n<switch inventory command>\n<BMC inventory/sensor export>\n<container digest and package lock export>")
    add_heading(document, "3. NVIDIA environment templates", 1)
    add_command(document, "nvidia-smi -q\nnvidia-smi topo -m\nNCCL_DEBUG=INFO NCCL_DEBUG_SUBSYS=INIT,COLL,GRAPH <service-or-collective-test>\nnsys profile -o evidence/.../traces/service <service-command>\n<DCGM diagnostic/telemetry command approved for the platform>")
    add_heading(document, "4. AMD environment templates", 1)
    add_command(document, "rocm-smi --showproductname --showdriverversion --showmeminfo vram\nrocm-smi --showtopo\nNCCL_DEBUG=INFO <RCCL service-or-collective-test>\nrocprofv3 --output-format json --output-file evidence/.../traces/service -- <service-command>\n<OEM/ROCm diagnostics approved for the platform>")
    add_heading(document, "5. Kubernetes templates", 1)
    add_command(document, "kubectl version -o yaml\nkubectl get nodes -o wide\nkubectl get pods -A -o wide\nkubectl get crd,resourceclaims,resourceslices -A -o yaml\nkubectl get events -A --sort-by=.lastTimestamp\n<ATMOS preflight/operator/support-bundle commands>")
    add_heading(document, "6. Command-use rules", 1)
    add_bullets(document, (
        "Run read-only inventory before changes and archive it under the run ID.",
        "Never place passwords, tokens, private keys or customer payloads in command history or evidence bundles.",
        "Do not execute reset, hot-plug, firmware, power, thermal or fault commands without the exact OEM/vendor procedure and lab authorization.",
        "Capture command, version, working directory, environment, stdout, stderr, exit status and start/end timestamps.",
        "A command's success is not a gate result; validate resulting state, counters, outputs and negative behavior.",
    ))


def add_source_and_scope_appendix(document: Document) -> None:
    add_part_opener(document, "VIII", "Scope, terminology, and source boundary",
                    "Keep this handbook useful without confusing planning inputs, public references, simulation and qualified product evidence.")
    add_heading(document, "1. Core terms", 1)
    add_table(document, ["Term", "Meaning"], [
        ["ExpertPack", "Immutable compiled/signed ATMOS expert generation with weights/layout/scales/kernels/metadata/goldens/compatibility."],
        ["Placement epoch", "Immutable expert-to-owner mapping used for the full lifetime of an in-flight request."],
        ["Complete service group", "Devices, packs, placement, runtime and reserve that can execute every required expert under the declared policy."],
        ["Accepted work", "Request or output token completed within frozen quality and latency SLO; excludes error, timeout, overload cancellation and late completion."],
        ["Host-staged", "GPU/ATMOS transfer through registered pinned host buffers using stock supported APIs; mandatory functional baseline."],
        ["Peer DMA", "Documented configuration-specific GPU-visible/ATMOS route without an application CPU copy; optional until fully qualified."],
        ["Useful HBF bandwidth", "Memory bytes contributing to real expert execution; separate from nominal interface rate and physical staging traffic."],
        ["Evidence class", "Analytic, simulated, emulated, prototype, production-silicon or named-OEM-rack; determines permitted conclusion."],
        ["Configuration ID", "Exact supported tuple of hardware/FW/platform/software/model/topology/workload/failure/limits."],
    ], font_size=7.45)
    add_heading(document, "2. Evidence boundary", 1)
    add_bullets(document, (
        "The draft proposal supplies the 18-artifact serving-stack matrix, common Sandisk requirements, GPU seams, owners and four grouped product questions.",
        "Repository commands verify current simulator, data-plane and control-plane software only; they do not establish ATMOS silicon or rack behavior.",
        "Public NVIDIA, AMD, Tenstorrent, switch and OEM material establishes published interfaces/offering context, not ATMOS compatibility or commitment.",
        "Planning values such as HBF capacity/rate, device power, traffic, SLO and economics remain targets until replaced by signed engineering/customer evidence.",
        "Every platform, OEM, GPU generation, switch/board, driver/FW and software release is qualified by configuration ID; results do not transfer silently.",
    ))
    add_heading(document, "3. Handbook maintenance", 1)
    add_numbered(document, (
        "Assign a handbook owner and package owners. Review open dependencies weekly and stable sections monthly.",
        "When an interface, target, supplier, model, topology, SLO or support boundary changes, update affected work packages, POC matrices, manifests, economics, risks and approved claims together.",
        "Preserve superseded handbook, manifests, evidence and decisions. Do not edit old gate packages to match a new design.",
        "Promote proven procedures from TEMPLATE to supported only after implementation, security/safety review, execution and release ownership.",
        "After each POC, update lessons learned, common failure symptoms, run duration/resources and the next gate's entry checklist.",
    ))


def package_summary(identifier: str) -> WorkPackageSummary:
    return next(item for item in WORK_PACKAGES if item.identifier == identifier)


def add_work_package(document: Document, detail: WorkPackageDetail) -> None:
    summary = package_summary(detail.identifier)
    add_page_break(document)
    document.add_paragraph(f"{detail.identifier} | {summary.domain.upper()} | ACTION {detail.action}", style="Kicker")
    add_heading(document, detail.title, 1)
    add_table(document, ["Purpose", "Primary owner", "Consumed by", "Dependencies"], [[
        detail.purpose, detail.owner, summary.consumed_by, detail.dependencies,
    ]], font_size=7.2, header_fill=TEAL, banded=False)
    add_heading(document, "Definition of ready", 2)
    add_bullets(document, detail.ready)
    add_heading(document, "Ordered execution", 2)
    add_numbered(document, detail.steps)
    if detail.commands:
        add_heading(document, "Command and automation examples", 2)
        add_body(document, "Commands marked TEMPLATE require the corresponding implementation, approved "
                 "configuration, credentials, and safety procedure. Replace placeholders and archive stdout/stderr.")
        for command in detail.commands:
            add_command(document, command)
    add_heading(document, "Verification matrix", 2)
    add_table(document, ["Check", "Method", "Pass evidence"],
              [[a, b, c] for a, b, c in detail.verification], font_size=7.15)
    add_heading(document, "Required deliverables", 2)
    add_bullets(document, detail.deliverables)
    add_heading(document, "Definition of done", 2)
    add_bullets(document, detail.done)
    add_heading(document, "Failure and pivot path", 2)
    add_bullets(document, detail.failure)
    add_heading(document, "NVIDIA / AMD platform seam", 2)
    add_body(document, detail.platform_seam)


def add_software_packages(document: Document) -> None:
    add_part_opener(document, "II", "Software work-package runbooks",
                    "Execute common Sandisk contracts first; keep NVIDIA and AMD differences behind explicit, independently qualified adapters.")
    add_callout(document, "PACKAGE RULE", "A work package is not complete when code merges. It is complete "
                "when its supported configuration, positive and negative tests, lifecycle, evidence, owner, "
                "and failure behavior satisfy its definition of done.", fill=PALE_BLUE, accent=BLUE)
    for detail in SOFTWARE_PACKAGES:
        add_work_package(document, detail)


def add_hardware_packages(document: Document) -> None:
    add_part_opener(document, "III", "Hardware and platform work-package runbooks",
                    "Qualify direct endpoint first, then transport, switch and one exact OEM server/rack envelope.")
    add_callout(document, "SAFETY", "Physical fault, reset, hot-plug, power, thermal, firmware and "
                "signal-integrity tests require approved lab hardware, disposable data, OEM/vendor procedure, "
                "trained personnel and defined stop conditions. Generic commands in this guide are not authorization.",
                fill=PALE_RED, accent=RED)
    for detail in HARDWARE_PACKAGES:
        add_work_package(document, detail)


def add_product_packages(document: Document) -> None:
    add_part_opener(document, "IV", "Productization work-package runbooks",
                    "Bind qualified components into one reproducible release and one equal-service economic/support contract.")
    for detail in PRODUCT_PACKAGES:
        add_work_package(document, detail)


def add_cover(document: Document, assets: dict[str, Path]) -> None:
    document.add_paragraph("IMPLEMENTATION | QUALIFICATION | EVIDENCE | RELEASE", style="Kicker")
    document.add_paragraph("ATMOS Rack-Scale MVP\nEngineering and POC Execution Handbook", style="Title")
    document.add_paragraph(
        "One-stop, step-by-step guide for verifying, modifying, creating, integrating, and qualifying "
        "the complete ATMOS MoE expert tier from primitive execution to a named rack-scale product.",
        style="Subtitle",
    )
    add_callout(
        document,
        "USE THIS",
        "Start with the readiness checklist, select a work package, execute its ordered procedure, "
        "publish the required evidence, and advance only through the four POC gate reviews. Configuration-"
        "specific runbooks and hardware safety procedures always override generic examples in this guide.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )
    add_figure(document, assets["delivery"], "Figure 1. Delivery layers and ownership surfaces covered by this handbook.")
    add_table(document, ["Coverage", "Primary users", "Final outcome"], [[
        "18 work packages + 4 grouped POC runbooks",
        "Model/compiler, runtime, driver/FW, hardware/OEM, SRE, validation, product/support",
        "A reproducible named-platform qualification package or a documented no-fit decision",
    ]], font_size=8.2, header_fill=TEAL, banded=False)
    document.add_paragraph(
        "Prepared 30 September 2026. This is an execution handbook: planning targets, example commands, "
        "and provisional thresholds must be replaced by approved configuration-specific values before a gate review.",
        style="Source Note",
    )


def add_operating_model(document: Document, assets: dict[str, Path]) -> None:
    add_page_break(document)
    document.add_paragraph("HOW TO USE THE HANDBOOK", style="Kicker")
    add_heading(document, "Execution model", 1)
    add_body(document, "Each work package produces versioned artifacts consumed by one or more POCs. "
             "The POC is the integration and decision layer; it does not replace package-level tests. "
             "A package may be complete for the development fixture and still require platform-specific "
             "requalification for NVIDIA, AMD, Arm64, x86-64, or a different OEM route.")
    add_figure(document, assets["poc"], "Figure 2. Four grouped POCs and the evidence each gate must consume.")
    add_heading(document, "Action codes", 2)
    add_table(document, ["Code", "Meaning", "Completion expectation"], [
        ["V - Verify/reuse", "Preserve an existing vendor/model/platform behavior; prove it still holds after ATMOS integration.", "Baseline, invariant tests, version pin, and evidence of no regression."],
        ["M - Modify/configure", "Change a supported extension point, configuration, topology, or deployment without taking ownership of the vendor core.", "Bounded diff/config, owner, upgrade strategy, conformance and regression tests."],
        ["N - New artifact", "Create and support a Sandisk/partner-owned implementation or product artifact.", "Design/API, implementation, tests, packaging, documentation, security, lifecycle, and support owner."],
        ["M/N", "Modify a platform seam and create the ATMOS-specific component around it.", "Both modification and new-artifact obligations apply."],
    ], font_size=7.6)
    add_heading(document, "Mandatory evidence repository", 2)
    add_command(document, "evidence/<program>/<poc>/<run-id>/\n  manifest.yaml\n  environment/\n  configs/\n  raw/\n  traces/\n  metrics/\n  quality/\n  faults/\n  power-thermal/\n  analysis/\n  economics/\n  decision/\n  checksums.sha256")
    add_bullets(document, [
        "Use one immutable run ID across GPU, host, switch, ATMOS, service, and analysis data.",
        "Record offered, admitted, accepted, rejected, failed, timed-out, cancelled, and out-of-SLO work separately.",
        "Label evidence as analytic, simulated, emulated, prototype, production-silicon, or named-OEM-rack.",
        "Do not overwrite failed runs. Store deviations, missing telemetry, and unsupported combinations.",
        "Require checksums and a decision record signed by architecture, performance, reliability/security, and product/finance where applicable.",
    ])
    add_heading(document, "Repository commands available today", 2)
    add_command(document, "cmake -S atmos-sim -B atmos-sim/build -DCMAKE_BUILD_TYPE=Release\ncmake --build atmos-sim/build --parallel\nctest --test-dir atmos-sim/build --output-on-failure")
    add_command(document, "cargo test --manifest-path control-plane/Cargo.toml\ncmake -S data-plane -B data-plane/build -DCMAKE_BUILD_TYPE=Release -DBUILD_TESTING=ON -DBUILD_BENCHMARKS=ON\ncmake --build data-plane/build --parallel\nctest --test-dir data-plane/build -N\n# Run --output-on-failure only if tests are registered; otherwise provide GoogleTest and use the standalone transfer benchmark.")
    add_callout(document, "BOUNDARY", "The simulator and transfer harness can establish reproducibility, "
                "causality, analytical bounds, transaction behavior, and calibrated what-if results. "
                "They cannot pass silicon timing, physical DMA, switch containment, thermal, OEM support, "
                "or rack economics gates.", fill=PALE_RED, accent=RED)


def add_master_register(document: Document) -> None:
    add_part_opener(document, "I", "Master work-package register",
                    "Select an artifact, satisfy its prerequisites, execute its runbook, and publish its evidence before integration.")
    rows = [[w.identifier, w.domain, w.artifact, w.action, w.owner, w.purpose, w.consumed_by]
            for w in WORK_PACKAGES]
    add_table(document, ["ID", "Domain", "Artifact", "Action", "Primary owner", "Purpose", "Consumed by"],
              rows, font_size=6.8)
    add_heading(document, "Suggested execution order", 1)
    add_table(document, ["Wave", "Work packages", "Why this order"], [
        ["0 - Control and evidence", "SW-01, SW-10, PR-01 skeleton", "Freeze identities, baselines, trace schema, and configuration fields before implementation results exist."],
        ["1 - Primitive execution", "SW-02, SW-07, HW-01, HW-02", "Establish packs, endpoint/runtime contracts, correctness, queues, and real expert service curves."],
        ["2 - Heterogeneous path", "SW-03, SW-04, SW-05, HW-03", "Create the end-to-end GPU-to-ATMOS execution path and placement/fan-in semantics."],
        ["3 - Platform integration", "SW-06, HW-04, HW-05, HW-06", "Add switch/OEM topology while preserving the native GPU platform."],
        ["4 - Operations", "SW-08, SW-09, PR-01", "Make the qualified group installable, schedulable, updateable, recoverable, and supportable."],
        ["5 - Product evidence", "PR-02 + all packages", "Populate equal-SLO economics and support only from measured, named configurations."],
    ], font_size=7.4)


def validate(document: Document) -> None:
    text = "\n".join([p.text for p in document.paragraphs] +
                     [c.text for t in document.tables for r in t.rows for c in r.cells])
    required = ["ATMOS Rack-Scale MVP", "Master work-package register", "SW-01", "PR-02",
                "Four grouped POCs", "Mandatory evidence repository", "ctest --test-dir",
                "Software work-package runbooks", "Model contract and learned-router baseline",
                "ATMOS SDK and Linux driver", "Telemetry and evidence system",
                "Hardware and platform work-package runbooks", "PCIe switch and backplane subsystem",
                "Server and rack integration envelope", "Productization work-package runbooks",
                "Compatibility manifest and release package", "Economic and support contract",
                "Four grouped POC execution runbooks", "Vendor-neutral expert primitive",
                "Switched E3.S subsystem", "Eight-GPU plus eight-ATMOS development service",
                "Named rack-scale product qualification", "Common POC control procedure",
                "Execution toolkit and reusable templates", "Run manifest template",
                "Troubleshooting and next action", "Final release checklist",
                "Command index and environment preflight", "Scope, terminology, and source boundary"]
    for item in required:
        if item not in text:
            raise RuntimeError(f"Missing required handbook content: {item}")
    if len(document.inline_shapes) < 3:
        raise RuntimeError("Expected at least three original figures")
    identifiers = [w.identifier for w in WORK_PACKAGES]
    if len(identifiers) != 18 or len(set(identifiers)) != 18:
        raise RuntimeError("Expected 18 unique work packages")
    exact_paragraphs = [paragraph.text for paragraph in document.paragraphs]
    expected_heading_counts = {
        "Definition of ready": 18,
        "Ordered execution": 18,
        "Ordered execution procedure": 4,
        "Experiment matrix": 4,
        "Gate package": 4,
    }
    for heading, expected in expected_heading_counts.items():
        actual = exact_paragraphs.count(heading)
        if actual != expected:
            raise RuntimeError(f"Expected {expected} '{heading}' headings, found {actual}")


def main() -> None:
    assets = build_assets()
    document = Document()
    configure_document(document)
    add_cover(document, assets)
    add_operating_model(document, assets)
    add_master_register(document)
    add_software_packages(document)
    add_hardware_packages(document)
    add_product_packages(document)
    add_poc_section(document)
    add_execution_toolkit(document, assets)
    add_command_index(document)
    add_source_and_scope_appendix(document)
    validate(document)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    validate(reopened)
    print(
        f"Wrote {OUTPUT} | paragraphs={len(reopened.paragraphs)} "
        f"tables={len(reopened.tables)} figures={len(reopened.inline_shapes)}"
    )


if __name__ == "__main__":
    main()