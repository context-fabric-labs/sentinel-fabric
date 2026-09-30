from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import re
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCS_DIR = Path(__file__).resolve().parent
ASSET_DIR = DOCS_DIR / "assets" / "atmos_moe_proposal"
OUTPUT = DOCS_DIR / "ATMOS_MoE_RackScale_Hybrid_E3S_Proposal_v4.docx"

NAVY = "112B45"
INK = "233548"
BLUE = "2F6F95"
TEAL = "008C82"
GREEN = "4D8B61"
ORANGE = "D9823B"
RED = "B74D4D"
WHITE = "FFFFFF"
MIST = "EAF1F5"
PALE_TEAL = "E5F3F1"
PALE_ORANGE = "FCEFE4"
PALE_RED = "F8E8E8"
GRAY = "637282"
LIGHT_GRAY = "F4F6F7"
MID_GRAY = "D8E0E5"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def pil_rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def font_path() -> str:
    candidates = [
        "/System/Library/Fonts/Avenir Next.ttc",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No supported TrueType font found")


FONT_PATH = font_path()


def image_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    # The Avenir collection exposes a usable regular face at index 0 and a heavier face at 1.
    try:
        return ImageFont.truetype(FONT_PATH, size=size, index=1 if bold else 0)
    except OSError:
        return ImageFont.truetype(FONT_PATH, size=size)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90,
                     end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = MID_GRAY, size: int = 6) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def set_table_width(table, width_inches: float) -> None:
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(int(width_inches * 1440)))
    width.set(qn("w:type"), "dxa")


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = rgb(BLUE)
    subtitle.paragraph_format.space_after = Pt(14)

    for name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 14, 7),
        ("Heading 2", 13.5, BLUE, 11, 5),
        ("Heading 3", 10.5, TEAL, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = rgb(GRAY)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)

    bullet = styles["List Bullet"]
    bullet.font.name = "Aptos"
    bullet.font.size = Pt(9.25)
    bullet.paragraph_format.left_indent = Inches(0.22)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(3)

    numbered = styles["List Number"]
    numbered.font.name = "Aptos"
    numbered.font.size = Pt(9.25)
    numbered.paragraph_format.left_indent = Inches(0.26)
    numbered.paragraph_format.first_line_indent = Inches(-0.16)
    numbered.paragraph_format.space_after = Pt(3)

    for style_name, size, color, bold in (
        ("Kicker", 8.5, TEAL, True),
        ("Source Note", 7.5, GRAY, False),
        ("Small", 8, GRAY, False),
    ):
        if style_name not in styles:
            styles.add_style(style_name, 1)
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = bold
        style.paragraph_format.space_after = Pt(4)
        if style_name == "Kicker":
            style.font.all_caps = True


def configure_document(document: Document) -> None:
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("ATMOS E3.S  |  RACK-SCALE MoE INFERENCE")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = rgb(GRAY)

    footer = section.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Inches(7.1))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left = table.cell(0, 0).paragraphs[0]
    left.add_run("INTERNAL CONCEPT PROPOSAL").font.color.rgb = rgb(GRAY)
    left.runs[0].font.size = Pt(7.5)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right.add_run("PAGE ")
    run.font.size = Pt(7.5)
    run.font.color.rgb = rgb(GRAY)
    add_field(right, "PAGE")

    properties = document.core_properties
    properties.title = "ATMOS E3.S Expert Tier for Rack-Scale MoE Inference"
    properties.subject = "Product architecture, service stack, OEM integration, economics, and qualification"
    properties.author = "Sandisk / ATMOS Architecture"
    properties.last_modified_by = "Sandisk / ATMOS Architecture"
    properties.comments = "Planning targets and vendor screening hypotheses require measured qualification."
    properties.created = datetime.now(timezone.utc)
    properties.modified = datetime.now(timezone.utc)


def add_kicker(document: Document, text: str) -> None:
    document.add_paragraph(text, style="Kicker")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_paragraph(text, style=f"Heading {level}")


def add_body(document: Document, text: str, *, lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if lead and text.startswith(lead):
        run = paragraph.add_run(lead)
        run.bold = True
        paragraph.add_run(text[len(lead):])
    else:
        paragraph.add_run(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK,
                  size: float = 8.2, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]], *,
              widths: list[float] | None = None, font_size: float = 8.1,
              header_fill: str = NAVY, banded: bool = True) -> object:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, sum(widths) if widths else 7.0)
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        set_cell_text(header.cells[index], text, bold=True, color=WHITE, size=font_size)
        set_cell_shading(header.cells[index], header_fill)
        set_cell_border(header.cells[index], header_fill)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value, size=font_size)
            set_cell_border(row.cells[index])
            if banded and row_index % 2:
                set_cell_shading(row.cells[index], LIGHT_GRAY)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(document: Document, label: str, text: str, *, fill: str = PALE_TEAL,
                accent: str = TEAL) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.75)
    set_cell_text(table.cell(0, 0), label, bold=True, color=WHITE, size=8.6,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_border(table.cell(0, 0), accent)
    set_cell_text(table.cell(0, 1), text, size=8.8)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_border(table.cell(0, 1), accent)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_source_note(document: Document, text: str) -> None:
    document.add_paragraph(text, style="Source Note")


def add_page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_part_opener(document: Document, number: str, title: str, promise: str) -> None:
    add_page_break(document)
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.05)
    table.columns[1].width = Inches(5.95)
    set_cell_text(table.cell(0, 0), number, bold=True, color=WHITE, size=20,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.cell(0, 0), TEAL)
    set_cell_border(table.cell(0, 0), TEAL)
    cell = table.cell(0, 1)
    cell.text = ""
    set_cell_shading(cell, NAVY)
    set_cell_border(cell, NAVY)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = rgb(WHITE)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(promise)
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(MIST)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def draw_wrapped(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
                 font: ImageFont.FreeTypeFont, fill: tuple[int, int, int],
                 align: str = "center", spacing: int = 5) -> None:
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 24
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_height = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y1 + max(0, (y2 - y1 - total_height) // 2)
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        if align == "left":
            x = x1 + 12
        elif align == "right":
            x = x2 - width - 12
        else:
            x = x1 + (x2 - x1 - width) // 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + spacing


def draw_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
             fill: str, outline: str = WHITE, text_color: str = WHITE,
             radius: int = 18, font_size: int = 27, bold: bool = True) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=pil_rgb(fill),
                           outline=pil_rgb(outline), width=3)
    draw_wrapped(draw, box, text, image_font(font_size, bold), pil_rgb(text_color))


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
               color: str = BLUE, width: int = 6) -> None:
    draw.line([start, end], fill=pil_rgb(color), width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - direction * 18, y2 - 11),
                  (x2 - direction * 18, y2 + 11)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 11, y2 - direction * 18),
                  (x2 + 11, y2 - direction * 18)]
    draw.polygon(points, fill=pil_rgb(color))


def make_architecture_diagram(path: Path) -> None:
    width, height = 1800, 1040
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "REFERENCE PRODUCT ARCHITECTURE", font=image_font(36, True),
              fill=pil_rgb(WHITE))
    draw.text((1190, 31), "Native GPU fabric preserved", font=image_font(24),
              fill=pil_rgb(MIST))

    draw.rounded_rectangle((45, 128, 1755, 850), radius=24, fill=pil_rgb(LIGHT_GRAY),
                           outline=pil_rgb(MID_GRAY), width=3)
    draw.text((80, 150), "GPU COMPUTE TRAY", font=image_font(30, True), fill=pil_rgb(NAVY))

    draw_box(draw, (90, 220, 400, 345), "Host CPU / PCIe root", NAVY, font_size=27)
    draw_box(draw, (490, 200, 760, 320), "GPU 0", BLUE, font_size=30)
    draw_box(draw, (790, 200, 1060, 320), "GPU 1", BLUE, font_size=30)
    draw_box(draw, (1090, 200, 1360, 320), "GPU 2", BLUE, font_size=30)
    draw_box(draw, (1390, 200, 1660, 320), "GPU 3", BLUE, font_size=30)
    draw.line((625, 360, 1525, 360), fill=pil_rgb(BLUE), width=10)
    draw.text((935, 375), "NVLink / NVSwitch or UALink GPU domain",
              font=image_font(24, True), fill=pil_rgb(BLUE))

    draw_box(draw, (670, 475, 1150, 595), "Balanced PCIe switch domain", TEAL,
             font_size=29)
    draw_arrow(draw, (400, 282), (670, 535), TEAL, 7)
    draw.text((435, 395), "Gen5/Gen6 x16 uplink", font=image_font(22, True),
              fill=pil_rgb(TEAL))

    device_boxes = [(210, 675, 500, 800), (560, 675, 850, 800),
                    (910, 675, 1200, 800), (1260, 675, 1550, 800)]
    for index, box in enumerate(device_boxes):
        draw_box(draw, box, f"ATMOS E3.S {index}\nHBF + NPU", GREEN, font_size=25)
        center = ((box[0] + box[2]) // 2, box[1])
        draw_arrow(draw, (910, 595), center, GREEN, 5)

    draw_box(draw, (85, 895, 520, 995), "Scale-out Ethernet / InfiniBand", INK,
             font_size=24)
    draw_box(draw, (680, 895, 1120, 995), "Companion ATMOS tray (fallback)", ORANGE,
             font_size=24)
    draw_box(draw, (1280, 895, 1715, 995), "Rack control, power, cooling", GRAY,
             font_size=24)
    draw_arrow(draw, (300, 850), (300, 895), INK, 5)
    draw_arrow(draw, (900, 850), (900, 895), ORANGE, 5)
    draw_arrow(draw, (1495, 850), (1495, 895), GRAY, 5)
    image.save(path, quality=95)


def make_service_stack_diagram(path: Path) -> None:
    width, height = 1800, 1180
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "ATMOS-AWARE SERVICE STACK", font=image_font(36, True),
              fill=pil_rgb(WHITE))

    layers = [
        ("SERVICE API AND SLO", "OpenAI-compatible API | admission | quality | TTFT / TPOT | tenant policy", NAVY),
        ("MODEL AND COMPILATION", "Checkpoint + tokenizer | router semantics | ExpertPack compiler | numerical oracle", BLUE),
        ("LLM EXECUTION", "vLLM / PyTorch | continuous batching | KV cache | attention | GPU collectives", TEAL),
        ("HETEROGENEOUS EXPERT LAYER", "expert ownership | token grouping | placement epoch | fan-out / fan-in | fallback", GREEN),
        ("TRANSPORT AND DEVICE", "registered buffers | DMA queues | fences | timeout / cancel | driver | firmware | NPU", ORANGE),
        ("PLATFORM OPERATIONS", "Kubernetes allocation | operator | rollout / rollback | health | recovery | compatibility manifest", GRAY),
    ]
    top = 130
    layer_height = 130
    for index, (title, description, color) in enumerate(layers):
        y1 = top + index * (layer_height + 20)
        y2 = y1 + layer_height
        draw.rounded_rectangle((70, y1, 430, y2), radius=18, fill=pil_rgb(color))
        draw_wrapped(draw, (80, y1 + 5, 420, y2 - 5), title, image_font(26, True),
                     pil_rgb(WHITE))
        draw.rounded_rectangle((460, y1, 1730, y2), radius=18, fill=pil_rgb(LIGHT_GRAY),
                               outline=pil_rgb(color), width=4)
        draw_wrapped(draw, (490, y1 + 5, 1700, y2 - 5), description,
                     image_font(25), pil_rgb(INK), align="left")
        if index < len(layers) - 1:
            draw_arrow(draw, (900, y2), (900, y2 + 20), color, 5)

    draw.rounded_rectangle((70, 1050, 1730, 1145), radius=18, fill=pil_rgb(PALE_RED),
                           outline=pil_rgb(RED), width=3)
    draw_wrapped(
        draw,
        (90, 1055, 1710, 1140),
        "Cross-cutting contracts: request identity | telemetry | security | compatibility | reproducible evidence",
        image_font(25, True),
        pil_rgb(RED),
    )
    image.save(path, quality=95)


def make_request_lifecycle_diagram(path: Path) -> None:
    width, height = 1800, 1010
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "ONE MoE LAYER: SYNCHRONOUS REQUEST LIFECYCLE",
              font=image_font(34, True), fill=pil_rgb(WHITE))

    legend = [("GPU", BLUE), ("EXPERT LAYER", TEAL), ("TRANSPORT", ORANGE),
              ("ATMOS DEVICE", GREEN)]
    legend_x = 90
    for label, color in legend:
        draw.rounded_rectangle((legend_x, 120, legend_x + 330, 175), radius=12,
                               fill=pil_rgb(color))
        draw_wrapped(draw, (legend_x + 5, 122, legend_x + 325, 172), label,
                     image_font(20, True), pil_rgb(WHITE))
        legend_x += 410

    x_positions = [90, 675, 1260]
    y_positions = [220, 465, 710]
    box_width = 450
    box_height = 145
    step_boxes = {
        1: (x_positions[0], y_positions[0], x_positions[0] + box_width, y_positions[0] + box_height),
        2: (x_positions[1], y_positions[0], x_positions[1] + box_width, y_positions[0] + box_height),
        3: (x_positions[2], y_positions[0], x_positions[2] + box_width, y_positions[0] + box_height),
        4: (x_positions[2], y_positions[1], x_positions[2] + box_width, y_positions[1] + box_height),
        5: (x_positions[1], y_positions[1], x_positions[1] + box_width, y_positions[1] + box_height),
        6: (x_positions[0], y_positions[1], x_positions[0] + box_width, y_positions[1] + box_height),
        7: (x_positions[0], y_positions[2], x_positions[0] + box_width, y_positions[2] + box_height),
        8: (x_positions[1], y_positions[2], x_positions[1] + box_width, y_positions[2] + box_height),
        9: (x_positions[2], y_positions[2], x_positions[2] + box_width, y_positions[2] + box_height),
    }
    step_content = {
        1: ("1  Router + top-k", BLUE),
        2: ("2  Resolve placement epoch", TEAL),
        3: ("3  Group + pack", TEAL),
        4: ("4  DMA + admit", ORANGE),
        5: ("5  HBF tiles + NPU FFN", GREEN),
        6: ("6  Complete + status", GREEN),
        7: ("7  Return + fence", ORANGE),
        8: ("8  Restore + fan-in", TEAL),
        9: ("9  Combine + continue", BLUE),
    }
    for step in range(1, 10):
        label, color = step_content[step]
        draw_box(draw, step_boxes[step], label, color, font_size=25)

    for first, second in ((1, 2), (2, 3), (4, 5), (5, 6), (7, 8), (8, 9)):
        first_box = step_boxes[first]
        second_box = step_boxes[second]
        if second_box[0] > first_box[0]:
            start = (first_box[2], (first_box[1] + first_box[3]) // 2)
            end = (second_box[0] - 10, (second_box[1] + second_box[3]) // 2)
        else:
            start = (first_box[0], (first_box[1] + first_box[3]) // 2)
            end = (second_box[2] + 10, (second_box[1] + second_box[3]) // 2)
        draw_arrow(draw, start, end, step_content[first][1], 6)
    draw_arrow(draw, ((step_boxes[3][0] + step_boxes[3][2]) // 2, step_boxes[3][3]),
               ((step_boxes[4][0] + step_boxes[4][2]) // 2, step_boxes[4][1] - 10),
               ORANGE, 6)
    draw_arrow(draw, ((step_boxes[6][0] + step_boxes[6][2]) // 2, step_boxes[6][3]),
               ((step_boxes[7][0] + step_boxes[7][2]) // 2, step_boxes[7][1] - 10),
               GREEN, 6)

    draw.rounded_rectangle((90, 895, 1710, 975), radius=15, fill=pil_rgb(PALE_RED),
                           outline=pil_rgb(RED), width=3)
    draw_wrapped(
        draw,
        (110, 900, 1690, 970),
        "Critical path closes only after every required top-k contribution is valid or the request follows an explicit fallback policy.",
        image_font(22, True),
        pil_rgb(RED),
    )
    image.save(path, quality=95)


def make_integration_options_diagram(path: Path) -> None:
    width, height = 1800, 1030
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "PHYSICAL INTEGRATION PATTERNS", font=image_font(36, True),
              fill=pil_rgb(WHITE))

    panels = [
        (60, 145, 570, 915, "A  OEM-INTEGRATED", TEAL,
         "Preferred for synchronous MoE", [
             ("GPU compute tray", BLUE),
             ("Host PCIe root", NAVY),
             ("Balanced switch", TEAL),
             ("4 x ATMOS E3.S", GREEN),
         ], "Lowest path length; highest OEM NRE"),
        (645, 145, 1155, 915, "B  COMPANION TRAY", ORANGE,
         "Fallback for mechanics / cooling", [
             ("GPU compute tray", BLUE),
             ("Approved cabled PCIe", ORANGE),
             ("ATMOS tray switch", TEAL),
             ("4-8 x ATMOS E3.S", GREEN),
         ], "Independent service; cable/reset qualification"),
        (1230, 145, 1740, 915, "C  NETWORK SERVICE", GRAY,
         "Asynchronous services only", [
             ("GPU racks", BLUE),
             ("Ethernet / InfiniBand", GRAY),
             ("ATMOS service rack", GREEN),
             ("Retrieval / preprocessing", INK),
         ], "Easy placement; not layer-synchronous baseline"),
    ]
    for x1, y1, x2, y2, title, color, subtitle, boxes, footer in panels:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=pil_rgb(LIGHT_GRAY),
                               outline=pil_rgb(color), width=4)
        draw.rounded_rectangle((x1, y1, x2, y1 + 82), radius=22, fill=pil_rgb(color))
        draw_wrapped(draw, (x1 + 10, y1 + 5, x2 - 10, y1 + 77), title,
                     image_font(28, True), pil_rgb(WHITE))
        draw_wrapped(draw, (x1 + 20, y1 + 95, x2 - 20, y1 + 155), subtitle,
                     image_font(22, True), pil_rgb(color))
        box_y = y1 + 190
        for index, (label, box_color) in enumerate(boxes):
            box = (x1 + 65, box_y + index * 125, x2 - 65, box_y + 90 + index * 125)
            draw_box(draw, box, label, box_color, font_size=23)
            if index < len(boxes) - 1:
                draw_arrow(draw, ((x1 + x2) // 2, box[3]),
                           ((x1 + x2) // 2, box[3] + 35), box_color, 5)
        draw.rounded_rectangle((x1 + 25, y2 - 105, x2 - 25, y2 - 25), radius=14,
                               fill=pil_rgb(WHITE), outline=pil_rgb(color), width=2)
        draw_wrapped(draw, (x1 + 35, y2 - 100, x2 - 35, y2 - 30), footer,
                     image_font(20, True), pil_rgb(INK))
    image.save(path, quality=95)


def make_value_chain_diagram(path: Path) -> None:
    width, height = 1800, 880
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "HOW ATMOS CAN CREATE ECONOMIC VALUE",
              font=image_font(36, True), fill=pil_rgb(WHITE))

    columns = [
        (70, "ARCHITECTURE", [
            ("Experts persist in HBF", GREEN),
            ("NPU executes local FFNs", GREEN),
            ("Compact activations cross PCIe", TEAL),
        ]),
        (630, "MEASURED SYSTEM EFFECT", [
            ("GPU GDDR released", BLUE),
            ("Less remote expert traffic", BLUE),
            ("More accepted work at SLO", BLUE),
        ]),
        (1190, "REMOVABLE COST", [
            ("GPU trays / replicas", ORANGE),
            ("Network / rack / energy", ORANGE),
            ("Or more service per rack", ORANGE),
        ]),
    ]
    for x, title, items in columns:
        draw.text((x, 135), title, font=image_font(25, True), fill=pil_rgb(NAVY))
        for index, (label, color) in enumerate(items):
            y1 = 205 + index * 160
            draw_box(draw, (x, y1, x + 470, y1 + 110), label, color, font_size=24)
    for y in (260, 420, 580):
        draw_arrow(draw, (545, y), (620, y), TEAL, 7)
        draw_arrow(draw, (1105, y), (1180, y), ORANGE, 7)

    draw.rounded_rectangle((70, 730, 1660, 835), radius=18, fill=pil_rgb(PALE_RED),
                           outline=pil_rgb(RED), width=3)
    draw_wrapped(
        draw,
        (90, 740, 1640, 825),
        "Credit value only when the same checkpoint, quality, TTFT, TPOT, availability, and demand are preserved and a real cost or capacity constraint changes.",
        image_font(25, True),
        pil_rgb(RED),
    )
    image.save(path, quality=95)


def make_poc_roadmap_diagram(path: Path) -> None:
    width, height = 1800, 960
    image = Image.new("RGB", (width, height), pil_rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 92), fill=pil_rgb(NAVY))
    draw.text((55, 24), "EVIDENCE LADDER: FROM BASELINE TO RACK PRODUCT",
              font=image_font(34, True), fill=pil_rgb(WHITE))

    phases = [
        ("P0", "Freeze baseline", NAVY),
        ("P1", "Expert primitive", BLUE),
        ("P2", "GPU-ATMOS loop", TEAL),
        ("P3", "4-device tray", GREEN),
        ("P4", "8+8 service", ORANGE),
        ("P5", "Named OEM path", GRAY),
        ("P6", "Rack product", RED),
    ]
    x_positions = [70, 315, 560, 805, 1050, 1295, 1540]
    y1, y2 = 205, 390
    for index, ((code, label, color), x) in enumerate(zip(phases, x_positions)):
        draw_box(draw, (x, y1, x + 190, y2), f"{code}\n{label}", color, font_size=24)
        if index < len(phases) - 1:
            draw_arrow(draw, (x + 190, (y1 + y2) // 2),
                       (x_positions[index + 1] - 10, (y1 + y2) // 2), color, 6)

    draw.text((70, 465), "EACH PHASE RETURNS A DECISION PACKAGE",
              font=image_font(26, True), fill=pil_rgb(NAVY))
    evidence = [
        ("Configuration manifest", BLUE),
        ("Raw traces + analysis", TEAL),
        ("Quality + SLO result", GREEN),
        ("Fault / recovery result", ORANGE),
        ("Cost + support update", GRAY),
        ("Pass / conditional / pivot / stop", RED),
    ]
    for index, (label, color) in enumerate(evidence):
        row = index // 3
        column = index % 3
        x1 = 70 + column * 570
        y = 525 + row * 130
        draw_box(draw, (x1, y, x1 + 500, y + 90), label, color, font_size=22)

    draw.rounded_rectangle((70, 815, 1730, 915), radius=18, fill=pil_rgb(PALE_ORANGE),
                           outline=pil_rgb(ORANGE), width=3)
    draw_wrapped(
        draw,
        (90, 825, 1710, 905),
        "OEM discovery starts at P0, but custom platform NRE starts only after the P2 latency gate. P5 can then run in parallel with P3/P4 engineering.",
        image_font(24, True),
        pil_rgb(INK),
    )
    image.save(path, quality=95)


def build_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "architecture": ASSET_DIR / "reference_architecture.png",
        "service_stack": ASSET_DIR / "service_stack.png",
        "request_lifecycle": ASSET_DIR / "request_lifecycle.png",
        "integration_options": ASSET_DIR / "integration_options.png",
        "value_chain": ASSET_DIR / "value_chain.png",
        "poc_roadmap": ASSET_DIR / "poc_roadmap.png",
    }
    make_architecture_diagram(assets["architecture"])
    make_service_stack_diagram(assets["service_stack"])
    make_request_lifecycle_diagram(assets["request_lifecycle"])
    make_integration_options_diagram(assets["integration_options"])
    make_value_chain_diagram(assets["value_chain"])
    make_poc_roadmap_diagram(assets["poc_roadmap"])
    return assets


def add_figure(document: Document, path: Path, caption: str, width: float = 7.0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    paragraph = document.add_paragraph(caption, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_cover(document: Document, assets: dict[str, Path]) -> None:
    add_kicker(document, "PRODUCT ARCHITECTURE | SERVICE STACK | OEM PRODUCTIZATION")
    document.add_paragraph("ATMOS E3.S Expert Tier for\nRack-Scale MoE Inference", style="Title")
    document.add_paragraph(
        "A product proposal for moving routed experts from GPU memory to persistent, "
        "near-data compute while preserving the native accelerator platform",
        style="Subtitle",
    )
    add_callout(
        document,
        "DECISION",
        "Fund a bounded feasibility and OEM co-design program. Product commitment is earned "
        "only when complete GPU-to-ATMOS-to-GPU latency, service resilience, and rack economics "
        "pass on a named platform.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )
    add_figure(
        document,
        assets["architecture"],
        "Conceptual production direction: a balanced ATMOS PCIe branch adjacent to each GPU compute tray. "
        "The GPU scale-up fabric remains unchanged.",
        width=7.05,
    )
    add_table(
        document,
        ["Proposal scope", "Initial qualification", "Production target"],
        [[
            "MoE inference only; routed-expert FFNs",
            "One checkpoint, one FP8 format, host-staged baseline",
            "Named OEM rack with supported local or companion PCIe path",
        ]],
        widths=[2.25, 2.35, 2.40],
        font_size=8.4,
        header_fill=TEAL,
        banded=False,
    )
    add_source_note(
        document,
        "Planning basis: ATMOS device values, platform choices, performance targets, schedules, and "
        "vendor fit remain subject to silicon, software, OEM, and commercial qualification. "
        "Prepared 30 September 2026.",
    )


def add_executive_section(document: Document) -> None:
    add_page_break(document)
    add_kicker(document, "EXECUTIVE VIEW")
    add_heading(document, "The proposal in one page", 1)
    add_body(
        document,
        "Large Mixture-of-Experts models activate only a small fraction of their experts per token, "
        "yet conventional serving keeps the complete expert population in expensive accelerator memory "
        "or distributes it across accelerator ranks. ATMOS proposes a persistent expert tier: each E3.S "
        "module combines HBF, local working memory, an NPU, and a PCIe endpoint so complete routed experts "
        "can remain beside the compute that executes them. GPUs retain attention, KV cache, routing, shared "
        "layers, and the native collective domain."
    )
    add_heading(document, "Product thesis", 2)
    add_table(
        document,
        ["If ATMOS can...", "Then the service can...", "But value is earned only when..."],
        [[
            "Execute real routed FFNs at small batch with bounded p99 latency",
            "Release GPU memory for KV/concurrency or reduce expert-parallel scale-out",
            "Quality, TTFT, TPOT, availability, and accepted throughput match the frozen contract",
        ], [
            "Keep expert weights persistent and move compact activations/results",
            "Avoid request-time weight movement and reduce remote expert traffic",
            "Useful local bandwidth reaches compute and PCIe/queue overhead does not dominate",
        ], [
            "Operate as a supported tray-local or companion subsystem",
            "Deploy through a turnkey OEM with one lifecycle and support path",
            "A named rack beats its GPU-only control on three-year cost per accepted request/token",
        ]],
        widths=[2.33, 2.33, 2.34],
        font_size=8.0,
    )
    add_heading(document, "The decisive constraint", 2)
    add_callout(
        document,
        "FIRST PROOF",
        "Capacity fit is necessary but not sufficient. The first proof is whether dispatch, transfer, "
        "queueing, expert execution, return, fan-in, and GPU continuation fit the per-layer latency "
        "budget under realistic top-k routing, batching, skew, and failure reserve.",
        fill=PALE_RED,
        accent=RED,
    )
    add_heading(document, "Proposed first product path", 2)
    add_bullets(document, [
        "Development vehicle: 8 RTX PRO-class GPUs plus 8 ATMOS E3.S devices, split across two balanced four-device PCIe domains for fault and NUMA study.",
        "First model: one Qwen3-235B-A22B-class checkpoint, one qualified FP8 representation, one pinned vLLM/PyTorch/CUDA stack, and host-staged transport as the correctness baseline.",
        "Production geometry: four GPUs plus four ATMOS devices per compute tray, using an OEM-supported balanced PCIe branch and documented GPU peer DMA only where the complete platform supports it.",
        "Availability baseline: complete service groups across independent failure domains; unique expert shards are never called highly available merely because they occupy separate devices.",
        "OEM route: co-design first with Giga Computing, Supermicro, and Penguin Solutions; use Dell or Lenovo as an enterprise lifecycle control. Every candidate must pass the same attachment and economics evidence request.",
    ])
    add_heading(document, "Document map", 2)
    add_table(
        document,
        ["Part", "Decision addressed", "Primary output"],
        [
            ["I", "What exactly is the ATMOS product architecture?", "Chosen data path, topology, capacity, failure, and claim boundaries"],
            ["II", "What changes in the complete serving stack?", "Component contracts, request lifecycle, release ownership, and operations"],
            ["III", "How does this become a supportable OEM platform?", "Hardware ledger, rack integration patterns, vendor strategy, and RFI"],
            ["IV", "Where can the economics work?", "Break-even method, TCO boundary, and value mechanisms"],
            ["V", "How will the hypothesis be proven?", "Detailed POC ladder, experiments, exit criteria, artifacts, and 90-day action plan"],
            ["VI", "How is the program governed and de-risked?", "Ownership, risks, security, release rules, and decision request"],
        ],
        widths=[0.55, 3.05, 3.40],
        font_size=8.1,
    )


def add_part_i(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(
        document,
        "I",
        "Reference product architecture",
        "Define one coherent product path before comparing vendors, switches, or headline capacity.",
    )
    add_heading(document, "1. Product definition and boundaries", 1)
    add_body(
        document,
        "ATMOS is an accelerator endpoint in an E3.S serviceable form factor, not an SSD and not a "
        "member of the GPU scale-up fabric. The product boundary includes the E3.S module, firmware, "
        "Linux driver, transport runtime, expert execution contract, orchestration integration, and "
        "platform compatibility manifest. A supported deployment also includes the selected switch or "
        "backplane, OEM BIOS/BMC behavior, power/cooling envelope, service procedure, and support route."
    )
    add_table(
        document,
        ["In scope", "Explicitly out of scope", "Future option after evidence"],
        [
            ["Layer-synchronous routed-expert FFNs for decoder MoE inference", "Replacing attention, KV cache, the learned router, or GPU collectives", "Grouped-layer or router/combine offload if fine-grained latency fails"],
            ["Host-staged transport on stock CUDA or ROCm stacks", "Proprietary GPU firmware/driver patches", "Documented direct GPU peer DMA on qualified platforms"],
            ["OEM-integrated or adjacent PCIe subsystem", "Treating ATMOS as an NVLink, NVSwitch, or UALink endpoint", "Native accelerator-fabric endpoint in a different silicon generation"],
            ["One qualified checkpoint and precision first", "A universal claim across all MoE models and topologies", "Additional checkpoints, AMD, Gen6, and larger groups after the baseline passes"],
        ],
        widths=[2.33, 2.33, 2.34],
        font_size=7.9,
    )
    add_heading(document, "2. Division of work", 1)
    add_table(
        document,
        ["Layer", "GPU platform retains", "ATMOS tier owns", "Invariant"],
        [
            ["Model semantics", "Tokenizer, checkpoint graph, learned router, top-k IDs and gate weights", "Compiled copies of assigned routed FFNs", "Same expert function and gate application as the reference"],
            ["Execution", "Attention, normalization, shared/dense layers, KV, GPU-local experts", "Assigned expert FFN matmuls, activation, and result production", "No silent approximation or missing expert"],
            ["Communication", "Native GPU collectives and scale-up fabric", "Activation dispatch/result return through registered PCIe buffers", "ATMOS is outside NCCL/RCCL rank identity"],
            ["Operations", "GPU operator and vendor stack", "ATMOS health, placement generation, packs, queues, reset, and telemetry", "One compatibility manifest binds the supported combination"],
        ],
        widths=[1.05, 2.05, 2.20, 1.70],
        font_size=7.8,
    )
    add_figure(
        document,
        assets["architecture"],
        "Figure 1. Preferred rack building block. Four ATMOS devices form a balanced PCIe branch adjacent "
        "to a four-GPU compute tray. A companion tray is the fallback when local mechanics or thermals do not fit.",
        width=7.05,
    )
    add_heading(document, "3. Topology decision", 1)
    add_table(
        document,
        ["Stage", "Topology", "Why this stage uses it", "Promotion condition"],
        [
            ["Primitive bring-up", "One direct ATMOS E3.S endpoint", "Removes switch effects and exposes endpoint/driver behavior", "Real expert correctness and complete round-trip service curve"],
            ["Tray emulator", "4 ATMOS behind one x16-to-4x4 switch domain", "Balanced Gen5 lane arithmetic and realistic four-GPU sharing", "Bounded p99, fault containment, all-device load, and serviceable reset"],
            ["8+8 development", "Two independent 4-device domains aligned to NUMA/GPU locality", "Avoids a single eight-device switch failure and 2:1 Gen5 uplink oversubscription", "Full-model accepted rate exceeds cost-adjusted break-even"],
            ["Production", "OEM-integrated 4+4 tray or approved PCIe companion tray", "Creates a named lifecycle, thermal design, and supported route", "Independent rack qualification against its own GPU-only control"],
        ],
        widths=[1.10, 1.75, 2.45, 1.70],
        font_size=7.8,
    )
    add_heading(document, "4. Critical execution path", 1)
    add_numbered(document, [
        "The GPU completes attention/shared work and the learned router emits expert IDs and gate weights.",
        "The heterogeneous expert layer resolves each expert against an immutable placement epoch and groups tokens by destination.",
        "Activations are packed into registered buffers; dependencies are recorded before DMA submission.",
        "ATMOS receives work, admits it to a bounded queue, stages the needed HBF tiles, and executes the expert FFN on the NPU.",
        "Results and completion status return to GPU-visible or host-staged buffers with explicit ordering and error state.",
        "The GPU restores token order, combines all top-k contributions, applies each learned gate exactly once, and continues the dependent layer.",
    ])
    add_callout(
        document,
        "LATENCY RULE",
        "Measure the complete continuation path, not PCIe bandwidth in isolation. At decode, small "
        "messages, queue submission, synchronization, slowest-destination fan-in, and GPU resumption can "
        "dominate even when aggregate link bandwidth appears sufficient.",
        fill=PALE_RED,
        accent=RED,
    )
    add_heading(document, "5. Capacity, placement, and availability", 1)
    add_body(
        document,
        "Nominal aggregate HBF is not usable service capacity. The placement plan reserves space for "
        "packing/scales, firmware, scratch, fragmentation, two pack generations during update, and the "
        "chosen failure policy. The first placement study must therefore publish three numbers: nominal "
        "capacity, application-usable capacity, and survivable capacity after the declared failure."
    )
    add_table(
        document,
        ["Policy", "Correctness behavior", "Capacity / latency consequence", "Initial posture"],
        [
            ["Two complete service groups", "Either group can serve the full model", "Highest capacity cost; simplest failover", "Conservative production baseline"],
            ["N+1 with rebuild", "Spare receives lost placement before readiness", "Requires spare capacity and measured rebuild/RTO", "Evaluate in tray and full-service POCs"],
            ["GPU fallback experts", "Missing ATMOS expert executes on GPU", "Consumes reclaimed GDDR and changes tail behavior", "Use only with explicit reserved copies"],
            ["Cross-node replica", "Remote healthy replica completes request", "Adds network path to layer-critical latency", "Not default for synchronous decode"],
            ["Fail closed", "No incorrect output when an expert is missing", "Complete service may become unavailable", "Mandatory safety behavior, not an HA strategy"],
        ],
        widths=[1.25, 2.00, 2.20, 1.55],
        font_size=7.7,
    )


def add_part_ii_foundation(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(
        document,
        "II",
        "Service stack impact",
        "Treat ATMOS as a maintained serving subsystem with explicit APIs, ownership, recovery, and release rules.",
    )
    add_heading(document, "1. Stack-wide change", 1)
    add_body(
        document,
        "The service-stack change is larger than a single vLLM callback. ATMOS introduces a second "
        "execution class, persistent expert ownership, asynchronous device queues, a new failure domain, "
        "and model artifacts that must be built, distributed, activated, rolled back, and observed with "
        "the GPU service. The design therefore separates model semantics, heterogeneous scheduling, "
        "transport, device execution, and platform operations so each boundary has an owner and test suite."
    )
    add_figure(
        document,
        assets["service_stack"],
        "Figure 2. ATMOS-aware service stack. The serving API remains stable while new model, scheduling, "
        "transport, device, and operational contracts are introduced below it.",
        width=7.0,
    )
    add_heading(document, "2. Component impact and ownership", 1)
    add_table(
        document,
        ["Component", "New responsibility", "Primary owner", "Definition of done"],
        [
            ["Expert Execution ABI", "Versioned tensor shapes, dtype/scales, expert identity, gate semantics, status, and golden outputs", "Sandisk model/runtime + NPU", "Bit/metric tolerance passes for every supported expert pack"],
            ["ExpertPack compiler", "Compile and sign resident expert artifacts with memory/scratch budget and compatibility metadata", "Sandisk + NPU partner", "Complete pack is reproducible, verified, and safely rollable"],
            ["vLLM integration", "Skip ATMOS-owned GPU weights; expose prepare/dispatch/finalize; preserve batching and APIs", "Sandisk serving team", "Pinned upstream/plugin build passes conformance and service tests"],
            ["Placement manager", "Map model/layer/expert to healthy owners; publish immutable epochs; control fan-out and replicas", "Sandisk runtime", "No request mixes epochs; skew and failure policies are measurable"],
            ["Transport runtime", "Registered buffers, queue submission, fences, cancel, timeout, retry, backpressure, and cleanup", "Sandisk runtime", "No stale DMA or ambiguous completion across error/reset"],
            ["Linux driver + firmware", "Discovery, isolation, DMA, MSI-X, telemetry, FLR/reset, secure update, and invalidation", "Sandisk silicon/firmware", "Kernel/platform qualification and negative tests pass"],
            ["NPU execution runtime", "Queue admission, tile movement, kernels, completion, and useful-bandwidth counters", "NPU partner + Sandisk", "Measured service curve meets the frozen expert envelope"],
            ["Kubernetes allocation", "Topology-aware GPU+ATMOS allocation and health without replacing the GPU operator", "Sandisk platform", "Qualified resource group is schedulable and fails closed"],
            ["ATMOS operator", "Pack staging, readiness, rollout, drain, rollback, repair, and compatibility enforcement", "Sandisk platform/SRE", "Generation-safe lifecycle and declared recovery objectives pass"],
            ["Observability", "Correlated request/layer/token/expert IDs; queue, DMA, NPU, quality, power, errors, and SLOs", "Performance + SRE", "One trace explains p99 across GPU, host, switch, and ATMOS"],
        ],
        widths=[1.25, 2.85, 1.25, 1.65],
        font_size=7.15,
    )
    add_heading(document, "3. Sandisk and Tenstorrent integration boundary", 1)
    add_body(
        document,
        "Tenstorrent's public TT-Metal/TT-NN stack, standalone vLLM TT plugin, and inference-server "
        "workflows are useful references for NPU programming, model enablement, serving integration, "
        "benchmarking, and evaluation. The legacy Tenstorrent vLLM fork is deprecated in favor of the "
        "plugin route. None of those public projects establishes support for the ATMOS HBF controller, "
        "memory map, PCIe endpoint, E3.S packaging, ExpertPack format, or hybrid GPU-plus-ATMOS execution. "
        "The proposed vLLM layer is also different from a full Tenstorrent platform backend: the GPU "
        "continues to own attention and KV while only selected expert FFNs execute on ATMOS."
    )
    add_table(
        document,
        ["Boundary", "Sandisk owns", "Tenstorrent / NPU partner delivers", "Joint acceptance"],
        [
            ["Functional / performance emulation", "ATMOS resource/topology contract, evidence adapter, workload traces, evidence labels", "Versioned functional and timing interfaces with stated accuracy and unsupported behavior", "Golden expert results, service-curve schema, calibration plan, pinned versions"],
            ["Compiler and packs", "Expert ABI, ExpertPack schema/signing, model ownership map, compatibility manifest", "Lowering, layouts, kernels, scratch/tile requirements, compiler diagnostics for the selected NPU", "Reproducible complete packs, shape/precision matrix, rollback and numerical conformance"],
            ["NPU runtime", "ATMOS queue/buffer contract, deadlines, cancellation, reset generation, cross-layer telemetry", "Device execution, local tile movement, kernel dispatch, completion and NPU counters", "Ordering, timeout ambiguity, safe retry/cleanup, measured useful HBF/NPU service"],
            ["Host and PCIe path", "HBF/controller firmware, endpoint/DMA, Linux driver, IOMMU/isolation, switch/OEM qualification", "NPU-facing command and memory requirements; no implied ownership of PCIe platform support", "One interface agreement spanning visibility, fences, completion, reset and fault ownership"],
            ["Serving integration", "Hybrid vLLM/PyTorch layer, GPU semantics, placement, batching, fan-in, fallback and release support", "Reusable plugin patterns or reviewed hooks where applicable; NPU model/runtime expertise", "Pinned upstream/plugin commits, bounded patch set, feature envelope and conformance tests"],
            ["Operations", "Kubernetes allocation, ATMOS operator, packs, readiness, recovery, support and security response", "NPU diagnostics and escalation artifacts under the signed scope", "Compatibility manifest, diagnostics bundle, release cadence and named escalation owners"],
        ],
        widths=[1.20, 2.10, 2.20, 1.50],
        font_size=6.9,
    )
    add_callout(
        document,
        "PARTNER RULE",
        "Confirm two separately accepted partner scopes: emulator/evidence interfaces and "
        "compiler-kernel-runtime enablement. Neither scope silently implies the other, and public model "
        "support on Tenstorrent products is not evidence that the model or operator works on ATMOS.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )
    add_heading(document, "4. Request lifecycle and concurrency model", 1)
    add_figure(
        document,
        assets["request_lifecycle"],
        "Figure 3. One layer-synchronous expert operation. Queueing and fan-in are part of the "
        "dependent path; successful DMA submission is not successful layer completion.",
        width=7.05,
    )
    add_body(
        document,
        "Every request carries a stable request ID, sequence ID, model/pack generation, placement "
        "epoch, layer, token range, expert IDs, gate weights, deadline, and cancellation generation. "
        "The expert layer may batch compatible work across requests, but it must never combine work "
        "from different model generations or restore results against a changed placement epoch."
    )
    add_table(
        document,
        ["Stage", "Required state", "Backpressure / failure behavior", "Key telemetry"],
        [
            ["Prepare", "Router output, deadline, placement epoch, output slots", "Reject or retain GPU-local path before ownership changes", "Tokens, selected experts, destinations, pack time"],
            ["Group", "Compatible model/layer/expert/precision and queue budget", "Cap group wait; preserve deadline ordering and fairness", "Tokens per expert, co-activation, group age"],
            ["Submit", "Registered buffers, dependency fence, queue credits", "No credit means bounded wait or declared fallback; never overwrite live buffers", "Queue depth, credit wait, bytes, submit latency"],
            ["Execute", "Validated pack generation and device readiness", "Timeout marks work indeterminate until firmware/driver resolves completion", "Admission, HBF bytes, NPU time, stalls, thermals"],
            ["Complete", "Status, result length, integrity tag, completion generation", "Late/stale completion is discarded; retry only when operation is declared idempotent", "DMA return, completion delay, errors/retries"],
            ["Finalize", "All required top-k contributions or explicit fallback result", "No partial expert set can silently continue", "Slowest contribution, fan-in, combine, continuation"],
        ],
        widths=[0.80, 2.10, 2.65, 1.45],
        font_size=7.35,
    )
    add_heading(document, "5. Contracts that keep the stack maintainable", 1)
    add_table(
        document,
        ["Contract", "Minimum fields and semantics", "Versioning rule", "Conformance evidence"],
        [
            ["Expert Execution ABI", "Hidden/output shape, dtype, scales, activation, expert ID, gate ownership, accumulation, scratch, status", "Major for semantic/layout break; minor for optional capability", "Golden vectors across batch/shape/precision and error paths"],
            ["ExpertPack", "Checkpoint hash, layer/expert map, weight layout, scales, kernel IDs, memory/scratch budget, target firmware", "Immutable content-addressed artifact signed before staging", "Compiler reproducibility, pack integrity, complete coverage, rollback"],
            ["Transport API", "Register/unregister, submit, dependencies, completion, cancel, timeout, retry class, reset invalidation, telemetry", "No buffer survives a generation/reset without re-registration", "Stress, race, reset-during-DMA, stale completion, isolation"],
            ["Placement API", "Owner set, replica role, affinity, epoch, readiness, fallback, drain state", "Epoch is immutable for an in-flight request", "Trace replay, skew, failure, rolling update, no mixed epoch"],
            ["Compatibility manifest", "Device/FW/driver, switch, BIOS, kernel, container, vLLM/PyTorch, pack, model, topology, known limits", "One release ID for every supported combination", "Fresh-install reproduction plus upgrade/rollback suite"],
        ],
        widths=[1.15, 2.75, 1.60, 1.50],
        font_size=7.2,
    )
    add_heading(document, "6. vLLM and PyTorch integration impact", 1)
    add_body(
        document,
        "The preferred implementation is a maintained plugin/backend and narrow reviewed upstream "
        "changes, not an indefinite private fork. The integration owns heterogeneous expert allocation "
        "and dispatch; it does not replace vLLM admission, attention, KV management, or distributed GPU "
        "execution. If a required hook is absent, the team must choose among upstreaming it, maintaining "
        "a bounded patch set with release tests, or narrowing the supported feature envelope."
    )
    add_table(
        document,
        ["Serving feature", "ATMOS impact", "Initial support posture", "Test obligation"],
        [
            ["Continuous batching", "Tokens from several requests can improve expert batch reuse but create deadline/fairness pressure", "Required", "Mixed prompt/decode, churn, cancellation, saturation, tenant fairness"],
            ["Prefix caching", "KV remains GPU-side; cache hits change arrival bursts into the expert tier", "Supported after baseline", "Hit/miss mixes and burst service curves; no cache semantic change"],
            ["Speculative decoding", "Draft/verify changes token bursts and cancellation of rejected work", "Disabled in first baseline", "Enable only after cancellation and wasted-work accounting pass"],
            ["CUDA/HIP graph capture", "External asynchronous queues and dynamic fan-out may break static capture assumptions", "Exclude ATMOS path from capture initially", "Correct synchronization and measured graph/non-graph delta"],
            ["Tensor/pipeline parallelism", "GPU ranks remain native; each stage needs explicit local/remote ATMOS ownership", "One fixed topology first", "Rank placement, collective ordering, no deadlock, stage imbalance"],
            ["Expert parallelism", "Some GPU EP ranks may disappear or retain hot experts; collective groups must be rebuilt", "Hybrid placement experiment", "Equivalent model semantics and lower total dispatch/combine cost"],
            ["Quantization", "ATMOS packs must match the chosen scale/layout and GPU combination precision", "One FP8 format first", "Layer/expert error, end-task quality, accumulation and overflow"],
            ["LoRA / adapters", "Expert weight overlays complicate resident pack identity and update", "Deferred unless customer-critical", "No mixed base/adapter generation; memory and compile impact"],
        ],
        widths=[1.25, 2.45, 1.45, 1.85],
        font_size=7.15,
    )
    add_heading(document, "7. Expert placement and batching policy", 1)
    add_body(
        document,
        "Placement is a deployment policy informed by measured router traces, not a second learned "
        "router. The neural router remains unchanged. The policy minimizes expected destination count, "
        "slowest-device fan-in, queue imbalance, and failure exposure while respecting capacity, replica, "
        "thermal, and locality constraints. It is recomputed offline or during a controlled drain; an "
        "in-flight request never follows a partially published placement."
    )
    add_table(
        document,
        ["Metric", "Why it matters", "Decision it drives"],
        [
            ["Mean and p95 ATMOS destinations per token/layer", "Top-k experts on many devices increase fan-out and slowest-device dependency", "Co-locate commonly co-activated experts; retain selected GPU-local experts"],
            ["Expert co-activation matrix", "Individual popularity misses pair/group routing structure", "Place groups, choose replicas, and predict switch traffic"],
            ["Queue imbalance and oldest-work age", "Average utilization can hide head-of-line blocking", "Change hashing, queue policy, replicas, or admission"],
            ["Slowest contribution by device/expert", "GPU continuation waits for the last required result", "Identify hot expert, thermal, link, or kernel tail"],
            ["GPU-local layer/expert percentage", "Hybrid value and latency depend on what stays local", "Choose full, cold-only, prefill-only, or grouped-layer offload"],
            ["Surviving complete-service capacity", "Separate devices do not imply a complete model after failure", "Size replicas/spares and fleet reserve"],
        ],
        widths=[2.05, 2.55, 2.40],
        font_size=7.4,
    )
    add_heading(document, "8. Model and ExpertPack lifecycle", 1)
    add_numbered(document, [
        "Build: pin checkpoint, tokenizer, model code, precision recipe, compiler, kernel library, and target device/firmware; generate signed content-addressed packs and golden outputs.",
        "Stage: distribute the complete next generation to every required replica without changing serving ownership; verify capacity, integrity, temperature, and firmware compatibility.",
        "Activate: publish readiness only when the full service group has the pack, then create one new placement epoch and route new requests to it.",
        "Drain: old-epoch requests complete or expire; cancellation and retries preserve their original generation semantics.",
        "Retire: remove the old pack only after no live request, fallback, or rollback policy references it; retain audit metadata and reproducible build inputs.",
        "Rollback: stop new activation, restore the prior placement/pack generation, and prove no partially updated expert bank can report ready.",
    ])
    add_callout(
        document,
        "READINESS RULE",
        "A healthy device is not necessarily a ready model service. Readiness requires a complete, "
        "compatible expert generation, valid placement, sufficient surviving capacity, and a tested "
        "fallback/recovery policy.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )
    add_heading(document, "9. Error, cancellation, and recovery semantics", 1)
    add_table(
        document,
        ["Event", "Immediate behavior", "Recovery owner", "Service consequence"],
        [
            ["Client cancellation", "Mark request generation cancelled; stop unsent work; discard late completions safely", "vLLM integration + runtime", "Capacity is reclaimed; no result enters another request"],
            ["Queue deadline exceeded", "Do not submit stale work; choose declared GPU/replica fallback or fail request", "Expert layer", "Deadline outcome is visible in accepted-rate accounting"],
            ["Device timeout", "Quarantine affected queue; resolve completion ambiguity before buffer reuse", "Driver/firmware + runtime", "Retry only if operation class is safe and deadline permits"],
            ["Endpoint reset / FLR", "Invalidate registrations, queues, pack readiness, and placement ownership", "Driver/operator", "Device rejoins only after discovery, self-test, pack verify, and new epoch"],
            ["Switch containment event", "DPC/AER isolates the path; unaffected domain remains available if topology permits", "OEM + switch + driver", "Availability policy decides failover, degradation, or service stop"],
            ["Pack integrity failure", "Fail closed and withdraw device/group readiness", "Operator + release", "No request can use an incomplete or unsigned generation"],
            ["Thermal/power throttle", "Advertise reduced service curve or withdraw readiness before queues diverge", "Firmware/BMC + scheduler", "Admission follows qualified degraded envelope, not nominal capacity"],
        ],
        widths=[1.25, 2.45, 1.45, 1.85],
        font_size=7.1,
    )
    add_heading(document, "10. Kubernetes and platform operations", 1)
    add_body(
        document,
        "The first release allocates whole ATMOS devices or prequalified complete service groups. "
        "Fractional sharing is deferred until DMA isolation, queue QoS, memory protection, accounting, "
        "and noisy-neighbor behavior are proven. The ATMOS device plugin or Dynamic Resource Allocation "
        "driver publishes health and topology; the ATMOS operator owns model readiness and recovery. "
        "The existing NVIDIA or AMD GPU operator remains the authority for GPUs."
    )
    add_table(
        document,
        ["Operational object", "Published state", "Controller action", "Failure boundary"],
        [
            ["ATMOS device", "Hardware/FW, NUMA/root, switch domain, health, thermals, capacity", "Discover, allocate, quarantine, reset, update", "Endpoint"],
            ["ATMOS service group", "Complete model generation, placement epoch, replicas, degraded capacity", "Stage, activate, drain, failover, rebuild", "Switch/tray/node/group"],
            ["Workload claim", "Requested GPU/ATMOS topology, model generation, SLO class", "Co-schedule and admit only compatible resources", "Pod/service"],
            ["Compatibility manifest", "Approved OS/kernel/driver/FW/switch/runtime/model tuple", "Block unsupported combination and report drift", "Release"],
        ],
        widths=[1.35, 2.45, 1.85, 1.35],
        font_size=7.45,
    )
    add_heading(document, "11. Observability and evidence schema", 1)
    add_body(
        document,
        "One correlated trace must explain a slow token from API admission through GPU execution, "
        "expert dispatch, PCIe, queues, HBF/NPU service, fan-in, and continuation. Metrics without "
        "request/layer/expert context are useful for health but insufficient for architectural decisions."
    )
    add_table(
        document,
        ["Plane", "Required measures", "Decision use"],
        [
            ["Service", "Offered/accepted/rejected RPS, active sequences, TTFT, TPOT, output rate, quality", "Prove equal-SLO comparison and economic numerator"],
            ["Routing", "Top-k, destinations/token, tokens/expert, co-activation, GPU-local share", "Placement and batching policy"],
            ["Transport", "Logical/physical bytes, registration, pack, submit, DMA, fence, completion, retries", "Separate protocol/queue cost from link utilization"],
            ["Device", "Queue age/depth, useful HBF bytes, LPDDR/SRAM traffic, NPU/kernel time, stalls, power/thermal", "Build expert service curves and detect throttling"],
            ["Reliability", "AER/DPC, timeouts, resets, pack failures, rebuild, degraded duration, lost/retried work", "Validate containment, RTO, and failure reserve"],
            ["Cost", "GPU/ATMOS utilization, energy, rack/network occupancy, operator effort, spares", "Populate three-year TCO per accepted token/request"],
        ],
        widths=[1.05, 3.85, 2.10],
        font_size=7.4,
    )
    add_heading(document, "12. Security and tenancy boundary", 1)
    add_bullets(document, [
        "Use bounded registered memory; ATMOS must never DMA to arbitrary host or GPU addresses. Address translation, IOMMU policy, ACS, ATS/PASID use, and peer apertures are platform-qualified capabilities.",
        "Bind every queue, buffer, pack, key, tenant, and completion to an ownership generation. Reset and process exit revoke access before memory is reused.",
        "Verify firmware and ExpertPack authenticity, secure update/rollback policy, measured boot evidence where available, and auditable compatibility-manifest identity.",
        "Zero or cryptographically invalidate tenant activation/result buffers and device scratch according to the data classification. Persistent expert weights follow model-IP custody policy.",
        "Test malformed descriptors, stale handles, out-of-range DMA, completion spoofing, queue exhaustion, reset races, cross-tenant access, and denial-of-service limits.",
        "Do not disable IOMMU, ACS, DPC, secure boot, or isolation controls to obtain a benchmark result. A faster unsupported route is not a product path.",
    ])
    add_heading(document, "13. NVIDIA and AMD portability", 1)
    add_table(
        document,
        ["Layer", "Common ATMOS contract", "NVIDIA path", "AMD path"],
        [
            ["Model", "Same checkpoint semantics, ExpertPack identity, placement schema", "Pinned PyTorch/CUDA/vLLM build", "Pinned PyTorch/ROCm/vLLM build; independently qualified"],
            ["GPU collectives", "ATMOS remains outside GPU rank groups", "Stock NCCL / NVLink-NVSwitch", "Stock RCCL / UALink or Infinity Fabric as platform provides"],
            ["Baseline transport", "Registered host-pinned buffers and explicit synchronization", "CUDA copy/stream path", "HIP copy/stream path"],
            ["Optional peer path", "Same transport semantics; platform-specific registration backend", "Documented NVIDIA peer-memory/GPUDirect route only", "Documented ROCm/amdgpu peer-DMA route only"],
            ["Host architecture", "Same wire/queue contract and telemetry schema", "x86-64 or Grace Arm64 build as selected", "EPYC x86-64 build as selected"],
            ["Qualification", "Same workload, quality, SLO, evidence and TCO method", "Named NVIDIA OEM configuration", "Named AMD/OCP configuration"],
        ],
        widths=[1.05, 2.45, 1.75, 1.75],
        font_size=7.35,
    )
    add_heading(document, "14. Release and upstream strategy", 1)
    add_table(
        document,
        ["Artifact", "Release model", "Upstream / support posture", "Owner"],
        [
            ["Expert ABI + pack schema", "Sandisk versioned specification and conformance kit", "Stable public/partner boundary; backward compatibility declared", "Model/runtime architecture"],
            ["vLLM integration", "Plugin/backend plus smallest possible reviewed patch set", "Upstream generic hooks; maintain bounded backports for supported releases", "Serving team"],
            ["Transport library", "Sandisk SDK with stable C/C++/Python bindings as needed", "Independent of CUDA/ROCm behind platform adapters", "Runtime team"],
            ["Driver/firmware", "Signed matched release with kernel/OEM qualification", "Supported distro kernels; security response and rollback policy", "Silicon/firmware"],
            ["Kubernetes components", "Device/DRA driver and operator tied to manifest", "Use upstream APIs; own ATMOS recovery semantics", "Platform/SRE"],
            ["Evidence tooling", "Reproducible runner, schemas, dashboards, and raw result bundle", "Open internal contract; evidence class always labeled", "Performance/validation"],
        ],
        widths=[1.30, 2.10, 2.35, 1.25],
        font_size=7.35,
    )
    add_callout(
        document,
        "PART II EXIT",
        "The service stack is ready for full-model POC only when the ABI, pack, transport, placement, "
        "lifecycle, telemetry, and recovery contracts have named owners and automated conformance tests. "
        "A demo that returns correct tensors under one happy-path load is not this exit.",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def add_part_iii(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(
        document,
        "III",
        "Hardware and OEM productization",
        "Turn an accelerator concept into a named, serviceable platform with explicit electrical, thermal, firmware, and commercial ownership.",
    )
    add_heading(document, "1. ATMOS hardware planning ledger", 1)
    add_body(
        document,
        "The values below are program inputs for architecture and POC sizing, not achieved product "
        "specifications. The silicon team must replace each target with a revisioned measured capability. "
        "Aggregate HBF bandwidth is local to independent modules and cannot be advertised as one shared "
        "memory pool. Application-usable capacity is lower than nominal after firmware, pack metadata, "
        "scales, alignment, scratch, update generations, fragmentation, and failure reserve."
    )
    add_table(
        document,
        ["Parameter", "Per E3.S planning direction", "4-device tray", "8-device development system", "Qualification evidence"],
        [
            ["HBF capacity", "128 GB nominal", "512 GB nominal", "1,024 GB nominal", "Usable and survivable bytes by pack/revision"],
            ["Local HBF delivery", "200 GB/s target", "800 GB/s arithmetic sum", "1.6 TB/s arithmetic sum", "Useful NPU-delivered bytes by real expert shape"],
            ["Local working memory", "8-16 GB LPDDR plus SRAM", "32-64 GB LPDDR sum", "64-128 GB LPDDR sum", "Read/write contention, tiling, scratch, staging limit"],
            ["Host endpoint", "PCIe Gen5 x4 proposed baseline", "4 x x4 downstream", "8 x x4 across two domains", "Negotiated rate, TLP efficiency, DMA, reset, errors"],
            ["Module power", "30-40 W planning range", "120-160 W plus switch", "240-320 W plus two switches", "Sustained mixed-load power and throttle curve"],
            ["Compute", "NPU capability is shape/precision dependent", "Independent workers", "Independent workers", "Kernel service curve; no peak-TOPS substitution"],
            ["Form factor", "E3.S 2T direction to confirm", "Four supported positions", "Eight supported positions", "Mechanics, connector, retention, insertion, airflow"],
        ],
        widths=[1.10, 1.50, 1.25, 1.45, 1.70],
        font_size=6.9,
    )
    add_heading(document, "2. Bandwidth and latency hierarchy", 1)
    add_table(
        document,
        ["Boundary", "Planning ceiling", "What it does not prove", "Required measurement"],
        [
            ["HBF -> NPU", "200 GB/s local target/device", "Useful FFN rate at small batch or a workable LPDDR/SRAM tile path", "Weight bytes, reuse, stalls, kernel time, tokens/expert"],
            ["LPDDR staging", "Combined read/write path is implementation-specific", "That each HBF byte reaches compute without extra movement", "Directional traffic and contention under exact kernels"],
            ["PCIe Gen5 x4 endpoint", "About 15.8 GB/s per direction before protocol overhead", "Small-message latency, peer support, queue cost, or GPU continuation", "Payload efficiency and p50/p95/p99 by message and concurrency"],
            ["PCIe Gen5 x16 uplink", "About 63 GB/s per direction before protocol overhead", "Eight x4 devices are balanced; they are roughly 2:1 oversubscribed", "All-device traffic, fairness, congestion, root/NUMA path"],
            ["Scale-out network", "Platform-specific 400/800 Gb/s class", "Suitability for a dependency at every MoE layer", "End-to-end remote expert round trip and tail"],
        ],
        widths=[1.40, 1.55, 2.45, 1.60],
        font_size=7.2,
    )
    add_callout(
        document,
        "DESIGN CHOICE",
        "Use one balanced x16-to-4x4 domain for a four-device tray. Use two independent domains for "
        "eight devices unless a qualified Gen6 uplink and fault model demonstrably outperform them. "
        "Do not pay for Gen6 to compensate for an unmeasured endpoint or kernel bottleneck.",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def add_part_iv(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(
        document,
        "IV",
        "Economic case and product threshold",
        "Replace capacity-only savings claims with equal-service measurements, complete failure reserve, and quote-backed cost per accepted work unit.",
    )
    add_heading(document, "1. Value must follow a measured causal chain", 1)
    add_figure(
        document,
        assets["value_chain"],
        "Figure 5. ATMOS creates business value only when architectural effects survive end-to-end "
        "measurement and remove a real infrastructure cost or increase accepted service capacity.",
        width=7.05,
    )
    add_table(
        document,
        ["Potential value mechanism", "Technical proof", "Economic proof", "Do not count"],
        [
            ["Release GPU memory", "Measured routed-expert GDDR removed; remaining model/KV/workspace fits", "More admitted sequences or fewer complete GPU service units", "Free bytes with no accepted-capacity change"],
            ["Increase throughput", "Bottleneck decomposition shows ATMOS adds usable parallel expert capacity", "Higher accepted requests/tokens within the same quality/TTFT/TPOT", "Completions after SLO, offered load, or synthetic kernel TOPS"],
            ["Change model replica geometry", "One hybrid group matches the complete GPU-only group at SLO", "Fewer GPU trays/racks including equivalent reserve", "Capacity fit without compute/latency equivalence"],
            ["Reduce expert-parallel traffic", "Measured network ports/bytes/tail decrease for the fixed service", "Avoided NIC/switch/port/rack cost or released capacity", "Traffic reduction that cannot remove cost or improve accepted load"],
            ["Improve energy / density", "Wall/rack energy and cooling measured at accepted load", "Lower energy/cooling/facility cost per accepted unit", "Lower ATMOS watts while total service energy rises"],
            ["Operational simplification", "Fewer GPU units/fabrics without worse incident or upgrade behavior", "Measured labor, spares, downtime, or support reduction", "Narrative benefit without an owned operating process"],
        ],
        widths=[1.35, 2.10, 2.05, 1.50],
        font_size=7.15,
    )
    add_heading(document, "2. Model capacity screen", 1)
    add_body(
        document,
        "The capacity screen decides whether a checkpoint deserves performance work; it does not predict "
        "service viability. Raw routed-expert bytes exclude scales, packing, metadata, shared/dense tensors, "
        "scratch, two update generations, fragmentation, and replicas. The first qualification uses one "
        "Qwen3-235B-A22B-class checkpoint because it offers a meaningful top-8 expert bank without making "
        "the initial service depend on a near-full eight-device capacity envelope."
    )
    add_table(
        document,
        ["Model screen", "Published / chosen geometry", "Raw routed-expert estimate", "Program role"],
        [
            ["Qwen3-235B-A22B class", "235B total / about 22B active; 94 layers; 128 routed experts; top-8", "About 227B parameters: about 227 GB FP8 or 454 GB at 16-bit before overhead", "First complete-model qualification; leaves room to test pack/update/failure policy"],
            ["DeepSeek-V3 class", "About 671B total / 37B active; 61 layers, first 3 dense; 256 routed + 1 shared; top-8", "About 654B routed parameters: about 654 GB FP8 or 1.31 TB at 16-bit before overhead", "Capacity and placement stress case only after first path; FP8 may fit nominally but reserve is tight"],
        ],
        widths=[1.35, 2.25, 2.00, 1.40],
        font_size=7.4,
    )
    add_heading(document, "3. Equal-service comparison contract", 1)
    add_body(
        document,
        "GPU-only and hybrid arms use the same checkpoint and license, tokenizer, model code, precision "
        "quality tolerance, prompt/output distribution, arrival process, TTFT, TPOT, availability target, "
        "failure reserve, warmup, measurement window, and accepted-work definition. Each arm may be tuned "
        "only through a documented procedure available to that architecture. Unsupported settings or a "
        "different failure policy invalidate the comparison."
    )
    add_callout(
        document,
        "PRIMARY METRIC",
        "Three-year TCO per accepted request or output token within the frozen quality and latency SLO. "
        "Report both units when output lengths vary. Accepted work excludes errors, timeouts, cancellations "
        "caused by overload, and completions outside the SLO.",
        fill=PALE_TEAL,
        accent=TEAL,
    )
    add_heading(document, "4. Break-even equations", 1)
    add_table(
        document,
        ["Question", "Equation / method", "Interpretation"],
        [
            ["Does one hybrid unit beat one GPU unit?", "C_hybrid / R_hybrid < C_gpu / R_gpu", "C is fully burdened comparable cost; R is accepted service rate at equal SLO"],
            ["What rate uplift is required?", "R_hybrid / R_gpu > C_hybrid / C_gpu", "A 25% cost premium requires more than 25% accepted-rate uplift unless another cost disappears"],
            ["How many units serve demand D?", "N = reserve_policy + ceil(D / R_unit)", "Reserve is explicit complete-service capacity, not an arbitrary extra server"],
            ["Does a changed replica geometry win?", "TCO_hybrid_groups(D,SLO,HA) < TCO_gpu_groups(D,SLO,HA)", "Compare complete groups and surviving capacity, not raw device count"],
            ["What is ROI?", "ROI = (TCO_control - TCO_hybrid - migration_cost) / (hybrid investment + migration_cost)", "Report NPV/payback separately when timing and financing are known"],
        ],
        widths=[1.65, 2.65, 2.70],
        font_size=7.35,
    )
    add_heading(document, "5. Fully burdened three-year cost boundary", 1)
    add_table(
        document,
        ["Cost domain", "GPU-only control", "Hybrid additions / changes", "Evidence source"],
        [
            ["Compute hardware", "GPU trays/cards, host CPUs, DRAM, local storage", "ATMOS modules, switch/backplane, cables/retimers, displaced storage/NIC", "Comparable OEM quotes and BOM revisions"],
            ["Rack / network", "Racks, PDUs, NICs, fabric switches, optics/cables", "Companion tray or changed port/rack count; management network", "Named topology and quote"],
            ["Power / cooling", "Measured wall/rack energy and facility cooling", "ATMOS/switch energy, changed fan/CDU/PUE behavior", "Accepted-load power and facility factors"],
            ["Software / support", "GPU software, OS/platform, support and orchestration", "ATMOS runtime/operator, switch FW, model compile, OEM support SKU", "License/support terms and owned effort"],
            ["Availability", "Complete spare groups, spares, repair and downtime", "ATMOS replicas/spares/rebuild capacity and service stock", "Failure policy, RTO, failure/repair assumptions"],
            ["Engineering / NRE", "Baseline integration and qualification", "Silicon/board/OEM NRE, vLLM work, validation, security, upstream/backport", "Approved labor/NRE estimate amortized over units"],
            ["Operations", "Deployment, upgrades, monitoring, incidents, capacity management", "Additional firmware/pack lifecycle and support handoffs; possible fleet reduction", "Measured time, incident drills, support model"],
            ["Migration / risk", "No architecture migration", "Application qualification, rollout, dual-running, contingency, decommission", "Deployment plan and sensitivity"],
        ],
        widths=[1.15, 1.85, 2.70, 1.30],
        font_size=6.95,
    )
    add_heading(document, "6. Sensitivity and feasible-region reporting", 1)
    add_body(
        document,
        "A single favorable point is not an investment case. The economic workbook sweeps the variables "
        "that can reverse the result and shows the region where hybrid remains better. Each chart labels "
        "measured values, quotes, contractual values, and assumptions separately."
    )
    add_table(
        document,
        ["Sweep", "Minimum range", "Why it can reverse the result"],
        [
            ["Accepted service rate", "Measured confidence interval plus degraded/failure rate", "Fleet count changes discretely and p99 queueing can collapse rate"],
            ["ATMOS configured cost", "Prototype, volume, support, spares, NRE allocation", "A low module-only price omits the subsystem and lifecycle"],
            ["GPU/rack baseline", "Current quote, committed discount, next platform scenario", "ATMOS competes with improving GPU memory and inference throughput"],
            ["Failure reserve", "Two complete groups, N+1 rebuild, GPU fallback, remote replica", "Reserve can consume the apparent capacity advantage"],
            ["Demand and utilization", "Low/expected/peak growth with burst and seasonality", "Savings disappear when reduced unit count cannot be realized"],
            ["Energy / cooling", "Measured power and local facility rates/PUE", "Incremental tray power may trigger facility or cooling changes"],
            ["Support / operations", "Best/expected/worst owned effort and downtime", "A new device class can add operational cost even when hardware falls"],
        ],
        widths=[1.45, 2.50, 3.05],
        font_size=7.35,
    )
    add_heading(document, "7. Commercial thresholds and claim discipline", 1)
    add_bullets(document, [
        "Set the minimum product hurdle before hybrid results are visible. The sponsor may choose positive ROI, a minimum payback/NPV, a minimum TCO reduction, or strategic capacity headroom; record the selected hurdle in the baseline manifest.",
        "Separate feasibility, design win, funded POC, supported release, purchase commitment, and recognized revenue. None implies the next without its own evidence and owner.",
        "Publish platform-specific results. An RTX development result does not size NVIDIA NVL, AMD Helios, or another turnkey rack. Each named rack receives its own control, hybrid result, quote, reserve, and support assessment.",
        "Do not monetize theoretical HBF bandwidth, nominal capacity, peak NPU operations, avoided network bytes, or free GDDR until the complete service converts them into accepted work or removable cost.",
        "Retain negative outcomes. A checkpoint, topology, OEM, transport, or workload that misses the feasible region is a valid no-fit result and protects the program from repeated unsupported investment.",
    ])
    add_callout(
        document,
        "PART IV EXIT",
        "The economic case is ready for product review only after the full-service POC populates the "
        "accepted-rate denominator, the OEM returns a configured quote and support boundary, the chosen "
        "failure reserve is included, and sensitivity remains positive over the approved operating range.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )


def add_part_iii_tail(document: Document, assets: dict[str, Path]) -> None:
    add_heading(document, "3. PCIe switch and fabric selection", 1)
    add_body(
        document,
        "The switch is selected after compatibility, not by lane count alone. The first board should be "
        "an orderable engineering fixture where possible; custom board design begins only when an existing "
        "board fails a documented requirement. Microchip Switchtec PM50052-class and Broadcom PEX89048-class "
        "devices are the initial Gen5 inquiries because their public lane counts can support the target maps. "
        "Neither is ATMOS-qualified by publication."
    )
    add_table(
        document,
        ["Selection gate", "Pass evidence", "Why it matters"],
        [
            ["Port map", "Written x16 upstream + 4x4 or 8x4 downstream configuration and board routing", "Arithmetic and physical lanes must both exist"],
            ["Generic endpoint support", "ATMOS enumeration, BAR/address-window sizing, MSI-X, non-NVMe handling", "NVMe switch marketing is not accelerator compatibility"],
            ["DMA and peer route", "Counters prove payload path, simultaneous directions, ordering, completion, and supported security settings", "Enumeration alone does not prove useful transport"],
            ["Containment", "ACS/IOMMU policy, AER/DPC, downstream reset, malformed endpoint, link degradation, switch reset", "One endpoint cannot corrupt or reset the entire service silently"],
            ["Management", "Firmware profile, diagnostics, telemetry, update/rollback, secure access, ownership", "The switch becomes part of the released lifecycle"],
            ["Commercial", "Board/chip quote, FAE scope, NRE, lead time, spares, lifecycle/EOL, production owner", "A lab success without supply/support is not a product"],
        ],
        widths=[1.40, 3.45, 2.15],
        font_size=7.35,
    )
    add_heading(document, "4. Physical integration patterns", 1)
    add_figure(
        document,
        assets["integration_options"],
        "Figure 4. Three deployment patterns. Local OEM integration is preferred for layer-synchronous "
        "MoE; a PCIe companion tray is the mechanical fallback; a network service is reserved for "
        "asynchronous retrieval, preprocessing, or offline work unless measurements change the boundary.",
        width=7.05,
    )
    add_table(
        document,
        ["Pattern", "When to choose it", "Required proof", "Disqualifier"],
        [
            ["OEM-integrated local", "Named tray exposes supported root path, mechanics, power, cooling, and firmware ownership", "Lane/NUMA map, thermal design, BMC/reset, peer route, service procedure", "Assumed spare bays/lanes or unowned platform change"],
            ["Companion PCIe tray", "Local tray cannot fit ATMOS but OEM approves cabled expansion", "Cable/retimer SI, hot-plug/reset, latency, failure domains, power/cooling, rack service", "Unsupported generic expansion or ambiguous warranty boundary"],
            ["Network-attached service", "Work can be asynchronous, coarse grained, prefill/offline, retrieval, or feature processing", "Complete-service latency, retries, network reserve, replica and economics", "Layer-synchronous decode misses the signed budget"],
        ],
        widths=[1.35, 2.45, 2.25, 0.95],
        font_size=7.25,
    )
    add_heading(document, "5. OEM platform design record", 1)
    add_body(
        document,
        "A platform name is not a configuration. Every qualified target returns one signed design record "
        "covering the exact tray, board revisions, device population, firmware, software, cooling, and "
        "support route. Replacing the GPU, host CPU, PCIe generation, switch, retimer, BIOS, device power, "
        "or E3.S population triggers impact review and selective requalification."
    )
    add_table(
        document,
        ["Domain", "Mandatory returned artifact"],
        [
            ["Electrical / topology", "Root/NUMA map; connector and lane ownership; generations/widths; switch ports; retimers/cables; oversubscription; peer/root crossings"],
            ["Mechanical", "E3.S 2T locations; retention; insertion/removal; keep-out; cable routing; service access; effect on boot/cache/storage bays"],
            ["Power / thermal", "Per-module/switch allocation; connector/PDB limits; cold plate/heatsink/airflow; fan/CDU behavior; sensors; throttle/shutdown curve"],
            ["Firmware / security", "BIOS windows; IOMMU/ACS/ATS/PASID policy; AER/DPC; BMC inventory/telemetry; FLR/reset domains; secure update/rollback"],
            ["Software", "Supported OS/kernel/container; driver/runtime; CUDA/ROCm; vLLM/PyTorch; Kubernetes; model/pack; diagnostics and manifest"],
            ["Reliability / service", "Failure matrix; degraded mode; RTO/RPO; field replacement; spares; diagnostics; escalation; warranty; support ownership"],
            ["Commercial / supply", "Prototype and production BOM; NRE; qualification cost; MOQ; lead time; lifecycle/EOL notice; geographies; support SKU"],
        ],
        widths=[1.35, 5.65],
        font_size=7.55,
    )
    add_heading(document, "6. Turnkey vendor strategy", 1)
    add_body(
        document,
        "The ranking below measures preliminary ATMOS partnership fit, not general vendor quality or GPU "
        "performance. Scores are hypotheses based on public offerings as of 30 September 2026 and must be "
        "replaced by a common RFI, engineering workshop, quote, lab access, and measured POC. NVIDIA NVL "
        "and AMD Helios are accelerator/platform ecosystems; the ranked companies are potential owners of "
        "the surrounding integration, deployment, and lifecycle boundary."
    )
    add_table(
        document,
        ["Factor", "Weight", "A score of 5 requires"],
        [
            ["ROI potential", "35%", "A credible path to reduce accelerator/rack/network/energy/support cost for the frozen service, with quoteable components"],
            ["Integration ease", "30%", "A supported PCIe/mechanical/thermal/firmware path, engineering owner, and accessible lab or configurable platform"],
            ["Flexibility", "25%", "Choice of accelerator, CPU, network, cooling, rack/form factor, and co-design without proprietary GPU-stack modification"],
            ["Lifecycle readiness", "10%", "Factory/rack validation, deployment, firmware, spares, field service, and accountable global escalation"],
        ],
        widths=[1.25, 0.70, 5.05],
        font_size=7.5,
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.50)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    add_heading(document, "Pre-RFI ATMOS partner priority", 2)
    add_table(
        document,
        ["Rank", "Provider / offer", "ROI", "Ease", "Flex", "Life", "Score", "Why it is in this position", "Primary uncertainty"],
        [
            ["1", "Giga Computing GIGAPOD", "4.5", "4.5", "5.0", "4.0", "91.5", "Explicit turnkey rack-scale offer across NVIDIA, AMD, Intel; ORV3; air/DLC; bespoke facility integration", "ATMOS tray/device engineering and global field-support terms"],
            ["2", "Supermicro DCBBS / SuperCluster", "4.5", "4.5", "4.5", "4.5", "90.0", "Broad system/rack portfolio, rack integration, liquid cooling, networking/cabling, L11/L12 validation and deployment", "Willingness to modify a validated GPU design and own ATMOS lifecycle"],
            ["3", "Penguin Solutions OriginAI", "4.0", "4.0", "5.0", "4.5", "86.0", "Multi-vendor hardware choice with tailored design, factory test, deploy, manage, software, spares, and AMD/NVIDIA support", "Relies on selected server OEM for deep tray electrical change"],
            ["4", "Lenovo Hybrid AI Factory / Neptune", "4.0", "3.5", "4.0", "4.5", "78.0", "Validated hybrid AI platforms, services, global enterprise route, and mature liquid-cooling capability", "Exact PCIe/E3.S accelerator path and change-control flexibility"],
            ["5", "Dell AI Factory", "4.0", "3.5", "3.5", "5.0", "76.5", "NVIDIA all-in-one path, AMD open architecture, broad lifecycle/services, financing, and enterprise support", "Validated-design modification may increase NRE and qualification friction"],
            ["6", "Cisco Secure AI Factory", "3.5", "3.0", "4.0", "4.5", "71.5", "Full-stack AI PODs, UCS/network/security/observability, build-your-own option, Supermicro-enabled rack systems", "Compute-tray ownership is shared; Cisco is strongest above the endpoint layer"],
            ["7", "HPE Private Cloud AI / AI Factory", "3.5", "3.0", "3.0", "5.0", "67.5", "Turnkey full-stack operations, production lab, enterprise governance, GreenLake and global lifecycle", "Appliance validation and NVIDIA-centric stack may constrain low-level co-design"],
        ],
        widths=[0.45, 1.70, 0.42, 0.42, 0.42, 0.42, 0.55, 3.05, 2.25],
        font_size=6.2,
    )
    add_source_note(
        document,
        "Score = (35 x ROI + 30 x Ease + 25 x Flexibility + 10 x Lifecycle) / 5. Public evidence "
        "establishes vendor offering scope only; it does not establish ATMOS compatibility, price, "
        "availability, or commitment. A failed attachment gate overrides the paper score.",
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    add_heading(document, "7. Engagement waves and action", 1)
    add_table(
        document,
        ["Wave", "Who", "Purpose", "Required exit"],
        [
            ["1 - co-design learning", "Giga Computing, Supermicro, Penguin Solutions", "Find the quickest credible local or companion PCIe route and price its NRE/lifecycle", "One returned design record and lab plan from at least two providers"],
            ["2 - enterprise control", "Dell and Lenovo", "Compare support, qualification, installed-base, procurement, and global service friction", "Quote-backed enterprise path or documented disqualifier"],
            ["3 - stack integrators", "Cisco and HPE", "Assess secure AI factory, operations, financing/consumption, and companion-service integration", "Clear ownership boundary with the compute OEM"],
            ["Watchlist", "QCT/Quanta, Foxconn/Ingrasys, Wiwynn, ASUS", "Activate on geography, customer pull, existing relationship, or superior orderable fixture", "Direct procurement/support owner plus the same RFI evidence"],
            ["Alternative accelerators", "Intel Gaudi through OEM; Cerebras, Groq, SambaNova", "Test only a funded customer use case or documented extension point", "Separate software baseline and no assumption of synchronous ATMOS insertion"],
        ],
        widths=[1.15, 1.80, 2.35, 1.70],
        font_size=7.35,
    )
    add_heading(document, "8. Common OEM RFI and workshop package", 1)
    add_numbered(document, [
        "Send a controlled device envelope: form factor, link modes, BAR/address needs, DMA model, power/thermal range, firmware/update assumptions, debug needs, and prohibited claims.",
        "Provide two target geometries: a four-GPU/four-ATMOS local tray and a PCIe companion tray. Ask the vendor to return exact orderable host/rack configurations, not a family brochure.",
        "Request the complete platform design record defined above, including displaced storage/NIC resources and the use of any lanes claimed to be saved by switching.",
        "Ask for prototype and production quotes separately: hardware, cables/retimers/backplane, engineering, firmware, qualification, lab access, deployment, support, spares, lead time, and lifecycle.",
        "Run a joint architecture workshop ending with a configuration ID, named technical/support owners, unresolved-interface register, POC access plan, dates, and explicit no-fit items.",
        "Advance only vendors that accept the evidence method: frozen workload/SLO, security controls enabled, measured route, fault injection, fully burdened economics, and platform-specific claims.",
    ])
    add_callout(
        document,
        "PART III EXIT",
        "A vendor advances to rack POC only with a signed or formally returned design record, named "
        "engineering and field-support owners, an accessible platform, a measurable route, and a quote "
        "whose power/cooling/failure reserve still fits the product economics.",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def add_part_v(document: Document, assets: dict[str, Path]) -> None:
    add_part_opener(
        document,
        "V",
        "Integrated POC and qualification plan",
        "Test the architecture in the order that retires the largest risk before custom hardware, platform NRE, or product claims.",
    )
    add_figure(
        document,
        assets["poc_roadmap"],
        "Figure 6. Evidence ladder. Each phase returns a decision package; OEM discovery begins early, "
        "while custom platform spending remains gated by the complete round-trip latency result.",
        width=7.05,
    )
    add_heading(document, "1. Evidence classes and permitted conclusions", 1)
    add_table(
        document,
        ["Evidence class", "Useful for", "Cannot establish", "Promotion rule"],
        [
            ["A - analytic bound", "Capacity/link arithmetic, impossible-case screening, first latency budget", "Queueing, kernels, software overhead, p99, fault behavior", "Inputs and units reviewed; uncertainty visible"],
            ["B - discrete-event simulation", "Trace replay, placement, queues, topology, what-if sensitivity", "Real endpoint/driver/NPU timing or platform support", "Calibrate each modeled component against higher evidence when available"],
            ["C - partner functional/performance emulation", "ABI, pack, kernels, approximate service curves, software integration", "Production silicon thermal, full DMA, reset, OEM behavior", "Versioned partner evidence and error bars; label provisional"],
            ["D - prototype / FPGA / engineering endpoint", "Driver/DMA/queue/fault semantics and closer timing", "Production power, final memory/NPU, supply, complete rack", "Matched manifest and known deltas to production"],
            ["E - production-intent ATMOS silicon", "Real service curves, latency, power, reset, full service POC", "Named OEM rack lifecycle and rack economics", "Silicon/FW/driver revision frozen for qualification"],
            ["F - named OEM production-intent rack", "Platform support, thermals, faults, operations, accepted service, TCO", "Other vendors, generations, or configurations", "Independent control, support owner, quote, release manifest"],
        ],
        widths=[1.35, 1.85, 2.30, 1.50],
        font_size=7.25,
    )
    add_callout(
        document,
        "EVIDENCE RULE",
        "A lower evidence class may justify the next experiment or stop an impossible path. It cannot pass "
        "a gate whose decision depends on real silicon, physical routing, thermals, fault containment, "
        "support, or quote-backed economics.",
        fill=PALE_RED,
        accent=RED,
    )
    add_heading(document, "2. Rules for every POC run", 1)
    add_bullets(document, [
        "Assign a run ID and immutable manifest covering model/tokenizer, precision, packs, placement, hardware/firmware, switch, BIOS, kernel, drivers, container, vLLM/PyTorch, workload, SLO, fault policy, and evidence class.",
        "Freeze business quality and service SLO before viewing hybrid results. Report offered, admitted, accepted, rejected, failed, timed-out, cancelled, and out-of-SLO work separately.",
        "Use low, target, burst, and saturation loads. Separate request rate, outstanding clients, active sequences, batch size, and tokens per expert; they are not interchangeable concurrency measures.",
        "Run a warmup defined by stable caches/clocks/temperatures, then at least three independent measured repetitions. Report distributions and confidence intervals; preserve raw traces and all failures.",
        "Measure full-service p50/p95/p99 directly. Do not add component p99 values to claim end-to-end p99. Distinguish overlapping wait from the dependency critical path.",
        "Keep quality, security, IOMMU/ACS/DPC, error reporting, and failure reserve enabled. Unsupported tuning cannot become the winning configuration.",
        "Compare GPU-only and hybrid at equivalent quality, prompt/output distribution, SLO, availability, and tuning process. A changed checkpoint or failure policy is a new experiment.",
        "Record logical and physical bytes, useful and nominal bandwidth, energy at the wall/rack, queue backlog, retries, temperature/throttling, and resource occupancy throughout the run.",
    ])
    add_heading(document, "3. First qualified configuration", 1)
    add_table(
        document,
        ["Dimension", "Frozen starting point", "Deferred expansion"],
        [
            ["Checkpoint", "One Qwen3-235B-A22B-class checkpoint and tokenizer", "DeepSeek-V3 class and additional customer checkpoints"],
            ["Precision", "One partner/Sandisk-qualified FP8 representation and accumulation policy", "BF16, other FP8 variants, mixed expert precisions"],
            ["Serving stack", "Pinned vLLM + PyTorch + CUDA + Linux/container build", "ROCm/AMD, additional vLLM releases, alternate serving engines"],
            ["Transport", "Host-staged registered buffers", "Documented NVIDIA/AMD GPU peer DMA"],
            ["Endpoint / fabric", "Gen5 x4 device; one direct endpoint then balanced x16-to-4x4 domain", "Gen6, eight-device single domain, native accelerator fabric"],
            ["Features", "Continuous batching; fixed distributed topology; no speculative decoding or LoRA", "Speculation, adapters, broader graph capture, fractional tenancy"],
            ["Availability", "Fail closed plus complete replica/control policy declared", "N+1 rebuild, GPU fallback, remote replicas after measurement"],
        ],
        widths=[1.20, 3.35, 2.45],
        font_size=7.4,
    )
    add_heading(document, "4. First workload card", 1)
    add_body(
        document,
        "This card makes the first experiment executable without turning planning values into a customer "
        "or product promise. P0 replaces each provisional item with the sponsor-approved checkpoint, trace, "
        "quality set, SLO, demand, availability policy, and economic hurdle before comparative results are viewed."
    )
    add_table(
        document,
        ["Field", "Starting profile", "P0 freeze / evidence"],
        [
            ["Service", "Enterprise knowledge/claims assistant with document-grounded generation; no named customer architecture implied", "Authorized prompts/documents or approved synthetic equivalent; data/IP and publication rules"],
            ["Checkpoint", "Qwen3-235B-A22B Instruct-class candidate; exact 2507/base variant and license to be selected", "Checkpoint/tokenizer/model-code hashes; non-thinking/thinking mode; vLLM and NPU operator support"],
            ["Expert scope", "Routed FFNs only; learned router and gate unchanged; shared/dense layers, attention and KV stay on GPU", "Layer/expert geometry, top-k normalization, GPU-local exceptions, numerical oracle"],
            ["Precision", "One qualified FP8 weight/activation/accumulation recipe", "Scale/layout/rounding/overflow rules and task-quality tolerance"],
            ["Prompt / output", "4K, 8K, 16K and 32K prompts; 128, 256 and 1,024 output tokens", "Trace distribution and weighting; longer contexts are a later qualified extension"],
            ["Concurrency / load", "Active sequences 1, 4, 8, 16, 32, 64; low/target/1.2x burst/saturation", "Sponsor target RPS and burst duration; separate offered rate, clients, sequences, batch and tokens/expert"],
            ["Starting SLO objectives", "TTFT <= 2.5 s and TPOT <= 45 ms as evaluation objectives, not commitments", "Business-approved percentile, error/reject budget, quality threshold and jitter reserve"],
            ["Quality", "GPU-only outputs and expert golden vectors plus task metrics for grounded answer correctness", "Datasets, seeds, judge/version, tolerance, safety and citation/factuality checks"],
            ["Availability", "Fail closed; compare complete replica groups first; declare one device/switch/tray failure target", "RTO/RPO, surviving accepted rate, fallback policy, reserve/spare assumptions"],
            ["Controls", "GPU-only service; ATMOS host-staged hybrid; optional documented peer arm", "Same model/quality/SLO/workload/tuning method; route and synchronization evidence"],
            ["Economic hurdle", "No default savings percentage", "Sponsor-approved TCO/ROI/payback or strategic-capacity threshold before P4"],
        ],
        widths=[1.20, 3.15, 2.65],
        font_size=7.05,
    )
    add_heading(document, "5. Phase summary and decision sequence", 1)
    add_table(
        document,
        ["Phase", "Primary question", "Minimum evidence", "Indicative window", "Spending unlocked"],
        [
            ["P0", "What exactly is the control, SLO, budget, and economic hurdle?", "A/B plus GPU control", "3-4 weeks", "Primitive software/emulator work and no-NRE OEM discovery"],
            ["P1", "Can one ATMOS device execute the real expert correctly and fast enough?", "C; E required for final", "4-6 weeks", "Endpoint/driver optimization and round-trip work"],
            ["P2", "Can the complete GPU-to-ATMOS-to-GPU continuation fit each MoE layer?", "D/E; E required for final", "4-6 weeks", "Four-device switch fixture and bounded OEM POC planning"],
            ["P3", "Can four devices handle realistic fan-out, skew, faults, and shared GPUs?", "E", "6-8 weeks", "8+8 full-service fixture and production topology design"],
            ["P4", "Does the complete hybrid service beat break-even at equal SLO and availability?", "E", "8-12 weeks", "Named OEM prototype/NRE subject to P5 record"],
            ["P5", "Is there a supported, serviceable, quoteable turnkey platform path?", "F design/prototype", "8-16 weeks, parallel after P2", "Production-intent rack qualification"],
            ["P6", "Does a named rack earn the product and ROI claim?", "F", "12-16+ weeks", "Platform-specific release and launch decision"],
        ],
        widths=[0.52, 2.42, 1.30, 1.10, 1.66],
        font_size=6.95,
    )
    add_source_note(
        document,
        "Windows begin after entry artifacts and required hardware are available; they are planning "
        "ranges, not silicon, supplier, or launch commitments. P5 discovery starts at P0, but custom "
        "NRE is not released until P2 passes on qualifying evidence.",
    )

    add_heading(document, "6. P0 - Definition, GPU control, and latency budget", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Create the immutable comparison contract and determine whether any plausible expert budget remains after measured GPU-only work."],
            ["Setup", "Pinned first checkpoint/precision/serving stack on the 8-GPU development host or closest available control; representative sponsor trace or approved synthetic mix."],
            ["Actions", "Verify model quality; profile prefill/decode; decompose attention/router/shared/collective/KV/queue time; export real router/expert co-activation traces; inventory GDDR/KV/workspace; freeze TTFT/TPOT, accepted-rate, availability, power, and TCO methods."],
            ["Latency budget", "B_expert = (TPOT target - measured non-expert critical path - jitter/reserve) / MoE-layer count. Replay end-to-end because per-layer averages alone do not prove p99."],
            ["Simulation", "Build/calibrate the Bridge model for endpoint, DMA, queue, HBF/NPU service, fan-out/fan-in, switch, and placement. Preserve analytic and simulation evidence labels."],
            ["Pass", "GPU control is reproducible across three runs; quality/SLO/economic hurdle are signed; trace and memory ledger are complete; expert budget is positive and testable."],
            ["Pivot / stop", "No stable control, checkpoint/license/operator cannot be supported, or non-expert path already consumes the SLO. Fix/narrow baseline before ATMOS timing claims."],
            ["Artifacts / owner", "Baseline manifest, raw traces, golden quality set, latency-budget worksheet, router trace, memory ledger, TCO skeleton, decision record. Owner: performance lead + product/SLO owner."],
        ],
        widths=[1.25, 5.75],
        font_size=7.4,
    )

    add_heading(document, "7. P1 - Single-device expert execution", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Prove the Expert ABI, ExpertPack, numerical behavior, local memory path, NPU kernels, and one-device service curve before adding PCIe fan-out."],
            ["Setup", "One real ATMOS device when available; before silicon, use the versioned partner functional/performance emulator with explicit timing uncertainty. Compare to the GPU expert oracle."],
            ["Sweep", "Representative FFN shapes; FP8 starting format; tokens/expert 1, 2, 4, 8, 16, 32, 64; warm/cold packs; queue depths; low/target/burst arrival; temperature/power states."],
            ["Measure", "Output error/end-task delta, compile/pack reproducibility, pack bytes/scratch, useful HBF and LPDDR/SRAM traffic, read amplification, NPU/kernel time, queueing, expert/s and token/s, p50/p95/p99, power/thermal/throttle."],
            ["Negative tests", "Malformed pack/descriptor, unsupported shape/precision, buffer bounds, queue exhaustion, cancel, timeout, firmware mismatch, reset, pack corruption, stale generation."],
            ["Pass", "Quality meets P0 tolerance; the service curve sustains trace-derived target arrival plus the approved burst reserve without unbounded queue growth; no unsafe error/reset behavior. Emulator pass is provisional until E evidence."],
            ["Pivot / stop", "Small-batch service is below demand or staging prevents useful HBF delivery. Optimize layout/kernels/tiling, increase grouping, retain hot experts on GPU, narrow to prefill/offline, or stop the target shape."],
            ["Artifacts / owner", "ABI/pack conformance suite, golden results, service-curve dataset/model, power/thermal curve, issue register, evidence-class statement. Owner: NPU/compiler + firmware/runtime leads."],
        ],
        widths=[1.25, 5.75],
        font_size=7.35,
    )

    add_heading(document, "8. P2 - Complete GPU-to-ATMOS expert round trip", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Retire the largest architectural risk: serial dispatch, transfer, queue, execution, return, fan-in, synchronization, and GPU continuation at every MoE layer."],
            ["Setup", "One GPU and one ATMOS endpoint first, then 1/4/8 logical destinations as fixtures permit; synchronized tracing; host-staged path mandatory; documented peer path is an optional compared arm."],
            ["Sweep", "Top-1/top-2/top-8; tokens/expert 1-64; representative hidden/output sizes; warm and loaded queues; one and concurrent requests; CPU/NUMA-local and crossing routes; low/target/burst/saturation."],
            ["Decomposition", "Router-to-pack, pack, registration/cache, submit, DMA out, device admission, expert, DMA back, completion/fence, fan-in/combine, and measured GPU continuation. Track slowest destination and actual stream synchronization."],
            ["Pass", "On production-intent evidence, complete p99 continuation fits the signed P0 expert budget at target accepted load and end-to-end trace replay preserves TPOT with queue reserve. Host-staged is viable or a qualified peer route has a named platform dependency."],
            ["Pivot / stop", "Both supported transports miss the budget. Evaluate grouped consecutive layers, router/combine relocation, larger batches, prefill/offline-only service, or cold-expert-only offload before switch/OEM NRE."],
            ["Artifacts / owner", "Complete latency distributions/decomposition, route matrix, synchronized traces, budget replay, peer-support dependency record, signed go/pivot/stop. Owner: performance + serving/runtime leads."],
        ],
        widths=[1.25, 5.75],
        font_size=7.35,
    )
    add_callout(
        document,
        "PROGRAM STOP POINT",
        "If P2 cannot fit a credible layer budget, more devices, switch lanes, HBF capacity, or rack "
        "integration do not rescue the same fine-grained FFN-offload architecture. Change the execution "
        "boundary before funding the scale-out version of a failing critical path.",
        fill=PALE_RED,
        accent=RED,
    )

    add_heading(document, "9. P3 - Four-device tray emulator", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Prove the balanced switch domain, multi-GPU sharing, placement, top-8 fan-out/fan-in, skew, containment, reset, and degraded service expected in one compute tray."],
            ["Setup", "Four GPUs (or four controlled producers) and four ATMOS devices behind one Gen5 x16-to-4x4 domain; exact lane/NUMA map; representative switch firmware/security; BMC/power/thermal instrumentation."],
            ["Traffic tests", "One/all GPUs; one/all ATMOS devices; independent and correlated expert traffic; host-staged and any qualified peer path; target/burst/saturation; pack loading and inference overlap; upstream saturation/fairness."],
            ["Placement tests", "Replay real P0 traces; compare random/hash, co-activation-aware, affinity, hot-expert replica, and selected GPU-local policies. Measure mean/p95 destinations, queue imbalance, slowest contribution, head-of-line blocking, and fan-in."],
            ["Fault tests", "Endpoint/link loss, malformed completion, DPC/AER, downstream FLR, switch reset/firmware failure, re-enumeration, pack corruption, throttling, power event, cancellation/reset races. Repeat enough times to expose nondeterminism."],
            ["Availability", "Compare one eight-device domain conceptually/physically with two four-device domains; test complete replica, N+1 rebuild, and GPU fallback policies against signed RTO/SLO and retained capacity."],
            ["Pass", "No silent quality error or stale DMA; bounded queues/p99 at target and approved burst; failed endpoint/domain is contained; unaffected complete service follows the declared policy; reset/rejoin and RTO pass."],
            ["Pivot / stop", "Re-place/replicate, split domains, reduce destinations, change queue/admission, or narrow workload when tail, contention, reset blast radius, thermal envelope, or reserve cost fails."],
            ["Artifacts / owner", "Qualified topology/route manifest, placement report, fault matrix/logs, recovery/RTO result, switch decision, availability-capacity model. Owner: hardware/OEM + validation/reliability leads."],
        ],
        widths=[1.25, 5.75],
        font_size=7.2,
    )

    add_heading(document, "10. P4 - Complete 8+8 development service", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Determine whether a complete checkpoint and maintained service stack create end-to-end capacity or cost headroom, not merely correct kernels or nominal fit."],
            ["Setup", "Eight RTX PRO-class GPUs plus eight ATMOS devices across two balanced four-device domains; full ExpertPack; versioned vLLM integration/operator; chosen failure reserve; GPU-only control."],
            ["A/B runs", "Same checkpoint/quality/trace/TTFT/TPOT/HA. Low/target/burst/saturation offered loads; contexts 4K-32K; outputs 128/256/1,024; active sequences 1-64; prefill/decode mixes; three independent runs."],
            ["Feature/ops runs", "Continuous batching, cancellation, timeout, backpressure, prefix hit/miss, fixed distributed topology, pack rollout/rollback, drain, endpoint and switch fault, degraded service, restart/recovery, 24-hour soak."],
            ["Pass", "Quality and SLO pass; accepted service rate exceeds the predeclared cost-adjusted hurdle; no unbounded queue; chosen failure policy preserves correctness and reserve; compatibility manifest and maintenance plan are reproducible."],
            ["Pivot / stop", "If capacity fits but accepted rate, p99, quality, availability, maintainability, or fully burdened economics fail, do not infer rack value. Narrow to prefill/offline/cold experts or stop fine-grained offload."],
            ["Artifacts / owner", "Reproducible deployment, A/B raw dataset/report, load curves, power/thermal, quality, fault/update/soak results, populated TCO workbook, compatibility manifest. Owner: service lead + product/finance."],
        ],
        widths=[1.25, 5.75],
        font_size=7.3,
    )

    add_heading(document, "11. P5 - Named OEM physical and lifecycle path", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Prove that the architecture can be bought, powered, cooled, managed, serviced, secured, and supported in one exact turnkey platform."],
            ["Entry / parallelism", "RFI discovery starts at P0; detailed lab planning follows promising P2; custom NRE requires P2 pass plus approved estimate. Select one co-design leader and one enterprise control where feasible."],
            ["Engineering actions", "Review and prototype local 4+4 first, companion PCIe second; close lane/NUMA, SI/retimer, mechanics, displaced storage/NIC, power/cooling, BIOS/BMC, IOMMU/ACS/DPC, reset, peer route, service, spares, warranty, support, and supply."],
            ["Physical tests", "Enumeration/address windows; all-device DMA; SI/error counters; sustained mixed workload thermal/power; fan/CDU response; service replacement; firmware/BIOS/driver update/rollback; endpoint/switch/root faults."],
            ["Pass", "Vendor returns and owns a supported design record, measurable route, field lifecycle, production-intent BOM/quote, qualification plan, supply/lifecycle terms, and platform model that remains within P2/P4 latency, power, availability, and economics envelopes."],
            ["Pivot / stop", "Use approved companion tray or next ranked provider when local integration fails. Stop a platform with no supported route, no thermal/mechanical envelope, no isolation/reset ownership, or no accountable field support."],
            ["Artifacts / owner", "Configuration/design record, CAD/lane/NUMA, SI/thermal/power, firmware/service contracts, BOM/quote/NRE, risk register, lab result, qualification schedule. Owner: OEM/hardware lead + procurement."],
        ],
        widths=[1.25, 5.75],
        font_size=7.3,
    )

    add_heading(document, "12. P6 - Independent rack-scale product qualification", 1)
    add_table(
        document,
        ["Element", "Action and required result"],
        [
            ["Purpose", "Earn a platform-specific product, support, and ROI statement against the actual rack control. Nothing from the development vehicle is transferred without remeasurement."],
            ["Setup", "Production-intent hybrid rack and comparable GPU-only control; same model/service/SLO/HA; full monitoring; production network, power/cooling, orchestration, support, security, and cost boundary."],
            ["Qualification", "Repeat workload/load sweep, placement, quality, latency, power, thermal, all declared faults, update/rollback, replacement/rebuild, security/tenancy, fresh install, and 72-hour representative soak. Rehearse support escalation."],
            ["Economics", "Populate three-year TCO per accepted request/token with quotes, measured service/power, complete reserve, spares, support, licenses, network/rack, amortized NRE, migration, and sensitivity."],
            ["Pass", "Hybrid beats the predeclared hurdle at equal quality/SLO/availability; support and supply are ready; no critical unresolved safety/security/reliability issue; new operators reproduce installation and operation from the package."],
            ["Pivot / stop", "No product claim on the failing platform. Release only a successful named platform, or define a new baseline for a narrower prefill/offline/asynchronous service."],
            ["Artifacts / owner", "Independent qualification report, raw benchmark package, TCO/ROI workbook, security/reliability/support readiness, production compatibility manifest, release/no-release decision. Owner: joint qualification board."],
        ],
        widths=[1.25, 5.75],
        font_size=7.3,
    )

    add_heading(document, "13. Standard sweep matrix", 1)
    add_table(
        document,
        ["Axis", "Initial points", "Applied phases", "Reporting rule"],
        [
            ["Tokens per expert", "1, 2, 4, 8, 16, 32, 64", "P1-P3", "Service curve by real shape/precision; no average-only result"],
            ["Top-k / destinations", "Top-1, top-2, top-8; 1, 2, 4, 8 destinations", "P2-P4", "Report slowest contribution and fan-in separately"],
            ["Context / output", "4K, 8K, 16K, 32K; 128, 256, 1,024 output", "P0, P4, P6", "Add 64K/128K only after model/runtime qualification"],
            ["Active sequences", "1, 4, 8, 16, 32, 64", "P0, P4, P6", "Independent from offered RPS and batch size"],
            ["Load", "Low, target, 1.2x burst, increasing to saturation", "All performance phases", "Stop unsafe overload; preserve rejects/timeouts/backlog"],
            ["Queue state", "Empty/warm, steady target, loaded, recovery after burst", "P1-P4", "Report oldest-work age and stability, not depth alone"],
            ["Transport", "Host-staged baseline; documented peer arm; root/NUMA local/crossing", "P2-P6", "Route counters and synchronization required"],
            ["Temperature", "Cold start, thermal steady state, qualified high inlet, throttle", "P1, P3-P6", "Accepted rate and errors at each state"],
            ["Failure", "Endpoint, link, switch, reset, pack, thermal/power, operator/service", "P1, P3-P6", "Repeat, record lost/retried work and RTO"],
        ],
        widths=[1.25, 2.20, 1.35, 2.20],
        font_size=7.25,
    )
    add_heading(document, "14. Decision package returned by every phase", 1)
    add_table(
        document,
        ["Artifact", "Required content"],
        [
            ["Configuration manifest", "All hardware/software/model/workload/SLO/failure/evidence revisions and known deviations"],
            ["Raw evidence bundle", "Traces, logs, counters, telemetry, outputs, power/thermal, errors, run timing, and analysis scripts"],
            ["Quality and service report", "Task/operator accuracy, offered/admitted/accepted work, TTFT/TPOT distributions, queue stability, confidence intervals"],
            ["Reliability report", "Fault cases, containment, ambiguity handling, retries, lost work, RTO/RPO, degraded capacity, update/rollback"],
            ["Economics / support update", "Current quotes/assumptions, complete reserve, TCO sensitivity, support/supply owners and unresolved commercial risk"],
            ["Decision record", "Pass, conditional pass with dated closure, pivot with funded next experiment, or stop; signers and dissent/open questions"],
            ["Compatibility delta", "What newly became supported, remained unsupported, changed, or invalidated from the prior phase manifest"],
        ],
        widths=[1.45, 5.55],
        font_size=7.55,
    )
    add_heading(document, "15. First 90 days", 1)
    add_table(
        document,
        ["Window", "Technical actions", "Partner / business actions", "Exit"],
        [
            ["Days 0-30", "Name owners; freeze first checkpoint/precision/SLO; reproduce GPU control; export traces; define ABI/pack/transport schemas; create manifest/evidence repository", "Issue common RFI to Giga, Supermicro, Penguin and one enterprise control; request switch FAE maps/boards; confirm NPU emulator/silicon evidence and dates", "P0 draft package, reviewed interfaces, RFI acknowledgement, blocker/assumption register"],
            ["Days 31-60", "Complete GPU baseline; implement golden expert harness and initial pack compiler; build transport queue/buffer test API; run analytic/simulation placement and latency; instrument end-to-end IDs", "Hold first OEM and switch workshops; obtain preliminary lane/power/cooling/support returns and budgetary quotes; define lab access", "P0 pass/conditional decision; P1 provisional service curves; candidate topology and vendor gaps"],
            ["Days 61-90", "Run P1 and provisional/real P2 as evidence permits; execute error/reset tests; calibrate simulator; choose host-staged implementation; publish latency go/pivot/stop", "Down-select two OEM routes only if P2 remains credible; close prototype configuration, NRE range, support owner, and P3 fixture plan", "Signed P2 decision; funded P3/P5 plan or documented architecture pivot/stop"],
        ],
        widths=[0.95, 2.75, 2.45, 0.85],
        font_size=6.8,
    )
    add_heading(document, "16. Gate review format", 1)
    add_table(
        document,
        ["Decision", "Meaning", "Required action"],
        [
            ["Pass", "All mandatory criteria pass at the required evidence class", "Release only the explicitly unlocked scope and archive package"],
            ["Conditional pass", "Decision is not affected by bounded closures with owner/date", "Track closure; no product claim until mandatory product evidence closes"],
            ["Pivot", "Hypothesis boundary changes but a specific alternative remains plausible", "Freeze new hypothesis, budget, experiment, and invalidated prior evidence"],
            ["Stop", "No credible route within service, support, schedule, or economics constraints", "Close spend, retain evidence, state what future change could reopen"],
        ],
        widths=[1.20, 3.05, 2.75],
        font_size=7.5,
    )
    add_callout(
        document,
        "PART V EXIT",
        "The POC program is complete only when P6 returns a reproducible platform-specific service, "
        "support, and economics package, or an earlier phase returns a documented no-fit result. A demo, "
        "capacity fit, vendor meeting, or simulated timing result is not program completion.",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def add_part_vi(document: Document) -> None:
    add_part_opener(
        document,
        "VI",
        "Program operation, risk, and release",
        "Run one cross-functional product program with phase-specific evidence, controlled claims, and an explicit no-fit outcome.",
    )
    add_heading(document, "1. Operating model", 1)
    add_body(
        document,
        "The program is organized around one versioned product contract and one evidence repository, not "
        "independent hardware, compiler, runtime, OEM, and benchmark projects. Each workstream owns a "
        "release artifact and participates in the phase whose decision depends on it. The architecture "
        "lead integrates decisions but cannot self-approve performance, security, economics, or release."
    )
    add_table(
        document,
        ["Workstream", "Accountable output", "Primary phases", "Required partners / functions"],
        [
            ["Product / service contract", "Checkpoint/use case, quality, TTFT/TPOT, HA, economic hurdle, approved claims", "P0, P4, P6", "Sponsor, customer/account, finance, legal"],
            ["Model / compiler / NPU", "Expert ABI, compiler, ExpertPack, kernels, golden outputs, service curves", "P1-P4", "NPU partner, model owners, quality"],
            ["Endpoint / firmware / driver", "PCIe/DMA/queues, isolation, telemetry, reset, update, secure lifecycle", "P1-P6", "Silicon/FW, OS, security, OEM"],
            ["Serving integration", "vLLM/PyTorch backend, placement, batching, fan-in, fallback, feature envelope", "P0-P6", "vLLM/PyTorch route, GPU platform"],
            ["Platform / SRE", "Kubernetes allocation, operator, packs, rollout/rollback, health, recovery, support runbook", "P3-P6", "Linux/Kubernetes vendor, customer SRE"],
            ["Hardware / OEM", "Switch/board, lane/mechanics/power/cooling, BIOS/BMC, design record, BOM and supply", "P2-P6", "Switch FAE, OEM/ODM, thermal/SI, procurement"],
            ["Performance / validation", "Evidence runner, synchronized telemetry, controls, statistics, faults, independent reports", "All", "Lab, reliability, quality, finance"],
            ["Security / reliability", "Threat model, DMA/isolation tests, failure policy, RTO/RPO, readiness signoff", "P1-P6", "Security, privacy, safety, OEM support"],
            ["Economics / commercial", "Quotes, TCO/ROI, NRE allocation, support/supply terms, feasible region", "P0, P4-P6", "Finance, product, procurement, OEM"],
        ],
        widths=[1.35, 2.70, 1.00, 1.95],
        font_size=6.95,
    )
    add_heading(document, "2. Minimum staffed functions", 1)
    add_body(
        document,
        "These are functions that need committed capacity and backups, not an approved headcount or a "
        "promise that each row is one person. Before P3, the sponsor should review a bottom-up plan based "
        "on actual silicon, partner scopes, OEM NRE, and feature commitments."
    )
    add_table(
        document,
        ["Function", "Why it cannot be part-time only", "Peak engagement"],
        [
            ["Program product owner / chief architect", "Owns service contract, tradeoffs, decision record, and cross-workstream closure", "All phases"],
            ["NPU/compiler/model engineers", "Real expert layouts, kernels, quality, and local-memory bottlenecks are core feasibility", "P1-P4"],
            ["Firmware/driver/transport engineers", "DMA, ordering, reset ambiguity, isolation, and telemetry are product-critical", "P1-P6"],
            ["Serving/vLLM engineers", "Continuous batching, placement, fan-in, fallback, and upstream maintenance define service viability", "P0-P6"],
            ["Hardware/platform/OEM engineers", "Lane, SI, mechanics, thermal, BMC/BIOS, switch, and service require direct ownership", "P2-P6"],
            ["Performance/validation/reliability", "Independent controls, tracing, statistics, fault injection, and soak cannot be deferred", "All phases"],
            ["Platform/SRE/security", "Operator, recovery, manifest, tenancy, updates, runbooks, and release response are required", "P3-P6"],
            ["Product finance/procurement/account", "Hurdle, quotes, NRE, demand, support, and commercial milestones prevent paper ROI", "P0, P4-P6"],
        ],
        widths=[1.55, 4.35, 1.10],
        font_size=7.35,
    )
    add_heading(document, "3. Critical path and parallel work", 1)
    add_table(
        document,
        ["Critical dependency", "Must finish before", "Can proceed in parallel", "Do not pre-commit"],
        [
            ["Frozen control, trace, SLO, expert budget", "Any pass/fail performance decision", "ABI/schema design and no-NRE RFIs", "Headline service rate or economics"],
            ["P1 useful expert service curve", "Confidence in device count and placement", "Driver/runtime harness and OEM discovery", "Large switch/card or rack design"],
            ["P2 complete round-trip latency", "P3 fixture and custom OEM NRE", "RFI workshops, simulation, software conformance", "Production topology or peer-DMA dependency"],
            ["P3 fault/placement/switch result", "P4 availability and production topology lock", "P5 physical design within bounded options", "Single/dual-domain and reserve policy"],
            ["P4 end-to-end service and economics", "Product investment beyond qualification", "P5 closure and support planning", "Fleet savings, release SKU, launch claim"],
            ["P5 named design/quote/support", "P6 rack build and rack TCO", "Control-rack readiness and release tooling", "Vendor selection from paper score"],
            ["P6 independent qualification", "Platform-specific release", "Sales/support enablement draft", "Cross-platform or future-generation claims"],
        ],
        widths=[1.65, 1.65, 2.25, 1.45],
        font_size=7.25,
    )
    add_heading(document, "4. Program risk register", 1)
    add_table(
        document,
        ["Risk", "Leading trigger", "Impact", "Response / decision", "Owner"],
        [
            ["Per-layer latency fails", "P2 p99 exceeds signed expert budget or queues grow at target", "Architecture not viable for synchronous decode", "Pivot execution boundary before any scale/rack NRE", "Performance + architecture"],
            ["Small-batch NPU/HBF efficiency", "P1 tokens/expert 1-8 show low useful HBF/NPU utilization", "Cannot meet enterprise low-rate/tail workloads", "Kernel/layout/tiling; larger grouping; prefill/offline or no-fit", "NPU/compiler"],
            ["Top-8 fan-out and skew", "High destinations/token, slowest-device tail, queue imbalance", "p99 and buffer pressure erase value", "Co-activation placement, replicas, GPU-local hot experts, admission", "Runtime/placement"],
            ["vLLM integration becomes a fork", "Growing patch set, release lag, unsupported features", "High maintenance/security cost", "Upstream generic hooks; narrow feature envelope; funded backport owner", "Serving lead"],
            ["Incomplete service after failure", "Unique expert loss or surviving group below demand", "Outage or incorrect quality", "Complete replicas, N+1/GPU fallback only after measured RTO/cost", "Reliability/product"],
            ["Switch blast radius / contention", "DPC/reset affects all devices; uplink queues saturate", "Availability and tail failure", "Two four-device domains, containment fixes, or direct topology", "Hardware/platform"],
            ["No supported OEM route", "No lanes/mechanics/thermal/BIOS owner or warranty support", "Cannot productize despite lab success", "Companion tray, next provider, or asynchronous service", "OEM lead"],
            ["Thermal/power envelope fails", "Throttle, connector/PDB/fan/CDU limit under mixed load", "Reduced rate or platform redesign", "Lower population/power, better cooling, companion tray, reject platform", "Hardware/OEM"],
            ["Peer DMA unavailable", "IOMMU/topology/driver blocks supported route", "Host staging may miss latency", "Keep host-staged baseline; only depend on peer after platform qualification", "Transport/OEM"],
            ["Security/isolation gap", "Out-of-range DMA, stale handle, cross-tenant access, unsigned pack/FW", "Release blocker", "Fail closed; redesign buffers/queues/update; no benchmark bypass", "Security + FW/driver"],
            ["BOM/support premium erases ROI", "Configured quote/reserve/NRE exceeds feasible region", "No business case", "Optimize subsystem/volume, different OEM, narrower value mechanism, stop", "Product/finance"],
            ["GPU platform improves faster", "New GPU memory/rate/price changes control", "ATMOS advantage shrinks", "Rebaseline at P4/P6 and every major platform generation", "Product/performance"],
            ["Silicon/partner schedule slips", "Emulator, samples, compiler, or firmware milestones miss", "Idle team or premature claims", "Evidence-class roadmap, reversible tooling, dated dependencies, scope reset", "Program owner"],
            ["Insufficient customer pull", "No trace, funded POC, owner, or buying path", "Technically valid but no product", "One sponsor/use case; separate interest, evaluation, design intent, and order", "Product/account"],
        ],
        widths=[1.30, 2.10, 1.45, 2.20, 0.95],
        font_size=6.45,
    )
    add_heading(document, "5. Security, compliance, and support release review", 1)
    add_table(
        document,
        ["Review", "Mandatory closure before release"],
        [
            ["Threat model", "Trusted boundaries, DMA/IOMMU, firmware, model IP, pack supply chain, management plane, tenancy, denial of service, physical service"],
            ["Secure lifecycle", "Signed firmware/packs, key/certificate ownership, update/rollback, vulnerability response, SBOM, support duration, EOL"],
            ["Data handling", "Activation/result sensitivity, scratch/buffer clearing, logs/traces, customer model custody, crash dump and RMA policy"],
            ["Platform hardening", "Secure boot, BIOS/BMC access, IOMMU/ACS/DPC, least privilege, network segmentation, audit, drift detection"],
            ["Reliability", "Declared failure policy, RTO/RPO, containment, rebuild, spares, replacement, degraded capacity, operational drills"],
            ["Support", "Single intake route, triage matrix across Sandisk/OEM/GPU/switch/NPU/software, diagnostics bundle, escalation and SLA"],
            ["Compliance", "Applicable industry/customer controls, export/licensing, open-source obligations, privacy, accessibility of operational tooling"],
        ],
        widths=[1.40, 5.60],
        font_size=7.45,
    )
    add_heading(document, "6. Change control and claim management", 1)
    add_bullets(document, [
        "Treat checkpoint, precision, expert operator, placement/failure policy, topology, device/FW, switch, BIOS, kernel/driver, serving stack, SLO, workload, and cost baseline as controlled configuration fields.",
        "A change creates an impact record identifying which conformance, performance, fault, security, thermal, economics, and support evidence is invalidated. Reuse is explicit; identical API names are not proof of compatibility.",
        "Label every result as analytic, simulated, emulated, prototype, production-silicon, or OEM-rack evidence. Label vendor statements, targets, quotes, measurements, and assumptions distinctly.",
        "Approve claims by configuration ID and evidence package. Do not say 'NVIDIA supported,' 'AMD supported,' or 'turnkey OEM supported' when only one named configuration passed.",
        "Archive superseded manifests, raw data, analyses, decisions, and approved wording. Restricted partner/customer data remains in controlled storage linked by evidence ID.",
    ])
    add_heading(document, "7. Product definition of done", 1)
    add_table(
        document,
        ["Dimension", "A release exists only when..."],
        [
            ["Function / quality", "The named checkpoint/precision/operator passes golden and end-task quality, including fallback and update paths"],
            ["Performance", "The named service meets accepted-throughput, TTFT, TPOT, queue stability, power/thermal, and degraded-mode envelopes"],
            ["Availability", "Complete-service capacity survives the declared failure, and reset/rebuild/replacement meet RTO/RPO without silent error"],
            ["Security", "DMA/isolation, firmware/pack lifecycle, tenancy, hardening, vulnerability response, and negative tests are approved"],
            ["Operations", "Fresh install, discovery, allocation, packs, rollout/rollback, monitoring, incident, drain, recovery, and upgrade are reproducible"],
            ["Platform", "One exact OEM configuration has lane/mechanical/power/cooling/BIOS/BMC/service evidence and a maintained compatibility manifest"],
            ["Commercial", "BOM, quote, NRE, supply/lifecycle, spares, support SKU, TCO/ROI, demand, and approved claims close the product hurdle"],
            ["Support", "Customer-facing docs, diagnostics, ownership/escalation, training, release notes, limitations, and EOL policy are ready"],
        ],
        widths=[1.35, 5.65],
        font_size=7.45,
    )
    add_heading(document, "8. Decision requested", 1)
    add_callout(
        document,
        "APPROVE NOW",
        "Approve P0-P2 engineering, the common software/evidence contracts, and parallel no-NRE OEM/switch "
        "discovery. Name accountable leads for product/SLO, NPU/compiler, firmware/driver, serving/runtime, "
        "performance, OEM hardware, platform/SRE, security, and economics/procurement.",
        fill=PALE_TEAL,
        accent=TEAL,
    )
    add_callout(
        document,
        "HOLD",
        "Do not approve a production rack direction, custom eight-device switch/card, peer-DMA dependency, "
        "platform NRE, fleet-savings claim, or broad NVIDIA/AMD/OEM product statement until its preceding "
        "phase returns the required evidence and owners.",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )
    add_body(
        document,
        "The immediate management return is a signed P2 go/pivot/stop package and two comparable OEM design "
        "responses. A positive P2 unlocks the four-device tray and bounded OEM prototype. A negative P2 is "
        "also a useful result: it redirects the architecture before switch, tray, and rack investment."
    )

    add_page_break(document)
    add_kicker(document, "APPENDIX A | DECISIONS TO FREEZE")
    add_heading(document, "Open decision register", 1)
    add_table(
        document,
        ["Decision", "Starting position", "Owner / needed by", "Evidence that may change it"],
        [
            ["First checkpoint", "One Qwen3-235B-A22B-class checkpoint", "Product/model / P0", "Sponsor workload, license, operator/precision support, baseline stability"],
            ["First precision", "One qualified FP8 format", "Model/NPU / P0", "Quality and compiler/kernel evidence"],
            ["Endpoint", "Gen5 x4 proposed", "Silicon/FW / P0", "Actual silicon capability and board signal-integrity plan"],
            ["Transport", "Host-staged required; peer optional", "Runtime/OEM / P2", "Complete supported-route latency and isolation"],
            ["Switch topology", "x16-to-4x4; two domains for eight devices", "Hardware / P3", "All-device traffic, containment, Gen6 availability/value"],
            ["Availability", "Complete service groups as conservative baseline", "Product/reliability / P3", "N+1/GPU fallback/remote replica RTO, retained rate, capacity, cost"],
            ["Production integration", "OEM-local first; companion PCIe second", "OEM lead / P5", "Returned mechanics/thermal/lane/service and P2 budget"],
            ["Primary OEM", "Common RFI to top co-design and enterprise candidates", "Procurement/OEM / P5", "Design record, lab access, quote, support, supply, POC"],
            ["AMD expansion", "After NVIDIA development baseline", "Product/runtime / after P4", "Customer pull, ROCm/vLLM readiness, named AMD/OCP fixture"],
            ["Product hurdle", "Set before hybrid result; no default percentage", "Sponsor/finance / P0", "Strategic capacity need, demand, approved financial policy"],
        ],
        widths=[1.35, 2.55, 1.35, 1.75],
        font_size=7.2,
    )

    add_kicker(document, "APPENDIX B | WORKING VOCABULARY")
    add_heading(document, "Key terms", 1)
    add_table(
        document,
        ["Term", "Meaning in this proposal"],
        [
            ["Accepted work", "Request or output token completed within the frozen quality and latency SLO; excludes error, timeout, overload cancellation, and late completion"],
            ["ATMOS service group", "Devices, switch/root path, packs, placement, runtime, and reserve that together provide one complete declared service"],
            ["Complete service capacity", "Capacity and rate that can execute every required expert after the declared failure; not nominal bytes across unique shards"],
            ["ExpertPack", "Immutable compiled and signed set of ATMOS-owned expert weights, layouts, scales, kernels, metadata, compatibility, and golden outputs"],
            ["Placement epoch", "Immutable mapping of model/layer/expert to ready owners and replicas used for the full lifetime of an in-flight request"],
            ["Host-staged", "GPU/ATMOS exchange through registered pinned host buffers using stock supported APIs; the correctness baseline"],
            ["Peer DMA", "Documented platform route that moves between GPU-visible and ATMOS buffers without an application CPU copy; optional until qualified"],
            ["Useful HBF bandwidth", "Bytes that contribute to real expert execution at the NPU, distinguished from nominal interface rate and staging amplification"],
            ["Turnkey", "Bound hardware, software, model artifacts, management, deployment, security, support, spares, supply, and qualification under a release manifest"],
            ["Compatibility manifest", "Configuration ID binding all supported hardware, firmware, platform, software, model, topology, workload envelope, and limitations"],
        ],
        widths=[1.45, 5.55],
        font_size=7.55,
    )

    add_kicker(document, "APPENDIX C | PUBLIC SOURCE REGISTER")
    add_heading(document, "Primary public references", 1)
    add_body(
        document,
        "Accessed 30 September 2026. Public references establish published model/platform/software/vendor "
        "context only. They do not establish ATMOS compatibility, measured performance, price, availability, "
        "partnership, or purchase commitment. Internal device values remain planning inputs pending "
        "revisioned engineering evidence."
    )
    add_table(
        document,
        ["Ref", "Source / scope", "URL", "Use and boundary"],
        [
            ["S1", "QwenLM - Qwen3", "https://github.com/QwenLM/Qwen3", "Model family and public configuration context; pin exact checkpoint/license"],
            ["S2", "DeepSeek - DeepSeek-V3", "https://github.com/deepseek-ai/DeepSeek-V3", "Stress-model architecture context; not a first-release commitment"],
            ["S3", "vLLM documentation", "https://docs.vllm.ai/", "Serving behavior/integration reference; no ATMOS backend implied"],
            ["S4", "Kubernetes device plugins", "https://kubernetes.io/docs/concepts/extend-kubernetes/compute-storage-net/device-plugins/", "Allocation/health API; application recovery remains ATMOS-owned"],
            ["S5", "Kubernetes Dynamic Resource Allocation", "https://kubernetes.io/docs/concepts/scheduling-eviction/dynamic-resource-allocation/", "Topology/resource evolution reference; qualify chosen release"],
            ["S6", "NVIDIA GPUDirect RDMA", "https://docs.nvidia.com/cuda/gpudirect-rdma/", "Peer-memory requirements; not automatic platform/ATMOS support"],
            ["S7", "AMD ROCm documentation", "https://rocm.docs.amd.com/", "AMD software/peer-path reference; named platform qualification required"],
            ["S8", "Microchip Switchtec PFX", "https://www.microchip.com/en-us/products/interface-and-connectivity/pcie-switches/switchtec-pfx-pcie-fanout-switches", "PCIe switch family inquiry; no ATMOS compatibility claim"],
            ["S9", "Broadcom PEX89000 family", "https://www.broadcom.com/products/pcie-switches-bridges/pcie-switches/pex89000", "Alternative Gen5 switch inquiry; exact part/board/firmware required"],
            ["S10", "NVIDIA DGX GB300", "https://docs.nvidia.com/dgx/dgxgb300-user-guide/introduction-to-dgx-gb300.html", "Published rack context; does not expose a generic ATMOS bay/path"],
            ["S11", "AMD Advancing AI / Helios", "https://www.amd.com/en/corporate/events/advancing-ai.html", "AMD rack-scale direction; ATMOS remains a separately qualified PCIe endpoint"],
            ["S12", "Giga Computing GIGAPOD", "https://www.gigabyte.com/Solutions/gigapod", "Turnkey/multi-accelerator/ORV3/cooling/customization public scope"],
            ["S13", "Supermicro NVIDIA systems / DCBBS", "https://www.supermicro.com/en/accelerators/nvidia", "Rack integration, cooling, validation, services and portfolio public scope"],
            ["S14", "Penguin Solutions OriginAI", "https://www.penguinsolutions.com/solutions/ai/originai", "Multi-vendor design-build-deploy-manage and lifecycle public scope"],
            ["S15", "Lenovo Hybrid AI", "https://www.lenovo.com/us/en/servers-storage/solutions/ai/", "Validated hybrid AI platforms, services and Neptune cooling public scope"],
            ["S16", "Dell AI solutions / AI Factory", "https://www.dell.com/en-us/dt/solutions/artificial-intelligence/index.htm", "NVIDIA turnkey and AMD open-architecture public scope"],
            ["S17", "Cisco Secure AI Factory", "https://www.cisco.com/site/us/en/solutions/artificial-intelligence/secure-ai-factory/index.html", "Full-stack AI POD, compute/network/security/observability public scope"],
            ["S18", "HPE AI Factory", "https://www.hpe.com/us/en/ai-factory.html", "Turnkey production foundation, lab, operations and lifecycle public scope"],
            ["S19", "Tenstorrent TT-Metal / TT-NN", "https://github.com/tenstorrent/tt-metal", "Public NPU programming/model reference; deprecated legacy vLLM fork noted; no ATMOS memory/controller support implied"],
            ["S20", "Tenstorrent vLLM TT plugin", "https://github.com/tenstorrent/vllm-tt-plugin", "Current standalone vLLM plugin route and feature constraints; full TT backend is not the proposed hybrid integration"],
            ["S21", "Tenstorrent inference server", "https://github.com/tenstorrent/tt-inference-server", "Public deployment, benchmarking and evaluation workflows; not an ATMOS release or support statement"],
        ],
        widths=[0.45, 1.85, 2.90, 1.80],
        font_size=6.15,
    )
    add_heading(document, "Evidence boundary", 2)
    add_callout(
        document,
        "FINAL CLAIM TO EARN",
        "For one named checkpoint, workload, availability policy, and OEM rack configuration, the qualified "
        "ATMOS E3.S expert tier increases accepted MoE service capacity or reduces fully burdened comparable "
        "infrastructure cost while preserving model quality, TTFT, TPOT, security, support, and lifecycle.",
        fill=PALE_TEAL,
        accent=TEAL,
    )


def validate(document: Document) -> None:
    content = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    required = [
        "ATMOS E3.S Expert Tier",
        "Reference product architecture",
        "Service stack impact",
        "Hardware and OEM productization",
        "Economic case and product threshold",
        "Integrated POC and qualification plan",
        "Program operation, risk, and release",
        "Expert Execution ABI",
        "Sandisk and Tenstorrent integration boundary",
        "Tenstorrent vLLM TT plugin",
        "Giga Computing",
        "Common OEM RFI",
        "First workload card",
        "Starting SLO objectives",
        "P2 - Complete GPU-to-ATMOS expert round trip",
        "First 90 days",
        "Product definition of done",
        "Primary public references",
        "FINAL CLAIM TO EARN",
        "complete continuation path",
    ]
    for item in required:
        if item not in content:
            raise RuntimeError(f"Missing required content: {item}")
    forbidden_patterns = {
        "proposal V2/V3": r"(?:proposal|document|revision|version)\s+(?:V2|V3|2|3)\b",
        "derived from an earlier version": r"(?:derived|extended|revised)\s+from\s+(?:V2|V3|an earlier|a previous)",
        "review comment": r"review comment",
        "assessment disposition": r"assessment disposition",
        "previous proposal": r"previous proposal",
    }
    for item, pattern in forbidden_patterns.items():
        if re.search(pattern, content, flags=re.IGNORECASE):
            raise RuntimeError(f"Standalone proposal contains forbidden history reference: {item}")
    if len(document.inline_shapes) < 7:
        raise RuntimeError("Expected at least seven original diagrams")


def main() -> None:
    assets = build_assets()
    document = Document()
    configure_document(document)
    add_cover(document, assets)
    add_executive_section(document)
    add_part_i(document, assets)
    add_part_ii_foundation(document, assets)
    add_part_iii(document, assets)
    add_part_iii_tail(document, assets)
    add_part_iv(document, assets)
    add_part_v(document, assets)
    add_part_vi(document)
    validate(document)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    validate(reopened)
    print(
        f"Wrote {OUTPUT} | paragraphs={len(reopened.paragraphs)} "
        f"tables={len(reopened.tables)} figures={len(reopened.inline_shapes)} "
        f"sections={len(reopened.sections)}"
    )


if __name__ == "__main__":
    main()